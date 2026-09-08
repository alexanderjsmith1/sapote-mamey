#!/usr/bin/env python3
"""insert_tree_into_report.py — append the phylogenomic tree to a strain's Word report.

"Each word doc can have the MLSA trees at a minimum." This reads the strain->tree map
(strain_to_MLSA_tree_map.csv) and appends, to a COPY of each strain report, the best
rendered tree for that strain's family (the pruned 138-SCG core-genome backbone, which is
derived from the full-pool 5-locus MLSA screen — the two-tier method) with a claim-safe
caption. It never overwrites the canonical thesis doc; it writes to --outdir.

Why the pruned core-genome and not the raw MLSA image: the full-pool MLSA (155-292 tips)
is not yet rendered and is unreadable at page size; the pruned backbone IS rendered,
readable, and derived from that MLSA. Swap the map's PNG column to use a different image.

CLAIM SAFETY (baked into every caption): the tree strengthens topology; whole-genome ANI
delimits species; "candidate novel" is a prior pending polyphasic work; judgment deferred.

Usage:
  python insert_tree_into_report.py --strain AS-XXX [--outdir DIR]
  python insert_tree_into_report.py --all           [--outdir DIR]
"""
from __future__ import annotations

import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402
import argparse, csv, glob, os, re, shutil
import os

ROOT = os.environ.get("SAPOTE_WORKSPACE_ROOT", os.getcwd())
MAP = f"{ROOT}/sessions/Amber/Deliverables/strain_to_MLSA_tree_map.csv"
REPORTS = f"{ROOT}/July 25 the Developer or User thesis bee paper/Strain_Level_Capture_2026-07-25/reports/docx"
OUTDIR_DEFAULT = f"{ROOT}/sessions/Amber/Deliverables/strain_reports_with_tree"

_TREE_KIND = {"pruned": "pruned 138-SCG core-genome backbone (derived from the full-pool MLSA screen)",
              "family": "138-SCG core-genome tree"}


def load_map():
    rows = {}
    with open(MAP) as fh:
        for r in csv.DictReader(fh):
            rows[r["strain"]] = r
    return rows


def report_for(strain):
    hits = glob.glob(f"{REPORTS}/*{strain}_*.docx") + glob.glob(f"{REPORTS}/{strain}_*.docx")
    return hits[0] if hits else None


def tree_kind(png):
    b = os.path.basename(png).lower()
    return _TREE_KIND["pruned"] if "pruned" in b else _TREE_KIND["family"]


def insert(strain, row, outdir):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    rep = report_for(strain)
    if not rep:
        return f"{strain}: no strain report docx"
    png = row.get("best_rendered_tree_png", "")
    if not png or png.startswith("<"):
        return f"{strain}: no rendered tree (family={row.get('family_authoritative_MLSA')})"
    png_abs = os.path.join(ROOT, png)
    if not os.path.exists(png_abs):
        return f"{strain}: tree PNG missing on disk ({png})"

    os.makedirs(outdir, exist_ok=True)
    out = os.path.join(outdir, f"{strain}_strain_report_with_tree.docx")
    shutil.copy(rep, out)
    d = docx.Document(out)

    d.add_page_break()
    h = d.add_heading("Phylogenomic placement", level=1)
    fam = row.get("family_authoritative_MLSA", "?")
    sub = d.add_paragraph()
    run = sub.add_run(f"{strain} within {fam}")
    run.bold = True
    run.font.size = Pt(12)

    pic = d.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(png_abs, width=Inches(6.2))

    cap = d.add_paragraph()
    cr = cap.add_run(
        f"Figure. {tree_kind(png)} for {fam}; {strain} is placed within the sampled panel. "
        f"The tree strengthens topology only — whole-genome ANI delimits species (~95-96%), "
        f"'candidate novel' is a prior pending polyphasic work (dDDH/AAI/POCP + chemotaxonomy), "
        f"and any fragmented assembly's branch lengths are not trustworthy even where its tip "
        f"position holds. Class-level hypotheses; judgment deferred; no bioactivity/structure claims.")
    cr.italic = True
    cr.font.size = Pt(8)
    cr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    d.save(out)
    return f"{strain}: appended {os.path.basename(png)} -> {os.path.relpath(out, ROOT)}"


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--strain")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--outdir", default=OUTDIR_DEFAULT)
    a = ap.parse_args(argv)
    rows = load_map()
    targets = ([s for s in rows if rows[s].get("has_strain_report_docx") == "yes"]
               if a.all else [a.strain])
    for s in targets:
        if s not in rows:
            emit(f"  {s}: not in map"); continue
        emit("  " + insert(s, rows[s], a.outdir))


if __name__ == "__main__":
    main()
