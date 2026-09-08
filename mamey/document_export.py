"""Render governed Sapote Markdown to DOCX and PDF from one canonical model.

The module is intentionally post-analysis.  It parses and validates authored
Markdown with :mod:`mamey.sapote_markdown`, then presents the same immutable
``DocumentModel`` through both output backends.  It does not rewrite claims,
infer missing evidence, or change scientific content.
"""
from __future__ import annotations

import sys

import argparse
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import tempfile
from typing import Any, Iterable

from .sapote_markdown import Block, DocumentModel, parse_path, validate_model


NAVY = "082A45"
TEAL = "0B8793"
GOLD = "E6A623"
CREAM = "F6F1E7"
PALE_TEAL = "E8F4F3"
PALE_GOLD = "FFF4D6"
PALE_CORAL = "FDEBE6"
CORAL = "D45B42"
INK = "17242E"
MUTED = "506270"
WHITE = "FFFFFF"


class DocumentExportError(RuntimeError):
    """Raised when an output contract cannot be met completely."""


@dataclass(frozen=True)
class ExportReceipt:
    status: str
    schema_version: str
    input_path: str
    input_sha256: str
    model_sha256: str
    outputs: list[dict[str, Any]]
    warnings: list[str]
    created_utc: str


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _plain(text: str) -> str:
    """Remove the small inline Markdown subset without altering its words."""
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text


def _iter_inline(text: str) -> Iterable[tuple[str, bool, bool, bool]]:
    """Yield text, bold, italic, code runs for the governed inline subset."""
    token = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))")
    cursor = 0
    for match in token.finditer(text):
        if match.start() > cursor:
            yield text[cursor : match.start()], False, False, False
        value = match.group(0)
        if value.startswith("`"):
            yield value[1:-1], False, False, True
        elif value.startswith("**"):
            yield value[2:-2], True, False, False
        else:
            yield value[1:-1], False, True, False
        cursor = match.end()
    if cursor < len(text):
        yield text[cursor:], False, False, False


def _asset_path(model: DocumentModel, source: str) -> Path:
    return (Path(model.source_path).resolve().parent / source).resolve()


def _docx_dependencies() -> None:
    try:
        import docx  # noqa: F401
        import lxml  # noqa: F401
    except ImportError as exc:
        raise DocumentExportError(
            "DOCX_UNAVAILABLE: install the governed 'documents' add-on profile "
            "(python-docx and lxml are required)."
        ) from exc


def _pdf_dependencies() -> None:
    try:
        import reportlab  # noqa: F401
    except ImportError as exc:
        raise DocumentExportError("PDF_UNAVAILABLE: reportlab is required") from exc


def _add_docx_inline(paragraph: Any, text: str) -> None:
    from docx.shared import Pt, RGBColor

    for value, bold, italic, code in _iter_inline(text):
        run = paragraph.add_run(value)
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = "Menlo"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(NAVY)


def _set_cell_shading(cell: Any, color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is not None:
        tc_pr.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def _set_cell_margins(cell: Any, top: int = 70, start: int = 80, bottom: int = 70, end: int = 80) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    row_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_pr.append(header)


def _set_cell_text(cell: Any, text: str, *, header: bool = False, compact: bool = False) -> None:
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.shared import Pt, RGBColor

    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    run = paragraph.add_run(_plain(text))
    run.font.name = "Arial"
    run.font.size = Pt(6.5 if compact else 8.2)
    run.bold = header
    run.font.color.rgb = RGBColor.from_string(WHITE if header else INK)


def _set_section_geometry(section: Any, *, landscape: bool = False) -> None:
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches

    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11 if landscape else 8.5)
    section.page_height = Inches(8.5 if landscape else 11)
    section.top_margin = Inches(0.56)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.22)


def _configure_docx_header_footer(section: Any, model: DocumentModel) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(f"SAPOTE-MAMEY  |  {model.meta.document_type.replace('_', ' ').upper()}")
    run.font.name = "Arial"
    run.font.size = Pt(7)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(TEAL)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"{model.meta.claim_safety_footer}  ·  ")
    run.font.name = "Arial"
    run.font.size = Pt(6.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def _docx_callout(document: Any, block: Block) -> None:
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    color_map = {
        "LEAD": (PALE_TEAL, TEAL),
        "NOTE": (PALE_TEAL, TEAL),
        "DECISION": (PALE_GOLD, GOLD),
        "CAUTION": (PALE_CORAL, CORAL),
        "HOLD": (PALE_CORAL, CORAL),
        "CLAIM CEILING": (PALE_GOLD, GOLD),
    }
    fill, accent = color_map[block.data["type"]]
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.right_indent = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    title = paragraph.add_run(f"{block.data['title']}.  ")
    title.bold = True
    title.font.name = "Arial"
    title.font.size = Pt(9)
    title.font.color.rgb = RGBColor.from_string(accent)
    _add_docx_inline(paragraph, block.data["text"])


def _docx_table(document: Any, block: Block) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Inches

    rows = block.data["rows"]
    columns = len(rows[0])
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    compact = columns > 7
    available = 9.7 if block.data["layout"] == "landscape" else 7.2
    weights = block.data.get("widths") or [1.0] * columns
    total = sum(weights)
    widths = [available * weight / total for weight in weights]
    for row_index, row_data in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0 and block.data["repeat_header"]:
            _set_repeat_table_header(row)
        for column_index, value in enumerate(row_data):
            cell = row.cells[column_index]
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            _set_cell_text(cell, value, header=row_index == 0, compact=compact)
            if row_index == 0:
                _set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                _set_cell_shading(cell, CREAM)


def _docx_list(document: Any, block: Block) -> None:
    """Write renderer-stable list paragraphs.

    Built-in Word list styles can move a continuation marker to the end of the
    preceding line when LibreOffice paginates a list.  Explicit markers inside
    hanging-indent paragraphs retain the same readable structure in Word and
    in the sanctioned headless QA renderer.
    """
    from docx.shared import Inches, Pt, RGBColor

    for index, item in enumerate(block.data["items"], start=1):
        paragraph = document.add_paragraph(style="List Paragraph")
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        paragraph.paragraph_format.space_after = Pt(1.5)
        paragraph.paragraph_format.keep_together = True
        marker = f"{index}." if block.data["ordered"] else "•"
        marker_run = paragraph.add_run(f"{marker}\t")
        marker_run.font.name = "Arial"
        marker_run.font.size = Pt(9.3)
        marker_run.font.color.rgb = RGBColor.from_string(INK)
        _add_docx_inline(paragraph, item)


def export_docx(model: DocumentModel, output: str | Path) -> Path:
    _docx_dependencies()
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    target = Path(output)
    document = Document()
    _set_section_geometry(document.sections[0])

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color, before, after in (
        ("Title", 24, NAVY, 0, 4),
        ("Subtitle", 11, MUTED, 0, 10),
        ("Heading 1", 17, NAVY, 12, 5),
        ("Heading 2", 12.5, TEAL, 9, 3),
        ("Heading 3", 10.5, GOLD, 7, 2),
        ("Heading 4", 9.3, NAVY, 5, 2),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = document.add_paragraph(style="Title")
    title.add_run(model.meta.title)
    if model.meta.subtitle:
        document.add_paragraph(model.meta.subtitle, style="Subtitle")
    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(8)
    run = rule.add_run("━" * 54)
    run.font.color.rgb = RGBColor.from_string(GOLD)
    run.font.size = Pt(6)
    meta_line = document.add_paragraph()
    meta_line.paragraph_format.space_after = Pt(8)
    meta_run = meta_line.add_run(f"Audience: {model.meta.audience}  ·  Authority: {model.meta.authority}")
    meta_run.font.size = Pt(7.5)
    meta_run.font.color.rgb = RGBColor.from_string(MUTED)
    for section in document.sections:
        _configure_docx_header_footer(section, model)

    landscape = False
    for block_index, block in enumerate(model.blocks):
        next_block = model.blocks[block_index + 1] if block_index + 1 < len(model.blocks) else None
        heading_for_wide_table = (
            block.kind == "heading" and next_block is not None and next_block.kind == "table"
            and next_block.data["layout"] == "landscape"
        )
        needs_landscape = heading_for_wide_table or (
            block.kind == "table" and block.data["layout"] == "landscape"
        )
        if needs_landscape and not landscape:
            section = document.add_section(WD_SECTION.NEW_PAGE)
            _set_section_geometry(section, landscape=True)
            _configure_docx_header_footer(section, model)
            landscape = True
        elif landscape and not needs_landscape:
            section = document.add_section(WD_SECTION.NEW_PAGE)
            _set_section_geometry(section, landscape=False)
            _configure_docx_header_footer(section, model)
            landscape = False

        if block.kind == "heading":
            document.add_paragraph(_plain(block.data["text"]), style=f"Heading {block.data['level']}")
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            _add_docx_inline(paragraph, block.data["text"])
        elif block.kind == "list":
            _docx_list(document, block)
        elif block.kind == "callout":
            _docx_callout(document, block)
        elif block.kind == "table":
            _docx_table(document, block)
        elif block.kind == "figure":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            max_width = 8.8 if landscape else 6.6
            picture = paragraph.add_run().add_picture(
                str(_asset_path(model, block.data["src"])), width=Inches(max_width)
            )
            picture._inline.docPr.set("descr", block.data["alt"])
            picture._inline.docPr.set("title", f"Figure {block.data['id']}")
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run(f"Figure {block.data['id']}. {block.data['caption']}")
            run.bold = True
            run.font.size = Pt(8.5)
            notes = document.add_paragraph()
            notes.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = notes.add_run(f"Methods: {block.data['methods']}  Source: {block.data['source']}")
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor.from_string(MUTED)
        elif block.kind == "rule":
            paragraph = document.add_paragraph("―" * 48)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.kind == "page_break":
            document.add_page_break()

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def _pdf_text(text: str) -> str:
    from xml.sax.saxutils import escape

    escaped = escape(text)
    escaped = re.sub(r"`([^`]+)`", r"<font name='Courier'>\1</font>", escaped)
    escaped = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", escaped)
    return escaped


def export_pdf(model: DocumentModel, output: str | Path) -> Path:
    _pdf_dependencies()
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import LETTER, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        BaseDocTemplate, Frame, Image, KeepTogether, ListFlowable, ListItem,
        NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle,
    )

    target = Path(output)
    def pdf_color(value: str) -> Any:
        return colors.HexColor(f"#{value.lstrip('#')}")

    target.parent.mkdir(parents=True, exist_ok=True)
    portrait = LETTER
    wide = landscape(LETTER)
    margin_x, margin_top, margin_bottom = 0.58 * inch, 0.55 * inch, 0.62 * inch
    doc = BaseDocTemplate(
        str(target), pagesize=portrait, leftMargin=margin_x, rightMargin=margin_x,
        topMargin=margin_top, bottomMargin=margin_bottom,
        title=model.meta.title, author="Sapote-Mamey",
    )

    def on_page(canvas: Any, document: Any) -> None:
        width, height = canvas._pagesize
        canvas.saveState()
        canvas.setStrokeColor(pdf_color(TEAL))
        canvas.setLineWidth(0.7)
        canvas.line(margin_x, height - 0.33 * inch, width - margin_x, height - 0.33 * inch)
        canvas.setFont("Helvetica-Bold", 6.5)
        canvas.setFillColor(pdf_color(TEAL))
        canvas.drawRightString(width - margin_x, height - 0.25 * inch, f"SAPOTE-MAMEY  |  {model.meta.document_type.replace('_', ' ').upper()}")
        canvas.setFont("Helvetica", 6.1)
        canvas.setFillColor(pdf_color(MUTED))
        footer = f"{model.meta.claim_safety_footer}  ·  {document.page}"
        canvas.drawCentredString(width / 2, 0.29 * inch, footer[:220])
        canvas.restoreState()

    portrait_frame = Frame(margin_x, margin_bottom, portrait[0] - 2 * margin_x, portrait[1] - margin_top - margin_bottom, id="portrait")
    wide_frame = Frame(margin_x, margin_bottom, wide[0] - 2 * margin_x, wide[1] - margin_top - margin_bottom, id="landscape")
    doc.addPageTemplates([
        PageTemplate(id="portrait", pagesize=portrait, frames=[portrait_frame], onPage=on_page),
        PageTemplate(id="landscape", pagesize=wide, frames=[wide_frame], onPage=on_page),
    ])

    styles = getSampleStyleSheet()
    pstyles = {
        "title": ParagraphStyle("SapoteTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=23, leading=25, textColor=pdf_color(NAVY), alignment=TA_LEFT, spaceAfter=3),
        "subtitle": ParagraphStyle("SapoteSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=10.5, leading=13, textColor=pdf_color(MUTED), spaceAfter=8),
        "meta": ParagraphStyle("SapoteMeta", parent=styles["Normal"], fontName="Helvetica", fontSize=7.2, leading=9, textColor=pdf_color(MUTED), spaceAfter=8),
        "body": ParagraphStyle("SapoteBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.7, leading=11.1, textColor=pdf_color(INK), spaceAfter=5),
        "caption": ParagraphStyle("SapoteCaption", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=7.7, leading=9.4, textColor=pdf_color(INK), alignment=TA_CENTER, spaceBefore=3, spaceAfter=1),
        "methods": ParagraphStyle("SapoteMethods", parent=styles["BodyText"], fontName="Helvetica", fontSize=6.8, leading=8.2, textColor=pdf_color(MUTED), alignment=TA_CENTER, spaceAfter=7),
    }
    for level, size, heading_color, before, after in ((1, 16, NAVY, 10, 4), (2, 12, TEAL, 8, 3), (3, 10, GOLD, 6, 2), (4, 9, NAVY, 5, 2)):
        pstyles[f"h{level}"] = ParagraphStyle(
            f"SapoteH{level}", parent=styles[f"Heading{min(level, 3)}"], fontName="Helvetica-Bold",
            fontSize=size, leading=size + 2, textColor=pdf_color(heading_color), spaceBefore=before,
            spaceAfter=after, keepWithNext=True,
        )

    story: list[Any] = [
        Paragraph(_pdf_text(model.meta.title), pstyles["title"]),
        Paragraph(_pdf_text(model.meta.subtitle), pstyles["subtitle"]) if model.meta.subtitle else Spacer(1, 1),
        Table([[""]], colWidths=[7.25 * inch], rowHeights=[3], style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), pdf_color(GOLD)), ("BOX", (0, 0), (-1, -1), 0, pdf_color(GOLD))])),
        Spacer(1, 5),
        Paragraph(_pdf_text(f"Audience: {model.meta.audience}  ·  Authority: {model.meta.authority}"), pstyles["meta"]),
    ]
    active_landscape = False
    for block_index, block in enumerate(model.blocks):
        next_block = model.blocks[block_index + 1] if block_index + 1 < len(model.blocks) else None
        heading_for_wide_table = (
            block.kind == "heading" and next_block is not None and next_block.kind == "table"
            and next_block.data["layout"] == "landscape"
        )
        wants_landscape = heading_for_wide_table or (
            block.kind == "table" and block.data["layout"] == "landscape"
        )
        if wants_landscape and not active_landscape:
            story.extend([NextPageTemplate("landscape"), PageBreak()])
            active_landscape = True
        elif active_landscape and not wants_landscape:
            story.extend([NextPageTemplate("portrait"), PageBreak()])
            active_landscape = False

        if block.kind == "heading":
            story.append(Paragraph(_pdf_text(block.data["text"]), pstyles[f"h{block.data['level']}"]))
        elif block.kind == "paragraph":
            story.append(Paragraph(_pdf_text(block.data["text"]), pstyles["body"]))
        elif block.kind == "list":
            items = [ListItem(Paragraph(_pdf_text(item), pstyles["body"]), leftIndent=11) for item in block.data["items"]]
            story.append(ListFlowable(items, bulletType="1" if block.data["ordered"] else "bullet", leftIndent=15, bulletFontSize=7, spaceAfter=4))
        elif block.kind == "callout":
            fill = PALE_TEAL if block.data["type"] in {"LEAD", "NOTE"} else PALE_GOLD if block.data["type"] in {"DECISION", "CLAIM CEILING"} else PALE_CORAL
            accent = TEAL if fill == PALE_TEAL else GOLD if fill == PALE_GOLD else CORAL
            content = Paragraph(_pdf_text(f"**{block.data['title']}.**  {block.data['text']}"), pstyles["body"])
            story.append(Table([[content]], colWidths=[9.55 * inch if active_landscape else 7.12 * inch], style=TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), pdf_color(fill)),
                ("LINEBEFORE", (0, 0), (0, -1), 4, pdf_color(accent)),
                ("BOX", (0, 0), (-1, -1), 0.4, pdf_color(fill)),
                ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ])))
            story.append(Spacer(1, 6))
        elif block.kind == "table":
            rows = []
            for row_index, row in enumerate(block.data["rows"]):
                cell_style = ParagraphStyle(
                    f"cell_{block.data['id']}_{row_index}", parent=pstyles["body"],
                    fontName="Helvetica-Bold" if row_index == 0 else "Helvetica",
                    textColor=colors.white if row_index == 0 else pdf_color(INK),
                    fontSize=6.2 if len(row) > 7 else 7.2,
                    leading=7.5 if len(row) > 7 else 8.7, spaceAfter=0,
                )
                rows.append([Paragraph(_pdf_text(cell), cell_style) for cell in row])
            available = 9.72 * inch if active_landscape else 7.22 * inch
            weights = block.data.get("widths") or [1.0] * len(rows[0])
            total = sum(weights)
            widths = [available * value / total for value in weights]
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), pdf_color(NAVY)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, pdf_color("B7C3C8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
            for row_index in range(2, len(rows), 2):
                style.append(("BACKGROUND", (0, row_index), (-1, row_index), pdf_color(CREAM)))
            story.append(Table(rows, colWidths=widths, repeatRows=1 if block.data["repeat_header"] else 0, style=TableStyle(style), hAlign="CENTER"))
            story.append(Spacer(1, 6))
        elif block.kind == "figure":
            asset = _asset_path(model, block.data["src"])
            image = Image(str(asset))
            max_width = 9.3 * inch if active_landscape else 6.65 * inch
            max_height = 5.9 * inch if active_landscape else 6.3 * inch
            scale = min(max_width / image.imageWidth, max_height / image.imageHeight)
            image.drawWidth = image.imageWidth * scale
            image.drawHeight = image.imageHeight * scale
            image.hAlign = "CENTER"
            story.append(KeepTogether([
                image,
                Paragraph(_pdf_text(f"Figure {block.data['id']}. {block.data['caption']}"), pstyles["caption"]),
                Paragraph(_pdf_text(f"Methods: {block.data['methods']}  Source: {block.data['source']}"), pstyles["methods"]),
            ]))
        elif block.kind == "rule":
            story.append(Table([[""]], colWidths=[9.7 * inch if active_landscape else 7.2 * inch], rowHeights=[1], style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), pdf_color(TEAL))])))
        elif block.kind == "page_break":
            story.append(PageBreak())

    doc.build(story)
    return target


def export_document(
    source: str | Path,
    *,
    outdir: str | Path | None = None,
    output_format: str = "both",
) -> ExportReceipt:
    """Validate once, then export requested formats atomically.

    ``both`` is strict: dependency checks run before rendering and final paths
    are populated only after both temporary artifacts have been written.
    """
    source_path = Path(source).resolve()
    model = parse_path(source_path)
    validation = validate_model(model, require_assets=True)
    validation.raise_for_errors()
    if output_format not in {"docx", "pdf", "both"}:
        raise DocumentExportError(f"FORMAT_INVALID:{output_format}")
    if output_format in {"docx", "both"}:
        _docx_dependencies()
    if output_format in {"pdf", "both"}:
        _pdf_dependencies()

    output_root = Path(outdir).resolve() if outdir else source_path.parent
    output_root.mkdir(parents=True, exist_ok=True)
    requested = ["docx", "pdf"] if output_format == "both" else [output_format]
    outputs: list[dict[str, Any]] = []
    with tempfile.TemporaryDirectory(prefix="sapote-render-", dir=output_root) as temp_dir:
        temp_root = Path(temp_dir)
        rendered: dict[str, Path] = {}
        for extension in requested:
            temporary = temp_root / f"{source_path.stem}.{extension}"
            rendered[extension] = export_docx(model, temporary) if extension == "docx" else export_pdf(model, temporary)
            if not rendered[extension].is_file() or rendered[extension].stat().st_size == 0:
                raise DocumentExportError(f"OUTPUT_NOT_WRITTEN:{extension}")
        for extension in requested:
            final = output_root / f"{source_path.stem}.{extension}"
            rendered[extension].replace(final)
            outputs.append({
                "format": extension,
                "path": str(final),
                "sha256": _sha256(final),
                "bytes": final.stat().st_size,
            })

    receipt = ExportReceipt(
        status="PASS",
        schema_version=model.meta.schema_version,
        input_path=str(source_path),
        input_sha256=model.source_sha256,
        model_sha256=model.canonical_sha256(),
        outputs=outputs,
        warnings=model.warnings,
        created_utc=datetime.now(timezone.utc).isoformat(),
    )
    receipt_path = output_root / f"{source_path.stem}.render_receipt.json"
    receipt_path.write_text(json.dumps(asdict(receipt), indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return receipt


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Render governed Sapote Markdown to DOCX and/or PDF")
    parser.add_argument("input", help="Governed Markdown document")
    parser.add_argument("--outdir", default=None, help="Output directory (default: beside input)")
    parser.add_argument("--format", dest="output_format", choices=["docx", "pdf", "both"], default="both")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        receipt = export_document(args.input, outdir=args.outdir, output_format=args.output_format)
    except Exception as exc:
        sys.stdout.write(str(json.dumps({"status": "FAIL", "error": str(exc)}, indent=2)) + "\n")
        return 1
    sys.stdout.write(str(json.dumps(asdict(receipt), indent=2)) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
