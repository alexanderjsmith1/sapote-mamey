"""Receipt producer for already-built GToTree/IQ-TREE phylogenies.

This module never runs alignment or tree inference. It binds artifacts and governance
evidence already present in an ``amber_phylo`` run directory for Figure Factory consumers.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import sys


PHYLO_EVIDENCE_SCHEMA = "sapote.phylogeny-evidence.v1"
REFERENCE_DEDUP_STATES = frozenset({
    "ONE_PER_SPECIES", "NO_OP_DETECTED", "NOT_RECORDED", "DEDUP_NOT_VERIFIED",
})
PHYLO_EVIDENCE_REQUIRED_KEYS = frozenset({
    "schema_version", "status", "tree_sha256", "alignment_sha256", "gtotree", "iqtree",
    "marker_set", "outgroup_tip", "outgroup_registry", "reference_dedup", "support",
    "n_tips", "assembly_quality_flags", "comparator_provenance", "package_tree_join",
})
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
SUPPORT_RE = re.compile(r"\)(\d+(?:\.\d+)?)/(\d+(?:\.\d+)?):")
# GToTree (and single-metric IQ-TREE / FastTree) runs annotate each internal node with ONE
# support value -- ")1.000:" or ")95:" -- not the SH-aLRT/UFBoot pair ")90/99:". Anchored on
# ")" and terminated by ":" so it never captures a branch length (which always follows ":",
# never ")") nor a slash-joined pair (whose first member is followed by "/", not ":").
SINGLE_SUPPORT_RE = re.compile(r"\)(\d+(?:\.\d+)?):")
# Directory tokens that mark an already-built, approved tree the receipt producer may bind.
# ``amber_phylo`` is the staged-run convention; ``phylo_tree`` is the real completed cohort
# GToTree/IQ-TREE output directory (workspace/cohort/phylo_tree/). Nothing here builds a tree.
COMPLETED_TREE_LOCATION_TOKENS = frozenset({"amber_phylo", "phylo_tree"})
OUTGROUP_COLUMNS = {
    "tree_scope", "ingroup_taxon", "family", "outgroup_genus",
    "outgroup_species_strain", "assembly_accession", "status", "rationale",
}


class PhyloEvidenceError(ValueError):
    """Typed refusal to publish an incomplete or contradictory phylogeny receipt."""

    def __init__(self, code: str, detail: str):
        self.code = code
        self.detail = detail
        super().__init__(f"{code}: {detail}")


def _fail(code: str, detail: str):
    raise PhyloEvidenceError(code, detail)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_object(path: Path, label: str) -> dict:
    try:
        value = json.loads(path.read_text(encoding="utf-8", errors="strict"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        _fail("PHYLO_EVIDENCE_INPUT_INVALID", f"{label}: {type(exc).__name__}")
    if not isinstance(value, dict):
        _fail("PHYLO_EVIDENCE_INPUT_INVALID", f"{label}: JSON object required")
    return value


def _tool(path: Path, expected_name: str) -> dict:
    value = _json_object(path, expected_name)
    name = str(value.get("tool_name", "")).strip()
    version = str(value.get("tool_version", "")).strip()
    digest = str(value.get("tool_sha256", "")).strip().lower()
    if name.casefold() != expected_name.casefold() or not version or not SHA256_RE.fullmatch(digest):
        _fail("PHYLO_TOOL_RECEIPT_INVALID", expected_name)
    return {"name": name, "version": version, "sha256": digest,
            "receipt_sha256": _sha256(path)}


def _marker_set(path: Path) -> dict:
    value = _json_object(path, "marker_set")
    name = str(value.get("name", "")).strip()
    count = value.get("count")
    digest = str(value.get("sha256", "")).strip().lower()
    if not name or not isinstance(count, int) or count <= 0 or not SHA256_RE.fullmatch(digest):
        _fail("PHYLO_MARKER_SET_INVALID", "name, positive count, and sha256 required")
    return {"name": name, "count": count, "sha256": digest,
            "receipt_sha256": _sha256(path)}


def _tree_tips(tree_text: str) -> list[str]:
    overlay_path = Path(__file__).resolve().parents[1] / "tools" / "tree_bgc_overlay.py"
    spec = importlib.util.spec_from_file_location("sapote_tree_overlay_for_receipt", overlay_path)
    if spec is None or spec.loader is None:
        _fail("PHYLO_TREE_PARSER_UNAVAILABLE", overlay_path.name)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    try:
        tips = module.tip_order(module.parse_newick(tree_text))
    except ValueError as exc:
        _fail("PHYLO_TREE_INVALID", str(exc))
    if not tips or len(tips) != len(set(tips)):
        _fail("PHYLO_TREE_INVALID", "non-empty unique tips required")
    return tips


def _canonical_tip(label: str) -> str:
    """Normalize a Newick-safe label to the display contract Genus species strain."""
    value = " ".join(str(label or "").replace("_", " ").split())
    parts = value.split()
    if (len(parts) < 3 or not re.fullmatch(r"[A-Z][a-z][A-Za-z.-]*", parts[0])
            or not re.fullmatch(r"(?:[a-z][A-Za-z.-]*|sp\.)", parts[1])):
        _fail("PHYLO_TIP_LABEL_INVALID", f"Genus species strain required: {label!r}")
    return value


def _tabular_rows(path: Path, *, sheet: str | None = None, header_row: int = 1) -> tuple[list[str], list[list[str]]]:
    if header_row < 1:
        _fail("PHYLO_HOST_TABLE_INVALID", "header_row must be positive")
    if path.suffix.lower() == ".xlsx":
        if not sheet:
            _fail("PHYLO_HOST_TABLE_INVALID", "XLSX requires an explicit sheet")
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            if sheet not in workbook.sheetnames:
                _fail("PHYLO_HOST_TABLE_INVALID", f"sheet absent: {sheet}")
            values = list(workbook[sheet].iter_rows(values_only=True))
            workbook.close()
        except PhyloEvidenceError:
            raise
        except Exception as exc:
            _fail("PHYLO_HOST_TABLE_INVALID", type(exc).__name__)
        if len(values) < header_row:
            _fail("PHYLO_HOST_TABLE_INVALID", "header row absent")
        header = [str(value or "").strip() for value in values[header_row - 1]]
        rows = [[str(value or "").strip() for value in row] for row in values[header_row:]]
        return header, rows
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open(newline="", encoding="utf-8-sig") as handle:
        values = list(csv.reader(handle, delimiter=delimiter))
    if len(values) < header_row:
        _fail("PHYLO_HOST_TABLE_INVALID", "header row absent")
    return [cell.strip() for cell in values[header_row - 1]], values[header_row:]


def _column_index(headers: list[str], selector: str | None, aliases: set[str], label: str) -> int:
    if selector and re.fullmatch(r"[A-Za-z]+", selector):
        index = 0
        for char in selector.upper():
            index = index * 26 + ord(char) - 64
        index -= 1
        if index < 0 or index >= len(headers):
            _fail("PHYLO_HOST_TABLE_INVALID", f"{label} column letter outside header")
        return index
    wanted = selector.casefold() if selector else None
    matches = [index for index, name in enumerate(headers)
               if (wanted is not None and name.casefold() == wanted)
               or (wanted is None and name.casefold() in aliases)]
    if len(matches) != 1:
        _fail("PHYLO_HOST_TABLE_AMBIGUOUS", f"{label} matched {len(matches)} columns; select one explicitly")
    return matches[0]


def _host_index(path: Path, *, sheet: str | None, header_row: int,
                strain_column: str | None, host_column: str | None) -> dict[str, str]:
    headers, rows = _tabular_rows(path, sheet=sheet, header_row=header_row)
    strain_index = _column_index(headers, strain_column, {"strain_id", "strain", "strain name"}, "strain")
    host_index = _column_index(headers, host_column, {"host", "source species"}, "host")
    result: dict[str, str] = {}
    for row in rows:
        strain = row[strain_index].strip() if strain_index < len(row) else ""
        host = row[host_index].strip() if host_index < len(row) else ""
        if not strain:
            continue
        key = strain.casefold()
        if key in result:
            _fail("PHYLO_HOST_TABLE_AMBIGUOUS", f"duplicate strain: {strain}")
        result[key] = host
    return result


def _package_tree_join(path: Path, tips: list[str], host_table: Path, *, host_sheet: str | None,
                       host_header_row: int, host_strain_column: str | None,
                       host_value_column: str | None) -> dict:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required = {"tree_tip", "role", "strain_id", "package_manifest", "package_manifest_sha256",
                    "host_label", "host_provenance"}
        if reader.fieldnames is None or not required <= set(reader.fieldnames):
            _fail("PHYLO_PACKAGE_TREE_JOIN_INVALID", "required columns missing")
        rows = list(reader)
    canonical_tips = [_canonical_tip(tip) for tip in tips]
    if len(canonical_tips) != len(set(canonical_tips)):
        _fail("PHYLO_TIP_LABEL_COLLISION", "tip normalization is not one-to-one")
    by_tip: dict[str, dict] = {}
    for row in rows:
        tip = _canonical_tip(row["tree_tip"])
        if tip in by_tip:
            _fail("PHYLO_PACKAGE_TREE_JOIN_INVALID", f"duplicate tip: {tip}")
        by_tip[tip] = row
    if set(by_tip) != set(canonical_tips):
        _fail("PHYLO_PACKAGE_TREE_JOIN_INCOMPLETE", "exactly one row per normalized tree tip required")
    hosts = _host_index(host_table, sheet=host_sheet, header_row=host_header_row,
                        strain_column=host_strain_column, host_column=host_value_column)
    joined: list[dict] = []
    for raw_tip, tip in zip(tips, canonical_tips):
        row = by_tip[tip]
        role = row["role"].strip().upper()
        strain = row["strain_id"].strip()
        if role not in {"PACKAGE", "REFERENCE"} or not strain or not tip.casefold().endswith(" " + strain.casefold()):
            _fail("PHYLO_PACKAGE_TREE_JOIN_INVALID", f"role/strain mismatch: {tip}")
        item = {"tree_tip": raw_tip, "canonical_tip": tip, "role": role, "strain_id": strain}
        if role == "PACKAGE":
            member = Path(row["package_manifest"].strip())
            if member.is_absolute() or ".." in member.parts:
                _fail("PHYLO_PACKAGE_MANIFEST_INVALID", f"portable relative path required: {strain}")
            manifest = (path.parent / member).resolve()
            expected = row["package_manifest_sha256"].strip().lower()
            if not manifest.is_file() or not SHA256_RE.fullmatch(expected) or _sha256(manifest) != expected:
                _fail("PHYLO_PACKAGE_MANIFEST_INVALID", strain)
            payload = _json_object(manifest, "package_manifest")
            if str(payload.get("strain_id") or payload.get("display_name") or "").strip() != strain:
                _fail("PHYLO_PACKAGE_MANIFEST_INVALID", f"strain mismatch: {strain}")
            authoritative_host = hosts.get(strain.casefold(), "").strip()
            if authoritative_host:
                host_label, provenance = authoritative_host, "AUTHORITATIVE_TABLE"
            else:
                host_label = row["host_label"].strip()
                provenance = row["host_provenance"].strip().upper()
                if not host_label or provenance != "PI_WORD_ONLY":
                    _fail("PHYLO_HOST_PROVENANCE_MISSING", strain)
            item.update({"manifest_sha256": expected, "host_label": host_label,
                         "host_provenance": provenance})
        else:
            item.update({"manifest_sha256": None, "host_label": "", "host_provenance": "NOT_APPLICABLE_REFERENCE"})
        joined.append(item)
    return {"crosswalk_sha256": _sha256(path), "host_table_sha256": _sha256(host_table),
            "host_table_sheet": host_sheet, "rows": joined}


def _outgroup_row(path: Path, ingroup_taxon: str, tree_scope: str, outgroup_tip: str) -> dict:
    lines = [line for line in path.read_text(encoding="utf-8").splitlines()
             if line.strip() and not line.lstrip().startswith("#")]
    reader = csv.DictReader(lines, delimiter="\t")
    if reader.fieldnames is None or not OUTGROUP_COLUMNS <= set(reader.fieldnames):
        _fail("PHYLO_OUTGROUP_REGISTRY_INVALID", "required columns missing")
    matches = [row for row in reader
               if row["ingroup_taxon"].casefold() == ingroup_taxon.casefold()
               and row["tree_scope"].casefold() == tree_scope.casefold()]
    locked = [row for row in matches if row["status"] == "LOCKED"]
    selected = locked if locked else matches
    if len(selected) != 1:
        _fail("PHYLO_OUTGROUP_REGISTRY_AMBIGUOUS", f"matched {len(selected)} rows")
    row = selected[0]
    scope = row["tree_scope"].strip().casefold()
    rationale = row["rationale"].strip().casefold()
    same_taxon = row["outgroup_genus"].strip().casefold() == row["ingroup_taxon"].strip().casefold()
    if scope == "genus":
        one_rank_out = bool(row["family"].strip()) and "sister genus" in rationale and not same_taxon
        expected_relation = "sister genus in the declared ingroup family"
    elif scope == "family":
        one_rank_out = (not same_taxon and
                        any(term in rationale for term in ("sister family", "sister order", "within-order sister")))
        expected_relation = "sister family/order outside the ingroup family"
    else:
        one_rank_out = False
        expected_relation = "supported genus or family scope"
    if not one_rank_out:
        _fail("OUTGROUP_TOO_DISTANT",
              f"registry does not declare one-rank-out relation ({expected_relation})")
    normalized_tip = re.sub(r"[^a-z0-9]+", "", outgroup_tip.casefold())
    species = re.sub(r"[^a-z0-9]+", "", row["outgroup_species_strain"].casefold())
    accession = re.sub(r"[^a-z0-9]+", "", row["assembly_accession"].casefold())
    if not ((species and species in normalized_tip) or (accession and accession in normalized_tip)):
        _fail("PHYLO_OUTGROUP_TIP_MISMATCH", outgroup_tip)
    canonical = json.dumps(row, sort_keys=True, separators=(",", ":"))
    return {"row": row, "rank_check": "ONE_RANK_OUT_REGISTRY_DECLARED",
            "row_sha256": hashlib.sha256(canonical.encode()).hexdigest(),
            "registry_sha256": _sha256(path)}


def _dedup(path: Path | None) -> dict:
    if path is None:
        return {"state": "NOT_RECORDED", "species_column": "", "rows": 0,
                "source_sha256": None}
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        if reader.fieldnames is None or "action" not in reader.fieldnames:
            _fail("PHYLO_REFERENCE_DEDUP_INVALID", "action column required")
        rows = list(reader)
    species_column = "species" if "species" in reader.fieldnames else ""
    kept_rows = [row for row in rows if str(row.get("action", "")).strip().upper() in {"KEEP", "KEPT"}]
    kept = [str(row.get(species_column, "")).strip().casefold()
            for row in kept_rows] if species_column else []
    if not species_column:
        state, reason = "DEDUP_NOT_VERIFIED", "species_column_missing"
    elif not kept:
        state, reason = "NO_OP_DETECTED", "no_kept_rows"
    elif any(not species for species in kept):
        state, reason = "DEDUP_NOT_VERIFIED", "kept_species_blank"
    elif len(kept) != len(set(kept)):
        state, reason = "DEDUP_NOT_VERIFIED", "kept_species_non_unique"
    else:
        state, reason = "ONE_PER_SPECIES", "one_unique_species_per_kept_row"
    return {"state": state, "reason": reason, "species_column": species_column, "rows": len(rows),
            "kept": len(kept_rows),
            "dropped": sum(str(row.get("action", "")).strip().upper().startswith(("DROP", "DELETE")) for row in rows),
            "source_sha256": _sha256(path)}


def _support(tree_text: str) -> dict:
    pairs = [(float(a), float(b)) for a, b in SUPPORT_RE.findall(tree_text)]
    if pairs:
        return {"label": "SH-aLRT/UFBoot", "node_count": len(pairs),
                "sh_alrt_min": min(a for a, _ in pairs), "ufboot_min": min(b for _, b in pairs),
                "weak_node_count": sum(a < 80 or b < 95 for a, b in pairs)}
    singles = [float(v) for v in SINGLE_SUPPORT_RE.findall(tree_text)]
    if singles:
        # One support metric per node (GToTree single-value branch support). The metric's
        # native scale is not carried in the Newick, so it is inferred conservatively from the
        # observed range and the applied weak-node cutoff is recorded rather than hidden: values
        # bounded by 1.0 are treated as a 0-1 scale (weak < 0.80), otherwise as a 0-100 scale
        # (weak < 95). Summary only; no biological interpretation of the tree is asserted.
        unit_scale = max(singles) <= 1.0
        weak_threshold = 0.80 if unit_scale else 95.0
        return {"label": "single_support", "support_format": "GTOTREE_SINGLE_VALUE",
                "node_count": len(singles), "support_min": min(singles),
                "support_scale": "UNIT_0_1" if unit_scale else "PERCENT_0_100",
                "weak_threshold": weak_threshold,
                "weak_node_count": sum(v < weak_threshold for v in singles)}
    _fail("PHYLO_SUPPORT_MISSING", "no SH-aLRT/UFBoot pair or single-value branch support in tree")


def _assembly_flags(path: Path, tips: list[str]) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        if reader.fieldnames is None or not {"tree_tip", "assembly_quality_flag"} <= set(reader.fieldnames):
            _fail("PHYLO_ASSEMBLY_FLAGS_INVALID", "tree_tip and assembly_quality_flag required")
        rows = [{"tree_tip": row["tree_tip"].strip(),
                 "assembly_quality_flag": row["assembly_quality_flag"].strip(),
                 "detail": row.get("detail", "").strip()} for row in reader]
    if {row["tree_tip"] for row in rows} != set(tips) or len(rows) != len(tips):
        _fail("PHYLO_ASSEMBLY_FLAGS_INCOMPLETE", "exactly one row per tree tip required")
    if any(not row["assembly_quality_flag"] for row in rows):
        _fail("PHYLO_ASSEMBLY_FLAGS_INVALID", "blank flag")
    return rows


def produce_phylo_evidence_receipt(*, tree: str | Path, alignment: str | Path,
                                   gtotree_receipt: str | Path, iqtree_receipt: str | Path,
                                   marker_set_receipt: str | Path, outgroup_tip: str,
                                   outgroup_registry: str | Path, ingroup_taxon: str,
                                   tree_scope: str, reference_dedup: str | Path | None,
                                   assembly_quality_flags: str | Path,
                                   package_tree_crosswalk: str | Path, host_table: str | Path,
                                   host_sheet: str | None = None, host_header_row: int = 1,
                                   host_strain_column: str | None = None,
                                   host_value_column: str | None = None) -> Path:
    tree_path = Path(tree).resolve()
    alignment_path = Path(alignment).resolve()
    if not (COMPLETED_TREE_LOCATION_TOKENS & set(tree_path.parts)) or not tree_path.is_file():
        _fail("PHYLO_TREE_LOCATION_INVALID",
              "existing tree under a completed-run directory (amber_phylo/ or phylo_tree/) required")
    if not alignment_path.is_file():
        _fail("PHYLO_ALIGNMENT_MISSING", alignment_path.name)
    output = tree_path.parent / "phylo_evidence_receipt.json"
    if output.exists():
        _fail("PHYLO_EVIDENCE_RECEIPT_EXISTS", output.name)
    tree_text = tree_path.read_text(encoding="utf-8", errors="strict")
    tips = _tree_tips(tree_text)
    # Enforce the display contract before any downstream table comparison can
    # obscure a malformed label with a secondary missing-row error.
    canonical_tips = [_canonical_tip(tip) for tip in tips]
    if len(canonical_tips) != len(set(canonical_tips)):
        _fail("PHYLO_TIP_LABEL_COLLISION", "tip normalization is not one-to-one")
    if outgroup_tip not in tips:
        _fail("PHYLO_OUTGROUP_TIP_MISMATCH", "outgroup is absent from tree tips")
    registry = _outgroup_row(Path(outgroup_registry), ingroup_taxon, tree_scope, outgroup_tip)
    dedup = _dedup(Path(reference_dedup) if reference_dedup else None)
    receipt = {
        "schema_version": PHYLO_EVIDENCE_SCHEMA,
        "status": "HOLD" if dedup["state"] == "DEDUP_NOT_VERIFIED" else "PASS",
        "tree_sha256": _sha256(tree_path), "alignment_sha256": _sha256(alignment_path),
        "gtotree": _tool(Path(gtotree_receipt), "GToTree"),
        "iqtree": _tool(Path(iqtree_receipt), "IQ-TREE"),
        "marker_set": _marker_set(Path(marker_set_receipt)), "outgroup_tip": outgroup_tip,
        "outgroup_registry": registry,
        "reference_dedup": dedup,
        "support": _support(tree_text), "n_tips": len(tips),
        "assembly_quality_flags": _assembly_flags(Path(assembly_quality_flags), tips),
        "package_tree_join": _package_tree_join(
            Path(package_tree_crosswalk), tips, Path(host_table), host_sheet=host_sheet,
            host_header_row=host_header_row, host_strain_column=host_strain_column,
            host_value_column=host_value_column,
        ),
        "comparator_provenance": {"outgroup_registry_row_sha256": registry["row_sha256"]},
        "claim_ceiling": "Existing-tree provenance and support summary only; no tree construction or biological interpretation.",
    }
    if receipt["status"] == "HOLD":
        receipt["hold_code"] = "DEDUP_NOT_VERIFIED"
    receipt["comparator_provenance"]["reference_dedup_state"] = receipt["reference_dedup"]["state"]
    output.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return output


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("tree", "alignment", "gtotree-receipt", "iqtree-receipt", "marker-set-receipt",
                 "outgroup-tip", "outgroup-registry", "ingroup-taxon", "tree-scope",
                 "assembly-quality-flags", "package-tree-crosswalk", "host-table"):
        parser.add_argument("--" + name, required=True)
    parser.add_argument("--reference-dedup")
    parser.add_argument("--host-sheet")
    parser.add_argument("--host-header-row", type=int, default=1)
    parser.add_argument("--host-strain-column", help="unique header name or XLSX column letter")
    parser.add_argument("--host-value-column", help="unique header name or XLSX column letter")
    args = parser.parse_args(argv)
    try:
        output = produce_phylo_evidence_receipt(**vars(args))
    except (OSError, UnicodeError, PhyloEvidenceError) as exc:
        sys.stderr.write(str(exc) + "\n")
        return 2
    sys.stdout.write(json.dumps({"status": "PASS", "receipt": str(output)}, sort_keys=True) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
