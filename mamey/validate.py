from __future__ import annotations
import warnings as _warnings

from mamey import __version__
import csv
import sys as _sys_for_csv
# v9.7.409 (BC hostile audit H12): a 5 MB annotation qualifier reached a package CSV and the read-back
# crashed with `_csv.Error: field larger than field limit (131072)`. The engine writes these fields itself;
# it must be able to read them back. Bounded to 2**31-1 (csv module ceiling on some platforms).
csv.field_size_limit(min(_sys_for_csv.maxsize, 2**31 - 1))
import json
import re  # v9.7.335: gold_completeness now parses BGC ids out of card filenames
from pathlib import Path
from .recovery_status import package_status_receipt, write_package_status_receipt
from . import exclusion_gate  # AMBER_EXCLUSION_HARDENING (.353): governed-output leak gate
from .bioactivity_metadata import BioactivityMetadataError, normalize_bioactivity

REQUIRED_SUFFIXES = [
    "manifest.json",
    "checksums_sha256.txt",
    "1_intake.json",
    "2_inventory.csv",
    "3_scan_states.json",
    "4A_RGGMCI_full.json",
    "4A_RGGMCI_ranked_pairs.csv",
    "4A_RGGMCI_evidence.csv",
    "4_triage_board.csv",
    "5_workbook.xlsx",
    "commit_receipt.json",
    "issue_log.md",
    "Project_Memory_Snapshot.json",
    "7_cell_provenance.csv",   # L3: cell_provenance is written on every run; validate must require it
]

# Enrichment files written AFTER the core validation gate (manifest_short, OPEN_ME_FIRST).
# These are checked by validate_package when enrichment_check=True (the post-seal pass).
ENRICHMENT_SUFFIXES = [
    "OPEN_ME_FIRST.html",      # C1: collaborator-facing package entry point
    "manifest_short.json",     # F2: compact LLM-consumable summary; regression-tested
]

REPORTING_V2_SUFFIXES = [
    "3_mibig_per_gene.json",
    "3_mibig_per_gene.csv",
    "3_mibig_convergence.json",
    "3_mibig_convergence.csv",
    "3_mibig_profile.csv",
    "3_antismash_structured.json",
    "3_antismash_modules.csv",
    "3_antismash_ripp_motifs.csv",
    "3_antismash_motifs.csv",
    "3_antismash_rrefinder.csv",
    "3_length_weighted_capacity.csv",
    "3_length_weighted_summary.json",
]


def _find(files, suffix):
    return next((p for p in files if p.name.endswith(suffix)), None)


CITATION_COMPACT_REQUIRED = [
    "Citation_Ledger.csv",
    "Citation_Ledger.json",
    "citation_compact/Literature_Search_WorkOrder.md",
    "citation_compact/Literature_Search_WorkOrder.json",
    "citation_compact/CITATION_COMPACT_QA.json",
]


def _rel(root: Path, path: Path) -> str:
    return str(path.relative_to(root)).replace("\\", "/")


def _csv_rows(path: Path) -> tuple[list[str], int]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        headers = list(reader.fieldnames or [])
        count = sum(1 for _ in reader)
    return headers, count


def validate_reporting_v2_outputs(package_dir: str | Path) -> dict:
    """Validate P-MPG, structured antiSMASH, RRE, P-LWC, and workbook parity."""
    root = Path(package_dir)
    files = [path for path in root.rglob("*") if path.is_file()]
    missing = [
        suffix for suffix in REPORTING_V2_SUFFIXES
        if _find(files, suffix) is None
    ]
    errors: list[str] = []
    counts: dict[str, int] = {}
    if missing:
        return {
            "status": "FAIL",
            "missing": missing,
            "errors": [],
            "counts": counts,
        }

    try:
        mpg = json.loads(_find(files, "3_mibig_per_gene.json").read_text(encoding="utf-8"))
        if mpg.get("schema_version") != "mibig_per_gene_v3":
            errors.append("P-MPG JSON schema_version is not mibig_per_gene_v3")
        expected = sum(
            len(rows) for rows in (mpg.get("per_gene_mibig") or {}).values()
        )
        headers, actual = _csv_rows(_find(files, "3_mibig_per_gene.csv"))
        counts["mibig_per_gene_rows"] = actual
        if expected != actual:
            errors.append(f"P-MPG JSON/CSV row mismatch: {expected} != {actual}")
        required = {
            "bgc_id", "query_gene", "subject_gene", "mibig_accession",
            "pct_identity", "pct_coverage", "pct_coverage_interpretation",
            "coverage_qc_flag", "reference_rank",
        }
        if not required.issubset(headers):
            errors.append(f"P-MPG CSV missing columns: {sorted(required - set(headers))}")
    except Exception as exc:
        errors.append(f"P-MPG parse error: {type(exc).__name__}: {exc}")

    try:
        convergence = json.loads(
            _find(files, "3_mibig_convergence.json").read_text(encoding="utf-8")
        )
        if convergence.get("schema") != "mibig_pathway_convergence_v2":
            errors.append(
                "MIBiG convergence JSON schema is not "
                "mibig_pathway_convergence_v2"
            )
        expected = len(convergence.get("rows") or [])
        headers, actual = _csv_rows(_find(files, "3_mibig_convergence.csv"))
        counts["mibig_convergence_rows"] = actual
        if expected != actual:
            errors.append(f"MIBiG convergence JSON/CSV row mismatch: {expected} != {actual}")
        required = {
            "bgc_id", "mibig_accession", "distinct_query_genes",
            "query_gene_share", "recognizable_gene_share",
            "class_concordance", "convergence_tier",
            "median_pct_coverage_interpretation", "coverage_qc_flag",
            "convergence_rank", "dominant_reference", "dominance_status",
        }
        if not required.issubset(headers):
            errors.append(f"MIBiG convergence CSV missing columns: {sorted(required - set(headers))}")
    except Exception as exc:
        errors.append(f"MIBiG convergence parse error: {type(exc).__name__}: {exc}")

    try:
        profile_headers, profile_rows = _csv_rows(_find(files, "3_mibig_profile.csv"))
        counts["mibig_profile_rows"] = profile_rows
        if "recognizable_gene_fraction" not in profile_headers:
            errors.append("MIBiG profile lacks recognizable_gene_fraction")
    except Exception as exc:
        errors.append(f"MIBiG profile parse error: {type(exc).__name__}: {exc}")

    try:
        from .antismash_tables import TABLE_COLUMNS
        structured = json.loads(
            _find(files, "3_antismash_structured.json").read_text(encoding="utf-8")
        )
        if structured.get("schema_version") != "antismash_structured_tables_v2":
            errors.append(
                "structured antiSMASH JSON schema_version is not antismash_structured_tables_v2"
            )
        for key, suffix in (
            ("modules", "3_antismash_modules.csv"),
            ("ripp_motifs", "3_antismash_ripp_motifs.csv"),
            ("motifs", "3_antismash_motifs.csv"),
            ("rrefinder", "3_antismash_rrefinder.csv"),
        ):
            headers, actual = _csv_rows(_find(files, suffix))
            expected = len(structured.get(key) or [])
            counts[f"antismash_{key}_rows"] = actual
            if expected != actual:
                errors.append(
                    f"structured antiSMASH {key} JSON/CSV row mismatch: {expected} != {actual}"
                )
            missing_columns = set(TABLE_COLUMNS[key]) - set(headers)
            if missing_columns:
                errors.append(
                    f"structured antiSMASH {key} CSV missing columns: "
                    f"{sorted(missing_columns)}"
                )
    except Exception as exc:
        errors.append(f"structured antiSMASH parse error: {type(exc).__name__}: {exc}")

    try:
        _, lwc_rows = _csv_rows(_find(files, "3_length_weighted_capacity.csv"))
        counts["p_lwc_rows"] = lwc_rows
        lwc_summary = json.loads(
            _find(files, "3_length_weighted_summary.json").read_text(encoding="utf-8")
        )
        if lwc_summary.get("status") != "TRIAL_ONLY_CAPACITY_METRIC":
            errors.append(f"P-LWC summary status is {lwc_summary.get('status')}")
        if int(lwc_summary.get("raw_regions", -1)) != lwc_rows:
            errors.append(
                f"P-LWC summary/CSV row mismatch: "
                f"{lwc_summary.get('raw_regions')} != {lwc_rows}"
            )
    except Exception as exc:
        errors.append(f"P-LWC parse error: {type(exc).__name__}: {exc}")

    workbook_path = _find(files, "5_workbook.xlsx")
    if workbook_path is not None:
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(
                workbook_path, read_only=True, data_only=True
            )
            expected_sheets = {
                "MIBiG_Per_Gene": "mibig_per_gene_rows",
                "MIBiG_Convergence": "mibig_convergence_rows",
                "MIBiG_Profile": "mibig_profile_rows",
                "antiSMASH_Modules": "antismash_modules_rows",
                "antiSMASH_RiPP_Motifs": "antismash_ripp_motifs_rows",
                "antiSMASH_Motifs": "antismash_motifs_rows",
                "antiSMASH_RREfinder": "antismash_rrefinder_rows",
                "P_LWC_Trial": "p_lwc_rows",
            }
            for sheet_name, count_key in expected_sheets.items():
                if sheet_name not in workbook.sheetnames:
                    errors.append(f"workbook missing sheet {sheet_name}")
                    continue
                data_rows = max(0, workbook[sheet_name].max_row - 1)
                expected_rows = counts.get(count_key)
                if expected_rows is not None and data_rows != expected_rows:
                    errors.append(
                        f"workbook/CSV row mismatch for {sheet_name}: "
                        f"{data_rows} != {expected_rows}"
                    )
            if "P_LWC_Summary" not in workbook.sheetnames:
                errors.append("workbook missing sheet P_LWC_Summary")
            workbook.close()
        except Exception as exc:
            errors.append(f"reporting workbook parse error: {type(exc).__name__}: {exc}")

    return {
        "status": "FAIL" if errors else "PASS",
        "missing": missing,
        "errors": errors,
        "counts": counts,
    }


def validate_citation_compact_outputs(package_dir: str | Path) -> dict:
    """Validate citation-compact outputs when present.

    This gate is optional-by-presence:
      - NOT_REQUESTED when no citation-compact artifacts exist;
      - PASS when compact artifacts are internally parseable and tracked;
      - FAIL when a partial/broken compact output set is present.

    The gate is intentionally package-level. It does not validate literature truth;
    it validates that missing citations were turned into explicit ledger/work-order
    tasks rather than silently omitted.
    """
    import os  # local: validate.py imports os nowhere else.

    root = Path(package_dir)

    # RGLOB-SEAL-COVERAGE (v9.7.409): ``Path.rglob`` swallows a per-directory OSError, so an
    # unreadable/unlistable subtree would vanish from the enumeration. Because this gate is
    # optional-by-presence, a silently-truncated scan would report NOT_REQUESTED for a package
    # whose citation-compact artifacts live under exactly the subtree that could not be read --
    # asserting absence from an incomplete scan. Walk fail-closed: any unreadable directory is
    # surfaced as an explicit FAIL (never NOT_REQUESTED), so an under-scan cannot be mistaken
    # for "not requested".
    walk_errors: list[str] = []

    def _onerror(exc: OSError) -> None:
        walk_errors.append(f"{type(exc).__name__}: {exc}")

    files = []
    for dirpath, _dirnames, filenames in os.walk(root, onerror=_onerror, followlinks=False):
        base = Path(dirpath)
        for name in filenames:
            p = base / name
            if p.is_file():
                files.append(p)
    if walk_errors:
        return {
            "status": "FAIL",
            "note": "citation-compact scan incomplete: unreadable subtree "
                    "(presence could not be determined)",
            "errors": sorted(walk_errors),
            "missing": [],
        }
    rels = {_rel(root, p) for p in files}
    compact_requested = (
        "Citation_Ledger.csv" in rels
        or "Citation_Ledger.json" in rels
        or any(r.startswith("citation_compact/") for r in rels)
    )
    if not compact_requested:
        return {
            "status": "NOT_REQUESTED",
            "note": "citation-compact outputs not present",
            "missing": [],
        }

    missing = [p for p in CITATION_COMPACT_REQUIRED if p not in rels]
    compact_reports = [
        r for r in rels
        if r.startswith("citation_compact/")
        and r.endswith("_citation_compact.md")
    ]
    if not compact_reports:
        missing.append("citation_compact/*_citation_compact.md")

    errors: list[str] = []

    # Parse ledger JSON.
    ledger_records = []
    ledger_json = root / "Citation_Ledger.json"
    if ledger_json.exists():
        try:
            payload = json.loads(ledger_json.read_text(encoding="utf-8"))
            if payload.get("schema_version") != "citation_ledger_v1":
                errors.append("Citation_Ledger.json schema_version is not citation_ledger_v1")
            ledger_records = payload.get("records", [])
            if not isinstance(ledger_records, list):
                errors.append("Citation_Ledger.json records is not a list")
        except Exception as exc:
            errors.append(f"Citation_Ledger.json parse error: {type(exc).__name__}: {exc}")

    # Parse ledger CSV.
    ledger_csv = root / "Citation_Ledger.csv"
    if ledger_csv.exists():
        try:
            with ledger_csv.open(newline="", encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle))
            if not rows:
                errors.append("Citation_Ledger.csv has no records")
            required_cols = {"citation_id", "scope", "source_status", "citation_label"}
            missing_cols = required_cols - set(rows[0].keys() if rows else [])
            if missing_cols:
                errors.append(f"Citation_Ledger.csv missing columns: {sorted(missing_cols)}")
        except Exception as exc:
            errors.append(f"Citation_Ledger.csv parse error: {type(exc).__name__}: {exc}")

    # Parse literature work-order JSON.
    workorder_tasks = []
    workorder_json = root / "citation_compact" / "Literature_Search_WorkOrder.json"
    if workorder_json.exists():
        try:
            payload = json.loads(workorder_json.read_text(encoding="utf-8"))
            if payload.get("schema_version") != "literature_search_workorder_v1":
                errors.append("Literature_Search_WorkOrder.json schema_version is not literature_search_workorder_v1")
            workorder_tasks = payload.get("tasks", [])
            if not isinstance(workorder_tasks, list):
                errors.append("Literature_Search_WorkOrder.json tasks is not a list")
        except Exception as exc:
            errors.append(f"Literature_Search_WorkOrder.json parse error: {type(exc).__name__}: {exc}")

    # If the ledger marks citations as needed, the work order must contain tasks.
    needed = [
        r for r in ledger_records
        if isinstance(r, dict) and r.get("source_status") == "citation_needed"
    ]
    if needed and not workorder_tasks:
        errors.append("ledger has citation_needed rows but Literature_Search_WorkOrder.json has no tasks")

    # Parse QA file and require it to know about compact outputs.
    qa_path = root / "citation_compact" / "CITATION_COMPACT_QA.json"
    if qa_path.exists():
        try:
            qa = json.loads(qa_path.read_text(encoding="utf-8"))
            if qa.get("status") not in {"PASS", "PASS_STRUCTURE", "WARN"}:
                errors.append(f"CITATION_COMPACT_QA status is {qa.get('status')}")
            outputs = set(qa.get("outputs", []))
            for req in ["Citation_Ledger.csv", "Citation_Ledger.json"]:
                if req not in outputs:
                    errors.append(f"CITATION_COMPACT_QA outputs missing {req}")
        except Exception as exc:
            errors.append(f"CITATION_COMPACT_QA.json parse error: {type(exc).__name__}: {exc}")

    # Package-level generic caveat policy. Citation-compact should use at most
    # one global genome-mining caveat across all compact Markdown reports.
    try:
        from .citation_compact import validate_global_caveat_count
        compact_text = "\n\n".join(
            (root / r).read_text(encoding="utf-8", errors="ignore")
            for r in compact_reports
            if (root / r).exists()
        )
        for err in validate_global_caveat_count(compact_text, max_count=1):
            errors.append(f"package caveat policy: {err.get('error')} count={err.get('count')}")
    except Exception as exc:
        errors.append(f"package caveat policy check error: {type(exc).__name__}: {exc}")

    # Check manifest/checksum tracking after seal.
    manifest_path = root / "manifest.json"
    checksum_path = root / "checksums_sha256.txt"
    if manifest_path.exists():
        try:
            mdata = json.loads(manifest_path.read_text(encoding="utf-8"))
            manifest_paths = {f.get("path") for f in mdata.get("files", []) if isinstance(f, dict)}
            for req in CITATION_COMPACT_REQUIRED:
                if req not in manifest_paths:
                    errors.append(f"manifest.json does not track {req}")
            for report in compact_reports:
                if report not in manifest_paths:
                    errors.append(f"manifest.json does not track {report}")
        except Exception as exc:
            errors.append(f"manifest.json compact tracking parse error: {type(exc).__name__}: {exc}")

    if checksum_path.exists():
        try:
            checksums = checksum_path.read_text(encoding="utf-8")
            for req in CITATION_COMPACT_REQUIRED:
                if req not in checksums:
                    errors.append(f"checksums_sha256.txt does not track {req}")
            for report in compact_reports:
                if report not in checksums:
                    errors.append(f"checksums_sha256.txt does not track {report}")
            # v9.7.145b: actually verify SHA256 values, not just presence
            import hashlib as _hashlib
            checksum_errors = []
            for line_no, line in enumerate(checksums.splitlines(), 1):
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                try:
                    expected_hash, rel_path = line.split(None, 1)
                except ValueError:
                    checksum_errors.append(f"line {line_no}: malformed checksum line")
                    continue
                file_path = root / rel_path.strip()
                if not file_path.exists():
                    # Only flag missing files that are not mutable receipts
                    from .packaging import MUTABLE_RECEIPT_NAMES as _mut, MUTABLE_RECEIPT_SUFFIXES as _muts
                    if file_path.name not in _mut and not file_path.name.endswith(_muts):  # v9.7.409 A10
                        checksum_errors.append(f"line {line_no}: missing {rel_path.strip()}")
                    continue
                actual_hash = _hashlib.sha256(file_path.read_bytes()).hexdigest()
                if actual_hash != expected_hash:
                    # Only fail on non-mutable files (mutable receipts appended post-seal)
                    from .packaging import MUTABLE_RECEIPT_NAMES as _mut, MUTABLE_RECEIPT_SUFFIXES as _muts
                    if file_path.name not in _mut and not file_path.name.endswith(_muts):  # v9.7.409 A10
                        checksum_errors.append(
                            f"line {line_no}: checksum mismatch for {rel_path.strip()}"
                        )
            if checksum_errors:
                errors.append(
                    f"checksums_sha256.txt has {len(checksum_errors)} verification failure(s): "
                    + "; ".join(checksum_errors[:3])
                    + (" ..." if len(checksum_errors) > 3 else "")
                )
        except Exception as exc:
            errors.append(f"checksums_sha256.txt read error: {type(exc).__name__}: {exc}")

    status = "FAIL" if missing or errors else "PASS_STRUCTURE"
    return {
        "status": status,
        "missing": missing,
        "errors": errors,
        "compact_report_count": len(compact_reports),
        "citation_count": len(ledger_records),
        "literature_task_count": len(workorder_tasks),
    }


CLAIM_SAFETY_WAIVER_NAME = "CLAIM_SAFETY_WAIVER.json"


def claim_safety_status_blocks(root: Path) -> tuple[bool, str]:
    """v9.7.409 B1 (owner ruling 2026-09-04: "validate must FAIL on claim_safety_status:FAIL").

    packaging.write_manifest records the seal-time claim-safety gate as
    manifest["claim_safety_status"] in {NOT_REQUESTED, PASS, FAIL}, but the PASS/FAIL chain below
    never read it, so a package whose own seal said FAIL still validated PASS. Returns
    (blocks, status). A waiver is a FILE beside STRICT_HEALTH_WAIVER.json -- never a hand-edited
    manifest key (manifest.json is engine-written and identity-bound)."""
    try:
        m = json.loads((Path(root) / "manifest.json").read_text(encoding="utf-8"))
    except Exception:
        return False, "UNREADABLE"  # an unreadable manifest already fails first in the chain
    status = str(m.get("claim_safety_status") or "ABSENT")
    # v9.7.412 (round 5, R3): the stored value is a SEAL-TIME receipt. A card added to the
    # checksum-exempt judgment/ subtree AFTER the seal is never re-scanned, so an overclaiming
    # card ("BGC001 produces **venezuelin**") left claim_safety_status at PASS — reproduced on a
    # fresh sealed package. Re-derive the gate over the package's current text files and take the
    # WORSE of stored and recomputed. The scan is the same one the seal ran, so this is symmetric
    # work, and it fails CLOSED: an unrunnable gate is a FAIL, never a silent pass.
    recomputed = ""
    try:
        from .claim_safety_gate import run_claim_safety_gate as _cs_gate
        recomputed = str(_cs_gate(Path(root)).get("claim_safety_status") or "")
    except Exception as exc:
        recomputed = "FAIL"
        status = f"{status}/RECHECK_UNAVAILABLE({type(exc).__name__})"
    if recomputed == "FAIL" and status != "FAIL":
        status = f"{status}/RECHECK_FAIL"
    if not (status.startswith("FAIL") or recomputed == "FAIL"):
        return False, status
    if (Path(root) / CLAIM_SAFETY_WAIVER_NAME).is_file():
        return False, f"{status}_WAIVED"
    return True, status



def _json_evidence_visibility(root, files):
    """Project saved parser facts into an advisory gate without inferring missing biology.

    Use the validator's containment-checked file list. Multiple receipts are
    ambiguous, never glob-first selected. Missing historical fields stay UNKNOWN.
    """
    import hashlib
    paths = sorted((p for p in files if p.name.endswith("_AntiSMASH_Evidence_Parse.json")),
                   key=lambda p: p.as_posix())
    result = {
        "policy": "ADVISORY_ONLY_NO_OVERALL_STATUS_CHANGE",
        "status": "UNKNOWN", "findings": [], "sources": [],
        "claim_ceiling": "Parser visibility only; no biological completeness or acceptance",
    }
    if not paths:
        result["findings"].append("EVIDENCE_RECEIPT_MISSING")
    elif len(paths) > 1:
        result["findings"].append("EVIDENCE_RECEIPT_AMBIGUOUS")
    for path in paths:
        entry = {"path": path.relative_to(root).as_posix(),
                 "attempt_scope": "Final returned parser receipt; interrupted attempts are not reconstructed",
                 "json_mode_requested": "UNKNOWN", "json_mode_effective": "UNKNOWN",
                 "json_bounded_truncated": "UNKNOWN", "parse_budget": "UNKNOWN",
                 "main_walker": {"completeness": "UNKNOWN"},
                 "record_extras": {"requested": "UNKNOWN", "completeness": "UNKNOWN"},
                 "tigrfam": {"requested": "UNKNOWN", "completeness": "UNKNOWN"}}
        try:
            raw = path.read_bytes()
            entry["sha256"] = hashlib.sha256(raw).hexdigest()
            ev = json.loads(raw)
            if not isinstance(ev, dict):
                raise ValueError("Evidence receipt is not an object")
        except (OSError, ValueError) as exc:
            entry["read_state"] = "PARSE_FAILED"
            entry["error"] = str(exc)
            result["findings"].append("EVIDENCE_RECEIPT_UNREADABLE")
            result["sources"].append(entry)
            continue
        entry["read_state"] = "OBSERVED"
        entry["raw_mode_fields"] = {key: ev[key] for key in ("json_mode_requested", "json_mode_effective", "json_mode", "json_bounded_truncated") if key in ev}
        requested = ev.get("json_mode_requested", "UNKNOWN")
        effective = ev.get("json_mode_effective", ev.get("json_mode", "UNKNOWN"))
        entry["json_mode_requested"] = requested if requested in ("off", "bounded", "full") else "UNKNOWN"
        entry["json_mode_effective"] = effective if effective in ("off", "bounded", "full") else "UNKNOWN"
        if "json_mode_effective" in ev and "json_mode" in ev and ev["json_mode_effective"] != ev["json_mode"]:
            result["findings"].append("JSON_EFFECTIVE_MODE_CONFLICT")
        if requested in ("bounded", "full") and effective == "off":
            result["findings"].append("JSON_MODE_DOWNGRADE")
        truncated = ev.get("json_bounded_truncated")
        if type(truncated) is bool:
            entry["json_bounded_truncated"] = truncated
        entry["parse_budget"] = ev.get("json_parse_budget", "UNKNOWN")
        entry["main_walker"].update({
            "scope": "Main JSON walker only; matched-leaf cap is not an extras or TIGRFAM completeness measure",
            "json_errors": ev.get("json_errors", "UNKNOWN"),
            "json_skipped": ev.get("json_skipped", "UNKNOWN"),
            "json_files": ev.get("json_files", "UNKNOWN"),
            "schema_warnings": ev.get("schema_warnings", "UNKNOWN"),
        })
        if effective == "off":
            entry["main_walker"]["completeness"] = "NOT_RUN"
        elif truncated is True:
            entry["main_walker"]["completeness"] = "TRUNCATED"
            result["findings"].append("MAIN_JSON_WALKER_TRUNCATED")
        elif ev.get("json_errors") or ev.get("json_skipped"):
            entry["main_walker"]["completeness"] = "PARSE_HELD"
            result["findings"].append("MAIN_JSON_PARSE_HELD")
        # No positive completeness assertion: absence of an error/cap field is not a completion receipt.
        for section, key in (("record_extras", "record_extras_requested"), ("tigrfam", "tigrfam_requested")):
            if type(ev.get(key)) is bool:
                entry[section]["requested"] = ev[key]
                if ev[key] is False:
                    entry[section]["completeness"] = "NOT_RUN"
            entry[section]["limit"] = "Independent record path; legacy extractor suppresses some errors and supplies no completion receipt"
        result["sources"].append(entry)
    result["findings"] = sorted(set(result["findings"]))
    if result["findings"]:
        result["status"] = "ADVISORY_FINDINGS"
    elif paths:
        result["status"] = "REPORTED_NOT_COMPLETENESS_VALIDATED"
    return result

def validate_package(package_dir: str | Path, gold_aware: bool = True,
                     enrichment_check: bool = False,
                     manifest_contract_check: bool = False,
                     write_status_receipt: bool = True) -> dict:
    """Validate a Mamey package.

    File-presence check always runs. When gold_aware is True (default) and the
    manifest mode is 'gold', additionally require that every BGC in the locked
    inventory has either a Mode B card or an explicit ledger entry in the
    judgment output — this is the extraction-side half of the v1.2 completeness
    gate. Mamey itself only produces the inventory and the depth-floor
    assignment; the actual Mode B / ledger entries are written by the judgment
    layer (Mamey v1.2 prompt). So a fresh extraction package reports
    JUDGMENT_PENDING for the gold completeness dimension rather than PASS.

    enrichment_check: when True, also check ENRICHMENT_SUFFIXES (OPEN_ME_FIRST.html,
    manifest_short.json). These are written after the core gate so this pass runs
    post-seal. The standalone `mamey validate` command always runs enrichment_check.

    write_status_receipt: when True (default) the recomputed package_status.json is
    written back into the package — the documented behaviour of the `mamey validate`
    command and of the seal pipeline, where package_status.json is a MUTABLE_RECEIPT.
    Read-only callers that must NOT touch the package they inspect (v9.7.409: the
    `seal-package` advisory QC gate, whose own receipts are redirected with --out) pass
    write_status_receipt=False; the receipt is still returned in-memory under
    result["package_status_receipt"], only the on-disk write is skipped.
    """
    root = Path(package_dir)
    from .packaging import PackageContainmentError, _package_files_fail_closed
    try:
        files = _package_files_fail_closed(root)
    except (PackageContainmentError, OSError) as exc:
        # v9.7.410 (CLAUDE_410_seal_locus_maps_coverage; SEAL_INTEGRITY_REATTACK A10): the .409
        # RGLOB-SEAL-COVERAGE fix makes _package_files_fail_closed RAISE on an unreadable/unlistable
        # subtree (os.walk onerror re-raise) instead of silently under-returning. That is the right
        # fail-closed semantics, but only PackageContainmentError was caught here, so the
        # PermissionError escaped `mamey validate` as an uncaught traceback rather than a verdict.
        # Convert it into the same clean typed FAIL the symlink case already gets. The enumeration
        # is the seal's single chokepoint: if it cannot complete, integrity is UNVERIFIED -> FAIL.
        if isinstance(exc, PackageContainmentError):
            error = f"package containment failure: {exc}"
            enumeration = "PACKAGE_CONTAINMENT"
        else:
            error = (f"SEAL_UNREADABLE_SUBTREE: package tree could not be fully enumerated "
                     f"({type(exc).__name__}: {exc}); integrity is UNVERIFIED, not clean")
            enumeration = "SEAL_UNREADABLE_SUBTREE"
        result = {
            "validator": f"Mamey v{__version__} package validator",
            "file_presence": "FAIL",
            "missing_required_suffixes": [],
            "package_containment": "FAIL",
            "containment_errors": [error],
            "seal_enumeration": enumeration,
            "checksum_integrity": "FAIL",
            "checksum_errors": [error],
            "status": "FAIL",
        }
        status_receipt = package_status_receipt(root, manifest={}, validator_status="FAIL")
        result["package_status"] = status_receipt["package_status"]
        result["package_status_receipt"] = status_receipt
        if manifest_contract_check:
            from .manifest_schema import check_package_contract
            result["manifest_contract_advisory"] = check_package_contract(root)
        return result

    suffixes_to_check = list(REQUIRED_SUFFIXES)
    if enrichment_check:
        suffixes_to_check += ENRICHMENT_SUFFIXES

    missing = [s for s in suffixes_to_check
               if not any(p.name.endswith(s) for p in files)]

    result = {
        "validator": f"Mamey v{__version__} package validator",
        "file_presence": "PASS" if not missing else "FAIL",
        "missing_required_suffixes": missing,
        "package_containment": "PASS",
    }

    result["json_evidence_visibility"] = _json_evidence_visibility(root, files)

    # _4B PKS-KS fragment scan (.359, phylogenomics-lane P358) — ADVISORY, non-blocking. Records presence only;
    # never added to missing/errors and never flips the overall status. A gold run always emits _4B, but a
    # zip with no PKS/KS content legitimately has none, so absence is a note (PKS_KS_SCAN_MISSING), not a fail.
    result["pks_ks_scan"] = ("PASS" if any(p.name.endswith("_4B_pks_ks_fragment_scan.csv") for p in files)
                             else "PASS_WITH_ISSUES (PKS_KS_SCAN_MISSING)")

    # Determine mode from manifest
    mode = None
    manifest_path = _find(files, "manifest.json")
    locked_ids: list[str] = []
    manifest_unreadable = False  # v9.7.335: fail closed when manifest.json will not parse
    if manifest_path:
        try:
            mdata = json.loads(manifest_path.read_text(encoding="utf-8"))
            mode = mdata.get("mode")
            locked_ids = [b.get("bgc_id") for b in mdata.get("bgcs", [])]
            result["manifest_parse"] = "PASS"
            try:
                normalized = normalize_bioactivity(mdata.get("bioactivity"))
                result["bioactivity_metadata_gate"] = "PASS"
                result["bioactivity_metadata_state"] = normalized["metadata_state"]
            except BioactivityMetadataError as exc:
                result["bioactivity_metadata_gate"] = exc.code
        except Exception as _mexc:
            # v9.7.335: swallowing this made a CORRUPT manifest score BETTER than a good one —
            # mode stayed None so the gold/depth-floor block never ran, mdata never bound so the
            # reporting-v2 gate became LEGACY_NOT_APPLICABLE, and the status fell through to PASS.
            # manifest.json is excluded from the checksum set, so nothing else catches it. Fail closed.
            result["manifest_parse"] = f"FAIL ({type(_mexc).__name__})"
            manifest_unreadable = True

    # ── v9.7.409 (BC hostile audit H3): identity binding ─────────────────────────
    # manifest.json `strain_id` swapped to a foreign id on a TEST-01 package validated PASS / MAMEY_COMPLETE and
    # mode-b then authored cards under the file-name strain. The package's identity must be ONE value:
    # the manifest's strain_id and the `<strain>_1_intake.json` prefix have to agree. Fail closed.
    try:
        _intake = _find(files, "1_intake.json")
        _prefix = _intake.name[: -len("_1_intake.json")] if _intake and _intake.name.endswith("_1_intake.json") else None
        _msid = mdata.get("strain_id") if "mdata" in locals() and isinstance(mdata, dict) else None
        if _prefix and _msid and str(_msid) != _prefix:
            result["identity_binding"] = "FAIL"
            result["identity_binding_detail"] = f"manifest strain_id {_msid!r} != package file prefix {_prefix!r}"
        elif _prefix and _msid:
            result["identity_binding"] = "PASS"
        else:
            result["identity_binding"] = "NOT_EVALUABLE"
    except Exception as _ib_exc:  # fail closed: an unverifiable binding is a FAIL
        result["identity_binding"] = f"ERROR: {_ib_exc}"

    # ── AMBER_EXCLUSION_HARDENING (.353): governed-output exclusion gate ────────
    # SSOT: OFFICIAL_DATA/exclusions.json -> exclusions.governed_excluded() = {AS-XXX}.
    # (1) A sealed GOVERNED package must not BE a hard-excluded strain's package.
    # (2) Any master workbook bundled into the package (a cross-strain governed
    #     table) must not carry a hard-excluded strain row. Fail closed on both.
    try:
        excl_offenders: dict = {}
        pkg_strain = mdata.get("strain_id") if "mdata" in locals() and isinstance(mdata, dict) else None
        from .exclusions import governed_excluded as _gov_excl
        if pkg_strain and str(pkg_strain) in _gov_excl():
            excl_offenders["<package strain_id>"] = {"*": [str(pkg_strain)]}
        master_xlsx = [p for p in files if p.suffix.lower() in (".xlsx", ".xlsm")
                       and "master" in p.name.lower()]
        if master_xlsx:
            mg = exclusion_gate.gate_governed_outputs(master_xlsx)
            if mg["status"] == "FAIL":
                excl_offenders.update(mg["offenders"])
        result["exclusion_gate"] = "FAIL" if excl_offenders else "PASS"
        if excl_offenders:
            result["exclusion_gate_offenders"] = excl_offenders
    except Exception as _eg_exc:  # fail closed: an unverifiable gate is a FAIL
        result["exclusion_gate"] = f"ERROR: {_eg_exc}"

    reporting_schema = (
        (mdata.get("reporting_features") or {}).get("schema_version")
        if "mdata" in locals() and isinstance(mdata, dict)
        else None
    )
    if reporting_schema == "sapote_reporting_features_v2":
        reporting = validate_reporting_v2_outputs(root)
        result["reporting_v2_gate"] = reporting["status"]
        result["reporting_v2"] = reporting
        if reporting["missing"]:
            result["missing_required_suffixes"].extend(reporting["missing"])
            result["file_presence"] = "FAIL"
    else:
        result["reporting_v2_gate"] = "LEGACY_NOT_APPLICABLE"


    # RG-GMCI gate — Mamey must complete this before triage/reporting.
    rg_path = _find(files, "4A_RGGMCI_full.json")
    rg_state = "FAIL_MISSING"
    rg_note = "RG-GMCI output file is missing."
    if rg_path:
        try:
            rg = json.loads(rg_path.read_text(encoding="utf-8"))
            rg_status = rg.get("status")
            _rg_ref_status = rg.get("reference_map_status")
            _rg_parse_errs = rg.get("reference_parse_error_count", 0) or 0
            if rg_status == "NULL_NO_RGGMCI_PAIRS" and _rg_parse_errs > 0:
                # GATE-11 (v9.7.338 verdict-changing): compute_rggmci sets
                # status=NULL_NO_RGGMCI_PAIRS purely from "no pairs produced", so a package
                # whose clusterblast reference geometry FAILED TO PARSE previously validated as
                # rggmci_gate: PASS — an integrity gate greenlighting a result built on
                # unparseable input. Fail closed when references were present but threw during
                # parsing (reference_parse_error_count > 0): the no-pairs verdict then rests on
                # unparseable input and is an unreliable empty set, not a genuine one.
                #
                # DESIGN NOTE (why parse_error_count, not reference_map_status): in rggmci.py
                # reference_map_status == "NULL_NO_CLUSTERBLAST_REFERENCES_PARSED" simply means
                # zero *usable* references, which is the NORMAL state for any antiSMASH package
                # produced without a knownclusterblast/clusterblast stage (no files present ->
                # parse_error_count == 0). Firing on that status alone (as the source addendum's
                # OR-form did) would fail essentially every reference-free package — turning a
                # targeted integrity gate into a blanket rejection. parse_error_count > 0 is the
                # true "present but did not parse" signal and matches the addendum's own
                # description of the known-bad artifact ("both present: NULL status AND
                # reference_parse_error_count > 0").
                rg_state = "FAIL"
                rg_note = (f"RG-GMCI status is {rg_status} but {_rg_parse_errs} clusterblast "
                           f"reference file(s) failed to parse "
                           f"(reference_map_status={_rg_ref_status}, "
                           f"reference_parse_error_count={_rg_parse_errs}); the no-pairs result is "
                           f"unreliable, not a genuine empty set.")
            elif rg_status in {"PASS", "NULL_NO_RGGMCI_PAIRS"}:
                rg_state = "PASS"
                rg_note = f"{rg_status}; {rg.get('pairs_total', 0)} pairs; {rg.get('reference_record_count', 0)} reference records."
            else:
                rg_state = "FAIL"
                rg_note = f"RG-GMCI status is {rg_status}; expected PASS or NULL_NO_RGGMCI_PAIRS."
        except Exception as exc:
            rg_state = "FAIL"
            rg_note = f"Could not parse RG-GMCI output: {type(exc).__name__}: {exc}"
    result["rggmci_gate"] = rg_state
    result["rggmci_note"] = rg_note

    # v9.7.136 citation-compact output gate — optional by presence.
    citation_compact = validate_citation_compact_outputs(root)
    result["citation_compact_gate"] = citation_compact["status"]
    result["citation_compact"] = citation_compact

    # Gold completeness dimension
    if gold_aware and mode == "gold":
        inv_path = _find(files, "2_inventory.csv")
        depth_assigned = 0
        full_mode_b = 0
        ledger = 0
        bgc_rows = 0
        if inv_path:
            with open(inv_path, newline="", encoding="utf-8") as f:
                for row in csv.DictReader(f):
                    bgc_rows += 1
                    depth = row.get("Depth_floor", "")
                    if depth:
                        depth_assigned += 1
                    if depth == "full_mode_b":
                        full_mode_b += 1
                    elif depth == "abbreviated_ledger":
                        ledger += 1

        # Look for judgment output (Mode B cards / report). These are written by
        # the v1.2 prompt, not by Mamey extraction.
        # v9.7.335: this was a filename substring test — one ZERO-BYTE file named "mode_b.md"
        # flipped the package to gold_completeness PASS with the note "All N BGCs accounted for
        # with Mode B cards", having opened nothing and counted nothing. Count authored cards per
        # BGC id instead, and carry the numerator/denominator into the note.
        _card_bgcs = set()
        # v9.7.409 (CLAUDE gate-completeness audit): content-commitment thresholds for a card to
        # count toward gold completeness. A real §1–§30 Mode-B card carries dozens of recognised
        # §-headings and thousands of authored chars; these floors reject filler while staying well
        # below any genuine (even abbreviated_ledger) card.
        _MIN_MODEB_SECTIONS = 3
        _MIN_MODEB_AUTHORED_CHARS = 200
        # v9.7.412 (R2): lexical-variety floors — see the filler note below. The dominant signal is
        # the single-token share (filler = 1.0; the 1,559-card corpus maxes at 0.131). The distinct
        # floor is set at 20 so that a MINIMAL card which only just clears the 200-char floor above
        # (~30 words) is never rejected for being short; the corpus minimum is 653.
        _MIN_MODEB_DISTINCT_TOKENS = 20
        _MAX_MODEB_TOP_TOKEN_SHARE = 0.5
        _card_seals: dict = {}
        _card_read_errors: list[dict[str, str]] = []
        _card_structure_errors: list[dict[str, str]] = []
        _card_seal_errors: list[dict[str, str]] = []
        for p in files:
            nm = p.name
            if "mode_b" not in nm.lower() and "Mode_B" not in nm:
                continue
            if p.suffix.lower() not in (".md", ".markdown"):
                continue
            # GATE-12 (v9.7.338 verdict-changing): the .335 hardening matched card FILENAMES
            # against locked_ids but never opened the files, so a set of empty/stub per-BGC
            # files (one AS-XXX_BGCNNN_mode_b.md per locked id) still flipped gold_completeness
            # to PASS with "All N BGCs accounted for". A 200-byte SIZE floor was added then — but
            # v9.7.409 (audit F3, `gate_gold_stub_cards_pass`) reproduced it end-to-end: 60
            # ~240-byte files of literal "x x x" (no §-headings, pure filler) still flipped a gold
            # package to PASS "All 60 BGCs accounted for with Mode B cards". `validate` never opened
            # the card. CONTENT-COMMITMENT: a card counts only if it carries REAL §-section content
            # — at least _MIN_MODEB_SECTIONS recognised Mode-B §-headings AND
            # _MIN_MODEB_AUTHORED_CHARS of non-whitespace body. Reuse the structure gate's own
            # heading detector (extract_section_titles) rather than stat().st_size.
            try:
                _card_text = p.read_text(encoding="utf-8", errors="replace")
            except OSError as exc:
                # An authored card that exists but cannot be read is not equivalent to a card
                # that has not been written yet.  The former means the judgment evidence is
                # present but unverifiable and must fail the gold completeness dimension; the
                # latter remains the ordinary JUDGMENT_PENDING state.
                _card_read_errors.append({
                    "path": _rel(root, p),
                    "error_type": type(exc).__name__,
                    "error": str(exc),
                })
                continue
            try:
                from .modeb_structure_gate import extract_section_titles as _extract_sections
                _sections = {n for (n, _t) in _extract_sections(_card_text)}
            except Exception as exc:
                # A broken structure evaluator is not evidence that an authored card is a
                # stub or absent.  Preserve the distinction so gold validation fails closed
                # instead of silently converting an infrastructure failure into pending work.
                _card_structure_errors.append({
                    "path": _rel(root, p),
                    "error_type": type(exc).__name__,
                    "error": str(exc),
                })
                _sections = set()
            _authored_chars = len("".join(_card_text.split()))
            if len(_sections) < _MIN_MODEB_SECTIONS or _authored_chars < _MIN_MODEB_AUTHORED_CHARS:
                # stub / filler card: real §-section content absent — do not count it.
                continue
            # v9.7.412 (round 5, R2): the .409 content-commitment floors count HEADINGS and CHARS,
            # so a card of `## §1..§4` plus 480 chars of literal "x x x" still flipped
            # gold_completeness to PASS (reproduced on a fresh sealed package). A real card is
            # LEXICALLY varied; filler is not. Corpus receipt over the 1,559 finished cards in the
            # W4 campaign: minimum distinct alphabetic tokens 653, maximum single-token share
            # 0.131 — so these floors sit an order of magnitude inside every real card.
            _words = re.findall(r"[A-Za-z][A-Za-z'-]{2,}", _card_text)
            _distinct = len({w.lower() for w in _words})
            _top_share = 0.0
            if _words:
                _counts: dict[str, int] = {}
                for _w in _words:
                    _lw = _w.lower(); _counts[_lw] = _counts.get(_lw, 0) + 1
                _top_share = max(_counts.values()) / len(_words)
            if _distinct < _MIN_MODEB_DISTINCT_TOKENS or _top_share > _MAX_MODEB_TOP_TOKEN_SHARE:
                continue
            m = re.search(r"(BGC\d{3,})", nm)
            if m:
                _card_bgcs.add(m.group(1))
                # v9.7.409 (audit F2, `gate_card_never_sealed`): the Mode-B card is the judgment
                # deliverable, yet it sits in the checksum-exempt mode_b/ subtree (F1), so a
                # post-seal swap/fabrication is invisible to `validate` — the one artifact carrying
                # every scientific claim has zero integrity coverage. Bind the COUNTED card's
                # content into the validate receipt: record a SHA256 keyed to its BGC id, so the
                # receipt now carries a verifiable hash of exactly the card credited toward
                # completeness. A post-seal card swap changes this hash (detectable on re-validate).
                try:
                    import hashlib as _hl
                    _card_seals[m.group(1)] = _hl.sha256(
                        _card_text.encode("utf-8", "replace")).hexdigest()
                except Exception as exc:
                    # A counted card without a receipt hash is not sealed judgment evidence.
                    # Record the exact failure and fail the gold dimension below.
                    _card_seal_errors.append({
                        "path": _rel(root, p),
                        "bgc_alias": m.group(1),
                        "error_type": type(exc).__name__,
                        "error": str(exc),
                    })
        _locked = {b for b in locked_ids if b}
        has_mode_b_cards = bool(_locked) and _locked.issubset(_card_bgcs)

        every_bgc_assigned = (bgc_rows > 0 and depth_assigned == bgc_rows
                              and len(locked_ids) == bgc_rows)

        if not every_bgc_assigned:
            gold_state = "FAIL"
            gold_note = (f"Depth floor not assigned to all BGCs "
                         f"({depth_assigned}/{bgc_rows} assigned; "
                         f"{len(locked_ids)} locked).")
        elif _card_read_errors:
            gold_state = "FAIL"
            gold_note = (
                f"Mode B judgment evidence is unreadable for {len(_card_read_errors)} "
                "authored card file(s); gold completeness is unverifiable, not pending."
            )
        elif _card_structure_errors:
            gold_state = "FAIL"
            gold_note = (
                f"Mode B card structure could not be evaluated for "
                f"{len(_card_structure_errors)} authored card file(s); gold completeness "
                "is unverifiable, not pending."
            )
        elif _card_seal_errors:
            gold_state = "FAIL"
            gold_note = (
                f"Mode B judgment evidence could not be sealed for "
                f"{len(_card_seal_errors)} counted card file(s); gold completeness "
                "is unverifiable."
            )
        elif not has_mode_b_cards:
            gold_state = "JUDGMENT_PENDING"
            gold_note = (f"Extraction complete: all {bgc_rows} BGCs have a depth-floor "
                         f"assignment ({full_mode_b} full Mode B, {ledger} ledger). "
                         f"Mode B cards / reports are written by the Mamey v1.2 "
                         f"judgment layer — load manifest.json into the prompt to "
                         f"complete gold mode.")
        else:
            gold_state = "PASS"
            gold_note = (f"All {bgc_rows} BGCs accounted for with Mode B cards or "
                         f"ledger entries ({len(_card_bgcs)}/{len(_locked)} BGC ids matched "
                         f"an authored card file).")

        result["gold_completeness"] = gold_state
        result["gold_note"] = gold_note
        if _card_read_errors:
            result["mode_b_card_read_errors"] = _card_read_errors
        if _card_structure_errors:
            result["mode_b_card_structure_errors"] = _card_structure_errors
        if _card_seal_errors:
            result["mode_b_card_seal_errors"] = _card_seal_errors
        # v9.7.409 (audit F2, `gate_card_never_sealed`): expose the per-BGC SHA256 of every card
        # that passed content-commitment and was counted toward completeness, binding the judgment
        # deliverable's content into the (otherwise card-blind) validate receipt.
        result["mode_b_card_seals"] = _card_seals
        result["depth_floor_breakdown"] = {
            "total_bgcs": bgc_rows,
            "full_mode_b": full_mode_b,
            "abbreviated_ledger": ledger,
        }

    # Overall status
    # v9.7.335: an unreadable manifest is checked FIRST. Previously the parse error was swallowed,
    # which disabled the gold/depth-floor block and the reporting-v2 gate and let the status fall
    # through to PASS — a corrupt package scored strictly better than a good one.
    _cs_blocks, result["claim_safety_status"] = claim_safety_status_blocks(root)  # v9.7.409 B1
    if manifest_unreadable:
        result["status"] = "FAIL"
    elif result["file_presence"] == "FAIL":
        result["status"] = "FAIL"
    elif str(result.get("identity_binding", "PASS")).startswith(("FAIL", "ERROR")):
        result["status"] = "FAIL"   # v9.7.409: manifest identity must match the package files
    elif result.get("rggmci_gate") != "PASS":
        result["status"] = "FAIL"
    elif result.get("citation_compact_gate") == "FAIL":
        result["status"] = "FAIL"
    elif result.get("reporting_v2_gate") == "FAIL":
        result["status"] = "FAIL"
    elif result.get("exclusion_gate", "PASS") != "PASS":
        # AMBER_EXCLUSION_HARDENING (.353): a hard-excluded strain reached a governed output.
        result["status"] = "FAIL"
    elif _cs_blocks:
        # v9.7.409 B1: the seal-time claim-safety gate wrote FAIL; surface it, do not hide it.
        result["status"] = "FAIL"
    elif result.get("gold_completeness") == "FAIL":
        result["status"] = "FAIL"
    elif result.get("gold_completeness") == "JUDGMENT_PENDING":
        result["status"] = "MAMEY_COMPLETE"
    else:
        result["status"] = "PASS"

    # v9.7.145c: wire verify_checksums so corrupted non-mutable files cannot return PASS
    try:
        checksum_errors = verify_checksums(root)
        if checksum_errors:
            result["checksum_integrity"] = "FAIL"
            result["checksum_errors"] = checksum_errors[:5]
            result["status"] = "FAIL"
        else:
            result["checksum_integrity"] = "PASS"
    except Exception as _cs_exc:
        # Fail-closed: if the integrity check itself cannot complete, the package's integrity is
        # UNVERIFIED — it must not be allowed to return PASS/MAMEY_COMPLETE (v9.7.145c intent).
        result["checksum_integrity"] = f"ERROR: {_cs_exc}"
        result["status"] = "FAIL"

    # ── v9.7.409 r2 (CLAUDE_409_seal_integrity_r2): provenance-binding + post-seal anchor gates ──
    # These RAISE THE BAR against naive tampering of the otherwise-unsealed provenance surface; they
    # are NOT cryptographic authenticity (an insider who rewrites every in-package copy stays
    # self-consistent — external signing is required). Each fails CLOSED only on a positive
    # contradiction, and is NOT_EVALUABLE when its inputs are absent, so partial/fixture/legit
    # packages keep validating.
    _mdata_for_bind = mdata if 'mdata' in locals() and isinstance(mdata, dict) else {}
    # (R2-1) anchor: post_seal_checksums.txt must match its anchor folded into the core seal.
    try:
        from .packaging import verify_post_seal_anchor as _vpsa
        _anchor_errs = _vpsa(root)
    except Exception as _pa_exc:  # fail closed: an unverifiable anchor is a FAIL
        _anchor_errs = [f"post_seal anchor check errored: {_pa_exc}"]
    if _anchor_errs:
        result["post_seal_anchor"] = "FAIL"
        result["post_seal_anchor_errors"] = _anchor_errs[:5]
        result["status"] = "FAIL"
    elif (root / "post_seal_checksums.txt").exists():
        result["post_seal_anchor"] = "PASS"
    # (provenance A) recompute the determinism fingerprint from the files; FAIL a forged receipt.
    try:
        from .packaging import verify_repro_fingerprint_recompute as _vfp
        _fp_gate = _vfp(root)
        result["repro_fingerprint_recomputed"] = _fp_gate["state"]
        if _fp_gate.get("detail"):
            result["repro_fingerprint_recomputed_detail"] = _fp_gate["detail"]
        if _fp_gate["state"] == "FAIL":
            result["status"] = "FAIL"
    except Exception as _fp_exc:  # fail closed
        result["repro_fingerprint_recomputed"] = f"ERROR: {_fp_exc}"
        result["status"] = "FAIL"
    # (provenance C5/D, PROV-01) bind load-bearing manifest fields to the checksum-covered tables.
    try:
        from .packaging import verify_manifest_provenance_binding as _vmb
        _mb = _vmb(root, _mdata_for_bind)
        result["manifest_provenance_binding"] = _mb["state"]
        result["engine_version_binding"] = _mb.get("engine_version_binding", "NOT_EVALUABLE")
        if _mb.get("detail"):
            result["manifest_provenance_binding_detail"] = _mb["detail"]
        if _mb["state"] == "FAIL":
            result["status"] = "FAIL"
    except Exception as _mb_exc:  # fail closed
        result["manifest_provenance_binding"] = f"ERROR: {_mb_exc}"
        result["status"] = "FAIL"

    status_receipt = package_status_receipt(root, manifest=mdata if 'mdata' in locals() else {}, validator_status=result.get("status"))
    result["package_status"] = status_receipt["package_status"]
    result["package_status_receipt"] = status_receipt
    if manifest_contract_check:
        # Advisory in this cut: observable in the receipt, never included in
        # the status decision above and therefore never blocks validation.
        from .manifest_schema import check_package_contract
        result["manifest_contract_advisory"] = check_package_contract(root)
    # v9.7.409 (CLAUDE_409_sealpackage_mutation): the on-disk rewrite is the ONLY step in
    # validate_package that mutates the package. A read-only caller (seal-package's advisory
    # QC gate) sets write_status_receipt=False so its inspection cannot change a byte — or the
    # mtime — of the sealed package, honouring the non-mutation contract of `--out`. The
    # recomputed receipt is still available in-memory via result["package_status_receipt"].
    if write_status_receipt:
        try:
            write_package_status_receipt(root, manifest=mdata if 'mdata' in locals() else {}, validator_status=result.get("status"))
        except Exception as _status_write_exc:
            result["package_status_receipt_write"] = {
                "state": "FAILED",
                "error_type": type(_status_write_exc).__name__,
                "error": str(_status_write_exc),
            }

    return result


# ── v9.7.123 (SM-P1-008): Workbook populated-sheet gate ───────────────────────

# Logical workbook-content gates. Each gate accepts one or more physical sheet/header
# spellings so older node-first workbooks (which used ``BGC`` instead of ``BGC_ID`` or the
# pre-SM-P1-008 sheet name ``Triage_First_Board``) do not fail simply for naming. New
# workbooks should still emit the canonical sheet names as their first alias. (v9.7.128)
_REQUIRED_SHEETS = {
    "BGC_Inventory": [("BGC_Inventory", ("BGC_ID", "BGC"))],
    "Triage_Board": [("Triage_Board", ("BGC_ID", "BGC")),
                     ("Triage_First_Board", ("BGC_ID", "BGC"))],
    "WetLab_Decision_Matrix": [("WetLab_Decision_Matrix", ("BGC_ID", "BGC"))],
    "Mode_B_Summary": [("Mode_B_Summary", ("BGC_ID", "BGC"))],
}
_MAX_ROWS = 200  # cap per sheet to avoid memory issues on large workbooks


def _is_pending(value) -> bool:
    """True if the primary-key cell value is PENDING or empty."""
    if value is None:
        return True
    s = str(value).strip()
    return s == "" or s.upper() == "PENDING"


def validate_workbook_content(workbook_path: str | Path) -> dict:
    """Validate that required workbook sheets are present and have real primary-key rows.

    Per-sheet status:
      PASS            — sheet present, at least one non-PENDING primary-key row
      FAIL_empty      — sheet present, header only (no data rows)
      FAIL_all_pending — sheet present, all data-row primary keys are PENDING or empty
      MISSING         — sheet not in the workbook

    Overall status: PASS only when every sheet is PASS.

    Cap: at most 200 rows per sheet are inspected (openpyxl read_only=True).
    """
    try:
        import openpyxl
    except ImportError:
        return {"status": "FAIL", "error": "openpyxl not installed", "sheets": {}}

    path = Path(workbook_path)
    sheet_results: dict[str, dict] = {}

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        present = set(wb.sheetnames)
    except Exception as exc:
        return {"status": "FAIL", "error": str(exc), "sheets": {}}

    for gate_name, aliases in _REQUIRED_SHEETS.items():
        # resolve the first present physical sheet for this logical gate
        selected = None
        for physical_sheet, pk_cols in aliases:
            if physical_sheet in present:
                selected = (physical_sheet, tuple(pk_cols))
                break
        if selected is None:
            sheet_results[gate_name] = {"status": "MISSING", "pk_col": aliases[0][1][0],
                                        "sheet": None, "rows_checked": 0, "real_rows": 0}
            continue
        sheet_name, pk_cols = selected

        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)

        # Find the header row to locate the pk column index
        try:
            header = next(rows_iter)
        except StopIteration:
            sheet_results[gate_name] = {"status": "FAIL_empty", "pk_col": pk_cols[0],
                                        "sheet": sheet_name, "rows_checked": 0, "real_rows": 0}
            continue

        # pk column index (0-based): first header matching any accepted pk spelling
        header_lower = [str(h).strip().lower() if h else "" for h in header]
        pk_choices = [c.lower() for c in pk_cols]
        pk_idx = next(
            (i for i, h in enumerate(header_lower) if h in pk_choices),
            None  # SM-P1-008: None -> the required primary-key column is absent from the header
        )
        if pk_idx is None:
            # FAIL_missing_primary_key: the sheet exists but lacks its required pk column.
            # Silently reading column 0 would mask a malformed sheet, so flag it explicitly.
            sheet_results[gate_name] = {"status": "FAIL_missing_primary_key", "pk_col": pk_cols[0],
                                        "sheet": sheet_name, "rows_checked": 0, "real_rows": 0}
            continue

        real_rows = 0
        rows_checked = 0
        for row in rows_iter:
            if rows_checked >= _MAX_ROWS:
                break
            rows_checked += 1
            pk_val = row[pk_idx] if pk_idx < len(row) else None
            if not _is_pending(pk_val):
                real_rows += 1

        if rows_checked == 0:
            status = "FAIL_empty"
        elif real_rows == 0:
            status = "FAIL_all_pending"
        else:
            status = "PASS"

        sheet_results[gate_name] = {
            "status": status,
            "pk_col": pk_cols[0],
            "sheet": sheet_name,
            "rows_checked": rows_checked,
            "real_rows": real_rows,
        }

    wb.close()

    overall = "PASS" if all(v["status"] == "PASS" for v in sheet_results.values()) else "FAIL"
    return {"status": overall, "sheets": sheet_results, "path": str(path)}


def _checksum_reciprocal_exempt(rel: str) -> bool:
    """SEAL-03: True if a package-relative path is legitimately allowed to be PRESENT in a
    sealed package but ABSENT from checksums_sha256.txt.

    verify_checksums only re-hashes files LISTED in the manifest (the forward direction), so a
    file added/injected into a sealed tree after the seal — one that is *not* tracked — slips
    past integrity entirely. The reciprocal scan closes that hole, but it must NOT flag the
    package's many legitimate post-seal artifacts, which are written AFTER the writer
    (packaging.write_manifest) captures the file list and so never enter the checksum set:

      * manifest.json / checksums_sha256.txt themselves (writer excludes both);
      * the writer's MUTABLE_NAMES (run_phase_receipts.jsonl, package_status.json,
        claim_safety_status.json, repro_fingerprint.json) + the per-strain
        *_judgment_register.json (rewritten by `ingest-receipts`);
      * everything under the post-seal figure dirs (smoke_figures/, gold_figures/, locus_maps/)
        — re-rendered by the figure phase / `render-figures`; their png/svg are already excluded
        by packaging.is_checksum_excluded, and their fig_*_data.csv / *_locus_map_data.csv
        companions are emitted post-capture too, so the whole subtree is post-seal;
      * package-ROOT figure images and their *_fig_*_data.csv companions, plus the supplementary
        presentation docs (PRINT_FIGURE_PACK.md / .pdf, FIGURES_SUPPLEMENTARY.md, *_strain_brief.pdf)
        and *_cnbu.json — all rendered/emitted after the seal.

    The predicate is intentionally broader than the writer's own exclusion set (which only sees
    files that exist at capture time); it mirrors "what a clean re-seal legitimately leaves
    untracked." Anything NOT matching here that is present-but-untracked is a real integrity
    error (a stray or injected file)."""
    rel = rel.replace("\\", "/")
    name = rel.rsplit("/", 1)[-1]
    # Writer-excluded singletons + mutable post-seal receipts.
    from .packaging import MUTABLE_RECEIPT_NAMES as _mut
    if name in ({'manifest.json', 'checksums_sha256.txt'} | set(_mut)):
        return True
    from .packaging import MUTABLE_RECEIPT_SUFFIXES as _muts
    if name.endswith(_muts):  # v9.7.409 A10: judgment register + timing breakdown
        return True
    # Whole governed post-seal output subtrees. `render-all-figures` gathers into
    # figures/ and writes its native figure suite to figures_rendered/.  The
    # domain-level command is likewise explicitly post-seal and writes both its
    # evidence tables and figures beneath domain_level/.  These trees are not
    # members of the sealed core checksum set.
    # v9.7.410 (CLAUDE_410_seal_locus_maps_coverage; SEAL_INTEGRITY_REATTACK A9): locus_maps/ is
    # deliberately NOT in the blanket prefix list below any more. It used to be, while
    # packaging.is_post_seal_covered pinned only *_data.csv/*.json beneath it — so a fabricated
    # locus_maps/*.md, *.txt or *.pdf was neither hash-pinned (SEAL-05) nor flagged as untracked
    # (SEAL-03) and validated MAMEY_COMPLETE (verified by execution). The exemption is now exactly
    # the union of what SEAL-05 owns (is_post_seal_covered — every non-image file, enforced when
    # post_seal_checksums.txt is present) and the documented non-deterministic image re-renders
    # (is_checksum_excluded: .png/.svg). Anything else under locus_maps/ is a real untracked file.
    if rel.startswith('locus_maps/'):
        from .packaging import is_post_seal_covered as _lm_covered
        from .packaging import is_checksum_excluded as _lm_image
        return _lm_covered(rel) or _lm_image(rel)
    if rel.startswith((
        'smoke_figures/',
        'gold_figures/',
        'figures/',
        'figures_rendered/',
        'domain_level/',
        # BLASTP overlays are governed, non-scoring post-seal evidence. The ingester writes
        # channel-preserving stores plus immutable quarantine and receipt artifacts only after
        # the sealed core checksum manifest exists.
        'blastp_online/',
        'blastp_quarantine/',
        'blastp_ingest_receipts/',
        # Canonical Sapote workflow outputs are intentionally authored after Tier-1 sealing.
        # Their own gates/receipts govern content; they must not invalidate the immutable core.
        'mode_b_templates/',
        'judgment/',
        'guide/',
        # v9.7.371 fix: chatgpt_commands.py's ChatGPT-safe post-seal render-figures/mode-b
        # commands default to writing straight into the package (mamey_native_figures/,
        # cohort_figures/, mode_b/) -- exactly the module's own documented "post-seal,
        # ChatGPT-safe" use case -- but none of the three were in this exemption list, so a
        # `mamey validate` run afterward spuriously flagged every file they wrote.
        'mamey_native_figures/',
        'cohort_figures/',
        'mode_b/',
        # v9.7.371 fix (follow-on to the chatgpt_commands.py exemption fix above): assembly_line.py's
        # own docstring says it emits per-BGC assembly lines "for a sealed package" -- explicitly a
        # post-seal, non-blocking additive step (like render-figures), but ASSEMBLY_LINES/ was never
        # added here, so the documented workflow spuriously fails its own subsequent checksum
        # verification.
        'ASSEMBLY_LINES/',
        # v9.7.374 fix: same class as the ASSEMBLY_LINES/ fix directly above, found in TWO sibling
        # "FREEZE-SAFE ADDITIVE REPORT LAYER... post-seal, non-blocking, like render-figures" modules
        # that were never added here despite matching ASSEMBLY_LINES/'s exact documented shape and
        # having real, registered CLI commands (`mamey p450-tailoring`, `mamey compound-families`):
        # p450_tailoring.py writes P450_TAILORING/; compound_family_report.py writes
        # COMPOUND_FAMILIES/. Either command followed by `mamey validate <sealed_pkg>` (a supported,
        # documented workflow) spuriously flipped checksum_integrity to FAIL on a genuinely valid
        # package.
        'P450_TAILORING/',
        'COMPOUND_FAMILIES/',
    )):
        return True
    # packaging.is_checksum_excluded — figure png/svg + supplementary presentation PDFs.
    from .packaging import is_checksum_excluded as _fig_excluded
    if _fig_excluded(rel):
        return True
    # Package-ROOT figure images + their data companions (rendered post-capture).
    if '/' not in rel and name.endswith(('.png', '.svg')):
        return True
    if '_fig_' in name and name.endswith('_data.csv'):
        return True
    if name in {
        'PRINT_FIGURE_PACK.md',
        'FIGURES_SUPPLEMENTARY.md',
        'figure_manifest_print.csv',
        'render_all_figures_summary.json',
    }:
        return True
    if name.endswith(('_compiled_report.md', '_SAPOTE_WORKFLOW_LEDGER.md')):
        return True
    if name.endswith('_cnbu.json'):
        return True
    return False


def verify_checksums(package_dir) -> list[str]:
    """Recompute every line in checksums_sha256.txt and return a list of error strings.

    v9.7.145c: mutable post-seal files should NOT be in checksums_sha256.txt at all.
    If they are present, they are skipped with a warning. Non-mutable file mismatches
    are always errors. Self-entry (checksums_sha256.txt itself) is always skipped.

    SEAL-03: also performs the RECIPROCAL scan — every file present in the tree but ABSENT
    from checksums_sha256.txt is an integrity error unless it is a legitimate post-seal
    artifact (see _checksum_reciprocal_exempt). This catches files injected into a sealed
    package after the seal, which the forward (line-by-line) pass cannot see.
    """
    import hashlib
    from pathlib import Path, PurePosixPath, PureWindowsPath
    root = Path(package_dir)
    checksum_path = root / 'checksums_sha256.txt'
    errors = []
    if not checksum_path.exists():
        return ['checksums_sha256.txt missing']
    # Fail before hashing any tracked entry.  Path.is_file() follows symlinks, so without this
    # preflight a package-local link could make validation read and bless bytes outside the
    # package root.  Reuse the writer's exact admission rule so seal and validation cannot drift.
    from .packaging import PackageContainmentError, _package_files_fail_closed
    try:
        package_files = _package_files_fail_closed(root)
    except PackageContainmentError as exc:
        return [f'package containment failure: {exc}']
    except OSError as exc:
        # v9.7.410 (A10): unreadable subtree -> typed clean error string, same fail-closed shape
        # as the containment refusal above (see validate_package for the rationale).
        return [f'SEAL_UNREADABLE_SUBTREE: package tree could not be fully enumerated '
                f'({type(exc).__name__}: {exc}); integrity is UNVERIFIED, not clean']
    admitted = {p.relative_to(root).as_posix(): p for p in package_files}
    # Known mutable files that may be appended after checksum generation
    from .packaging import MUTABLE_RECEIPT_NAMES as mutable
    # v9.7.152: per-strain judgment register is rewritten post-seal by `ingest-receipts`
    # (its documented purpose), so exempt by suffix — it is prefixed per strain
    # (AS-XXX_judgment_register.json, SID-001_judgment_register.json, …).
    # See PATCH_AUDIT_v9.7.151c_Part2 Finding 1.
    from .packaging import MUTABLE_RECEIPT_SUFFIXES as mutable_suffixes  # v9.7.409 A10: single source, writer + validator in lockstep
    from .packaging import is_checksum_excluded as _fig_excluded
    tracked: set[str] = set()
    for line_no, line in enumerate(checksum_path.read_text(encoding="utf-8").splitlines(), 1):
        line = line.strip()
        if not line or line.startswith('#'):
            continue
        try:
            expected, rel = line.split(None, 1)
        except ValueError:
            errors.append(f'line {line_no}: malformed checksum line')
            continue
        # Check lexical containment BEFORE exemptions or filesystem reads. A checksum
        # record is untrusted: absolute/parent paths bypass the package symlink scan.
        # Normalize both separator styles for portable manifests and reciprocal parity.
        rel = rel.strip().replace("\\", "/")
        logical = PurePosixPath(rel)
        if ("\x00" in rel or logical.is_absolute() or PureWindowsPath(rel).drive
                or ".." in logical.parts or not logical.parts):
            errors.append(f'line {line_no}: unsafe checksum path {rel!r}')
            continue
        rel = logical.as_posix()
        tracked.add(rel)
        name = logical.name
        # Skip self-entry
        if name == 'checksums_sha256.txt':
            continue
        # Skip mutable files (warn but do not fail)
        if name in mutable or name.endswith(mutable_suffixes):
            continue
        # v9.7.160: skip post-seal figure images (rendered after the seal; excluded from the
        # checksum set by packaging.is_checksum_excluded). A stale figure entry from an older
        # checksums file must be skipped, not failed — mirrors the writer-side exclusion.
        if _fig_excluded(rel):
            continue
        path = admitted.get(rel)
        if path is None:
            errors.append(f'line {line_no}: missing {rel} (not an admitted regular package file)')
            continue
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual != expected:
            errors.append(f'line {line_no}: checksum mismatch for {rel}')
    # SEAL-04: a manifest that exists but lists NOTHING is not a clean package — it is an
    # unverified one. Both passes are silent on it: the forward loop has no lines to check, and
    # the reciprocal scan below is guarded on `tracked` being non-empty (correctly, so it does
    # not flag every file as untracked). The net effect was a zero-scan PASS — emptying
    # checksums_sha256.txt in a tampered sealed package returned no errors at all, so
    # validate_package reported checksum_integrity=PASS and status=MAMEY_COMPLETE, identical to
    # the pristine package. A missing manifest was already an error; an empty one must be too.
    # The wording deliberately avoids the substring "untracked" so the SEAL-03 guarantee tested
    # by test_empty_checksums_manifest_does_not_trigger_reciprocal (no per-file untracked
    # errors) is preserved exactly.
    if not tracked:
        errors.append('checksums_sha256.txt lists no files (empty or comment-only manifest); '
                      'package integrity is UNVERIFIED, not clean')
    # SEAL-03 reciprocal scan: files present in the tree but not tracked by the manifest.
    # Guard: only meaningful when the checksum manifest actually tracks something. A real
    # sealed package always lists 100+ files (packaging.write_manifest), so `tracked` is
    # never empty in practice; an empty checksums_sha256.txt tracks nothing and there is no
    # file set to be reciprocal *to* — skipping avoids flagging every file as "untracked" for
    # a degenerate/empty manifest (which the forward pass already treats as "nothing to
    # verify"), and keeps the reciprocal contract "present but not in a real manifest."
    if tracked:
        for p in package_files:
            rel = p.relative_to(root).as_posix()
            if rel in tracked:
                continue
            # v9.7.409 (CLAUDE gate-completeness audit F1, `gate_exemption_injection`):
            # _checksum_reciprocal_exempt exempts whole DELIVERABLE subtrees by prefix, so a
            # fabricated file dropped into one passes `validate` with arbitrary, unverified content
            # (audit probe: `judgment/fake_verdict.json {"verdict":"PASS"}` → MAMEY_COMPLETE / PASS).
            # The GOVERNED shape of judgment/ is authored Markdown: judgment_store and
            # compile_report write only `.md` there (*_mode_b.md, *_laypersons_section.md,
            # *_fermentation_section.md, *_<section>.md). A machine artifact injected there is not a
            # governed post-seal output — refuse an untracked non-`.md` file placed DIRECTLY under
            # judgment/ even though the broad prefix would otherwise exempt it. Governed mutable
            # receipt names (the SSOT for legitimately-untracked files) are still honoured.
            _low = rel.lower()
            _bn = rel.rsplit("/", 1)[-1]
            if (_low.startswith("judgment/") and _low.count("/") == 1
                    and not _low.endswith((".md", ".markdown"))
                    and _bn not in mutable and not _bn.endswith(mutable_suffixes)):
                errors.append(
                    "ungoverned file injected under exempt judgment/ subtree "
                    f"(not an authored .md deliverable, not a governed receipt): {rel}")
                continue
            if _checksum_reciprocal_exempt(rel):
                continue
            errors.append(f'untracked file present but absent from checksums_sha256.txt: {rel}')
    # SEAL-05 (v9.7.409, CLAUDE post_seal_checksums lane; audit N1/N2/N9): enforce the tracked
    # post-seal integrity manifest when the package carries one. The core seal above is a
    # point-in-time snapshot that leaves the post-seal DATA/deliverable subtrees (mode_b/,
    # blastp_online/, domain_level/, guide/, figures*/…, locus_maps data, root figure data)
    # reciprocal-EXEMPT, so injection/swap/delete there is invisible. post_seal_checksums.txt (written
    # by packaging.write_post_seal_checksums at end of run) hash-pins exactly those covered files;
    # ABSENCE of the manifest = the old behaviour (a package sealed before this patch still validates),
    # PRESENCE = enforced.
    errors.extend(_verify_post_seal_checksums(root, admitted))
    return errors


def _verify_post_seal_checksums(root, admitted: dict) -> list[str]:
    """SEAL-05 reciprocal integrity for post-seal DATA/deliverable files.

    Reads post_seal_checksums.txt (if present) and reports, for the CONSERVATIVE covered set
    (packaging.is_post_seal_covered — never the genuinely-mutable receipts or regenerated
    reports):
      * a covered file present on disk but absent from the manifest  -> INJECTION,
      * a covered file whose bytes no longer match the recorded hash -> SWAP,
      * a file listed in the manifest but missing on disk            -> DELETE.

    `admitted` is verify_checksums' {rel: Path} of admitted (symlink-refused) regular files, reused
    so this scan cannot read outside the package root. Absent manifest -> [] (backward-compatible)."""
    from pathlib import PurePosixPath, PureWindowsPath
    post_seal_path = root / 'post_seal_checksums.txt'
    if not post_seal_path.exists():
        return []
    import hashlib
    from .packaging import is_post_seal_covered
    errors: list[str] = []
    recorded: dict[str, str] = {}
    try:
        raw = post_seal_path.read_text(encoding="utf-8")
    except OSError as exc:
        return [f'post_seal_checksums.txt unreadable: {exc}']
    for line_no, line in enumerate(raw.splitlines(), 1):
        line = line.strip()
        if not line or line.startswith('#'):
            continue
        try:
            expected, rel = line.split(None, 1)
        except ValueError:
            errors.append(f'post_seal line {line_no}: malformed checksum line')
            continue
        rel = rel.strip().replace("\\", "/")
        logical = PurePosixPath(rel)
        if ("\x00" in rel or logical.is_absolute() or PureWindowsPath(rel).drive
                or ".." in logical.parts or not logical.parts):
            errors.append(f'post_seal line {line_no}: unsafe path {rel!r}')
            continue
        recorded[logical.as_posix()] = expected
    # Forward: every recorded covered file must be present and byte-identical (catches SWAP + DELETE).
    for rel, expected in recorded.items():
        path = admitted.get(rel)
        if path is None:
            errors.append(f'post-seal file deleted (listed in post_seal_checksums.txt but absent): {rel}')
            continue
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual != expected:
            errors.append(f'post-seal checksum mismatch for {rel}')
    # Reciprocal: every covered file on disk must be recorded (catches INJECTION into a covered subtree).
    for rel in admitted:
        if is_post_seal_covered(rel) and rel not in recorded:
            errors.append(f'injected file present in covered post-seal subtree but absent from '
                          f'post_seal_checksums.txt: {rel}')
    return errors
