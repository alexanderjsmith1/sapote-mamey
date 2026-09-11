"""Layered per-gene BGC Guide deliverable — lay→technical, Structure/Function/BLASTp per gene."""
from __future__ import annotations

try:  # pragma: no cover - import shape depends on package vs direct-script use
    from .console import emit
except ImportError:  # direct execution: no parent package to resolve against.
    # v9.7.407: a bare-script run (documented for workbook_schema_check.py) has neither a
    # parent package NOR the bundle root on sys.path, so put the root there first.
    import os, sys
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from mamey.console import emit

# The BLASTp similarity readout is rendered from blastp_evidence_store fields (pct_identity,
# pct_positive, query_coverage, evalue, bitscore, subject_accession, evidence_tier, claim_safety) —
# the Guide does NOT compute similarity; it renders the store's existing tiering. This is what keeps
# every per-gene claim capped at its BLASTp evidence_tier. All coordinate/domain content comes from
# the sealed package's gene_by_gene_all_bgcs.csv. The module renders banked data; it does not re-scan.

import csv, json, os, re
from pathlib import Path
from typing import Any

try:
    from .blastp_evidence_store import HIT_FIELDS  # REUSE, do not re-derive
except Exception:  # pragma: no cover - store is always present in-bundle
    HIT_FIELDS = [
        "round_id", "strain", "bgc_id", "query_gene", "node", "region", "query_id",
        "subject_id", "subject_accession", "subject_title", "subject_sciname", "subject_taxid",
        "pct_identity", "pct_positive", "align_len", "query_len", "query_coverage",
        "evalue", "bitscore", "evidence_tier", "next_action", "claim_safety",
    ]

GUIDE_PARTS = ["P1", "P2", "P3", "P4", "P5"]
_TIER_GLOSS = {
    "TIER_A_CLASS_DEFINING": "A (class-defining)",
    "TIER_B_NEIGHBORHOOD_SUPPORT": "B (neighborhood support)",
    "TIER_C_GENERIC_FUNCTION_CONFIRMED": "C (generic function)",
    "TIER_D_UNINFORMATIVE_OR_WEAK": "D (uninformative/weak)",
}
# role grouping for the gene catalogue (Part 4)
_ROLE_ORDER = ["core", "resistance", "regulation", "mobile_element", "accessory"]
CLAIM_SAFETY_LINE = (
    "Claim-safety: biosynthetic capacity consistent with the evidence below; not a production or "
    "bioactivity claim. KCB/BLASTp = similarity, not identity. Bioactivity metadata may be absent."
)


# --------------------------------------------------------------------------- data readers
def _find_gene_table(package: Path, strain: str) -> Path | None:
    for cand in (f"{strain}_gene_by_gene_all_bgcs.csv", "gene_by_gene_all_bgcs.csv"):
        p = package / cand
        if p.exists():
            return p
    hits = list(package.glob("*gene_by_gene_all_bgcs.csv"))
    return hits[0] if hits else None


def _strain_from_package(package: Path) -> str:
    mp = package / "manifest.json"
    if mp.exists():
        try:
            m = json.loads(mp.read_text(encoding="utf-8"))
            s = m.get("strain") or m.get("strain_id")
            if s:
                return str(s)
        except Exception:
            pass
    # fall back to a *_gene_by_gene_all_bgcs.csv prefix
    for p in package.glob("*_gene_by_gene_all_bgcs.csv"):
        return p.name.split("_gene_by_gene_all_bgcs.csv")[0]
    return package.name


def _role_of(row: dict[str, Any]) -> str:
    fn = (row.get("gene_function_inference") or "").lower()
    res = (row.get("resistance_tier") or "").lower()
    dom = (row.get("sec_met_domains") or "").lower()
    if "core biosynthetic" in fn or "core" == fn:
        return "core"
    # Read the TIER, don't truth-test the string. `resistance_tier` is a four-member
    # enum whose NULL member is itself a string — NULL_NO_SOURCE_DERIVED_RESISTANCE,
    # meaning "scanned, found nothing". That naming is deliberate and correct (it keeps
    # scanned-and-negative distinct from not-scanned), but it is TRUTHY, so a truthiness
    # test promoted an explicit negative into the resistance role. T3 stays accessory:
    # transporter-only routing is not self-protection evidence.
    if res.startswith(("t1", "t2")):
        return "resistance"
    if any(t in fn for t in ("regulat", "transcription")) or "tetr" in dom or "luxr" in dom:
        return "regulation"
    if any(t in fn for t in ("transpos", "integrase", "recombinase", "mobile")):
        return "mobile_element"
    return "accessory"


def _load_blastp_store(store_dir: Path | None) -> dict[str, dict[str, Any]]:
    """query_gene(lower) -> best store row. Empty dict if no store (graceful)."""
    if not store_dir or not store_dir.exists():
        return {}
    best: dict[str, dict[str, Any]] = {}
    for csvp in store_dir.glob("**/*.csv"):
        try:
            for r in csv.DictReader(open(csvp, encoding="utf-8")):
                # v9.7.245 (independently confirms the AS-XXX session's P10): the nr overlay written by
                # `ingest-blastp --package` — the artifact .239 created and .241 fixed — names its gene
                # column `locus_tag`, not `query_gene`. This reader skipped every such row, so the guide
                # could not see the nr evidence the rest of the pipeline had just gone to great lengths
                # to produce. Accept either column; `query_gene` still wins when both are present.
                k = str(r.get("query_gene") or r.get("locus_tag") or "").strip().lower()
                if not k:
                    continue
                try:
                    bit = float(r.get("bitscore") or 0)
                except Exception:
                    bit = 0.0
                prev = best.get(k)
                if prev is None or bit > float(prev.get("bitscore") or 0):
                    best[k] = r
        except Exception:
            continue
    return best


def _auto_store_dir(package: Path) -> Path | None:
    for cand in ("blastp_evidence_store", "blastp_store", "blastp"):
        p = package / cand
        if p.exists():
            return p
    return None


# --------------------------------------------------------------------------- renderers
def render_blastp_readout(hit_row: dict[str, Any] | None, has_translation: bool = True,
                          was_queried: bool = True) -> str:
    """One standardized similarity line from a store hit row. Handles no-hit / no-translation.

    was_queried=False means this gene was never submitted to BLASTp (not in the evidence store's
    query set). Absence from the store is NOT evidence of an orphan / no-hit — only a gene that was
    queried and returned nothing is a genuine no-hit. Rendering an un-queried gene as
    "orphan / fast-evolving" asserts a negative result that was never tested (claim-safety: absence
    of evidence rendered as evidence of absence). (v9.7.196 false-orphan fix)"""
    if not has_translation:
        return "**BLASTp:** not submitted (no translation — pseudogene/fragment)"
    if not was_queried:
        return "**BLASTp:** not run for this gene (not submitted to BLASTp; absence is not an orphan call)"
    if not hit_row:
        return "**BLASTp:** no hit at e-value \u22641e-5 (orphan / fast-evolving)"
    g = lambda k: hit_row.get(k, "")
    tier_raw = str(g("evidence_tier") or "")
    tier = _TIER_GLOSS.get(tier_raw, tier_raw or "unclassified")
    title = str(g("subject_title") or "").strip() or "uncharacterized"
    sci = str(g("subject_sciname") or "").strip()
    sci = f" [{sci}]" if sci else ""
    return (
        f"**BLASTp:** {title}{sci} \u00b7 **id** {g('pct_identity')}% \u00b7 **pos** {g('pct_positive')}% "
        f"\u00b7 **cov** {g('query_coverage')} \u00b7 **E** {g('evalue')} \u00b7 **bit** {g('bitscore')} "
        f"\u00b7 acc `{g('subject_accession')}` \u00b7 **tier** {tier} \u00b7 **claim** {g('claim_safety') or 'capacity-consistent'}"
    )


def build_guide(package, bgc_id, blastp_store=None, audience="both") -> dict[str, Any]:
    """Assemble the deterministic guide skeleton (dict) from banked data only.
    Fills all coordinate/domain/BLASTp content; leaves <!-- LAY: ... --> prose slots."""
    package = Path(package)
    strain = _strain_from_package(package)
    gene_table = _find_gene_table(package, strain)
    if not gene_table:
        raise FileNotFoundError(f"no gene_by_gene_all_bgcs.csv under {package}")
    store_dir = Path(blastp_store) if blastp_store else _auto_store_dir(package)
    store = _load_blastp_store(store_dir)

    genes: list[dict[str, Any]] = []
    products = ""
    for row in csv.DictReader(open(gene_table, encoding="utf-8")):
        if str(row.get("bgc_id", "")).strip() != str(bgc_id).strip():
            continue
        products = products or (row.get("bgc_products") or "")
        locus = (row.get("locus_tag") or "").strip()
        has_tx = bool((row.get("aa_length") or "").strip() and (row.get("aa_length") != "0"))
        hit = store.get(locus.lower())
        # queried = the locus is a key in the BLASTp store (whether or not it returned a hit).
        # Not a key = never submitted → must not be rendered as a no-hit orphan.
        was_queried = bool(store) and (locus.lower() in store)
        genes.append({
            "locus_tag": locus,
            "role": _role_of(row),
            "aa_length": row.get("aa_length", ""),
            "cds_start": row.get("cds_start", ""),
            "cds_end": row.get("cds_end", ""),
            "strand": row.get("strand", ""),
            "domains": row.get("sec_met_domains", ""),
            "function_inference": row.get("gene_function_inference", ""),
            "blastp_readout": render_blastp_readout(hit, has_translation=has_tx, was_queried=was_queried),
            "blastp_tier": str((hit or {}).get("evidence_tier", "") or ("NO_TRANSLATION" if not has_tx else ("NO_HIT" if was_queried else "NOT_QUERIED"))),
        })

    return {
        "schema_version": "bgc-guide-1.0",
        "strain": strain,
        "bgc_id": str(bgc_id),
        "products": products,
        "audience": audience,
        "parts_present": list(GUIDE_PARTS),
        "genes": genes,
        "gene_count": len(genes),
        "claim_safety_line": CLAIM_SAFETY_LINE,
    }


def render_markdown(guide: dict[str, Any]) -> str:
    """Deterministic .md (source of truth). Prose slots marked <!-- LAY: ... --> for a Sapote pass."""
    L: list[str] = []
    strain, bgc = guide["strain"], guide["bgc_id"]
    aud = guide.get("audience", "both")
    L.append(f"# BGC Guide — {strain} {bgc}")
    L.append("")
    L.append(f"*Products (antiSMASH): {guide.get('products','') or 'n/a'} \u00b7 {guide['gene_count']} genes*")
    L.append("")
    # Part 1
    L.append("## Part 1 — Plain-Language Story")
    L.append(f"> {guide['claim_safety_line']}")
    L.append("")
    L.append("<!-- LAY: one-paragraph identity, why-care, factory analogy, what-makes-it-special, honest caveats -->")
    L.append("")
    # Part 2 (suppressible)
    if aud != "technical":
        L.append("## Part 2 — Biology Primer")
        L.append("<!-- LAY: genes/proteins/domains, how to read the evidence tags, the enzyme types present in THIS BGC -->")
        L.append("")
    # Part 3
    L.append("## Part 3 — Technical Overview")
    L.append("")
    L.append(f"| field | value |")
    L.append(f"|---|---|")
    L.append(f"| strain | {strain} |")
    L.append(f"| BGC | {bgc} |")
    L.append(f"| products | {guide.get('products','') or 'n/a'} |")
    L.append(f"| genes | {guide['gene_count']} |")
    L.append("")
    L.append("<!-- LAY: core biosynthetic logic, open questions -->")
    L.append("")
    # Part 4 — gene catalogue, role-grouped
    L.append("## Part 4 — Complete Gene Catalogue")
    L.append("")
    by_role: dict[str, list[dict]] = {r: [] for r in _ROLE_ORDER}
    for g in guide["genes"]:
        by_role.setdefault(g["role"], []).append(g)
    for role in _ROLE_ORDER:
        grp = by_role.get(role) or []
        if not grp:
            continue
        L.append(f"### {role.replace('_',' ').title()}")
        L.append("")
        for g in grp:
            L.append(f"#### {g['locus_tag']} — <!-- LAY: plain-title -->")
            L.append(f"*{g['aa_length']} aa \u00b7 {g['cds_start']}\u2013{g['cds_end']} bp \u00b7 strand {g['strand']}*")
            L.append("")
            L.append(f"<!-- LAY: one to three sentence plain-language summary for {g['locus_tag']} -->")
            L.append("")
            L.append(f"**Structure:** {g['domains'] or 'no antiSMASH domain annotation'}")
            L.append("")
            L.append(f"**Function:** {g['function_inference'] or 'inference pending'}")
            L.append("")
            L.append(g["blastp_readout"])
            L.append("")
    # Part 5
    L.append("## Part 5 — Synthesis and Next Steps")
    L.append(f"> {guide['claim_safety_line']}")
    L.append("")
    L.append("<!-- LAY: whole-picture read, prioritized experiments -->")
    L.append("")
    return "\n".join(L)


def render_docx(guide: dict[str, Any], out_path) -> bool:
    """Thin renderer over the .md; graceful no-op if python-docx absent. Returns True if written."""
    try:
        from docx import Document  # type: ignore
    except Exception:
        return False
    try:
        md = render_markdown(guide)
        doc = Document()
        for line in md.splitlines():
            if line.startswith("# "):
                doc.add_heading(line[2:], level=0)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=1)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=2)
            elif line.startswith("#### "):
                doc.add_heading(line[5:], level=3)
            elif line.startswith("<!--"):
                continue  # LAY slot; not rendered in docx presentation layer
            else:
                doc.add_paragraph(line)
        doc.save(str(out_path))
        return True
    except Exception:
        return False


# --------------------------------------------------------------------------- quality gate
def _load_gene_count_crosscheck(package: Path | None, bgc_id: str) -> dict | None:
    """Read the seal-time independent CDS recount sidecar for one BGC, if present.

    The sidecar (`gene_count_crosscheck.json`) is written at seal time by walking the region GBKs
    directly — a DIFFERENT code path than the CSV the guide is built from — so it is an independent
    source of truth for how many CDS the region actually contains. Returns the per-BGC entry
    {gbk_cds_count, gbk_edge_cds_count, gbk_source} or None if the sidecar or entry is absent.
    """
    if not package:
        return None
    import json
    p = Path(package) / "gene_count_crosscheck.json"
    if not p.exists():
        return None
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return None
    return (data.get("per_bgc") or {}).get(str(bgc_id))


def guide_quality_gate(guide: dict[str, Any], package: Path | None = None
                       ) -> tuple[bool, list[str], list[str]]:
    """SKELETON gate — validates the deterministic `guide` dict (rebuilt from the sealed package),
    NOT authored prose. It confirms structural completeness (all 5 Parts, every CDS has a subsection
    with the three readouts, no per-gene claim exceeds its BLASTp tier, claim-safety line present) so
    the emitted skeleton is sound. It CANNOT see authored content: an empty template passes it. To
    verify a FINISHED guide (residual LAY slots, per-Part authored length, gene summaries) use
    `verify_authored_guide()` / `mamey verify-guide` — do not report this gate's pass as validation
    of authored work.

    v9.7.191: when a seal-time `gene_count_crosscheck.json` sidecar is present, cross-check the
    guide's gene_count against an INDEPENDENT CDS recount from the region GBKs (AS-XXX upstream-
    omission hardening). Errors ONLY on interior-gene omission (a gene the clipped GBK could not
    legitimately have dropped); edge/FC differences warn, not fail; an absent sidecar degrades to a
    warning. Pass `package` to enable the cross-check.

    Returns (ok, errors, warnings)."""
    errors: list[str] = []
    warnings: list[str] = []
    # 5 parts (P2 legitimately suppressed under audience=technical)
    required = set(GUIDE_PARTS)
    if guide.get("audience") == "technical":
        required.discard("P2")
    if not required.issubset(set(guide.get("parts_present", []))):
        errors.append(f"missing Parts: {sorted(required - set(guide.get('parts_present', [])))}")
    # every gene has a subsection with the three readouts (structure/function/blastp)
    if guide.get("gene_count", 0) == 0:
        errors.append("no gene subsections (AS-XXX-class omission check)")
    # v9.7.191: independent-recount cross-check (AS-XXX upstream-omission hardening)
    _xc = _load_gene_count_crosscheck(package, guide.get("bgc_id", ""))
    if package is None:
        pass  # cross-check not requested
    elif _xc is None:
        warnings.append("no gene-count cross-check sidecar; AS-XXX upstream-omission check is degraded")
    else:
        csv_n = int(guide.get("gene_count", 0))
        gbk_n = int(_xc.get("gbk_cds_count", 0))
        edge_n = int(_xc.get("gbk_edge_cds_count", 0))
        lower_ok = gbk_n - edge_n  # CSV may legitimately drop up to edge_n edge/contig-spanning CDS
        if csv_n < lower_ok:
            errors.append(
                f"gene omission: guide has {csv_n} genes but region GBK has {gbk_n} CDS "
                f"({edge_n} edge-droppable) -> at least {lower_ok - csv_n} interior gene(s) missing "
                f"(AS-XXX upstream-omission)")
        elif csv_n < gbk_n:
            warnings.append(
                f"gene count {csv_n} < GBK {gbk_n}; within edge-droppable budget ({edge_n}) — OK")
        elif csv_n > gbk_n:
            warnings.append(
                f"gene count {csv_n} > GBK {gbk_n}; guide carries genes not in the counted GBK — verify source")
    for g in guide.get("genes", []):
        if "blastp_readout" not in g or not g["blastp_readout"]:
            errors.append(f"{g.get('locus_tag','?')}: missing BLASTp readout")
        # claim-not-above-tier: a TIER_C/D (or NO_HIT) gene must not carry a class-level claim.
        # The deterministic skeleton never emits such a claim (prose is a LAY slot), so this guards
        # a filled guide: flag if a claim token appears against a weak tier.
        tier = str(g.get("blastp_tier", ""))
        if tier in ("TIER_C_GENERIC_FUNCTION_CONFIRMED", "TIER_D_UNINFORMATIVE_OR_WEAK", "NO_HIT"):
            fn = (g.get("function_inference") or "").lower()
            if "core biosynthetic" in fn and tier in ("TIER_D_UNINFORMATIVE_OR_WEAK", "NO_HIT"):
                warnings.append(f"{g.get('locus_tag','?')}: core-role gene with weak BLASTp tier ({tier}) — verify claim ceiling")
    # claim-safety line present (Parts 1 & 5 both carry it in the renderer)
    if not guide.get("claim_safety_line"):
        errors.append("claim-safety line absent")
    return (len(errors) == 0, errors, warnings)


# --------------------------------------------------------------------------- CLI
def guide_command(args) -> int:
    """CLI entrypoint: build → gate → write <pkg>/guide/<STRAIN>_<BGC>_Guide.{md,docx} + receipt."""
    package = Path(args.package)
    outdir = Path(args.outdir) if getattr(args, "outdir", None) else package / "guide"
    outdir.mkdir(parents=True, exist_ok=True)
    fmt = getattr(args, "format", "both")
    audience = getattr(args, "audience", "both")
    store = getattr(args, "blastp_store", None)

    # resolve which BGCs to cover
    bgc_ids: list[str]
    if getattr(args, "bgc", None):
        bgc_ids = [args.bgc]
    else:
        strain = _strain_from_package(package)
        gt = _find_gene_table(package, strain)
        seen: list[str] = []
        if gt:
            for r in csv.DictReader(open(gt, encoding="utf-8")):
                b = str(r.get("bgc_id", "")).strip()
                if b and b not in seen:
                    seen.append(b)
        bgc_ids = seen[: getattr(args, "top_n", 10)]

    rc = 0
    for bgc in bgc_ids:
        try:
            guide = build_guide(package, bgc, blastp_store=store, audience=audience)
        except Exception as e:
            emit(f"[guide] {bgc}: build failed — {e}")
            rc = 1
            continue
        # AUDIT_374: was `guide_quality_gate(guide)` -- the ONLY real production caller
        # never passed the sealed-package dir already sitting in scope (`package`, set above),
        # so the AS-XXX interior-gene-omission cross-check this function was specifically built
        # for (see its own docstring: "AS-XXX upstream-omission hardening") silently degraded to
        # "cross-check not requested" on every real `mamey guide` run, with no warning emitted.
        ok, errors, warnings = guide_quality_gate(guide, package=package)
        strain = guide["strain"]
        stem = f"{strain}_{bgc}_Guide"
        wrote = []
        if fmt in ("md", "both"):
            (outdir / f"{stem}.md").write_text(render_markdown(guide), encoding="utf-8")
            wrote.append("md")
        docx_ok = None
        if fmt in ("docx", "both"):
            docx_ok = render_docx(guide, outdir / f"{stem}.docx")
            if docx_ok:
                wrote.append("docx")
            elif fmt == "docx":
                emit(f"[guide] {bgc}: docx renderer unavailable (python-docx) — writing md fallback")
                (outdir / f"{stem}.md").write_text(render_markdown(guide), encoding="utf-8")
                wrote.append("md")
                rc = 1 if not wrote else rc
        # receipt (mirrors mode_b_receipt shape)
        receipt = {
            "schema_version": "bgc-guide-receipt-1.0",
            "strain": strain, "bgc_id": bgc, "gene_count": guide["gene_count"],
            "gate_ok": ok, "errors": errors, "warnings": warnings,
            "formats_written": wrote,
            "docx_rendered": bool(docx_ok) if docx_ok is not None else "not_requested",
        }
        (outdir / f"{stem}.receipt.json").write_text(json.dumps(receipt, indent=2), encoding="utf-8")
        status = "OK" if ok else "GATE_FAIL"
        emit(f"[guide] {bgc}: {status} ({guide['gene_count']} genes; wrote {','.join(wrote)})", f'        note: this gates the SKELETON only. After authoring the LAY slots, run `mamey verify-guide {stem}.md` to validate the finished guide.', sep="\n")
        if not ok:
            for e in errors:
                emit(f"        ERROR: {e}")
            rc = 1
    return rc


# ---------------------------------------------------------------------------------------------
# v9.7.192 — AUTHORED-output verification.
#
# WHY THIS EXISTS. guide_quality_gate() gates the deterministic SKELETON (a `guide` dict rebuilt
# from the sealed package). It CANNOT see authored prose — an empty template passes it identically
# to a finished guide. That let a green "guide gate: OK" be reported over unauthored work (a
# fabricated assurance, worse than none). verify_authored_guide() closes the hole by reading the
# FINISHED .md and checking the authoring actually happened: no residual LAY slots, every Part
# carries real prose, every gene subsection has a non-empty plain-language summary.
# ---------------------------------------------------------------------------------------------

_LAY_SLOT_RE = re.compile(r"<!--\s*LAY:", re.IGNORECASE)
_PART_HEADER_RE = re.compile(r"^##\s+Part\s+(\d+)\b.*$", re.MULTILINE)
_GENE_HEADER_RE = re.compile(r"^####\s+(\S+)\b.*$", re.MULTILINE)

# per-Part minimum authored body length (characters of prose, excluding headers/tables/readouts).
# Deliberately modest — this catches "blank" and "one-line stub", not a judgment of quality.
_PART_MIN_CHARS = {1: 300, 2: 200, 3: 200, 5: 300}
_GENE_SUMMARY_MIN_CHARS = 20


def verify_authored_guide(md_path, *, audience: str = "both") -> tuple[bool, list[str], list[str]]:
    """Read a FINISHED guide .md and verify the authoring happened. Unlike guide_quality_gate (which
    gates the skeleton), this opens the authored file. Returns (ok, errors, warnings).

    Checks (all ERROR unless noted):
      - file exists and is non-empty
      - ZERO residual `<!-- LAY: ... -->` slots (any remaining = unauthored section)
      - every required Part header present (P2 optional under audience=technical)
      - each Part 1/2/3/5 carries >= its minimum authored body length (real prose, not a stub)
      - every gene subsection (#### locus) has a non-empty plain-language summary line
    """
    errors: list[str] = []
    warnings: list[str] = []
    # LEAK-1 (CLAUDE_409_claimsafety_coverage): accumulate the AUTHORED lay prose so the same
    # claim-safety detector the seal scan / verify-modeb lane uses can be run over it below. The
    # skeleton readout lines (Structure/Function/BLASTp coord italics) are rendered from banked,
    # class-capped evidence and are excluded here — we lint only what a human/LLM author wrote into
    # the LAY slots, exactly the surface verify-guide previously certified without any claim check.
    authored_prose: list[str] = []
    p = Path(md_path)
    if not p.exists():
        return False, [f"authored guide not found: {p}"], []
    if p.is_dir():
        return False, [f"expected a guide file, got a directory: {p}"], []
    text = p.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        return False, [f"authored guide is empty: {p}"], []

    # 1. residual LAY slots = unauthored sections (the core check)
    residual = _LAY_SLOT_RE.findall(text)
    if residual:
        errors.append(f"{len(residual)} residual '<!-- LAY: -->' slot(s) — sections left unauthored")

    # 2. required Parts present
    parts_present = {int(m.group(1)) for m in _PART_HEADER_RE.finditer(text)}
    required_parts = {1, 2, 3, 4, 5}
    if audience == "technical":
        required_parts.discard(2)
    missing = sorted(required_parts - parts_present)
    if missing:
        errors.append(f"missing Part(s): {missing}")

    # 3. per-Part authored body length. Slice the text by Part headers and measure the prose in each.
    part_spans = [(int(m.group(1)), m.start()) for m in _PART_HEADER_RE.finditer(text)]
    part_spans.sort(key=lambda t: t[1])
    for i, (num, start) in enumerate(part_spans):
        end = part_spans[i + 1][1] if i + 1 < len(part_spans) else len(text)
        body = text[start:end]
        # strip headers, table rows, blockquotes, and BLASTp/Structure/Function readout lines —
        # what remains is authored prose. This is a floor for "was anything written", not a quality bar.
        prose_lines = []
        for ln in body.splitlines():
            s = ln.strip()
            if not s or s.startswith("#") or s.startswith("|") or s.startswith(">"):
                continue
            if s.startswith("**Structure:**") or s.startswith("**Function:**") or s.startswith("**BLASTp:**"):
                continue
            if s.startswith("*") and s.endswith("*") and "aa ·" in s:  # the coord italic line
                continue
            prose_lines.append(s)
        authored_prose.extend(prose_lines)  # LEAK-1: authored Part prose feeds the claim-safety lint
        prose_len = len(" ".join(prose_lines))
        floor = _PART_MIN_CHARS.get(num)
        if floor is not None and prose_len < floor:
            errors.append(f"Part {num}: only {prose_len} chars of authored prose (floor {floor}) — thin/stub")

    # 4. every gene subsection has a non-empty plain-language summary. A gene block runs from its
    # #### header to the next #### or ## header; the summary is the prose between the coord italic
    # line and the **Structure:** line.
    gene_headers = [(m.group(1), m.start(), m.end()) for m in _GENE_HEADER_RE.finditer(text)]
    for gi, (locus, gstart, ghead_end) in enumerate(gene_headers):
        gend = gene_headers[gi + 1][1] if gi + 1 < len(gene_headers) else len(text)
        block = text[ghead_end:gend]
        struct_idx = block.find("**Structure:**")
        summary_region = block[:struct_idx] if struct_idx >= 0 else block
        summary_prose = []
        for ln in summary_region.splitlines():
            s = ln.strip()
            if not s or s.startswith("*") or s.startswith("#") or s.startswith("|"):
                continue
            summary_prose.append(s)
        authored_prose.extend(summary_prose)  # LEAK-1: authored gene summaries feed the claim-safety lint
        if len(" ".join(summary_prose)) < _GENE_SUMMARY_MIN_CHARS:
            errors.append(f"gene {locus}: plain-language summary missing/too short (< {_GENE_SUMMARY_MIN_CHARS} chars)")
        # also catch an unfilled plain-title in the header itself
        if "LAY:" in text[gstart:ghead_end]:
            errors.append(f"gene {locus}: plain-title slot left unauthored")

    # 5. LEAK-1 (CLAUDE_409_claimsafety_coverage): claim-safety gate on the AUTHORED lay prose.
    # verify_authored_guide previously enforced slot-fill + prose-length + gene summaries but ran
    # ZERO claim-safety check at any severity (DEEP_AUDIT_claimsafety.md LEAK-1, HIGH) — a
    # structurally-complete guide whose LAY slots carry an unhedged compound-identity/bioactivity
    # claim ("makes the antibiotic streptomycin, active against MRSA") passed verify-guide GREEN,
    # and the _Guide.md filename is filtered out of the one seal scan. This mirrors the
    # finished-profile blocking semantics CLAUDE_409_claim_safety_enforce added to verify-modeb:
    # verify-guide certifies a FINISHED, authored deliverable, so a claim-safety finding is an
    # ERROR (verify_guide_command then returns rc 1 / FAIL). We reuse the exact same detector
    # (claim_safety_gate.lint_text) — the seal scan's detector — rather than a second copy, so
    # clean/hedged capacity prose that passes the package gate keeps passing here.
    try:
        from .claim_safety_gate import lint_text as _cs_lint_text
    except Exception:  # pragma: no cover - direct-script fallback (sys.path set at module top)
        from claim_safety_gate import lint_text as _cs_lint_text
    cs_findings = _cs_lint_text("\n".join(authored_prose))
    for f in cs_findings:
        errors.append(f"claim-safety: {f} — hedge to class-level capacity language before certifying")

    return (len(errors) == 0, errors, warnings)
