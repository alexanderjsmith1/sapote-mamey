#!/usr/bin/env python3
"""build_gcf_tags.py — map each BGC's KCB anchor to a curated GCF / product-family tag (Sapote annotation layer).

This is an ANNOTATION / DEREPLICATION layer, not a detection step: it does NOT find new clusters (antiSMASH does
that) and it does NOT identify a product. It normalizes the antiSMASH/KCB output into a curated controlled
vocabulary (resources/gcf_thesaurus.json) so the workbook carries two distinct columns:
  - Predicted_BGC_Class       : the antiSMASH operational protocluster label (detector output)
  - GCF_Product_Family_Tag    : the curated product-family tag, ANCHORED ON the KCB similarity hit (a hypothesis)

Claim-safety: the GCF tag inherits the KCB anchor's status — similarity, NOT identification. UNRESOLVED anchors
stay dark (no family). Unmatched named anchors fall to a class-level _orphan tag, never a false specific call.

Usage:
  python tools/build_gcf_tags.py --banked-dir cohort --thesaurus resources/gcf_thesaurus.json \
      --workbook Master.xlsx --out-dir analysis
"""
import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402
import argparse, os, json, csv
try:  # v9.7.410 CSV formula-cell guard (CLAUDE_v9.7.410_tools_csv_writer_coverage)
    from mamey.csv_safety import SafeDictWriter as _SafeDictWriter, SafeWriter as _SafeWriter
except ImportError:  # bare-script run: bundle root is one level up
    import os as _cs_os, sys as _cs_sys
    _cs_sys.path.insert(0, _cs_os.path.dirname(_cs_os.path.dirname(_cs_os.path.abspath(__file__))))
    from mamey.csv_safety import SafeDictWriter as _SafeDictWriter, SafeWriter as _SafeWriter
from collections import Counter
from _wbio import atomic_save, atomic_open


def _read_json(_path, *, encoding="utf-8"):
    """P3b: context-managed JSON read; closes the handle a bare open() leaked."""
    import json as _json
    with open(_path, encoding=encoding) as _fh:
        return _json.load(_fh)


# antiSMASH product (lowercased) -> thesaurus class bucket for the _orphan fallback
CLASS_BUCKET = {
 'phosphonate':'phosphonate','t1pks':'T1PKS','transat-pks':'transAT-PKS','t2pks':'T2PKS','hr-t2pks':'T2PKS',
 'nrps':'NRPS','nrps-like':'NRPS','ripp':'RiPP','ripp-like':'RiPP','lanthipeptide':'RiPP','lassopeptide':'RiPP',
 'thiopeptide':'RiPP','thioamitides':'RiPP','saccharide':'saccharide','oligosaccharide':'saccharide',
 'nucleoside':'nucleoside','terpene':'terpene','indole':'indole_alkaloid','siderophore':'siderophore',
 'ni-siderophore':'siderophore','nrp-metallophore':'siderophore','betalactone':'betalactone','cdps':'CDPS',
 'ectoine':'other','melanin':'other','butyrolactone':'other','hgle-ks':'other','pks':'T1PKS','pks-like':'T1PKS',
}

def load_aliases(thesaurus_path):
    th=_read_json(thesaurus_path)
    pairs=[]  # (alias, tag)
    for cls, tags in th.items():
        for t in tags:
            for a in t['aliases']:
                pairs.append((a.lower(), t['tag']))
    pairs.sort(key=lambda p:-len(p[0]))  # longest alias first = most specific match
    orphan={cls:next((t['tag'] for t in tags if t['tag'].endswith('_orphan') or t['tag'].endswith('_unknown')), None) for cls,tags in th.items()}
    return pairs, orphan

def primary_class(products):
    cl=[p.strip().lower() for p in (products or '').split(';') if p.strip()]
    for pref in ['transat-pks','t1pks','hr-t2pks','t2pks','nrps','nrps-like','phosphonate','lanthipeptide',
                 'lassopeptide','thiopeptide','ripp','terpene','saccharide','nucleoside','siderophore','ectoine',
                 'betalactone','cdps','indole','butyrolactone','melanin','pks']:
        if pref in cl: return pref
    return cl[0] if cl else 'unknown'

def tag_bgc(kcb, products, pairs, orphan):
    anchor=(kcb or '').strip().lower()
    pc=primary_class(products)
    if not anchor or anchor=='unresolved':
        return 'DARK_no_family', 'unresolved'  # the dark/novel pool
    for alias,tag in pairs:
        if alias and alias in anchor:
            return tag, f'anchor~"{alias}"'
    # named anchor but no thesaurus match -> class-level orphan
    bucket=CLASS_BUCKET.get(pc)
    ot=orphan.get(bucket) if bucket else None
    return (ot or 'OTHER_unknown'), 'unmatched_anchor'

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--banked-dir', default='cohort'); ap.add_argument('--thesaurus', default='resources/gcf_thesaurus.json')
    ap.add_argument('--workbook', default=None); ap.add_argument('--out-dir', default='analysis')
    a=ap.parse_args(); os.makedirs(a.out_dir, exist_ok=True)
    pairs, orphan = load_aliases(a.thesaurus)
    bgc=_read_json(os.path.join(a.banked_dir,'bgc_data.json')); bgcs=bgc['bgcs']
    rows=[]; tagc=Counter(); basis=Counter()
    for b in bgcs:
        pc=primary_class(b.get('products')); kcb=b.get('closest_kcb_product') or ''
        tag,why=tag_bgc(kcb, b.get('products'), pairs, orphan)
        tagc[tag]+=1; basis[why.split('~')[0].split('"')[0].strip()]+=1
        rows.append({'strain':b['sid'],'bgc':b.get('bgc_id',''),'Predicted_BGC_Class':pc,
                     'KCB_anchor':kcb,'GCF_Product_Family_Tag':tag,'match_basis':why})
    out_csv=os.path.join(a.out_dir,'GCF_Tags.csv')
    with atomic_open(out_csv,'w',newline='') as f:
        w=_SafeDictWriter(f,fieldnames=['strain','bgc','Predicted_BGC_Class','KCB_anchor','GCF_Product_Family_Tag','match_basis']); w.writeheader(); w.writerows(rows)
    # coverage
    n=len(bgcs); dark=tagc['DARK_no_family']; named=sum(v for k,v in tagc.items() if not k.endswith(('_orphan','_unknown','_no_family')))
    orphans=sum(v for k,v in tagc.items() if k.endswith(('_orphan','_unknown')))
    emit(f"  tagged {n} BGCs -> GCF_Tags.csv", f"  dark (UNRESOLVED, no family): {dark} ({100*dark/n:.0f}%)", f"  resolved to a named GCF family: {named} ({100*named/n:.0f}%)", f"  named anchor but class-orphan: {orphans} ({100*orphans/n:.0f}%)", sep="\n")
    anchored=n-dark
    emit(f"  of {anchored} anchored BGCs, {named} ({100*named/anchored:.0f}%) got a specific family tag", "  top GCF families:", sep="\n")
    for t,c in tagc.most_common(14):
        if t!='DARK_no_family': emit(f"    {c:4}  {t}")
    if a.workbook and os.path.exists(a.workbook):
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb=openpyxl.load_workbook(a.workbook)
        if 'GCF_Tags' in wb.sheetnames: del wb['GCF_Tags']
        ws=wb.create_sheet('GCF_Tags'); ws.append(['strain','bgc','Predicted_BGC_Class','KCB_anchor','GCF_Product_Family_Tag','match_basis'])
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='2E5E7E')
        for r in rows: ws.append([r['strain'],r['bgc'],r['Predicted_BGC_Class'],r['KCB_anchor'],r['GCF_Product_Family_Tag'],r['match_basis']])
        ws.freeze_panes='A2'; atomic_save(wb, a.workbook); emit(f"  GCF_Tags sheet added to {os.path.basename(a.workbook)}")

if __name__=='__main__': main()
