#!/usr/bin/env python3
"""placement_to_docx.py — assemble phylogenetic-placement figures + their methods captions into a Word doc.

Each placement output dir holds `<group>_placement_tree.png` + `<group>_placement_FIGURE_CAPTION.txt` (written
by phylo_place.py report). This tool drops each tree on its own page with the full methods caption underneath —
so a reader who has never heard of EPA-ng still understands the figure. The individual PNG + caption files stay
on disk untouched, so you can also mix-and-match trees into other documents.

Usage:
  Tools/phylo_place bin: Tools/bin/python3 Tools/placement_to_docx.py \
      --dir "strain_data/_PLACEMENT/nocardia/placements_enriched" \
      --dir "strain_data/_PLACEMENT/streptomyces/placements" \
      --out "strain_data/_PLACEMENT/PLACEMENT_TREES_<date>.docx" --title "16S placement trees"
  # or auto-discover every placement dir:
  Tools/bin/python3 Tools/placement_to_docx.py --auto --out combined.docx
"""
import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402
import argparse, glob, os, re, sys, datetime
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # bundle root for `import mamey` (v9.7.367 A10)
from mamey.workspace_root import workspace_root
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _wbio import atomic_save


def _find(dirs):
    pairs = []
    for d in dirs:
        for png in sorted(glob.glob(os.path.join(d, "*_placement_tree.png"))):
            cap = png.replace("_placement_tree.png", "_placement_FIGURE_CAPTION.txt")
            pairs.append((png, cap if os.path.exists(cap) else None, d))
    return pairs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", action="append", default=[], help="placement output dir (repeatable)")
    ap.add_argument("--auto", action="store_true", help="discover strain_data/_PLACEMENT/*/placements*")
    ap.add_argument("--out", required=True)
    ap.add_argument("--title", default="Phylogenetic placement trees")
    ap.add_argument("--root", default=str(workspace_root()))
    a = ap.parse_args()
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    dirs = list(a.dir)
    if a.auto:
        dirs += sorted(glob.glob(os.path.join(a.root, "strain_data/_PLACEMENT/*/placements*")))
    dirs = [d for i, d in enumerate(dirs) if d not in dirs[:i]]  # dedupe, keep order
    pairs = _find(dirs)
    if not pairs:
        sys.exit("no *_placement_tree.png found in the given dirs")

    doc = docx.Document()
    h = doc.add_heading(a.title, level=0)
    doc.add_paragraph(f"Generated {datetime.date.today().isoformat()} · {len(pairs)} figures · "
                      "each tree + methods caption; individual PNG/caption files remain separate for mix-and-match.")
    doc.add_paragraph("Claim-safety: placement = neighborhood hypothesis; 16S is an anchor, not a species call; "
                      "judgment deferred.").runs[0].italic = True

    for i, (png, cap, d) in enumerate(pairs):
        doc.add_page_break()
        grp = os.path.basename(png).replace("_placement_tree.png", "")
        run = os.path.basename(os.path.dirname(png))
        hd = doc.add_heading(f"{grp} — {run}", level=1)
        # image scaled to fit page width (~6.5")
        try:
            doc.add_picture(png, width=Inches(6.3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            doc.add_paragraph(f"[image failed: {e}]")
        # caption
        if cap:
            txt = open(cap).read().strip()
            for j, para in enumerate(txt.split("\n\n")):
                p = doc.add_paragraph()
                r = p.add_run(para.strip())
                r.font.size = Pt(9)
                if para.lower().startswith("figure"):
                    r.bold = True
                if para.startswith("Claim-safety"):
                    r.italic = True; r.font.color.rgb = RGBColor(0x80, 0x30, 0x30)
        else:
            doc.add_paragraph("[no caption file found — run phylo_place.py report to generate it]")

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    # v9.7.371 fix: was a bare doc.save(a.out) -- an interrupted write mid-save leaves a
    # truncated/corrupt .docx (same zip-container failure mode tools/_wbio.py's own docstring
    # warns about for .xlsx: "corrupts the workbook into an unreadable BadZipFile"). atomic_save
    # is format-agnostic (calls obj.save(tmp) then os.replace()) and works as a drop-in here --
    # python-docx's Document exposes the same .save(path) interface openpyxl's Workbook does.
    atomic_save(doc, a.out)
    emit(f"WROTE {a.out}  ({len(pairs)} figures)")
    for png, cap, d in pairs:
        emit(f"  - {os.path.basename(png)}  caption={'yes' if cap else 'MISSING'}")


if __name__ == "__main__":
    sys.exit(main())
