#!/usr/bin/env python3
"""manuscript_qa.py — Read-only quality check of a Word (.docx) manuscript; writes a Markdown report and a TSV.

It never changes the document: the file is read once into memory and opened from those bytes. Every finding is a
lead for a person to check, with the paragraph number (¶ = index in python-docx `Document(...).paragraphs`; `T<k>`
= the k-th top-level table), a quote of at most 15 words and a detail. Nothing is judged or fixed.

Checks (the `check` column):
 1. citation_unmatched, citation_year_mismatch, reference_uncited, reference_order, citation_hold:
    author-year citations against the reference list. Reference entries are read with the parser of
    tools/paper_reference_audit.py; citations are author-year pairs found inside any (...) or [...] group, so a
    mixed group such as "(PCoA; Gower 1966)" still counts. "2019a,b" cites both years; "Smith, Jones, and Brown
    2019" is a citation of Smith. Surnames match without particles ("van der Maaten" = "Maaten"). Order is
    letter by letter on the first author's surname as written in the list, then the given name; apostrophes,
    spaces, hyphens and accents are ignored. The list starts at a References, Reference List, References Cited,
    Bibliography, Literature Cited, Cited Literature or Works Cited heading; with none of them the report gets
    one citation_hold row saying the citations were not checked. A numbered list is left to
    paper_reference_audit.py. Author names in these docstrings are illustrations of citation forms.
 2. british_spelling: a word list (default below; --spelling-file replaces it). Reference entries are skipped.
 3. genus_upright, family_italic, panel_label_italic: italics read from each run, including runs inside
    hyperlinks (run, character style, paragraph style, document default). Genus names come from --genus-file
    or the default list. Family names (-aceae) should be upright; single-letter panel labels in captions too.
 4. short_sentence: sentences of 5 words or fewer in body text (not headings, captions, references or tables).
    A run-in label opening a paragraph (an all-bold or all-italic first sentence, or one of 1-2 words such as
    "Methods.") is not counted.
 5. internal_name: names that must not appear in paper text (default list below; --internal-names replaces it).
 6. hyperlink_internal_target, style_undefined: a hyperlink whose relationship points at a part inside the
    package (an image, say) rather than an external URL; a paragraph style the document does not define (one row
    per paragraph); a run style it does not define (one row per style, with the number of paragraphs).
 7. print_scale, image_taller_than_text_block, image_wider_than_text_block, image_no_dpi, effective_text_below_min:
    each embedded raster image's placed width against its drawn width (pixels / dpi). Below --min-scale (0.95)
    it is flagged. If the PNG declares its text size (tEXt key min_text_pt) or --assume-text-pt is given, the
    printed text size is reported, and flagged below --min-text-pt (8).
 8. superseded_render (optional, --superseded DIR): an embedded image byte-identical to a file in that folder.

Usage:
  python tools/manuscript_qa.py paper.docx [more.docx] --out REPORT_DIR [--genus-file F] [--spelling-file F]
      [--internal-names F] [--superseded DIR] [--min-scale 0.95] [--assume-text-pt 8] [--min-text-pt 8]
Exit 0 when every input was read (findings or not); 2 when an input is refused (not .docx, unreadable).
"""
from __future__ import annotations

import argparse
import hashlib
import io
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from mamey.csv_safety import SafeWriter  # noqa: E402
from mamey.logging_setup import get_logger  # noqa: E402
from tools.paper_reference_audit import NUMBER, YEAR, YEAR_RE, author_key  # noqa: E402

_LOG = get_logger("tools.manuscript_qa")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
EMU_PER_IN = 914400
QUOTE_WORDS = 15
COLUMNS = ["check", "¶", "quote", "detail"]

CHECKS = [  # report order, with the one-line description used in the Markdown counts table
    ("citation_unmatched", "1", "in-text citation with no reference of that author"),
    ("citation_year_mismatch", "1", "citation author is in the list, but not with that year"),
    ("reference_uncited", "1", "reference never cited (a citation with the wrong year does not count)"),
    ("reference_order", "1", "reference sorts before the entry above it (first author, letter by letter)"),
    ("citation_hold", "1", "citation or reference the parser could not read; check by hand"),
    ("british_spelling", "2", "British spelling outside the reference list"),
    ("genus_upright", "3", "genus name set upright"),
    ("family_italic", "3", "family name (-aceae) set italic"),
    ("panel_label_italic", "3", "single-letter panel label set italic in a caption"),
    ("short_sentence", "4", "body sentence of 5 words or fewer"),
    ("internal_name", "5", "internal name (file, path, tool or run label) in paper text"),
    ("hyperlink_internal_target", "6", "hyperlink points inside the package, not at a URL"),
    ("style_undefined", "6", "paragraph or run style not defined in the document"),
    ("print_scale", "7", "image placed below the minimum share of its drawn width"),
    ("image_taller_than_text_block", "7", "image taller than the page's text block"),
    ("image_wider_than_text_block", "7", "image wider than the page's text block"),
    ("image_no_dpi", "7", "raster image with no dpi, so its drawn size is unknown"),
    ("effective_text_below_min", "7", "declared figure text prints below the minimum size"),
    ("superseded_render", "8", "embedded image byte-identical to a superseded render"),
]

DEFAULT_GENERA = (
    # actinomycetes
    "Actinoallomurus Actinokineospora Actinomadura Actinophytocola Actinoplanes Actinosynnema Agromyces Allokutzneria "
    "Amycolatopsis Arthrobacter Blastococcus Brevibacterium Catenulispora Cellulomonas Cellulosimicrobium Clavibacter "
    "Corynebacterium Couchioplanes Crossiella Curtobacterium Dactylosporangium Dermacoccus Dietzia Embleya Frankia "
    "Geodermatophilus Glycomyces Gordonia Herbidospora Isoptericola Janibacter Jiangella Kitasatospora Kocuria "
    "Kribbella Kutzneria Lechevalieria Leifsonia Lentzea Microbacterium Microbispora Micrococcus Micromonospora "
    "Microtetraspora Modestobacter Mycobacterium Mycolicibacterium Nakamurella Nocardia Nocardioides Nocardiopsis "
    "Nonomuraea Planomonospora Plantactinospora Prauserella Promicromonospora Pseudonocardia Rathayibacter Rhodococcus "
    "Saccharomonospora Saccharopolyspora Saccharothrix Salinispora Sciscionella Streptacidiphilus Streptomyces "
    "Streptosporangium Thermobifida Thermomonospora Tsukamurella Umezawaea Verrucosispora Williamsia "
    # test organisms, hosts and other genera common in these papers
    "Acinetobacter Acromyrmex Apis Apterostigma Arabidopsis Ascosphaera Aspergillus Atta Bacillus Bombus Botrytis "
    "Caenorhabditis Candida Ceratodon Cladonia Cryptococcus Cyphomyrmex Drosophila Enterococcus Escherichia "
    "Escovopsis Fusarium Galleria Hypogymnia Klebsiella Leucoagaricus Leucocoprinus Marchantia Megachile Nosema "
    "Osmia Paenibacillus Philanthus Physcomitrium Polytrichum Pseudomonas Saccharomyces Salmonella Sphagnum "
    "Staphylococcus Trachymyrmex Trichophyton Usnea Xylocopa"
).split()

_ISE_STEMS = ("character recogn organ summar normal optim standard util minim maxim priorit visual stabil synthes "
              "hypothes emphas categor critic random harmon solubil mobil special local final general real colon immun "
              "steril homogen crystall polymer parameter initial serial digit author custom memor metabol neutral oxid "
              "ion lyophil linear marginal mineral fertil sensit capital symbol apolog").split()


def _re_sub(pattern, repl):
    return pattern, (lambda word: re.sub(pattern, repl, word, flags=re.I))


DEFAULT_SPELLING = [  # (regex matched as a whole word, case-insensitive; function giving the American form)
    _re_sub(r"(colo|behavio|neighbo|favo|flavo|harbo|humo|labo|tumo|vapo|odo|rumo|savo|endeavo|armo|valo|vigo|rigo"
            r"|hono)ur(\w*)", r"\1r\2"),
    _re_sub(r"(label|model|travel|cancel|signal|channel|tunnel|level|fuel|total|counsel|dial|equal|jewel)l"
            r"(ed|ing|er|ers)", r"\1\2"),
    _re_sub(r"grey(s|ish|er|est|ed|ing|scale)?", r"gray\1"),
    _re_sub(r"(cent|lit|met|fib|calib|spect|theat|lust|sab|somb|meag|millilit|microlit|nanolit|kilomet|centimet"
            r"|millimet|micromet|nanomet)re(s)?", r"\1er\2"),
    _re_sub(r"(cent)red", r"\1ered"),
    _re_sub(r"(para|ortho|homo|ana|xeno|cata)logue(s)?", r"\1log\2"),
    _re_sub(r"programme(s)?", r"program\1"),
    # -yse verbs: analyse -> analyze. "analyses" is also the American plural of "analysis", so -yses is never flagged.
    _re_sub(r"(ana|cata|hydro|dia|para|electro|pyro|photo)lys(e|ed|ing)", r"\1lyz\2"),
    _re_sub(r"(" + "|".join(_ISE_STEMS) + r")is(e|ed|es|ing|ation|ations|er|ers)", r"\1iz\2"),
    _re_sub(r"(an|leuk|septic|isch|bacter|tox)aem(\w*)", r"\1em\2"),
    _re_sub(r"haem(\w*)", r"hem\1"),
    _re_sub(r"oestr(\w*)", r"estr\1"),
    _re_sub(r"foet(us|uses|al)", r"fet\1"),
    _re_sub(r"paed(\w+)", r"ped\1"),
    _re_sub(r"oedema(s|tous)?", r"edema\1"),
    _re_sub(r"faec(es|al)", r"fec\1"),
    _re_sub(r"aetiolog(\w*)", r"etiolog\1"),
    _re_sub(r"diarrhoea(l)?", r"diarrhea\1"),
    _re_sub(r"caesium", r"cesium"),
    _re_sub(r"aluminium", r"aluminum"),
    _re_sub(r"sulph(\w+)", r"sulf\1"),
    _re_sub(r"mould(s|ed|y|ing)?", r"mold\1"),
    _re_sub(r"defence(s)?", r"defense\1"),
    _re_sub(r"licence(s)?", r"license\1"),
    _re_sub(r"ageing", r"aging"),
    _re_sub(r"artefact(s|ual)?", r"artifact\1"),
    _re_sub(r"enrolment(s)?", r"enrollment\1"),
    _re_sub(r"fulfil(s|ment)?", r"fulfill\1"),
    _re_sub(r"practise(s|d)?", r"practice\1"),
    _re_sub(r"whilst", r"while"),
    _re_sub(r"amongst", r"among"),
]

PATH_RE = re.compile(r"(?<![\w/.~-])(?:~|\.{1,2})?/?[\w.-]+(?:/[\w.-]+)+/?")
URL_RE = re.compile(r"(?:https?://|ftp://|www\.)\S+|\bdoi:\s*\S+|\b10\.\d{4,9}/\S+", re.I)


def _path_like(token):
    """A slash token is a path when rooted, or when a segment has an underscore or the last one a file extension.
    Units (mg/L, µg/mL) and word pairs (and/or) are none of these."""
    segments = [s for s in token.split("/") if s]
    return (token.startswith(("/", "~", "./", "../")) or any("_" in s for s in segments)
            or bool(segments and re.search(r"[A-Za-z]\w*\.[A-Za-z]{1,5}$", segments[-1])))


DEFAULT_INTERNAL = [  # (label, compiled regex, optional filter on the matched text)
    ("path", PATH_RE, _path_like),
    ("file extension .tsv/.py/.json", re.compile(r"[\w.-]*\.(?:tsv|py|json)\b"), None),
    ("Tier A", re.compile(r"(?<!\w)Tier A(?!\w)", re.I), None),
    ("run 7", re.compile(r"(?<!\w)run 7(?!\w)", re.I), None),
]

# Reference-list headings: those of paper_reference_audit.py plus the other common names. A document with none of
# them gets one citation_hold row, so an unrecognized heading can never read as a clean citation check.
REF_HEADING = re.compile(r"^(?:(?:supplementary|supporting)\s+)?(?:references(?:\s+cited)?|reference\s+list|bibliography"
                         r"|literature\s+cited|cited\s+literature|works\s+cited)\s*[:.]?$", re.I)
CAPTION_START = re.compile(r"^\s*(?:Supplementary\s+)?(?:Figs?\.?|Figures?|Tables?)\s*S?\d+[A-Za-z]?\s*[.:|]", re.I)
SUPP_END = re.compile(r"(?:supplementary|supporting) (?:information|material)|(?:figure|table) legends", re.I)
ABBREV = {"al", "fig", "figs", "e.g", "i.e", "eg", "ie", "sp", "spp", "subsp", "var", "vs", "cf", "ca", "approx", "no",
          "nos", "dr", "st", "eq", "eqs", "ref", "refs", "vol", "pp", "suppl", "etc", "inc", "ltd", "co", "prof", "resp",
          "ed", "eds", "est", "tab", "jan", "feb", "mar", "apr", "jun", "jul", "aug", "sep", "sept", "oct", "nov", "dec"}
BOUNDARY = re.compile(r"[.!?][\"”’')\]]*\s+(?=[\"“‘(\[]?[A-Z0-9])")


# --------------------------------------------------------------------------------------------- reading the document

class _Styles:
    """Style lookups from styles.xml: names, defined ids and inherited italic/bold."""

    def __init__(self, doc):
        el = doc.styles.element
        self.by_id = {s.get(W + "styleId"): s for s in el.iter(W + "style")}
        self.default_para = next((sid for sid, s in self.by_id.items()
                                  if s.get(W + "type") == "paragraph" and s.get(W + "default") in ("1", "true")), None)
        rpr = el.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr")
        self.doc_default = {t: _toggle(rpr.find(W + t)) if rpr is not None else None for t in ("i", "b")}
        self._cache = {}

    def name(self, sid):
        s = self.by_id.get(sid)
        n = s.find(W + "name") if s is not None else None
        return n.get(W + "val") if n is not None else (sid or "")

    def prop(self, sid, tag, depth=0):
        """Inherited on/off of a run property (i or b) through basedOn; None when no style in the chain sets it."""
        key = (sid, tag)
        if key in self._cache:
            return self._cache[key]
        s = self.by_id.get(sid)
        value = None
        if s is not None and depth < 20:
            rpr = s.find(W + "rPr")
            if rpr is not None and rpr.find(W + tag) is not None:
                value = _toggle(rpr.find(W + tag))
            else:
                based = s.find(W + "basedOn")
                if based is not None:
                    value = self.prop(based.get(W + "val"), tag, depth + 1)
        self._cache[key] = value
        return value


def _toggle(el):
    if el is None:
        return None
    return el.get(W + "val") not in ("0", "false", "off")


def _run_flag(run, tag, para_sid, styles):
    rpr = run.find(W + "rPr")
    if rpr is not None and rpr.find(W + tag) is not None:
        return _toggle(rpr.find(W + tag))
    cs = None
    if rpr is not None and rpr.find(W + "rStyle") is not None:
        cs = styles.prop(rpr.find(W + "rStyle").get(W + "val"), tag)
    ps = styles.prop(para_sid, tag)
    if cs is not None and ps is not None:
        return cs != ps          # style toggles combine by exclusive or
    if cs is not None or ps is not None:
        return bool(cs if cs is not None else ps)
    return bool(styles.doc_default.get(tag))


def _para_sid(p, styles):
    ps = p.find(f"{W}pPr/{W}pStyle")
    return ps.get(W + "val") if ps is not None else styles.default_para


def _read_paragraph(p, styles):
    """Visible text of one w:p with per-character italic and bold, hyperlinked and inserted runs included.
    Deleted text, text boxes and field instructions are skipped."""
    sid = _para_sid(p, styles)
    text, italic, bold = [], [], []

    def walk(el):
        for child in el:
            tag = child.tag
            if tag in (W + "del", W + "txbxContent", W + "pPr", W + "moveFrom", W + "rPr"):
                continue
            if tag == W + "r":
                it, bd = _run_flag(child, "i", sid, styles), _run_flag(child, "b", sid, styles)
                for rc in child:
                    s = {W + "t": rc.text or "", W + "tab": "\t", W + "br": " ", W + "cr": " ",
                         W + "noBreakHyphen": "-"}.get(rc.tag)
                    if s:
                        text.append(s)
                        italic.extend([it] * len(s))
                        bold.extend([bd] * len(s))
            else:
                walk(child)
    walk(p)
    return "".join(text), italic, bold, sid


def _load(path):
    import docx  # python-docx; named in README
    data = Path(path).read_bytes()
    return docx.Document(io.BytesIO(data)), hashlib.sha256(data).hexdigest()


def _units(doc, styles):
    """One unit per top-level paragraph (label = its index in doc.paragraphs) and per paragraph inside a top-level
    table (label T<k>), each with kind: heading, reference, caption, body or table."""
    units, in_refs = [], False
    for i, para in enumerate(doc.paragraphs):
        text, italic, bold, sid = _read_paragraph(para._p, styles)
        sname = styles.name(sid)
        low = sname.casefold()
        heading = low.startswith(("heading", "title", "subtitle", "toc")) or low in ("toc heading",)
        stripped = text.strip()
        if REF_HEADING.fullmatch(stripped):
            in_refs, kind = True, "heading"
        elif in_refs and (heading or SUPP_END.fullmatch(stripped)):
            in_refs, kind = False, "heading"
        elif heading:
            kind = "heading"
        elif in_refs:
            kind = "reference"
        elif CAPTION_START.match(text) or "caption" in low:
            kind = "caption"
        else:
            kind = "body"
        units.append(dict(label=i, el=para._p, text=text, italic=italic, bold=bold, kind=kind, style=sname, sid=sid))
    tables = [c for c in doc.element.body if c.tag == W + "tbl"]
    for k, tbl in enumerate(tables, 1):
        for p in tbl.iter(W + "p"):
            text, italic, bold, sid = _read_paragraph(p, styles)
            units.append(dict(label=f"T{k}", el=p, text=text, italic=italic, bold=bold, kind="table",
                              style=styles.name(sid), sid=sid))
    return units


def quote(text, start=0, end=None, limit=QUOTE_WORDS):
    """At most `limit` words of text around [start, end); an ellipsis is attached to the end word when cut."""
    words = list(re.finditer(r"\S+", text))
    if not words:
        return ""
    end = len(text) if end is None else end
    hit = [i for i, m in enumerate(words) if m.end() > start and m.start() < max(end, start + 1)] or [0]
    a, b = hit[0], min(hit[-1], hit[0] + limit - 1)
    while b - a + 1 < limit and (a > 0 or b < len(words) - 1):
        if a > 0:
            a -= 1
        if b - a + 1 < limit and b < len(words) - 1:
            b += 1
    q = " ".join(text[words[a].start():words[b].end()].split())
    return ("…" if a > 0 else "") + q + ("…" if b < len(words) - 1 else "")


def _finding(check, unit, q, detail):
    return dict(check=check, para=unit["label"] if unit is not None else "", quote=q, detail=detail)


def _all_set(flags, text, start, end):
    chars = [flags[i] for i in range(start, end) if not text[i].isspace()]
    return bool(chars) and all(chars), any(chars)


def _in_url(text, start):
    s = text.rfind(" ", 0, start) + 1
    e = text.find(" ", start)
    token = text[s:] if e < 0 else text[s:e]
    return "://" in token or token.startswith("www.")


# --------------------------------------------------------------------------------------------- check 1

def _letters(name):
    decomposed = unicodedata.normalize("NFKD", name)
    return "".join(c for c in decomposed if c.isalpha()).casefold()


def _given_name(prefix, surname):
    """The first author's given name or initials: the text after the surname up to the next author separator
    ("Smith, John, and B. Jones" -> "John"; "Smith JA, Lee B" -> "JA"). Empty for a corporate author."""
    lead = prefix.lstrip()
    if not unicodedata.normalize("NFKC", lead).casefold().startswith(surname):
        return ""
    rest = lead[len(surname):].lstrip(" ,")
    return re.split(r",|;|\s+(?:and|&)\s+|\(", rest, maxsplit=1)[0].strip(" .")


def _initials(given):
    return "".join(w[0] for w in re.findall(r"[^\W\d_]+", given)).casefold()


def _sorts_before(e, prev):
    """True when entry e belongs above prev: surname letter by letter, then given name. Two given names with the same
    initials ("C. R." and "Cameron R.") are taken as one person written two ways, not as an order error."""
    if e["order"][0] != prev["order"][0]:
        return e["order"][0] < prev["order"][0]
    if _initials(e["given"]) == _initials(prev["given"]):
        return False
    return e["order"][1] < prev["order"][1]


PARTICLES = {"van", "von", "de", "der", "den", "la", "le", "du", "da", "di", "dos", "del", "ten", "ter", "the"}


def _surname(name):
    """Matching key: letters only, leading particles and "The" dropped ("van der Maaten" = "Maaten, L. van der")."""
    words = name.casefold().split()
    while len(words) > 1 and words[0] in PARTICLES:
        words = words[1:]
    return _letters(" ".join(words))


_UP = "A-ZÀ-ÖØ-Þ"
_NAME = rf"[{_UP}][\w'’\-]*"
_PART = r"(?:(?:van|von|de|der|den|la|le|du|da|di|dos|del|ten|ter)\s+)"
_AUTHOR = rf"{_PART}*{_NAME}(?:\s+{_NAME}){{0,2}}"
_MORE = (rf"(?:(?:,\s+{_PART}*{_NAME})+,?\s+(?:and|&)\s+{_PART}*{_NAME}"   # "Smith, Jones, and Brown"
         rf"|(?:,\s+{_PART}*{_NAME})*,?\s+et\s+al\.?"                    # "Smith et al.", "Currie, Scott, et al."
         rf"|\s+(?:and|&)\s+{_PART}*{_NAME})?")                            # "Smith and Jones"
_YEARS = rf"({YEAR}(?:\s*[,;]\s*(?:{YEAR}|[a-z](?!\w)(?=\s*(?:[,;)\]]|$))))*)"   # also "2019a,b"
CITE_IN_GROUP = re.compile(rf"({_AUTHOR}){_MORE},?\s+{_YEARS}(?!\w)")
CITE_NARRATIVE = re.compile(rf"({_AUTHOR}){_MORE}\s+\({_YEARS}\)")
GROUP = re.compile(r"\(([^()]*)\)|\[([^\[\]]*)\]")
NOT_AUTHORS = {"january", "february", "march", "april", "may", "june", "july", "august", "september", "october",
               "november", "december", "spring", "summer", "autumn", "fall", "winter", "fig", "figs", "figure",
               "figures", "table", "tables", "version", "release", "since", "in", "from", "and", "until", "before",
               "after", "during"}


def _expand_years(text):
    """'2019a,b' -> 2019a, 2019b; '2018, 2020' -> 2018, 2020. A bare letter takes the year before it."""
    out = []
    for tok in re.findall(rf"{YEAR}|(?<![\w])[a-z](?!\w)", text):
        if len(tok) == 1:
            if out and out[-1][-1].isalpha():
                out.append(out[-1][:-1] + tok)
        else:
            out.append(tok)
    return out


def find_citations(text):
    """Author-year pairs: inside any (...) or [...] group, wherever they sit among other material
    ("PCoA; Gower 1966", "RAxML-NG, Kozlov et al. 2019"), and narrative "Author et al. (2019)".
    Yields (names, year, start, end): names are the comma-listed names before "et al." in order ("Currie, Scott,
    et al." gives Currie then Scott; "RAxML-NG, Kozlov et al." gives RAxML-NG then Kozlov), or the one first
    author. The caller takes the first name the reference list knows. The stricter grammar of
    paper_reference_audit.py holds such mixed groups for a person; this QA pass needs the pairs themselves."""
    found = []
    for g in GROUP.finditer(text):
        body, offset = (g.group(1), g.start(1)) if g.group(1) is not None else (g.group(2), g.start(2))
        if not YEAR_RE.search(body):
            continue
        found += [(m, offset) for m in CITE_IN_GROUP.finditer(body)]
    found += [(m, 0) for m in CITE_NARRATIVE.finditer(text)]
    out = []
    for m, offset in found:
        head = m.group(0)[:m.start(2) - m.start()]
        names = [m.group(1)]
        if "," in head[len(m.group(1)):]:          # the listed co-authors, not the one after the final "and"
            names += [n.strip() for n in re.findall(rf"\s+({_PART}*{_NAME})(?=,|\s+et\s+al|\s+(?:and|&)\s)",
                                                    head[len(m.group(1)):])]
        names = [n for n in names if n.split()[-1].casefold().strip(".") not in NOT_AUTHORS]
        if not names:
            continue                                   # "14 July 2026", "Fig. 2019": a date or a label
        for y in _expand_years(m.group(2)):
            row = (tuple(names), y, offset + m.start(), offset + m.end())
            if row not in out:
                out.append(row)
    return out


def check_citations(units):
    out, refs = [], [u for u in units if u["kind"] == "reference" and u["text"].strip()]
    if not refs:
        n = sum(len(find_citations(u["text"])) for u in units)
        first = next((u for u in units if find_citations(u["text"])), None)
        return [_finding("citation_hold", first, quote(first["text"]) if first else "",
                         f"no reference list found (heading not recognized); {n} citation{'' if n == 1 else 's'} "
                         "in the text not checked")]
    if all(NUMBER.match(u["text"]) for u in refs):
        return [_finding("citation_hold", refs[0], quote(refs[0]["text"]),
                         "numbered reference list; use tools/paper_reference_audit.py --style numbered")]
    entries, prev = [], None
    for u in refs:
        year = YEAR_RE.search(u["text"])
        name = author_key(u["text"][:year.start()]) if year else None
        if not name:
            out.append(_finding("citation_hold", u, quote(u["text"]),
                                "reference not parsed: one surname-first entry with a year per paragraph"))
            continue
        given = _given_name(u["text"][:year.start()], name)
        e = dict(unit=u, key=(name, year.group(1)), order=(_letters(name), _letters(given)), given=given, cited=False)
        entries.append(e)
        if prev is not None and _sorts_before(e, prev):
            this = f"{name.title()}, {given}".rstrip(", ")
            that = f'{prev["key"][0].title()}, {prev["given"]}'.rstrip(", ")
            out.append(_finding("reference_order", u, quote(u["text"]),
                                f'"{this}" sorts before "{that}" (¶{prev["unit"]["label"]}): letter by letter on the '
                                "surname, then the given name"))
        prev = e
    index, years, shown = defaultdict(list), defaultdict(set), {}
    for e in entries:
        canon = _surname(e["key"][0])
        index[(canon, e["key"][1])].append(e)
        years[canon].add(e["key"][1])
        shown.setdefault(canon, e["key"][0].title())
    for (canon, year), group in index.items():
        if len(group) > 1:
            out.append(_finding("citation_hold", group[1]["unit"], quote(group[1]["unit"]["text"]),
                                f"{len(group)} references share the key {shown[canon]} {year}"))
    for u in units:
        if u["kind"] == "reference":
            continue
        text = u["text"]
        for names, year, start, end in find_citations(text):
            # first listed name the reference list knows; a multi-word name may carry a leading capitalized word
            # ("However Marsh"), so its shorter tails are tried too
            tails = [" ".join(n.split()[k:]) for n in names for k in range(len(n.split()))]
            canon = next((_surname(t) for t in tails if _surname(t) in years), _surname(names[0]))
            author = names[0] + (" et al." if len(names) > 1 else "")
            if (canon, year) in index:
                for e in index[(canon, year)]:
                    e["cited"] = True
            elif canon in years:
                out.append(_finding("citation_year_mismatch", u, quote(text, start, end),
                                    f"{shown[canon]} {year} cited; the reference list has {shown[canon]} "
                                    + ", ".join(sorted(years[canon]))))
            else:
                out.append(_finding("citation_unmatched", u, quote(text, start, end), f"no reference for {author} {year}"))
    for e in entries:
        if not e["cited"]:
            out.append(_finding("reference_uncited", e["unit"], quote(e["unit"]["text"]),
                                f"{e['key'][0].title()} {e['key'][1]} is not cited in the text"))
    return out


# --------------------------------------------------------------------------------------------- check 2

def load_spelling(path):
    rules = []
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        brit, _, amer = line.partition("\t")
        brit, amer = brit.strip(), amer.strip()
        if brit.startswith("re:"):
            rules.append(_re_sub(brit[3:], amer))
        else:
            rules.append((re.escape(brit), (lambda a: (lambda w: a))(amer)))
    return rules


def check_spelling(units, rules, genera=DEFAULT_GENERA):
    """Italic words (Latin names, foreign terms) and a species epithet after a genus or its initial ("S. sulphureus")
    are skipped: a deposited name keeps its spelling."""
    out, genera = [], set(genera)
    compiled = [(re.compile(rf"\b(?:{pat})\b", re.I), fix) for pat, fix in rules]
    for u in units:
        if u["kind"] == "reference":
            continue
        seen = set()
        for rx, fix in compiled:
            for m in rx.finditer(u["text"]):
                if m.span() in seen or _in_url(u["text"], m.start()) or _all_set(u["italic"], u["text"], *m.span())[0]:
                    continue
                before = re.search(r"(\b[A-Z][a-z]*\.?)\s+$", u["text"][:m.start()])
                if before and (re.fullmatch(r"[A-Z]\.", before.group(1)) or before.group(1) in genera):
                    continue
                seen.add(m.span())
                word = m.group(0)
                us = fix(word)
                us = us[:1].upper() + us[1:] if word[:1].isupper() else us
                out.append(_finding("british_spelling", u, quote(u["text"], *m.span()), f'"{word}" (American: "{us}")'))
    return sorted(out, key=lambda f: (str(f["para"]), f["detail"]))


# --------------------------------------------------------------------------------------------- check 3

def load_genera(path):
    return [ln.strip() for ln in Path(path).read_text(encoding="utf-8").splitlines()
            if ln.strip() and not ln.lstrip().startswith("#")]


PANEL_GROUP = re.compile(r"\(([^()]{1,20})\)")
PANEL_LETTERS = re.compile(r"^\s*[A-Za-z](?:\s*(?:[,–-]|and)\s*[A-Za-z])*\s*$")
FAMILY = re.compile(r"\b[A-Z][a-z]+aceae\b")


def check_italics(units, genera):
    out = []
    genus_rx = re.compile(r"\b(" + "|".join(sorted(map(re.escape, genera), key=len, reverse=True)) + r")\b") \
        if genera else None
    for u in units:
        text, it = u["text"], u["italic"]
        if genus_rx is not None:
            for m in genus_rx.finditer(text):
                if _in_url(text, m.start()):
                    continue
                every, some = _all_set(it, text, *m.span())
                if not every:
                    how = "partly upright" if some else "set upright"
                    out.append(_finding("genus_upright", u, quote(text, *m.span()), f'"{m.group(0)}" {how} (genus names are italic)'))
        for m in FAMILY.finditer(text):
            if _all_set(it, text, *m.span())[0]:
                out.append(_finding("family_italic", u, quote(text, *m.span()), f'"{m.group(0)}" set italic (family names are upright)'))
        if u["kind"] == "caption":
            for m in PANEL_GROUP.finditer(text):
                if not PANEL_LETTERS.match(m.group(1)):
                    continue
                letters = [m.start(1) + j for j, c in enumerate(m.group(1)) if c.isalpha() and
                           (j == 0 or not m.group(1)[j - 1].isalpha()) and
                           (j + 1 == len(m.group(1)) or not m.group(1)[j + 1].isalpha())]
                if any(it[i] for i in letters):
                    out.append(_finding("panel_label_italic", u, quote(text, *m.span()),
                                        f'"({m.group(1)})" panel label set italic (panel labels are upright)'))
    return out


# --------------------------------------------------------------------------------------------- check 4

def sentences(text):
    spans, start = [], 0
    for m in BOUNDARY.finditer(text):
        before = re.search(r"(\S+)$", text[:m.start() + 1])
        token = before.group(1).lstrip("([\"“‘").rstrip(".").casefold() if before else ""
        if token in ABBREV or re.fullmatch(r"[a-z]", token) or re.fullmatch(r"(?:[a-z]\.)+[a-z]", token):
            continue
        spans.append((start, m.start() + 1))
        start = m.end()
    if text[start:].strip():
        spans.append((start, len(text.rstrip())))
    return spans


def check_short_sentences(units, max_words=5):
    out = []
    for u in units:
        text = u["text"]
        if u["kind"] != "body" or not text.strip():
            continue
        letters = [b for b, c in zip(u["bold"], text) if c.isalnum()]
        if letters and all(letters):
            continue                                       # a bold run-in heading
        if not re.search(r"[.!?][\"”’')\]]*\s*$", text) and len(text.split()) <= max_words:
            continue                                       # a label, not a sentence
        spans = sentences(text)
        for k, (s, e) in enumerate(spans):
            sent = text[s:e].strip()
            n = len([w for w in sent.split() if any(c.isalnum() for c in w)])
            if k == 0 and (any(_all_set(flags, text, s, e)[0] for flags in (u["bold"], u["italic"]))
                           or (n <= 2 and len(spans) > 1)):
                continue                                   # a run-in label: a bold "Assembly." or "Methods. Counts…"
            if 0 < n <= max_words:
                out.append(_finding("short_sentence", u, quote(sent), f"{n} words"))
    return out


# --------------------------------------------------------------------------------------------- check 5

def load_internal(path):
    entries = []
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        s = line.strip()
        if not s or s.startswith("#"):
            continue
        if s.startswith("re:"):
            entries.append((s[3:], re.compile(s[3:]), None))
        else:
            entries.append((s, re.compile(rf"(?<!\w){re.escape(s)}(?!\w)", re.I), None))
    return entries


def check_internal_names(units, entries):
    out = []
    for u in units:
        if u["kind"] == "reference":
            continue
        masked = URL_RE.sub(lambda m: " " * len(m.group(0)), u["text"])
        taken = []
        for label, rx, keep in entries:
            for m in rx.finditer(masked):
                hit = m.group(0).rstrip(".,;:")
                if not hit or (keep is not None and not keep(hit)):
                    continue
                span = (m.start(), m.start() + len(hit))
                if any(a < span[1] and span[0] < b for a, b in taken):
                    continue
                taken.append(span)
                out.append(_finding("internal_name", u, quote(u["text"], *span), f'"{hit}" ({label})'))
    return out


# --------------------------------------------------------------------------------------------- check 6

def check_links_and_styles(units, doc, styles):
    """Undefined paragraph styles: one row per paragraph. Undefined run styles: one row per style, at its first use,
    with the number of paragraphs (Word falls back to default formatting for them)."""
    out, rels, run_styles = [], doc.part.rels, {}
    for u in units:
        for hl in u["el"].iter(W + "hyperlink"):
            rid = hl.get(R_NS + "id")
            if not rid:
                continue                                  # w:anchor links go to a bookmark in the document
            link_text = "".join(t.text or "" for t in hl.iter(W + "t"))
            pos = max(u["text"].find(link_text), 0) if link_text else 0
            rel = rels.get(rid)
            if rel is None:
                out.append(_finding("hyperlink_internal_target", u, quote(u["text"], pos, pos + len(link_text)),
                                    f'link "{link_text}" names {rid}, which has no relationship'))
            elif not rel.is_external:
                out.append(_finding("hyperlink_internal_target", u, quote(u["text"], pos, pos + len(link_text)),
                                    f'link "{link_text}" points at {rel.target_ref} inside the document '
                                    f'({rel.reltype.rsplit("/", 1)[-1]}), not a URL'))
        ps = u["el"].find(f"{W}pPr/{W}pStyle")
        if ps is not None and ps.get(W + "val") not in styles.by_id:
            out.append(_finding("style_undefined", u, quote(u["text"]),
                                f'paragraph style "{ps.get(W + "val")}" is not defined in the document'))
        for sid in {rs.get(W + "val") for rs in u["el"].iter(W + "rStyle")} - set(styles.by_id):
            first, n = run_styles.get(sid, (u, 0))
            run_styles[sid] = (first, n + 1)
    for sid, (u, n) in sorted(run_styles.items()):
        out.append(_finding("style_undefined", u, quote(u["text"]),
                            f'run style "{sid}" is not defined in the document (used in {n} paragraphs; first shown)'))
    return out


# --------------------------------------------------------------------------------------------- checks 7 and 8

def _sections_by_element(doc):
    """Top-level body element -> section index (a paragraph holding w:sectPr ends its section)."""
    out, k = {}, 0
    for child in doc.element.body:
        out[child] = k
        if child.tag == W + "p" and child.find(f"{W}pPr/{W}sectPr") is not None:
            k += 1
    return out


def _top(el, body):
    while el is not None and el.getparent() is not body:
        el = el.getparent()
    return el


def _caption_near(units, i):
    for j in list(range(i + 1, min(i + 4, len(units)))) + list(range(i - 1, max(i - 3, -1), -1)):
        if units[j]["kind"] == "caption" and units[j]["label"] != "" and isinstance(units[j]["label"], type(units[i]["label"])):
            return units[j]["text"]
    return ""


def _declared_pt(info):
    for key, value in info.items():
        if str(key).casefold().replace(" ", "_") in ("min_text_pt", "text_pt"):
            try:
                return float(value)
            except (TypeError, ValueError):
                return None
    return None


def check_images(units, doc, *, min_scale=0.95, assume_text_pt=None, min_text_pt=8.0, superseded_dir=None):
    from PIL import Image  # Pillow; named in README
    out, body = [], doc.element.body
    sec_of, sections = _sections_by_element(doc), list(doc.sections)
    old = {}
    if superseded_dir:
        base = Path(superseded_dir)
        for f in sorted(base.rglob("*")):
            if f.is_file() and f.stat().st_size <= 200 * 1024 * 1024:
                old.setdefault(hashlib.sha256(f.read_bytes()).hexdigest(), f.relative_to(base).as_posix())
    n = 0
    for i, u in enumerate(units):
        for drawing in u["el"].iter(W + "drawing"):
            frame = next((c for c in drawing if c.tag in (WP + "inline", WP + "anchor")), None)
            extent = frame.find(WP + "extent") if frame is not None else None
            blip = next(drawing.iter(A + "blip"), None)
            rid = blip.get(R_NS + "embed") if blip is not None else None
            if extent is None or not rid or rid not in doc.part.related_parts:
                continue
            n += 1
            blob = doc.part.related_parts[rid].blob
            pw, ph = int(extent.get("cx")) / EMU_PER_IN, int(extent.get("cy")) / EMU_PER_IN
            q = quote(_caption_near(units, i) or u["text"]) or f"image {n}"
            sec = sections[min(sec_of.get(_top(u["el"], body), 0), len(sections) - 1)] if sections else None
            if sec is not None and sec.page_width is not None and sec.page_height is not None:
                tw = (sec.page_width - (sec.left_margin or 0) - (sec.right_margin or 0)) / EMU_PER_IN
                th = (sec.page_height - (sec.top_margin or 0) - (sec.bottom_margin or 0)) / EMU_PER_IN
                if ph > th + 0.01:
                    out.append(_finding("image_taller_than_text_block", u, q, f"placed {ph:.2f} in tall; text block {th:.2f} in"))
                if pw > tw + 0.01:
                    out.append(_finding("image_wider_than_text_block", u, q, f"placed {pw:.2f} in wide; text block {tw:.2f} in"))
            digest = hashlib.sha256(blob).hexdigest()
            if digest in old:
                out.append(_finding("superseded_render", u, q, f"embedded image is byte-identical to {old[digest]}"))
            try:
                img = Image.open(io.BytesIO(blob))
                px_w, info, fmt = img.size[0], dict(img.info), img.format
            except Exception:                              # vector formats (EMF, SVG) have no drawn pixel size
                continue
            dpi = info.get("dpi")
            if not dpi or not dpi[0] or dpi[0] <= 1:
                out.append(_finding("image_no_dpi", u, q, f"{fmt} {img.size[0]}x{img.size[1]} px with no dpi; "
                                                          f"drawn size unknown; placed {pw:.2f} in"))
                continue
            drawn = px_w / float(dpi[0])
            scale = pw / drawn if drawn else 1.0
            declared = _declared_pt(info) or assume_text_pt
            text_note = f"; {declared:g} pt text prints at {declared * scale:.1f} pt" if declared else ""
            if scale < min_scale:
                out.append(_finding("print_scale", u, q, f"placed {pw:.2f} in, drawn {drawn:.2f} in ({px_w} px at "
                                                         f"{float(dpi[0]):.0f} dpi): prints at {scale:.1%}{text_note}"))
            if declared and declared * scale < min_text_pt - 1e-9:
                out.append(_finding("effective_text_below_min", u, q, f"{declared:g} pt text at {scale:.1%} prints at "
                                                                      f"{declared * scale:.1f} pt (minimum {min_text_pt:g} pt)"))
    return out


# --------------------------------------------------------------------------------------------- driver and reports

def run_checks(path, *, genus_file=None, spelling_file=None, internal_names_file=None, superseded_dir=None,
               min_scale=0.95, assume_text_pt=None, min_text_pt=8.0, return_meta=False):
    doc, digest = _load(path)
    styles = _Styles(doc)
    units = _units(doc, styles)
    findings = (check_citations(units)
                + check_spelling(units, load_spelling(spelling_file) if spelling_file else DEFAULT_SPELLING,
                                 load_genera(genus_file) if genus_file else DEFAULT_GENERA)
                + check_italics(units, load_genera(genus_file) if genus_file else DEFAULT_GENERA)
                + check_short_sentences(units)
                + check_internal_names(units, load_internal(internal_names_file) if internal_names_file else DEFAULT_INTERNAL)
                + check_links_and_styles(units, doc, styles)
                + check_images(units, doc, min_scale=min_scale, assume_text_pt=assume_text_pt, min_text_pt=min_text_pt,
                               superseded_dir=superseded_dir))
    order = {c: k for k, (c, _, _) in enumerate(CHECKS)}
    findings.sort(key=lambda f: (order.get(f["check"], 99), isinstance(f["para"], str),
                                 f["para"] if isinstance(f["para"], int) else str(f["para"])))
    if return_meta:
        n_images = sum(1 for u in units for _ in u["el"].iter(W + "drawing"))
        return findings, dict(sha256=digest, paragraphs=len(doc.paragraphs), images=n_images,
                              references=sum(u["kind"] == "reference" for u in units))
    return findings


def _md_cell(s):
    return str(s).replace("|", "\\|").replace("\n", " ")


def write_reports(path, findings, meta, out_dir, settings):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^\w.-]+", "_", Path(path).stem).strip("_") or "document"
    tsv, md = out_dir / f"{stem}_manuscript_qa.tsv", out_dir / f"{stem}_manuscript_qa.md"
    with tsv.open("w", encoding="utf-8", newline="") as fh:
        w = SafeWriter(fh, delimiter="\t", lineterminator="\n")
        w.writerow(COLUMNS)
        w.writerows([[f["check"], f["para"], f["quote"], f["detail"]] for f in findings])
    counts = defaultdict(int)
    for f in findings:
        counts[f["check"]] += 1
    lines = [f"# Manuscript QA: {Path(path).name}", "",
             f"- File: `{Path(path).resolve()}`", f"- SHA-256: `{meta['sha256']}`",
             f"- {meta['paragraphs']} paragraphs, {meta['references']} reference entries, {meta['images']} images.",
             "- ¶ is the index in python-docx `Document(...).paragraphs`; `T<k>` is the k-th top-level table.",
             f"- Settings: {settings}.",
             "- The check reads the file and never changes it. Each row is a lead for a person to check.", "",
             "## Counts", "", "| # | Check | What it looks for | Findings |", "|---|---|---|---|"]
    lines += [f"| {num} | `{c}` | {desc} | {counts.get(c, 0)} |" for c, num, desc in CHECKS]
    lines += ["", f"Total: {len(findings)} findings.", ""]
    for c, num, desc in CHECKS:
        rows = [f for f in findings if f["check"] == c]
        if not rows:
            continue
        lines += [f"## {num}. `{c}` ({len(rows)})", "", desc[:1].upper() + desc[1:] + ".", "",
                  "| ¶ | Quote | Detail |", "|---|---|---|"]
        lines += [f"| {_md_cell(f['para'])} | {_md_cell(f['quote'])} | {_md_cell(f['detail'])} |" for f in rows]
        lines.append("")
    md.write_text("\n".join(lines), encoding="utf-8")
    return md, tsv, dict(counts)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("docx", type=Path, nargs="+", help="Word documents to check (read only)")
    ap.add_argument("--out", type=Path, required=True, help="folder for <name>_manuscript_qa.md and .tsv")
    ap.add_argument("--genus-file", type=Path, help="genus names, one per line (replaces the default list)")
    ap.add_argument("--spelling-file", type=Path, help="TSV british<TAB>american, re:<regex> allowed (replaces the default)")
    ap.add_argument("--internal-names", type=Path, help="one name per line, re:<regex> allowed (replaces the default)")
    ap.add_argument("--superseded", type=Path, help="folder of superseded renders; embedded images matching their bytes are flagged")
    ap.add_argument("--min-scale", type=float, default=0.95, help="flag images placed below this share of drawn width")
    ap.add_argument("--assume-text-pt", type=float, help="figure text size to assume when a PNG declares none")
    ap.add_argument("--min-text-pt", type=float, default=8.0, help="minimum printed text size (pt)")
    a = ap.parse_args(argv)
    for p in a.docx:
        if p.suffix.casefold() != ".docx" or not p.is_file():
            _LOG.error("manuscript_qa refused: not a readable .docx file: %s", p)
            return 2
    for opt in (a.genus_file, a.spelling_file, a.internal_names):
        if opt is not None and not opt.is_file():
            _LOG.error("manuscript_qa refused: list file not found: %s", opt)
            return 2
    if a.superseded is not None and not a.superseded.is_dir():
        _LOG.error("manuscript_qa refused: superseded folder not found: %s", a.superseded)
        return 2
    settings = (f"min scale {a.min_scale:g}, min text {a.min_text_pt:g} pt"
                + (f", assumed text {a.assume_text_pt:g} pt" if a.assume_text_pt else "")
                + f", genera {'from ' + a.genus_file.name if a.genus_file else 'default list'}"
                + f", spelling {'from ' + a.spelling_file.name if a.spelling_file else 'default list'}"
                + f", internal names {'from ' + a.internal_names.name if a.internal_names else 'default list'}"
                + (f", superseded renders from {a.superseded}" if a.superseded else ""))
    for p in a.docx:
        try:
            findings, meta = run_checks(p, genus_file=a.genus_file, spelling_file=a.spelling_file,
                                        internal_names_file=a.internal_names, superseded_dir=a.superseded,
                                        min_scale=a.min_scale, assume_text_pt=a.assume_text_pt,
                                        min_text_pt=a.min_text_pt, return_meta=True)
        except Exception as exc:                           # a corrupt or non-Word zip
            _LOG.error("manuscript_qa refused: %s could not be read: %s", p, exc)
            return 2
        md, tsv, counts = write_reports(p, findings, meta, a.out, settings)
        _LOG.info("%s: %d findings -> %s", p.name, len(findings), md)
        for c, _, _ in CHECKS:
            if counts.get(c):
                _LOG.info("  %-30s %d", c, counts[c])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
