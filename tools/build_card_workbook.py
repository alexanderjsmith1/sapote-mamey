#!/usr/bin/env python3
"""build_card_workbook.py (v9.7.199) — cross-strain card workbook.

Demonstrates/produces a single Excel workbook that embeds per-BGC deliverables as sheets for
cross-strain analysis, across N sealed packages:

  B10_BGC_Report_Cards  — one row per BGC (deterministic triage export: products, boundary,
                          AB/AF, novelty, CCTT, lead tier, KCB-top).
  E5_ModeB_Cards        — one row per Mode B §-section (AUTHORED cards if present in the package,
                          else the emit-modeb-template SKELETON, tagged in Provenance).

This is an ADDITIVE tool. It does NOT modify the FROZEN master workbook (MASTER_SCHEMA_FROZEN_v1_1,
validated by mamey/workbook_schema_check.py). Folding these sheets into the master's writer must be
append-only (new sheet codes B10/E5, never repurposing a frozen code) and is a separate, schema-
gated change.

Binding Excel constraint handled here: a cell holds at most 32,767 chars. A dense §4 gene-by-gene
section can exceed that, so any over-limit section body is CHUNKED into continuation rows
("§4 (cont.1)") rather than silently truncated. File size is not a constraint — the derived
workbook is small (KB–low MB) even though the antiSMASH inputs are 20–120 MB.

Usage:
  python tools/build_card_workbook.py --packages runs_*/  --out cards.xlsx
  python tools/build_card_workbook.py --packages runs/A/package runs/B/package --out cards.xlsx
"""
from __future__ import annotations

import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402
import argparse, csv, glob, importlib.util, re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

try:
    from _wbio import atomic_save
except ModuleNotFoundError:  # supports spec-based import used by the hermetic regression test
    _wbio_path = Path(__file__).resolve().with_name("_wbio.py")
    _wbio_spec = importlib.util.spec_from_file_location("_wbio_build_card_workbook", _wbio_path)
    if not _wbio_spec or not _wbio_spec.loader:
        raise
    _wbio_module = importlib.util.module_from_spec(_wbio_spec)
    _wbio_spec.loader.exec_module(_wbio_module)
    atomic_save = _wbio_module.atomic_save

CELL_LIMIT = 32767
CHUNK_AT = 32000  # leave headroom
FONT = "Arial"
HDR = Font(name=FONT, bold=True, color="FFFFFF", size=11)
HDRFILL = PatternFill("solid", fgColor="2F5496")
BODY = Font(name=FONT, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")

SECTION_RE = re.compile(r"(?m)^#{1,3}\s*§(\d+[a-z]?)\s+([^\n]+)\n(.*?)(?=^#{1,3}\s*§\d|\Z)", re.S)


def _first(package_dir: Path, suffix: str) -> Path | None:
    m = sorted(package_dir.glob(f"*{suffix}"))
    return m[0] if m else None


def _strain_id(package_dir: Path) -> str:
    cw = _first(package_dir, "_2b_bgc_crosswalk.csv")
    if cw:
        return cw.name[:-len("_2b_bgc_crosswalk.csv")]
    return package_dir.parent.name


def bgc_report_rows(package_dir: Path, strain: str) -> list[list]:
    """One row per BGC from the triage board (deterministic)."""
    tb = _first(package_dir, "_4_triage_board.csv")
    rows = []
    if not tb:
        return rows
    with tb.open(newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            rows.append([
                strain, r.get("BGC_ID", ""),
                f"{r.get('Contig') or r.get('Node_ID','')} · {r.get('antiSMASH_Region','')}",
                r.get("Products", ""), r.get("Boundary_Status") or r.get("Boundary", ""),
                # v9.7.371 fix: real triage-board headers (mamey/cli.py's triage_headers, verified
                # directly) are AB_auto/AF_auto/Lead_tier_auto -- AB_score/AB/AF_score/AF/Lead_tier
                # never existed, so .get() silently returned "" for every row (blank AB/AF columns
                # cohort-wide), and the Lead_tier fallback fell through to Corrected_rank -- a
                # different field (a numeric rank correction, not a tier label) -- mislabeling the
                # column with numbers instead of tier names like Exceptional/High/Medium.
                r.get("AB_auto", ""), r.get("AF_auto", ""),
                r.get("Novelty_auto") or r.get("novelty_auto", ""),
                r.get("CCTT_triggers", ""), r.get("Lead_tier_auto", ""),
                r.get("KCB_top") or r.get("kcb_top", ""),
            ])
    return rows


def _find_modeb_cards(package_dir: Path) -> list[Path]:
    """Authored Mode B cards, if any (ingest-receipts writes judgment/*_mode_b.md)."""
    cards = []
    for pat in ("judgment/*_mode_b.md", "*_ModeB.md", "*_Mode_B*.md", "mode_b_templates/*.md"):
        cards += list(package_dir.glob(pat))
    # de-dup by BGC_ID (authored `{strain}_{bgc}_mode_b.md` and template `{bgc}_template.md` differ by
    # filename but are the same BGC → were double-counted), preferring judgment/ (authored) over templates.
    def _bgc_of(card):
        txt = card.read_text(encoding="utf-8", errors="replace")[:400]
        mm = re.search(r"bgc:\s*(BGC\d+)", txt) or re.search(r"(BGC\d+)", card.name)
        return mm.group(1) if mm else card.stem
    seen, out = set(), []
    for c in sorted(cards, key=lambda p: ("mode_b_templates" in str(p), p.name)):
        bid = _bgc_of(c)
        if bid not in seen:
            seen.add(bid); out.append(c)
    return out


def modeb_section_rows(package_dir: Path, strain: str) -> list[list]:
    rows = []
    for card in _find_modeb_cards(package_dir):
        text = card.read_text(encoding="utf-8", errors="replace")
        m = re.search(r"bgc:\s*(BGC\d+)", text) or re.search(r"(BGC\d+)", card.name)
        bgc = m.group(1) if m else card.stem
        provenance = ("AUTHORED" if (re.search(r"<!--\s*MODE B:\s*BGC", text[:400]) and "mode_b_templates" not in str(card))
                      else "SKELETON (emit-modeb-template)")
        for sm in SECTION_RE.finditer(text):
            num, title, bodytext = sm.group(1), sm.group(2).strip(), sm.group(3).strip()
            # chunk over-limit bodies into continuation rows
            if len(bodytext) <= CHUNK_AT:
                rows.append([strain, bgc, f"§{num}", title, bodytext, len(bodytext), provenance])
            else:
                parts = [bodytext[i:i+CHUNK_AT] for i in range(0, len(bodytext), CHUNK_AT)]
                for i, part in enumerate(parts):
                    tag = f"§{num}" if i == 0 else f"§{num} (cont.{i})"
                    rows.append([strain, bgc, tag, title, part, len(part), provenance])
    return rows


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--packages", nargs="+", required=True,
                    help="Package dirs or globs (e.g. 'runs_*/*/package')")
    ap.add_argument("--out", required=True, help="Output .xlsx path")
    args = ap.parse_args()

    # resolve packages: accept dirs, globs, or run roots (find */package under them)
    pkgs: list[Path] = []
    for spec in args.packages:
        for hit in glob.glob(spec):
            p = Path(hit)
            if (p / "manifest.json").exists() or list(p.glob("*_2b_bgc_crosswalk.csv")):
                pkgs.append(p)
            else:
                pkgs += [q.parent for q in p.rglob("*_2b_bgc_crosswalk.csv")]
    pkgs = sorted({p.resolve() for p in pkgs})
    if not pkgs:
        emit("no packages resolved from", args.packages); return 1

    wb = openpyxl.Workbook()
    # README
    ws = wb.active; ws.title = "README"
    for line in [
        ["Cross-strain card workbook (build_card_workbook.py, v9.7.199)"],
        [f"strains: {len(pkgs)}"],
        [""],
        ["B10_BGC_Report_Cards: one row per BGC (deterministic triage export)."],
        ["E5_ModeB_Cards: one row per Mode B §-section; AUTHORED if present, else SKELETON."],
        ["Over-limit section bodies (>32000 chars) are chunked into '§N (cont.k)' rows."],
        ["ADDITIVE tool — does not modify the FROZEN master (append-only fold-in is a separate change)."],
    ]:
        ws.append(line)
    ws["A1"].font = Font(name=FONT, bold=True, size=13); ws.column_dimensions["A"].width = 96

    # B10
    ws = wb.create_sheet("B10_BGC_Report_Cards")
    hdr = ["Strain", "BGC_ID", "Node_Region", "Products", "Boundary", "AB", "AF",
           "Novelty_auto", "CCTT_triggers", "Lead_tier", "KCB_top"]
    ws.append(hdr)
    n_bgc = 0
    for p in pkgs:
        for row in bgc_report_rows(p, _strain_id(p)):
            ws.append(row); n_bgc += 1
    for c in range(1, len(hdr)+1):
        ws.cell(1, c).font = HDR; ws.cell(1, c).fill = HDRFILL
    ws.freeze_panes = "A2"; ws.column_dimensions["A"].width = 26; ws.column_dimensions["D"].width = 38

    # E5
    ws = wb.create_sheet("E5_ModeB_Cards")
    hdr = ["Strain", "BGC_ID", "Section_No", "Section_Title", "Body", "Char_Count", "Provenance"]
    ws.append(hdr)
    n_sec = 0
    for p in pkgs:
        for row in modeb_section_rows(p, _strain_id(p)):
            ws.append(row); n_sec += 1
    if n_sec == 0:
        ws.append(["—", "—", "—", "(no Mode B cards authored or templated in these packages yet)", "", 0, "NONE"])
    for c in range(1, len(hdr)+1):
        ws.cell(1, c).font = HDR; ws.cell(1, c).fill = HDRFILL
    ws.freeze_panes = "A2"; ws.column_dimensions["A"].width = 26; ws.column_dimensions["E"].width = 90
    for r in ws.iter_rows(min_row=2):
        r[4].alignment = WRAP
        for cell in r:
            cell.font = BODY

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    atomic_save(wb, args.out)
    wb.close()

    # verification receipts
    wb2 = openpyxl.load_workbook(args.out, read_only=True)
    maxlen, maxloc = 0, None
    try:
        for s in wb2.sheetnames:
            for row in wb2[s].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and len(cell.value) > maxlen:
                        maxlen, maxloc = len(cell.value), f"{s}!{cell.coordinate}"
    finally:
        wb2.close()
    size_kb = Path(args.out).stat().st_size / 1024
    emit(f'wrote {args.out} ({size_kb:.0f} KB) — strains={len(pkgs)} BGC_rows={n_bgc} ModeB_section_rows={n_sec}', f"max cell length {maxlen} at {maxloc} (Excel limit {CELL_LIMIT}) -> {('OK' if maxlen < CELL_LIMIT else 'EXCEEDS (chunking failed)')}", sep="\n")
    return 0 if maxlen < CELL_LIMIT else 2


if __name__ == "__main__":
    raise SystemExit(main())
