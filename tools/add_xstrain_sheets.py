#!/usr/bin/env python3
"""Append or refresh the three deterministic cross-strain workbook sheets.

The module is import-safe.  Call :func:`main` or execute the file after the base
workbook has been built.  Existing generated sheets are replaced idempotently;
similarly prefixed user sheets are preserved.
"""
from __future__ import annotations

import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402

import argparse
import os
import re
import subprocess
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from _wbio import atomic_dump_json, atomic_save


GENERATED_SHEETS = (
    "Cross_Strain_Class_Prevalence",
    "Cross_Strain_Findings",
    "Strain_Cohort_Context",
)
UNIVERSAL = {"saccharide", "fatty_acid", "other", "terpene"}
DIAGNOSTICS = {
    "TIGR01454": "ansamycin",
    "TIGR03604": "thiopeptide",
    "TIGR03828": "enediyne",
    "TIGR04186": "nucleoside",
}


def _read_json(path, *, encoding="utf-8"):
    import json

    with open(path, encoding=encoding) as handle:
        return json.load(handle)


def _generated_sheet_name(existing: str, canonical: str) -> bool:
    """Match the canonical sheet or openpyxl's numeric collision suffix only."""
    return existing == canonical or re.fullmatch(re.escape(canonical) + r"[0-9]+", existing) is not None


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--banked-dir",
        default=os.environ.get("MAMEY_BANKED_DIR", os.path.join(os.path.dirname(__file__), "..", "cohort")),
    )
    parser.add_argument("--out", default=os.path.join(os.getcwd(), "Sapote-Mamey_Master_Workbook.xlsx"))
    return parser


def main(argv=None) -> int:
    args = _parser().parse_args(argv)
    workbook_path = os.path.abspath(args.out)
    banked_dir = os.path.abspath(args.banked_dir)

    bgc_data = _read_json(os.path.join(banked_dir, "bgc_data.json"))
    gene_data = _read_json(os.path.join(banked_dir, "gene_data.json"))
    tigrfam = _read_json(os.path.join(banked_dir, "tigrfam.json"))
    # These inputs are part of the bank contract even though this builder does not
    # currently project their fields into its three sheets.
    _read_json(os.path.join(banked_dir, "deep_data.json"))
    _read_json(os.path.join(banked_dir, "rggmci_full.json"))

    strains = bgc_data["strains"]
    bgcs = bgc_data["bgcs"]
    if not strains:
        raise SystemExit("add_xstrain_sheets: bgc_data.json contains zero strains; cohort denominator is undefined")
    if not isinstance(bgcs, list):
        raise SystemExit("add_xstrain_sheets: bgc_data.json field 'bgcs' must be a list")
    if not isinstance(gene_data, dict):
        raise SystemExit("add_xstrain_sheets: gene_data.json must contain a JSON object")

    denominator = len(strains)

    def tier(strain):
        n50 = strains[strain]["n50"]
        return "GOOD" if n50 >= 1e6 else "MOD" if n50 >= 1e5 else "POOR"

    strain_order = sorted(strains, key=lambda strain: -strains[strain]["corrected_bgcs"])
    rank = {strain: number for number, strain in enumerate(strain_order, 1)}

    strain_classes = defaultdict(set)
    class_bgc_count = Counter()
    for row in bgcs:
        for product in (row["products"] or "").split(";"):
            product = product.strip()
            if product:
                strain_classes[row["sid"]].add(product)
                class_bgc_count[product] += 1
    prevalence = Counter()
    for classes in strain_classes.values():
        for product in classes:
            prevalence[product] += 1

    def band(count):
        if count == denominator:
            return "CORE"
        if count >= denominator * 0.5:
            return "COMMON"
        if count == 1:
            return "UNIQUE"
        return "ACCESSORY"

    zero_kcb_total = sum(1 for row in bgcs if not row["kcb_top"])
    context = {}
    for strain in strain_order:
        classes = strain_classes[strain]
        shared = sum(1 for product in classes if prevalence[product] > 1)
        unique = sum(1 for product in classes if prevalence[product] == 1)
        zero_kcb = sum(1 for row in bgcs if row["sid"] == strain and not row["kcb_top"])
        diagnostics = list(tigrfam.get(strain, {}).get("present", {}))
        context[strain] = {
            "rank": rank[strain],
            "tier": tier(strain),
            "n_classes": len(classes),
            "shared": shared,
            "unique": unique,
            "zero_kcb": zero_kcb,
            "zk_share": f"{zero_kcb} of {zero_kcb_total}",
            "diagnostics": ",".join(sorted(diagnostics)) if diagnostics else "none",
        }

    diagnostic_strains = defaultdict(list)
    for strain in strain_order:
        for accession in tigrfam.get(strain, {}).get("present", {}):
            diagnostic_strains[accession].append(strain)

    enediyne_leads = []
    for strain in diagnostic_strains.get("TIGR03828", []):
        candidates = [row for row in bgcs if row["sid"] == strain and "PKS" in (row["products"] or "")]
        if candidates:
            row = sorted(candidates, key=lambda item: -(item.get("kcb_cumulative") or 0))[0]
            identity = f"{strain} / {row['contig']} / {row['region']} / {row['bgc_id']}"
            enediyne_leads.append(f"{identity} [{row['products']}] edge={row['edge_status']}")

    workbook = openpyxl.load_workbook(workbook_path)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")

    for canonical in GENERATED_SHEETS:
        for existing in list(workbook.sheetnames):
            if _generated_sheet_name(existing, canonical):
                del workbook[existing]

    def style_header(sheet, columns):
        for column in range(1, columns + 1):
            cell = sheet.cell(1, column)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"

    sheet = workbook.create_sheet("Cross_Strain_Class_Prevalence")
    sheet.append(["product_class", "n_strains", "pct_strains", "n_BGCs", "band", "informative_for_comparison"])
    for product, count in sorted(prevalence.items(), key=lambda item: (-item[1], -class_bgc_count[item[0]])):
        sheet.append([
            product,
            count,
            round(100 * count / denominator, 1),
            class_bgc_count[product],
            band(count),
            "NO (universal)" if product in UNIVERSAL else "yes",
        ])
    style_header(sheet, 6)
    for column, width in zip("ABCDEF", [28, 10, 12, 9, 12, 26]):
        sheet.column_dimensions[column].width = width

    sheet = workbook.create_sheet("Cross_Strain_Findings")
    sheet.append(["#", "finding", "metric", "value", "claim_status", "note"])
    tier_counts = Counter(tier(strain) for strain in strain_order)
    raw_total = sum(strains[strain].get("raw_bgcs", 0) for strain in strain_order)
    corrected_total = round(sum(strains[strain].get("corrected_bgcs", 0) for strain in strain_order), 1)
    universal_present = sorted(product for product, count in prevalence.items() if count == denominator and product in UNIVERSAL)
    hgle_classes = sorted(product for product in prevalence if "hgl" in product.lower())
    hgle_summary = "; ".join(f"{product}={prevalence[product]}/{denominator}" for product in hgle_classes) or "none detected"
    napaa_count = prevalence.get("NAPAA", 0)
    zero_kcb_pct = round(100 * zero_kcb_total / len(bgcs)) if bgcs else 0
    findings = [
        (1, "Assembly-quality distribution", "tiers (GOOD/MOD/POOR)", f"{tier_counts.get('GOOD', 0)}/{tier_counts.get('MOD', 0)}/{tier_counts.get('POOR', 0)}", "GROUNDED", "computed from banked n50 thresholds; interpret fragmentation-sensitive comparisons cautiously"),
        (2, "BGC totals", "raw -> corrected", f"{raw_total} -> {corrected_total}", "GROUNDED", "corrected count = Interior + 0.5*Edge + 0.25*Full-contig"),
        (3, "Universal non-discriminating classes", f"universal classes ({denominator}/{denominator})", ",".join(universal_present) if universal_present else "none", "GROUNDED", "classes universal in this cohort are not suitable for shared/novelty claims"),
        (4, "Core ecological signal", "ectoine / NI-siderophore", f"{prevalence.get('ectoine', 0)}/{denominator} / {prevalence.get('NI-siderophore', 0)}/{denominator}", "GROUNDED, presence-only", "presence-only genome-mining signal; not an assay of ecological function"),
        (5, "hglE-like class prevalence", "computed class names containing hgl", hgle_summary, "GROUNDED, collection-specific", "computed from the current class-prevalence table; no hard-coded strain list"),
        (6, "NAPAA exclusion basis", "strains with NAPAA (this set)", f"{napaa_count}/{denominator}", "GROUNDED, collection-specific", "exclude from headline ecological/comparative claims when treated as housekeeping-adjacent/non-discriminating"),
        (7, "Novelty floor", "fully KCB-dark BGCs", f"{zero_kcb_total}/{len(bgcs)} ({zero_kcb_pct}%)", "GROUNDED floor", "KCB-dark is a lower-bound novelty signal and can be inflated by fragmented assemblies"),
        (8, "KCB field caveat", "kcb_cumulative", "cumulative SCORE, not % identity", "METHOD CAVEAT", "do not label KCB score as percent identity or product identity"),
        (9, "Thiopeptide diagnostic", "TIGR03604 strains", f"{len(diagnostic_strains.get('TIGR03604', []))}/{denominator}", "GROUNDED", "diagnostic panel presence only; validate exact regions before product claims"),
        (10, "Enediyne diagnostic", "TIGR03828 strains", f"{len(diagnostic_strains.get('TIGR03828', []))}/{denominator}", "PRIORITY/VERIFY", "warhead-scaffold candidates require exact-region and accessory-gene review before assertion"),
        (11, "Ansamycin diagnostic", "TIGR01454 strains", f"{len(diagnostic_strains.get('TIGR01454', []))}/{denominator}", "GROUNDED", "diagnostic panel presence only"),
        (12, "RG-GMCI HIGH caveat", "HIGH-pair count", "not a monotonic fragmentation metric", "METHOD CAVEAT", "reflects reference-anchored homologous-pair richness and BGC content, not physical contig joining"),
    ]
    for row in findings:
        sheet.append(list(row))
    sheet.append([])
    sheet.append(["", "enediyne candidate regions (locators):", "", "", "", ""])
    for locator in enediyne_leads:
        sheet.append(["", "", locator, "", "PRIORITY/verify", ""])
    style_header(sheet, 6)
    for column, width in zip("ABCDEF", [4, 30, 26, 30, 20, 60]):
        sheet.column_dimensions[column].width = width
    for row in range(2, sheet.max_row + 1):
        for column in range(1, 7):
            sheet.cell(row, column).alignment = Alignment(wrap_text=True, vertical="top")

    sheet = workbook.create_sheet("Strain_Cohort_Context")
    sheet.append(["strain", "corrected_rank", "tier", "n_classes", "shared_classes", "unique_classes", "zero_KCB_regions", "zero_KCB_share_of_cohort", "diagnostics_present"])
    for strain in strain_order:
        row = context[strain]
        sheet.append([strain, row["rank"], row["tier"], row["n_classes"], row["shared"], row["unique"], row["zero_kcb"], row["zk_share"], row["diagnostics"]])
    style_header(sheet, 9)
    for column, width in zip("ABCDEFGHI", [10, 14, 6, 11, 14, 14, 16, 22, 30]):
        sheet.column_dimensions[column].width = width

    required = {"A2_Strain_Registry", "B4_Cross_Strain_Scans", "D1_RGGMCI_All_Strains", "C1_DAPR_Antibacterial", "C2_DAPR_Antifungal"}
    missing = sorted(required - set(workbook.sheetnames))
    atomic_save(workbook, workbook_path)
    workbook.close()
    atomic_dump_json(context, os.path.join(banked_dir, "cohort_context.json"), indent=0)

    emit(f"added 3 cross-strain sheets | total generated sheets {len(GENERATED_SHEETS)}")
    emit("class bands:", dict(Counter(band(count) for count in prevalence.values())))
    emit("cohort context written for", len(context), "strains")
    if missing:
        emit("auto: DAPR rescue refresh SKIPPED — workbook lacks cohort/DAPR sheets: " + ",".join(missing))
    else:
        subprocess.run([sys.executable, os.path.join(os.path.dirname(__file__), "build_dapr_rescue_sheets.py"), workbook_path], check=True)
        emit("auto: D5 Fragment_Rescue_Tiers + DAPR Activity_Ref regenerated")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
