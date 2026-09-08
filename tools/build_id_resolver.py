#!/usr/bin/env python3
"""build_id_resolver.py — emit the BGC ID-resolver table for a banked cohort.

    BGC_ID | bgc_uid | contig/NODE | region | antiSMASH file | workbook row

Pass --workbook for authoritative bgc_uid + row from the BGC_Master sheet; otherwise uids are derived
and the row column is left blank.
"""
from __future__ import annotations

import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402
import argparse, csv, json, os, sys
try:  # v9.7.410 CSV formula-cell guard (CLAUDE_v9.7.410_tools_csv_writer_coverage)
    from mamey.csv_safety import SafeDictWriter as _SafeDictWriter, SafeWriter as _SafeWriter
except ImportError:  # bare-script run: bundle root is one level up
    import os as _cs_os, sys as _cs_sys
    _cs_sys.path.insert(0, _cs_os.path.dirname(_cs_os.path.dirname(_cs_os.path.abspath(__file__))))
    from mamey.csv_safety import SafeDictWriter as _SafeDictWriter, SafeWriter as _SafeWriter


def _read_json(_path, *, encoding="utf-8"):
    """P3b: context-managed JSON read; closes the handle a bare open() leaked."""
    import json as _json
    with open(_path, encoding=encoding) as _fh:
        return _json.load(_fh)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _wbio import atomic_write_text
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from mamey.id_resolver import resolver_rows, COLUMNS

def _workbook_map(path):
    wm = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = next((wb[n] for n in ("BGC_Master", "C1_BGC_Master", "BGC_Inventory") if n in wb.sheetnames), None)
        if ws is None: return wm
        rows = list(ws.iter_rows(values_only=True)); hdr = [str(c) for c in rows[0]]; ix = {c: i for i, c in enumerate(hdr)}
        sidc = next((ix[c] for c in ("strain", "sid", "cohort_id") if c in ix), None)
        for ri, r in enumerate(rows[1:], start=2):
            bid = r[ix["bgc_id"]] if "bgc_id" in ix else None
            if bid:
                wm[(r[sidc] if sidc is not None else None, bid)] = {
                    "bgc_uid": r[ix["bgc_uid"]] if "bgc_uid" in ix else None, "row": ri}
    except Exception as e:
        sys.stderr.write(f"  [warn] workbook read failed ({e}); using derived uids\n")
    return wm

def main(argv=None):
    ap = argparse.ArgumentParser(description="Emit BGC ID-resolver table for a banked cohort.")
    ap.add_argument("--banked-dir", default="cohort"); ap.add_argument("--workbook", default=None); ap.add_argument("--out", default=None)
    a = ap.parse_args(argv)
    bgcs = _read_json(os.path.join(a.banked_dir, "bgc_data.json"))["bgcs"]
    rows = resolver_rows(bgcs, _workbook_map(a.workbook) if a.workbook else None)
    # Build the CSV in memory, then write atomically (file) or stream to stdout. Building first means
    # a file output is never left half-written if the process dies mid-loop (v9.7.115).
    import io as _io
    _buf = _io.StringIO()
    w = _SafeDictWriter(_buf, fieldnames=COLUMNS); w.writeheader()
    for r in rows: w.writerow(r)
    if a.out:
        atomic_write_text(a.out, _buf.getvalue())
        emit(f"  wrote {len(rows)} resolver rows -> {a.out}")
    else:
        sys.stdout.write(_buf.getvalue())
    return 0

if __name__ == "__main__": raise SystemExit(main())
