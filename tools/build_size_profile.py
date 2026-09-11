#!/usr/bin/env python3
"""build_size_profile.py — per-strain BGC nucleotide content by class (Mamey standard module).

Counting BGCs inflates with assembly fragmentation (a split megasynthase is counted several times); summing the
nucleotide content (kb) of each strain's BGCs by class is far more fragmentation-robust, because a split cluster
still contributes ~its true length. This module makes size-by-type a first-class Mamey output: it writes a
Size_By_Type workbook sheet and emits the size figures, so every cohort (incl. the AS merge) gets it.

Honest caveat carried in the outputs: size is biased slightly DOWN by edge truncation (clusters running off a
contig end lose nucleotides), the opposite direction to count inflation — so count (upper) and size (lower)
bracket the truth. See docs/BGC_size_by_type.md.

Outputs (to --out-dir): fig_bgc_size_by_type.png (+ _data.csv), fig_count_vs_size_fragmentation.png,
fig_strain_class_size_heatmap.png. Adds the Size_By_Type sheet to --workbook when given.

Usage:
  python tools/build_size_profile.py --banked-dir cohort --workbook Master.xlsx --out-dir figures
  python tools/build_size_profile.py --out-dir figures --replot   # re-draw from fig_bgc_size_by_type_data.csv
"""
import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402
import argparse, os, json, csv, math, statistics, sys
try:  # v9.7.410 CSV formula-cell guard (CLAUDE_v9.7.410_tools_csv_writer_coverage)
    from mamey.csv_safety import SafeDictWriter as _SafeDictWriter, SafeWriter as _SafeWriter
except ImportError:  # bare-script run: bundle root is one level up
    import os as _cs_os, sys as _cs_sys
    _cs_sys.path.insert(0, _cs_os.path.dirname(_cs_os.path.dirname(_cs_os.path.abspath(__file__))))
    from mamey.csv_safety import SafeDictWriter as _SafeDictWriter, SafeWriter as _SafeWriter


def _read_json(_path, *, encoding="utf-8"):
    """P3b: context-managed JSON read; closes the handle a bare open() leaked."""
    import json as _json
    with open(_path, encoding=encoding) as _fh:
        return _json.load(_fh)

_TOOL_DIR = os.path.dirname(os.path.abspath(__file__))
_REPO_ROOT = os.path.dirname(_TOOL_DIR)
for _p in (_TOOL_DIR, _REPO_ROOT):
    if _p not in sys.path:
        sys.path.insert(0, _p)
from collections import defaultdict, Counter
from _wbio import atomic_save, atomic_open

SACC={'saccharide','oligosaccharide','amglyccycl','aminoglycoside'}
BACKBONE={'t1pks','t2pks','t3pks','transat-pks','pks','pks-like','hr-t2pks','nrps','nrps-like','terpene','ripp',
 'ripp-like','lanthipeptide','lassopeptide','thiopeptide','siderophore','ectoine','betalactone','indole',
 'butyrolactone','phosphonate','nucleoside','aminocoumarin','melanin','arylpolyene','redox-cofactor','cdps',
 'other','fatty_acid','napaa'}
_PREF=['transat-pks','t1pks','hr-t2pks','t2pks','nrps','nrps-like','pks','terpene','ripp','lanthipeptide',
 'lassopeptide','thiopeptide','siderophore','ectoine','phosphonate','ni-siderophore','butyrolactone']
_NORM={'nrps-like':'NRPS','t1pks':'T1PKS','t2pks':'T2PKS','hr-t2pks':'HR-T2PKS','transat-pks':'transAT',
 'pks':'PKS','pks-like':'PKS','ni-siderophore':'Siderophore'}

def _classes(b): return [p.strip().lower() for p in (b.get('products') or '').split(';') if p.strip()]
def _primary(b):
    cl=_classes(b)
    if not cl: return 'unclassified'
    bb=[c for c in cl if c in BACKBONE and c not in ('other','fatty_acid')]
    if bb:
        for pref in _PREF:
            if pref in cl: return pref
        return bb[0]
    if any(c in SACC for c in cl): return 'saccharide-only'
    return cl[0]
def _norm(c): return _NORM.get(c, c.upper() if len(c)<=5 else c.capitalize())
def _pearson(x,y):
    if len(x)<3: return 0.0
    mx=statistics.mean(x); my=statistics.mean(y); cov=sum((a-mx)*(b-my) for a,b in zip(x,y))
    sx=math.sqrt(sum((a-mx)**2 for a in x)); sy=math.sqrt(sum((b-my)**2 for b in y)); return cov/(sx*sy) if sx and sy else 0.0

def compute(banked_dir, include_saccharide=False):
    bgc=_read_json(os.path.join(banked_dir,'bgc_data.json'))
    from mamey.figure_policy import omit_saccharides
    strains=bgc['strains']; bgcs=omit_saccharides(bgc['bgcs'])  # figure policy: pure-saccharide omitted
    siz=defaultdict(Counter); cnt=defaultdict(Counter)
    for b in bgcs:
        p=_primary(b)
        if p=='saccharide-only' and not include_saccharide: continue
        c=_norm(p); siz[b['sid']][c]+=(b.get('length_kb') or 0); cnt[b['sid']][c]+=1
    tot=Counter()
    for s in siz:
        for c,v in siz[s].items(): tot[c]+=v
    cols=[c for c,_ in tot.most_common(8)]
    sids=sorted(strains, key=lambda s: strains[s].get('contigs',0) or 0)
    return strains, sids, cols, siz, cnt

def write_csv(path, strains, sids, cols, siz):
    with atomic_open(path,'w',newline='') as f:
        w=_SafeWriter(f); w.writerow(['strain','contigs']+cols)
        for s in sids: w.writerow([s, strains[s].get('contigs','')]+[round(siz[s][c],1) for c in cols])

def write_sheet(workbook, strains, sids, cols, siz, cnt):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb=openpyxl.load_workbook(workbook)
    if 'Size_By_Type' in wb.sheetnames: del wb['Size_By_Type']
    ws=wb.create_sheet('Size_By_Type')
    hdr=['strain','contigs']+[f'{c}_kb' for c in cols]+['total_kb']+[f'{c}_n' for c in cols]
    ws.append(hdr)
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='2E5E7E')
    for s in sids:
        row=[s, strains[s].get('contigs','')]+[round(siz[s][c],1) for c in cols]
        row+=[round(sum(siz[s].values()),1)]+[cnt[s][c] for c in cols]
        ws.append(row)
    ws.freeze_panes='C2'
    atomic_save(wb, workbook)

def plot_all(out_dir, strains, sids, cols, siz, cnt):
    import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt, numpy as np
    PAL=['#2E86AB','#E59866','#7FB069','#B07AA1','#C0392B','#F1C40F','#16A085','#7F8C8D']
    lc=[math.log10(max(strains[s].get('contigs',1) or 1,1)) for s in sids]
    # 1) stacked size-by-type
    fig,ax=plt.subplots(figsize=(13,5)); x=np.arange(len(sids)); bottom=np.zeros(len(sids))
    for i,c in enumerate(cols):
        v=np.array([siz[s][c] for s in sids]); ax.bar(x,v,bottom=bottom,width=0.84,color=PAL[i%len(PAL)],label=c); bottom+=v
    ax.set_xticks(x); ax.set_xticklabels(sids,rotation=90,fontsize=5); ax.set_xlim(-0.6,len(sids)-0.4)
    ax.set_ylabel('total BGC size (kb)'); ax.spines[['top','right']].set_visible(False)
    ax.set_title('Total BGC nucleotide content per strain by type\nstrains ordered by contig count (most contiguous at left)',fontsize=11,weight='bold')
    ax.legend(ncol=8,fontsize=7.5,frameon=False,loc='upper center',bbox_to_anchor=(0.5,-0.18))
    plt.tight_layout(); plt.savefig(os.path.join(out_dir,'fig_bgc_size_by_type.png'),dpi=160,bbox_inches='tight'); plt.close()
    # 2) count vs size fragmentation
    rows=[(c,_pearson(lc,[cnt[s][c] for s in sids]),_pearson(lc,[siz[s][c] for s in sids])) for c in cols]; rows.sort(key=lambda r:r[1])
    fig,(a1,a2)=plt.subplots(1,2,figsize=(11,4.6)); y=np.arange(len(rows)); h=0.38
    a1.barh(y+h/2,[r[1] for r in rows],h,color='#C0392B',label='count vs fragmentation')
    a1.barh(y-h/2,[r[2] for r in rows],h,color='#2E86AB',label='size (kb) vs fragmentation')
    a1.set_yticks(y); a1.set_yticklabels([r[0] for r in rows],fontsize=9); a1.axvline(0,color='#444',lw=0.9)
    a1.set_xlabel('correlation with fragmentation (log10 contigs)'); a1.spines[['top','right']].set_visible(False)
    a1.set_title('Count inflates (+), size deflates (−):\nopposite biases bracket the truth',fontsize=10,weight='bold'); a1.legend(fontsize=8,frameon=False,loc='lower right')
    tc=[sum(cnt[s].values()) for s in sids]; ts=[sum(siz[s].values()) for s in sids]; a2b=a2.twinx()
    a2.scatter(lc,tc,s=22,c='#C0392B',alpha=0.7); a2b.scatter(lc,ts,s=22,c='#2E86AB',alpha=0.7,marker='s')
    for v,c,axx in [(tc,'#C0392B',a2),(ts,'#2E86AB',a2b)]:
        m,b=np.polyfit(lc,v,1); xs=np.array([min(lc),max(lc)]); axx.plot(xs,m*xs+b,c=c,lw=1.2)
    a2.set_xlabel('log10 contig count'); a2.set_ylabel(f'BGC count (r={_pearson(lc,tc):+.2f})',color='#C0392B'); a2b.set_ylabel(f'BGC size kb (r={_pearson(lc,ts):+.2f})',color='#2E86AB')
    a2.set_title('Whole-genome: count vs size vs fragmentation',fontsize=10,weight='bold'); a2.spines[['top']].set_visible(False); a2b.spines[['top']].set_visible(False)
    plt.tight_layout(); plt.savefig(os.path.join(out_dir,'fig_count_vs_size_fragmentation.png'),dpi=160); plt.close()
    # 3) size heatmap (strain x class)
    M=np.array([[siz[s][c] for c in cols] for s in sids],float)
    fig,ax=plt.subplots(figsize=(8.5,11)); im=ax.imshow(M,aspect='auto',cmap='YlGnBu')
    ax.set_xticks(range(len(cols))); ax.set_xticklabels(cols,rotation=45,ha='right',fontsize=9)
    ax.set_yticks(range(len(sids))); ax.set_yticklabels(sids,fontsize=5.5)
    for i in range(len(sids)):
        for j in range(len(cols)):
            if M[i,j]>0: ax.text(j,i,int(M[i,j]),ha='center',va='center',fontsize=4.5,color='black' if M[i,j]<M.max()*0.6 else 'white')
    ax.set_title('Per-strain BGC size (kb) by class\nstrains ordered by contig count (most contiguous at top)',fontsize=10,weight='bold')
    plt.colorbar(im,ax=ax,shrink=0.4,label='BGC size (kb)')
    plt.tight_layout(); plt.savefig(os.path.join(out_dir,'fig_strain_class_size_heatmap.png'),dpi=160); plt.close()

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--banked-dir', default='cohort'); ap.add_argument('--workbook', default=None)
    ap.add_argument('--out-dir', default='figures'); ap.add_argument('--include-saccharide', action='store_true')
    ap.add_argument('--replot', action='store_true')
    a=ap.parse_args(); os.makedirs(a.out_dir, exist_ok=True)
    csv_path=os.path.join(a.out_dir,'fig_bgc_size_by_type_data.csv')
    if a.replot:
        rows=list(csv.reader(open(csv_path))); cols=rows[0][2:]
        strains={r[0]:{'contigs':int(r[1]) if r[1] else 0} for r in rows[1:]}
        sids=[r[0] for r in rows[1:]]
        siz=defaultdict(Counter); cnt=defaultdict(Counter)
        for r in rows[1:]:
            for j,c in enumerate(cols): siz[r[0]][c]=float(r[2+j])
        plot_all(a.out_dir, strains, sids, cols, siz, cnt); emit('  replotted from CSV (count panels need full data)'); return
    strains, sids, cols, siz, cnt = compute(a.banked_dir, a.include_saccharide)
    write_csv(csv_path, strains, sids, cols, siz)
    plot_all(a.out_dir, strains, sids, cols, siz, cnt)
    if a.workbook and os.path.exists(a.workbook):
        write_sheet(a.workbook, strains, sids, cols, siz, cnt); emit(f'  Size_By_Type sheet added to {os.path.basename(a.workbook)}')
    emit(f'  size-by-type: {len(sids)} strains x {len(cols)} classes; figures + CSV in {a.out_dir}')

if __name__=='__main__': main()
