#!/usr/bin/env python3
"""cohort_scoring_version_gate.py — fail-closed: a comparative build may only aggregate strains
scored under ONE engine version.

STATUS (v9.7.374 doc-fix): this module docstring previously read "DRAFT for v9.7.99 ...
intentionally NOT registered yet" — stale since v9.7.338. It has in fact been reclassified WIRED
since .338 (COH-01) and is hooked into `mamey.cohort_deliverable._check_engine_uniformity`, which
calls `assert_uniform_scoring_engine(...)` before any cross-strain aggregation in
`run_cohort_deliverable()`. `tools/gate_registry.tsv` correctly lists it WIRED, and
`tests/test_gate_wiring_invariant.py::_UNIT_TESTED_OPERATOR_GATES` already documents the .338
reclassification in a comment — only this file's own docstring had not caught up. See
`tests/test_cohort_scoring_version_gate.py` and `tests/test_cohort_deliverable.py` for the
regression coverage.

Why (COHORT_RESCORING_PLAN.md): scoring boundaries stacked 1.9.84→1.9.96. Cross-strain AB/AF
comparison is only valid when every aggregated strain was scored under the same engine. Before
this gate existed, the comparative builds (hub_merge, build_cross_strain_figures, lead_board,
cohort_concordance_summary, build_master*, merge_workbooks) aggregated strains with NO version
check — so a mixed-version cohort silently produced an invalid comparison. This converts "remember
to re-score" into a gate, the same fail-closed pattern as the leak audit and tier-derivation parity
gate.

Per-strain provenance already exists: the master workbook's A3_Run_Manifest sheet records
"version" = f"Mamey v{engine}" (mamey/master_workbook.py). The gate reads that field per strain.
"""
from __future__ import annotations

import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402
import argparse, re, sys
from pathlib import Path


class CohortVersionError(RuntimeError):
    """Raised when a comparative cohort spans more than one scoring engine version."""


_VER = re.compile(r"(\d+\.\d+\.\d+)")


def parse_engine_version(version_field: str) -> str | None:
    """'Mamey v1.9.96' -> '1.9.96'; tolerant of surrounding text. None if no version found."""  # version-sync-ok
    m = _VER.search(version_field or "")
    return m.group(1) if m else None


def assert_uniform_scoring_engine(strain_versions: dict[str, str], current_engine: str) -> None:
    """Raise CohortVersionError unless every strain's parsed engine == current_engine.

    strain_versions: {strain_id: raw version field, e.g. 'Mamey v1.9.96'}. Missing/unparseable  # version-sync-ok
    versions are themselves a failure (a strain with no recorded provenance can't be certified).
    """
    offenders, unstamped = {}, []
    # v9.7.115: an empty cohort is a FAILURE, not a vacuous pass. A fail-closed gate that says "OK"
    # when it found zero strains to certify almost always means a bad glob/path, not a clean cohort —
    # and a comparative build on zero strains is meaningless. Refuse rather than wave it through.
    if not strain_versions:
        raise CohortVersionError(
            f"no strains found to certify at engine {current_engine} — the comparative cohort is "
            "empty (likely a wrong workbooks path). A fail-closed gate does not pass on zero input."
        )
    for strain, raw in strain_versions.items():
        v = parse_engine_version(raw)
        if v is None:
            unstamped.append(strain)
        elif v != current_engine:
            offenders[strain] = v
    if unstamped or offenders:
        parts = []
        if offenders:
            parts.append("scored under a different engine: "
                         + ", ".join(f"{s}={v}" for s, v in sorted(offenders.items())))
        if unstamped:
            parts.append("no recorded engine version: " + ", ".join(sorted(unstamped)))
        raise CohortVersionError(
            f"cohort is not uniformly scored at engine {current_engine} — " + "; ".join(parts)
            + ". Re-score these strains under the current engine before building a comparative deliverable."
        )


def _read_workbook_version(path: Path) -> str:
    """Read A3_Run_Manifest 'version' from a strain master workbook (best-effort, read-only)."""
    import openpyxl  # local import: comparative builds already depend on openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in ("A3_Run_Manifest", "Run_Manifest"):
        if sheet in wb.sheetnames:
            ws = wb[sheet]
            header = [c.value for c in next(ws.iter_rows(max_row=1))]
            if "version" in header:
                col = header.index("version")
                row = next(ws.iter_rows(min_row=2, values_only=True), None)
                return str(row[col]) if row and col < len(row) else ""
    return ""


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbooks-dir", required=True, help="dir of per-strain master workbooks (*.xlsx)")
    ap.add_argument("--engine", required=True, help="current engine version, e.g. 1.9.96")
    a = ap.parse_args()
    strain_versions = {}
    for wbp in sorted(Path(a.workbooks_dir).glob("*.xlsx")):
        strain_versions[wbp.stem] = _read_workbook_version(wbp)
    try:
        assert_uniform_scoring_engine(strain_versions, a.engine)
    except CohortVersionError as e:
        emit(f"FAIL: {e}", file=sys.stderr)
        return 1
    emit(f"OK: all {len(strain_versions)} strains scored at engine {a.engine}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
