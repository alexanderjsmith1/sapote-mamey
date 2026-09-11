#!/usr/bin/env python3
"""build_chitinase_screen.py — add the Chitinase_Screen sheet to the master workbook.

Ecological screen (whole-genome GH18/GH19 + chitin-binding genes) is a default cohort analysis for
this insect-associated collection: chitinolytic load is the degradative arm of antifungal defensive
symbiosis. Reads cohort/chitinase.json (banked per-strain counts) and writes a sorted sheet with the
genus and host annotation. Counts are annotation-derived (GH18/GH19 family + chitin-binding/LPMO);
exact GH-family partition needs an HMM scan.

Usage: python tools/build_chitinase_screen.py --workbook <xlsx> --banked-dir <dir>
"""
import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402
import argparse, json, os


from _wbio import atomic_save


def _read_json(_path, *, encoding="utf-8"):
    """P3b: context-managed JSON read; closes the handle a bare open() leaked."""
    import json as _json
    with open(_path, encoding=encoding) as _fh:
        return _json.load(_fh)

_PLACEHOLDER={'not','unknown','unclassified','uncultured','unidentified','na','n/a','none',''}
def _genus(org):
    """First token of the organism string, but robust to placeholder taxonomies like 'not verified'
    (which naive split()[0] turns into 'not'). Returns 'unclassified' for placeholders."""
    org=(org or '').strip()
    if not org: return 'unclassified'
    first=org.split()[0]
    if first.lower() in _PLACEHOLDER: return 'unclassified'
    if first.lower()=='candidatus' and len(org.split())>1: return org.split()[1]
    return first

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--workbook', required=True); ap.add_argument('--banked-dir', default='cohort')
    a=ap.parse_args()
    cpath=os.path.join(a.banked_dir,'chitinase.json')
    if not os.path.exists(cpath):
        emit('  [chitinase] no chitinase.json in banked dir — skipping (ecological screen optional)'); return
    chit=_read_json(cpath)
    bgc=_read_json(os.path.join(a.banked_dir,'bgc_data.json'))
    strains=bgc['strains']; bgcs=bgc['bgcs']
    # fragmentation-robust normalizer: ectoine + NAPAA are compact, single-copy, ubiquitous BGCs that
    # are NOT inflated by assembly fragmentation (unlike raw/NRPS BGC counts). chitinase/(ecto+napaa)
    # gives a fragmentation-robust ratio (|corr with contig count| ~0.2 vs ~0.4 for raw chitinase) that
    # surfaces genuine chitinolytic outliers rather than assembly artifacts. See docs/HOW_TO_USE.md.
    from collections import Counter
    ecto=Counter(); napaa=Counter()
    for b in bgcs:
        p=(b.get('products') or '').lower()
        if 'ectoine' in p: ecto[b['sid']]+=1
        if 'napaa' in p: napaa[b['sid']]+=1
    import statistics
    def den_of(sid): return ecto.get(sid,0)+napaa.get(sid,0)
    # score the normalized ratio ONLY over strains that have a real ecto+napaa reference; strains
    # missing it (reference dropped by fragmentation) are flagged 'no_norm_ref', not scored as outliers.
    scored={s: chit[s].get('chitinase',0)/den_of(s) for s in chit if den_of(s)>=1}
    mu=statistics.mean(scored.values()) if scored else 0
    sd=statistics.pstdev(scored.values()) or 1
    def norm(sid):
        d=den_of(sid)
        return round(chit[sid].get('chitinase',0)/d,2) if d>=1 else None
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb=openpyxl.load_workbook(a.workbook)
    if 'Chitinase_Screen' in wb.sheetnames: del wb['Chitinase_Screen']
    ws=wb.create_sheet('Chitinase_Screen')
    hdr=['strain','genus','host','chitinases_GH18_GH19','chitin_binding_CBM_LPMO','total_chitinolytic',
         'ecto+napaa (norm ref)','chit_per_unit','z_score','outlier_flag']
    ws.append(hdr)
    for c in range(1,len(hdr)+1):
        cell=ws.cell(1,c); cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='2E5E7E')
    rows=[]
    for sid,c in chit.items():
        genus=_genus(strains.get(sid,{}).get('organism','')) if sid in strains else 'unclassified'
        den=den_of(sid); nv=norm(sid)
        if nv is None:
            z=None; flag='no_norm_ref (ref dropped)'
        else:
            z=(nv-mu)/sd
            flag='HIGH chitinolytic' if z>1.3 else ('LOW chitinolytic' if z<-1.3 else '')
        rows.append([sid,genus,c.get('host',''),c.get('chitinase',0),c.get('chitin_binding',0),
                     c.get('chitinase',0)+c.get('chitin_binding',0), den,
                     nv if nv is not None else 'n/a', round(z,2) if z is not None else 'n/a', flag])
    # v9.7.371 fix: was -(r[8] if isinstance(r[8],(int,float)) else -1) -- a "no_norm_ref"
    # row (z_score stored as the string 'n/a') got sort key -(-1)=1, landing it in the SAME
    # sort-order neighborhood as a genuinely-measured z=-1.0 (key=1) or z=-1.3 LOW-threshold row
    # (key=1.3) -- "never evaluated" silently interleaved with real low-outlier strains instead of
    # being visually distinct. float('inf') guarantees a no-reference row always sorts strictly
    # after every real z-score, regardless of how extreme that real score is.
    rows.sort(key=lambda r:(-r[8] if isinstance(r[8],(int,float)) else float('inf')))
    LOW=PatternFill('solid',fgColor='FDE7E7'); HI=PatternFill('solid',fgColor='E6F4EA')
    for r in rows:
        ws.append(r)
        if 'HIGH' in r[9]:
            for c in range(1,len(hdr)+1): ws.cell(ws.max_row,c).fill=HI
        elif 'LOW' in r[9] or r[3]==0:
            for c in range(1,len(hdr)+1): ws.cell(ws.max_row,c).fill=LOW
    for col,w in zip('ABCDEFGHIJ',[10,14,22,22,24,16,18,14,10,18]): ws.column_dimensions[col].width=w
    ws.freeze_panes='A2'
    vals=[r[3] for r in rows]; strep=[r[3] for r in rows if r[1]=='Streptomyces']
    ws.append([])
    # v9.7.374 fix (AUDIT_374): the printed formula claimed a "+0.5" smoothing term that
    # does not exist anywhere in the real computation -- den_of()/norm()/scored above are exactly
    # chit/(ecto+napaa), gated on den_of(sid)>=1 (rows below that threshold get 'no_norm_ref'
    # instead of a smoothed score). A reader trying to reproduce this sheet's chit_per_unit column
    # from raw chitinase.json using the printed "+0.5" formula would get different numbers than
    # what is actually in the sheet. Label corrected to describe the real, applied formula.
    ws.append(['SUMMARY',f'{len(rows)} strains',f'mean chit {sum(vals)/len(vals):.1f}',
               '', '', '', 'norm = chit/(ecto+napaa), scored only when ecto+napaa>=1',f'frag-robust (|r|~0.2 vs 0.4 raw)','',''])
    atomic_save(wb, a.workbook)
    nhi=sum(1 for r in rows if 'HIGH' in r[9]); nlo=sum(1 for r in rows if 'LOW' in r[9] or r[3]==0)
    emit(f"  Chitinase_Screen sheet written: {len(rows)} strains | normalized chit/(ecto+napaa) | {nhi} HIGH, {nlo} LOW outliers")

if __name__=='__main__': main()
