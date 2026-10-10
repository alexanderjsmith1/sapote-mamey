#!/usr/bin/env python3
"""Per-strain slide decks: every antiSMASH region of one strain, with every evidence channel held for it.

Usage:
  python tools/strain_slides.py build --sources <strain>_sources.json --out <folder> [--tag v2]
  python tools/strain_slides.py template > sources.json        # an empty sources file to fill

The sources file names every input; nothing is found by guessing paths. Only `strain` and `package` are required.
Every other channel is optional, and a slide says so when its channel is missing.

What a deck holds, in order:
1. overview: isolation and genome as recorded (metadata given in the sources file), nearest public genomes;
2. trees: the approved panels listed in the sources file (never a new tree);
3. BGC landscape: antiSMASH product types and how many regions each channel supports;
4. region tables: one row per antiSMASH region with every channel and its slide number;
5. a BiG-SCAPE figure, when one is given;
6. one slide per antiSMASH region, strongest evidence first:
   - a centred gap-rescue map (drawn by tools/gap_directed_rescue.redraw; the layout it used is saved next to it);
   - a gene strip in the map's own orientation: arrows by antiSMASH role, the genes matched to MIBiG in a band,
     GECCO's per-gene probability above them, and domain labels for core, matched and high-probability genes;
     for a gap rescue, every partner contig the map drew is added beside the core ("//" between contigs), each with
     its own GECCO bars, orientation and Pfam names, cut to its matched genes +/- 2.5 kb;
   - an "assembled across contigs" block: the reference genes found, grouped by biosynthetic step, each with the
     contig and antiSMASH region it sits in, its identity, the partner check's verdict, and the best KnownClusterBlast
     reference for that gene when it differs from the map's reference;
   - one line per channel: antiSMASH, core genes, KnownClusterBlast, per-gene MIBiG, gap rescue and its partner
     verdicts, RG-GMCI with the two-proof verdict, BiG-SCAPE, stored BLASTp, GECCO, and a link verdict when a rescue
     verdict table is given;
   - family figures listed for that region in the sources file;
7. GECCO-only candidates: GECCO clusters that overlap no antiSMASH region;
8. extra figures (for example protein PCoA panels from tools/strain_slides_pcoa.py);
9. where the rest lives (paths listed in the sources file).
A second file, <strain>_gene_tables_<tag>_<date>.pptx, lists every gene of every region: role, Pfam/TIGRFAM/rule
domains with descriptions, smCOG, GECCO probability, MIBiG gene match and stored BLASTp hit.

Rules the tool keeps:
- Each region is shown as `strain / full contig / region / BGC alias`, copied from the package inventory row.
- A gap-rescue partner is drawn on the map; an RG-GMCI link is not, and its line names the two-proof verdict.
  A link the two-proof check rates WEAK does not raise a region's rank.
- Until a rescue verdict table is supplied, every region slide carries "review draft: links not yet adjudicated"; with
  one, a region carries "review draft: n of m links not settled" while a shown split partner has no settled ruling.
- An existing deck is never overwritten; a new --tag writes new files.
- Similarity is not identity; capacity is not production; a family is shared architecture, not a compound.
Needs python-pptx, matplotlib and Pillow (pip install 'sapote-mamey[slides]'); without them it exits with a message.
"""
import argparse
import copy
import csv
import datetime
import json
import math
import re
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _console import emit  # noqa: E402

HERE = Path(__file__).resolve().parent
TAG_DEFAULT = "v2"
SW, SH, M = 13.333, 7.5, 0.45
INK, MUTED, ACCENT, WARN = "1F2937", "4B5563", "1F4E79", "B45309"
ROLE_COL = {"core biosynthetic": "#B91C1C", "tailoring / modification": "#F59E0B", "halogenation": "#F59E0B",
            "chain release / TE": "#7C3AED", "regulation": "#15803D", "self-resistance / export": "#2563EB",
            "biosynthetic context": "#C7CBD1", "unknown": "#E5E7EB"}
# Pfam families counted as a core biosynthetic enzyme in a GECCO-only cluster (the list the cohort GECCO table used).
CORE_PFAM = {"PF00501": "AMP-binding (A domain)", "PF00109": "KS N", "PF02801": "KS C", "PF00698": "AT",
             "PF00668": "condensation", "PF04055": "radical SAM", "PF04738": "LanB", "PF05147": "LanC"}
BAR = 74.0  # median identity for a close MIBiG match

TEMPLATE = {
    "strain": "", "package": "", "antismash_zip": "", "mibig_gbk_dir": "",
    "gap_rescue_dir": "", "gecco_dir": "", "bigscape_region_tsv": "", "bigscape_figure": "",
    "rescue_verdicts_tsv": "", "blastp_csv_globs": [], "pcoa_kit": "", "pcoa_drop_origins": "", "pfam_names_tsv": "", "rggmci_scorecard_tsv": "", "gene_labels_tsv": "", "locus_maps_dir": "", "multi_ref_dir": "", "ref_genome_dir": "", "typed_rules_json": "",
    "metadata": {"organism": "", "host": "", "location": "", "collected": "", "accession_16s": "", "bioassay": ""},
    "nearest": [{"name": "", "ani": "", "af": "", "method": ""}],
    "trees": [{"image": "", "label": "", "note": ""}],
    "family_figures": [{"bgc": "", "image": "", "caption": ""}],
    "extra_figures": [{"image": "", "heading": "", "caption": ""}],
    "pointers": [{"label": "", "path": "", "note": ""}],
}


def need_libs():
    try:
        import pptx  # noqa: F401
        import matplotlib  # noqa: F401
        import PIL  # noqa: F401
    except ImportError as e:
        sys.exit(f"strain_slides needs python-pptx, matplotlib and Pillow ({e}); install the 'slides' extra")


def rows(path, delim="\t"):
    with open(path, newline="") as fh:
        return [{k: (v or "").strip() for k, v in r.items() if k} for r in csv.DictReader(fh, delimiter=delim)]


def node_key(contig):
    """`NODE_n_length_L`: GECCO and some tools rewrite the coverage part of a SPAdes contig name."""
    return contig.split("_cov")[0]


def resolve(src, base):
    out = dict(src)
    for k, v in src.items():
        if isinstance(v, str) and v and k not in ("strain",) and not v.startswith("/"):
            out[k] = str((base / v).resolve()) if (base / v).exists() else v
    return out


# ---------------------------------------------------------------- data
def load(src):
    strain = src["strain"]
    pkg = Path(src["package"])
    D = {"strain": strain, "pkg": pkg, "src": src}
    D["manifest"] = json.load(open(pkg / "manifest.json")) if (pkg / "manifest.json").exists() else {}
    D["inv"] = {r["BGC_ID"]: r for r in rows(pkg / f"{strain}_2_inventory.csv", ",")}
    D["genes"] = defaultdict(list)
    for r in rows(pkg / f"{strain}_gene_by_gene_all_bgcs.csv", ","):
        D["genes"][r["bgc_id"]].append(r)
    opt = lambda name: pkg / f"{strain}_{name}"
    D["mibig"] = {r["bgc_id"]: r for r in rows(opt("3_mibig_profile.csv"), ",")} if opt("3_mibig_profile.csv").exists() else {}
    D["rgg"] = rows(opt("4A_RGGMCI_ranked_pairs.csv"), ",") if opt("4A_RGGMCI_ranked_pairs.csv").exists() else []
    D["twoproof"] = {tuple(sorted((r["bgc_a"], r["bgc_b"]))): r for r in rows(opt("4D_two_proof_rescue.csv"), ",")} \
        if opt("4D_two_proof_rescue.csv").exists() else {}
    D["dom"] = defaultdict(list)
    if opt("domains.csv").exists():
        for r in rows(opt("domains.csv"), ","):
            try:
                ev = float(r.get("evalue") or 1)
            except ValueError:
                ev = 1.0
            D["dom"][r["locus_tag"]].append((ev, r.get("domain", ""), r.get("description", ""), r.get("feature_type", "")))
        for k in D["dom"]:
            D["dom"][k].sort()
    D["cds"] = {r["locus_tag"]: r for r in rows(opt("cds_table.csv"), ",")} if opt("cds_table.csv").exists() else {}
    D["polymer"] = defaultdict(list)
    if opt("predicted_polymers.csv").exists():
        for r in rows(opt("predicted_polymers.csv"), ","):
            D["polymer"][(node_key(r["record_id"]), r["region_number"])].append(r)
    D["adom"] = defaultdict(list)
    if opt("nrps_prediction.csv").exists():
        for r in rows(opt("nrps_prediction.csv"), ","):
            if r.get("domain_class") == "NRPS_A":
                D["adom"][r["locus_tag"]].append(r)
    # gap rescue: SUMMARY.tsv plus one folder per run, and the adjudicator's per-run verdict file when present
    D["gap"], D["gap_root"] = {}, Path(src["gap_rescue_dir"]) if src.get("gap_rescue_dir") else None
    if D["gap_root"] and (D["gap_root"] / "SUMMARY.tsv").exists():
        for r in rows(D["gap_root"] / "SUMMARY.tsv"):
            D["gap"][r["bgc"].split(" / ")[-1]] = r
    # GECCO: genes, features, clusters; overlaps are computed here on NODE_n_length_L and coordinates
    D["ggenes"], D["gfeat"], D["gclusters"] = defaultdict(list), defaultdict(list), []
    gd = Path(src["gecco_dir"]) if src.get("gecco_dir") else None
    if gd and gd.is_dir():
        for f in gd.glob("*.genes.tsv"):
            for r in rows(f):
                D["ggenes"][node_key(r["sequence_id"])].append(r)
        for f in gd.glob("*.features.tsv"):
            for r in rows(f):
                D["gfeat"][r["protein_id"]].append(r)
        for f in gd.glob("*.clusters.tsv"):
            D["gclusters"] += rows(f)
    D["gover"], D["gonly"] = gecco_overlaps(D)
    D["big"] = {r["bgc"]: r for r in rows(src["bigscape_region_tsv"])} if src.get("bigscape_region_tsv") and \
        Path(src["bigscape_region_tsv"]).exists() else {}
    D["verdicts"] = defaultdict(list)
    if src.get("rescue_verdicts_tsv") and Path(src["rescue_verdicts_tsv"]).exists():
        for r in rows(src["rescue_verdicts_tsv"]):
            core = r.get("core identity") or r.get("core_identity", "")
            if r.get("strain") == strain and core:
                D["verdicts"][core.split(" / ")[-1]].append(r)
    D["blastp"] = blastp_hits(src.get("blastp_csv_globs") or [])
    D["pcoa_cache"] = {}
    D["pcoa_drop"] = frozenset(l.strip() for l in open(src["pcoa_drop_origins"])) \
        if src.get("pcoa_drop_origins") and Path(src["pcoa_drop_origins"]).exists() else frozenset()
    D["pfam"] = {}
    if src.get("pfam_names_tsv") and Path(src["pfam_names_tsv"]).exists():
        D["pfam"] = {r["accession"]: r["name"] for r in rows(src["pfam_names_tsv"])}
    D["glabel"] = gene_labels(src.get("gene_labels_tsv"), D)
    D["typed_rules"] = None
    if src.get("typed_rules_json") and Path(src["typed_rules_json"]).exists():
        sys.path.insert(0, str(HERE))
        import typed_class_rules
        D["typed_rules"] = typed_class_rules.load_rules(src["typed_rules_json"])
    D["region_of"] = defaultdict(list)  # NODE_n_length_L -> [(start, end, BGC, region)]
    for b, r in D["inv"].items():
        D["region_of"][node_key(r["Contig"])].append((int(float(r["Start"])), int(float(r["End"])), b, r["antiSMASH_Region"]))
    D["kcb"] = kcb_genes(src.get("antismash_zip"))
    # RG-GMCI pair scorecard (intact relative + biosynthetic logic): strain, pair, layers_verdict, relative, gap
    D["scorecard"] = {}
    if src.get("rggmci_scorecard_tsv") and Path(src["rggmci_scorecard_tsv"]).exists():
        for r in rows(src["rggmci_scorecard_tsv"]):
            if r.get("strain") != strain or "+" not in r.get("pair", ""):
                continue
            ids = {x.split(" / ")[-1]: x for x in (r.get("identity_a", ""), r.get("identity_b", "")) if x}
            ok = all(b in D["inv"] and ids.get(b, "").split(" / ")[1:3] == [D["inv"][b]["Contig"], D["inv"][b]["antiSMASH_Region"]]
                     for b in r["pair"].split("+"))
            if ok:  # a row from another run is used only when both regions' contig and region match this package
                D["scorecard"][frozenset(r["pair"].split("+"))] = r
    D["mibig_names"] = {}
    D["famfigs"] = defaultdict(list)
    for f in src.get("family_figures") or []:
        if f.get("bgc") and f.get("image") and Path(f["image"]).exists():
            D["famfigs"][f["bgc"]].append(f)
    return D


def gecco_overlaps(D):
    by_node = defaultdict(list)
    for b, r in D["inv"].items():
        by_node[node_key(r["Contig"])].append((int(float(r["Start"])), int(float(r["End"])), b))
    over, only = defaultdict(list), []
    for c in D["gclusters"]:
        k = node_key(c["sequence_id"])
        s, e = int(c["start"]), int(c["end"])
        hit = [b for (rs, re_, b) in by_node.get(k, []) if s <= re_ and e >= rs]
        for b in hit:
            over[b].append(c)
        if not hit:
            only.append(c)
    return over, only


def blastp_hits(globs):
    """Rank-1 hits per gene from stored BLASTp CSVs (columns gene, hit_rank, subject_acc, subject_organism, pct_identity)."""
    import glob as _g
    best = {}
    for pat in globs:
        for f in sorted(_g.glob(pat, recursive=True), reverse=True):
            try:
                rr = rows(f, ",")
            except (OSError, csv.Error, UnicodeDecodeError):
                continue
            for r in rr:
                if r.get("hit_rank") != "1" or not r.get("gene"):
                    continue
                org = r.get("subject_organism", "")
                if r["gene"] not in best or (org and not best[r["gene"]]["org"]):
                    best[r["gene"]] = dict(org=org, pid=r.get("pct_identity", ""), acc=r.get("subject_acc", ""))
    return best


def kcb_genes(zip_path):
    """Per-gene KnownClusterBlast hits from an antiSMASH result ZIP: locus tag -> [(MIBiG accession, cluster name,
    subject protein, identity)], best first. The ranked hit list is antiSMASH's; nothing is re-scored here."""
    import zipfile
    out = defaultdict(list)
    out["__rank__"] = {}  # (NODE_n_length_L, region number) -> MIBiG accessions in antiSMASH's rank order
    if not zip_path or not Path(zip_path).exists():
        return out
    try:
        z = zipfile.ZipFile(zip_path)
    except (OSError, zipfile.BadZipFile):
        return out
    try:
        for n in z.namelist():
            if "__MACOSX" in n or not re.search(r"(^|/)knownclusterblast/[^/]+\.txt$", n):
                continue  # the folder sits at the ZIP's top level or under one result folder
            acc = name = None
            in_tab = False
            m = re.match(r"(.+)_c(\d+)\.txt$", Path(n).name)
            ranked = out["__rank__"].setdefault((node_key(m.group(1)), int(m.group(2))), []) if m else []
            for line in z.read(n).decode("utf-8", "replace").splitlines():
                m = re.match(r"^\d+\. (BGC\d{7})", line)
                if m:
                    acc, in_tab = m.group(1), False
                    if acc not in ranked:
                        ranked.append(acc)
                    continue
                if line.startswith("Source: "):
                    name = line[8:].strip()
                elif line.startswith("Table of Blast hits"):
                    in_tab = True
                elif in_tab and line.strip():
                    f = line.split("\t")
                    if len(f) >= 3 and acc:
                        try:
                            identity = float(f[2])
                            if not math.isfinite(identity) or not 0 <= identity <= 100:
                                raise ValueError("identity must be a finite percentage")
                            out[f[0]].append((acc, name or acc, f[1], identity))
                        except ValueError as exc:
                            raise ValueError(f"STRAIN_SLIDE_KCB_UNVERIFIED: {n}: invalid identity {f[2]!r} for {f[0]!r}: {exc}") from exc
                elif in_tab:
                    in_tab = False
    finally:
        z.close()
    for k in out:
        if k != "__rank__":
            out[k].sort(key=lambda x: -x[3])
    return out


def kcb_top(D, bgc, n=3):
    """The region's first n KnownClusterBlast clusters, in antiSMASH's order."""
    inv = D["inv"][bgc]
    num = int(re.sub(r"\D", "", inv["antiSMASH_Region"]) or 0)
    return D["kcb"].get("__rank__", {}).get((node_key(inv["Contig"]), num), [])[:n]


def mibig_gene_name(D, acc, protein):
    """Gene name (or locus tag) of a MIBiG protein, read from the local MIBiG GenBank folder; the protein id otherwise."""
    if acc not in D["mibig_names"]:
        names = {}
        d = D["src"].get("mibig_gbk_dir")
        f = next(iter(sorted(Path(d).glob(f"{acc}*.gbk"))), None) if d and Path(d).is_dir() else None
        if f:
            for block in f.read_text(errors="replace").split("     CDS             ")[1:]:
                pid = re.search(r'/protein_id="([^"]+)"', block)
                gn = re.search(r'/gene="([^"]+)"', block)
                pr = re.search(r'/product="([^"]+)"', block)
                lt = re.search(r'/locus_tag="([^"]+)"', block)
                if pr and not (re.fullmatch(r"[A-Za-z]{2,5}\d*[A-Z]?\d*", pr.group(1)) and len(pr.group(1)) <= 8):
                    pr = None  # a product line names the gene only when it is short, like LooH
                nm = gn or pr or lt
                if pid and nm:
                    names[pid.group(1)] = nm.group(1)
        D["mibig_names"][acc] = names
    return D["mibig_names"][acc].get(protein, protein)


def pfam_names(D, tag):
    """Pfam names of one gene from GECCO's feature table, best E-value first (accessions when no name table is given)."""
    fs = sorted(D["gfeat"].get(tag, []), key=lambda f: float(f.get("i_evalue") or 1))
    return list(dict.fromkeys(D["pfam"].get(f["domain"], f["domain"]) for f in fs))


# Biosynthetic step of a reference gene, read from the reference's own product line (first match wins).
STEPS = [("halogenation", ["halogenase"]),
         ("glycosylation", ["glycosyltransferase", "glycosyl transferase", "glycosyl-transferase"]),
         ("methylation", ["methyltransferase", "methylase"]),
         ("sugar supply", ["sugar", "hexose", "epimerase", "dtdp", "udp-", "ndp-", "nucleotidyl", "glucose", "mannos"]),
         ("regulation", ["regulat", "transcription"]),
         ("transport / resistance", ["antiporter", "efflux", "transporter", "transmembrane", "permease", "abc "])]
STEP_COL = {"core and other biosynthetic": "#B91C1C", "halogenation": "#F59E0B", "glycosylation": "#F59E0B",
            "methylation": "#F59E0B", "sugar supply": "#FBBF24", "regulation": "#15803D", "transport / resistance": "#2563EB"}
FOUND = ("PRESENT_IN_CORE", "MISSING_FOUND_CLEAR", "MISSING_FOUND_AMBIGUOUS")
# The same steps read from the found gene's Pfam names, used when the reference's product line names no step.
STEP_PFAM = [("halogenation", ["halogenase"]), ("glycosylation", ["glycos_transf", "glyco_trans", "udpgt", "erycii"]),
             ("methylation", ["methyltransf"]),
             ("sugar supply", ["epimerase", "udpg_mgdp", "ntp_transf", "hexose_dehydrat", "aldo_ket_red", "rmld", "gdp_man",
                               "degt_dnrj"]),
             ("regulation", ["marr", "tetr", "luxr", "gere", "hth", "response_reg", "trans_reg"]),
             ("transport / resistance", ["mfs", "sugar_tr", "na_h_exchanger", "abc_tran", "abc2_membrane"])]


def step_of(r, D=None):
    p = (r.get("reference_product") or "").lower()
    for label, words in STEPS:
        if any(w in p for w in words):
            return label
    if D is not None and r.get("best_locus"):
        pf = " ".join(pfam_names(D, r["best_locus"])).lower()
        for label, words in STEP_PFAM:
            if any(w in pf for w in words):
                return label
    if p in ("", "unknown", "hypothetical protein"):
        return "unannotated in the reference"
    return "core and other biosynthetic" if r.get("reference_gene_kind", "").startswith("biosynthetic") else "other"


def where(D, r):
    """`NODE_n_length_L` plus the antiSMASH region the found gene sits in, from the gap-rescue row."""
    k = node_key(r.get("best_contig", ""))
    reg = (r.get("best_region_identity") or "").split(" / ")
    return k, (f"{reg[-1]}" if len(reg) == 4 else "no antiSMASH region")


def kcb_concordance(D, a, b):
    """Per-gene KnownClusterBlast agreement between two regions: MIBiG clusters both regions hit, with the reference
    genes each one hits. Complementary = mostly different reference genes, the pattern of one cluster split in two."""
    def subj(bgc):
        m = defaultdict(set)
        names = {}
        for g in D["genes"].get(bgc, []):
            for acc, name, prot, _ in D["kcb"].get(g["locus_tag"], []):
                m[acc].add(prot)
                names[acc] = name
        return m, names
    ma, na = subj(a)
    mb, nb = subj(b)
    out = []
    for acc in set(ma) & set(mb):
        ov = len(ma[acc] & mb[acc])
        out.append(dict(acc=acc, name=(na.get(acc) or nb.get(acc) or acc).split("/")[0], a=len(ma[acc]), b=len(mb[acc]),
                        both=ov, union=len(ma[acc] | mb[acc]), complementary=ov <= 1 and min(len(ma[acc]), len(mb[acc])) >= 2))
    return sorted(out, key=lambda x: (-x["complementary"], -x["union"]))


CONTRADICTED = ("APART_CLOSE", "CONFLICT")


SETTLED = ("HOLDS", "TWO_SIMILAR_LOCI", "REJECT")


def unsettled_links(D, bgc, st=None):
    """(links without a settled rescue verdict, links shown) for a region's split partners. A verdict row counts for
    the partner its `partner region` names, or failing that its `partner contig`."""
    shown, _ = split_partners(D, bgc, st)
    done = set()
    for v in D["verdicts"].get(bgc, []):
        if v.get("verdict") in SETTLED:
            done.add((v.get("partner region") or "").split(" / ")[-1])
            done.add(node_key(v.get("partner contig") or v.get("partner_contig", "")))
    return sum(x["other"] not in done and node_key(D["inv"][x["other"]]["Contig"]) not in done for x in shown), len(shown)


def split_partners(D, bgc, st=None):
    """Regions that may be the rest of this cluster, with every layer that speaks to the link.
    Candidates: RG-GMCI HIGH or MODERATE partners, and regions holding a SUPPORTED gap-rescue find.
    Layers: per-gene KnownClusterBlast (both regions hit mostly different genes of a MIBiG cluster in both top threes)
    and the RG-GMCI scorecard's intact-relative reading (CONSISTENT: the two loci lie together in the best public
    relative; APART_CLOSE: far apart in a complete relative; CONFLICT: a single-copy marker on both; APART_WEAK and
    NOT_READ: the relative is a draft or distant, or none is shared).
    Returns (shown, set_aside): shown when position says CONSISTENT, or KCB is complementary and position does not
    contradict; set aside when position contradicts. Benchmark (lineage benchmark): KCB alone kept 3 of 19 true
    splits and 2 of 17 false HIGH links; with position, the readable false link is rejected and the true ones kept."""
    cand = {}
    for r in D["rgg"]:
        if bgc in (r["bgc_a"], r["bgc_b"]) and r["rggmci_confidence"].startswith(("HIGH", "MODERATE")):
            o = r["bgc_b"] if r["bgc_a"] == bgc else r["bgc_a"]
            cand[o] = f"RG-GMCI {r['rggmci_confidence'].split('_')[0]} {r.get('rggmci_score', '')}, two-proof " \
                      f"{two_proof(D, r) or 'not checked'}"
    for r in (st or {}).get("table", []):
        reg = (r.get("best_region_identity") or "").split(" / ")
        if r["status"] in FOUND[1:] and r.get("partner_verdict") == "SUPPORTED" and len(reg) == 4 and reg[-1] != bgc:
            cand.setdefault(reg[-1], "gap-rescue SUPPORTED find")
    shown, aside = [], []
    for o, why in cand.items():
        if o not in D["inv"]:
            continue
        top = set(kcb_top(D, bgc)) & set(kcb_top(D, o))
        conc = [c for c in kcb_concordance(D, bgc, o) if c["complementary"] and c["acc"] in top]
        sc = D["scorecard"].get(frozenset((bgc, o)), {})
        pos = sc.get("layers_verdict", "")
        item = dict(other=o, why=why, conc=conc, pos=pos, sc=sc)
        if pos in CONTRADICTED:
            aside.append(item)
        elif pos == "CONSISTENT" or conc:
            shown.append(item)
    key = lambda x: (x["pos"] == "CONSISTENT" and bool(x["conc"]), x["pos"] == "CONSISTENT",
                     max([c["union"] for c in x["conc"]] or [0]))
    return sorted(shown, key=key, reverse=True), aside


def position_text(sc):
    """The intact-relative check as a plain sentence (a draft relative fits but cannot settle a split)."""
    v = sc.get("layers_verdict", "")
    if not v:
        return "Their order was not checked against a public relative."
    rel = (sc.get("relative_source") or sc.get("relative") or "").split(",")[0][:34]
    gap = sc.get("locus_gap_genes", "")
    asm = (sc.get("relative_assembly") or "").lower()
    done = asm.startswith("complete")
    return {"CONSISTENT": (f"In {rel}, a complete genome, the two loci sit together ({gap} genes apart)." if done else
                           f"In {rel}, a {asm or 'draft'} genome, the two loci sit together, but a draft cannot settle it."),
            "APART_WEAK": f"In {rel} ({asm or 'a distant relative'}) they lie {gap} genes apart; that relative cannot decide it.",
            "NOT_READ": "No shared public relative to check their order.",
            "APART_CLOSE": f"In {rel}, a complete genome, they lie {gap} genes apart: two separate clusters.",
            "CONFLICT": "Both hold the same single-copy core gene: two separate clusters."}.get(v, v)


VERDICT_WORDS = {"HOLDS": "the link holds", "TWO_SIMILAR_LOCI": "two similar loci, not one cluster",
                 "REJECT": "rejected", "UNRESOLVED": "unresolved"}


def short_rule(rule, n=110):
    rule = "" if rule.lstrip().startswith("{") else rule.strip()
    return rule if len(rule) <= n else rule[:n].rsplit(" ", 1)[0].rstrip(",;:") + "…"


def verdict_sentence(v):
    """A gene-level ruling on one link, in words."""
    rule = short_rule(v.get("rule", ""))
    w = v.get("verdict", "")
    return (f"Gene-level review: {VERDICT_WORDS.get(w, w.lower())}"
            + ((f"; next check: {rule[:1].lower() + rule[1:]}" if w == "UNRESOLVED" else f" ({rule})") if rule else "") + ".")


def verdict_for(D, bgc, other):
    """The ruling on the link between bgc and another region, if the verdict table has one."""
    node = node_key(D["inv"][other]["Contig"]) if other in D["inv"] else ""
    for v in D.get("verdicts", {}).get(bgc, []):
        if (v.get("partner region") or "").split(" / ")[-1] == other or \
                (node and node_key(v.get("partner contig") or v.get("partner_contig", "")) == node):
            return v
    return None


def split_paras(D, bgc, st=None):
    """Regions on other contigs that may be the rest of this cluster, one plain paragraph each: the link, the
    KnownClusterBlast pattern, the position in a public relative, and the gene-level ruling when there is one."""
    shown, aside = split_partners(D, bgc, st)
    out = []
    for x in shown[:3]:
        o = x["other"]
        inv = D["inv"][o]
        head = f"{o} on {node_key(inv['Contig']).split('_length')[0]} ({inv['Products'][:26]}): "
        m = re.match(r"RG-GMCI (\w+) ([\d.]+), two-proof (\w+)", x["why"])
        if m:
            body = f"linked by RG-GMCI, {m.group(1).lower()} confidence (score {m.group(2)})."
            if m.group(3) == "WEAK":
                body += " The two-proof check rates it weak: more like two similar loci."
        else:
            body = "genes the partner checks support were found here."
        if x["conc"]:
            k = x["conc"][0]
            body += (f" The two regions hit different genes of {k['name'][:30]} ({k['a']} here, {k['b']} there"
                     + (f", {k['both']} in both" if k["both"] else "") + ").")
        elif m:
            body += " KnownClusterBlast does not show them hitting different genes of one cluster."
        body += " " + position_text(x["sc"])
        v = verdict_for(D, bgc, o)
        if v:
            body += " " + verdict_sentence(v)
        out.append([(head, True, INK), (body, False, INK)])
    if aside:
        out.append([("Set aside: ", True, MUTED),
                    (", ".join(x["other"] for x in aside[:3]) + (" and others" if len(aside) > 3 else "")
                     + ", which a complete relative or a shared single-copy gene shows to be separate clusters.", False, MUTED)])
    return out


def pathway_paras(D, gr, st):
    """The gap rescue in plain sentences: how many reference genes are in the region, which turned up on which other
    contigs, which genes are closer to another MIBiG cluster, and what is missing. Class level; nothing is joined."""
    ref_acc = gr.get("reference", "")
    ref = gr.get("reference_name", "").split("/")[0][:40] or ref_acc
    found = [r for r in st["table"] if r["status"] in FOUND and r.get("best_locus")]
    anchored = {node_key(r.get("best_contig", "")) for r in found if r.get("partner_verdict") == "SUPPORTED"}
    here, away, aside = [], defaultdict(list), []
    for r in found:
        v = r.get("partner_verdict", "")
        if r["status"] == "PRESENT_IN_CORE":
            here.append(r)
        elif v and v != "SUPPORTED" and node_key(r.get("best_contig", "")) not in anchored:
            aside.append(r)
        else:
            away[where(D, r)].append(r)
    if not away:
        return []  # every kept find is on the core contig: nothing is assembled across contigs
    ids = [float(r["best_identity_pct"]) for r in here if r.get("best_identity_pct")]
    out = [[(f"{len(here)} of {st['n']} {ref} genes are in this region"
             + (f" (median {statistics.median(ids):.0f}% identity)" if ids else "") + ".", False, INK)]]
    for (k, reg), rs in away.items():
        sup = sum(r.get("partner_verdict") == "SUPPORTED" for r in rs)
        genes = ", ".join(f"{ref_label(D, r, 30)} {float(r['best_identity_pct']):.0f}%" for r in rs[:6]) + (", ..." if len(rs) > 6 else "")
        note = ("supported by the partner checks" if sup == len(rs) else
                f"{sup} supported by the partner checks, the others weaker matches beside them" if sup else
                "weaker matches only")
        out.append([(f"On {k.split('_length')[0]} ({reg}): ", True, INK), (f"{genes}; {note}.", False, INK)])
    alts = defaultdict(list)
    for r in here + sum(away.values(), []):
        alt = next((h for h in D["kcb"].get(r["best_locus"], []) if not h[0].startswith(ref_acc.split(".")[0])), None)
        name = alt[1].split("/")[0].strip()[:30] if alt else ""
        if alt and alt[3] > float(r["best_identity_pct"] or 0) and name.lower() != ref.lower():  # not another entry of the same compound
            alts[name].append(alt[3])
    for name, vals in list(alts.items())[:2]:
        out.append([(f"{len(vals)} of these genes {'is' if len(vals) == 1 else 'are'} closer to {name} (KnownClusterBlast, "
                     f"up to {max(vals):.0f}%) than to {ref}.", False, INK)])
    missing = [r for r in st["table"] if r["status"] not in FOUND]
    if aside:
        out.append([(f"{len(aside)} weaker match{'es' if len(aside) > 1 else ''} on contigs with nothing supported "
                     "set aside (paralog families or single genes).", False, MUTED)])
    if missing:
        names = ", ".join(ref_label(D, r, 20) for r in missing[:4]) + (", ..." if len(missing) > 4 else "")
        out.append([(f"Not found anywhere in the genome: {len(missing)} of {st['n']} ({names}).", False, MUTED)])
    return out


def gap_stats(D, bgc):
    r = D["gap"].get(bgc)
    if not r or r.get("check") != "run" or not D["gap_root"]:
        return r, None
    f = D["gap_root"] / r["folder"]
    if not (f / "gap_rescue.tsv").exists():
        return r, None
    sys.path.insert(0, str(HERE))
    from gap_directed_rescue import apply_adjudication
    g = apply_adjudication(rows(f / "gap_rescue.tsv"), f)  # the adjudicator's verdicts, when its file is present
    for x in g:  # a reference gene named only by its protein accession takes the MIBiG file's short name (LooH)
        if re.fullmatch(r"[A-Z]{1,3}_?\d{5,}\.\d+", x.get("name", "")) and r.get("reference"):
            x["name"] = mibig_gene_name(D, r["reference"], x["name"])
    found = [x for x in g if x["status"] in ("PRESENT_IN_CORE", "MISSING_FOUND_CLEAR")]
    ids = [float(x["best_identity_pct"]) for x in found if x["best_identity_pct"]]
    adj = D["gap_root"] / f"{r['folder']}_ADJUDICATION.tsv"
    av = Counter(x["verdict"] for x in rows(adj)) if adj.exists() else None
    return r, dict(folder=f, table=g, n=len(g), core=sum(x["status"] == "PRESENT_IN_CORE" for x in g),
                   clear=sum(x["status"] == "MISSING_FOUND_CLEAR" for x in g),
                   frac=len(found) / len(g) if g else 0, med=statistics.median(ids) if ids else None, adj=av)


NO_SLIDE_TYPES = {"ectoine", "NAPAA"}


def left_out(D, bgc):
    """Why a region gets no region slide, or "" when it gets one. The owner, 3 Oct: ectoine is near universal and "does not
    belong in our slide decks"; NAPAA and geosmin go with it; terpene and saccharide stay. antiSMASH has no geosmin
    type, so a geosmin region is a terpene-only region whose KnownClusterBlast best match is geosmin. A region with
    split-link evidence keeps its slide whatever its type: a HIGH RG-GMCI link the two-proof check does not reject, or a
    gap-rescue find on another contig that the partner checks support. The region keeps its region-table row."""
    inv = D["inv"][bgc]
    types = {t.strip() for t in inv["Products"].split(";") if t.strip()}
    kc = inv.get("KCB_top", "").split(" | ")
    if types and types <= NO_SLIDE_TYPES:
        why = " and ".join(sorted(types)) + " only"
    elif types == {"terpene"} and len(kc) > 1 and "geosmin" in kc[1].lower():
        why = "terpene, best KnownClusterBlast match geosmin"
    else:
        return ""
    if rgg_for(D, bgc)[0]:
        return ""
    _, st = gap_stats(D, bgc)
    core = node_key(inv["Contig"])
    if st and any(x.get("partner_verdict") == "SUPPORTED" and x.get("best_contig")
                  and node_key(x["best_contig"]) != core for x in st["table"]):
        return ""
    return why


def two_proof(D, r):
    return D["twoproof"].get(tuple(sorted((r["bgc_a"], r["bgc_b"]))), {}).get("verdict", "")


def rgg_for(D, bgc):
    """HIGH links count only when the engine's two-proof check does not rate them WEAK."""
    hits = [r for r in D["rgg"] if bgc in (r["bgc_a"], r["bgc_b"])]
    high = [r for r in hits if r["rggmci_confidence"] == "HIGH_RG_GMCI_RESCUE" and two_proof(D, r) != "WEAK"]
    mod = [r for r in hits if r["rggmci_confidence"].startswith("MODERATE")]
    return high, mod


def gecco_max(D, bgc):
    ps = [float(c["average_p"]) for c in D["gover"].get(bgc, []) if c.get("average_p")]
    return max(ps) if ps else None


def score(D, bgc):
    """Evidence order of the region slides: gap-rescue completeness x median identity, plus GECCO's probability,
    a HIGH RG-GMCI link the two-proof check does not reject, a BiG-SCAPE family, and a HOLDS rescue verdict."""
    _, st = gap_stats(D, bgc)
    s = (st["frac"] * (st["med"] or 0) / 100) if st else 0
    s += 0.3 * (gecco_max(D, bgc) or 0)
    s += 0.3 if rgg_for(D, bgc)[0] else 0
    s += 0.1 if D["big"].get(bgc, {}).get("gcf_tightest", "unplaced") not in ("", "unplaced") else 0
    s += 0.3 if any(v.get("verdict") == "HOLDS" for v in D["verdicts"].get(bgc, [])) else 0
    return s


def big_line(b):
    if not b:
        return "no row for this region in the family table"
    if not b.get("gcf_tightest") or b["gcf_tightest"] == "unplaced":
        return "a singleton at every cutoff: no other region in the run shares its family"
    cut = b.get("tightest_shared_cutoff", "")
    mem = b.get(f"members_c{cut}", "")
    layers = Counter(m.split(":")[0] for m in mem.split("; ") if ":" in m)
    own = sorted({m.split(":", 1)[1] for m in mem.split("; ") if m.startswith("AS:")})
    mib = sorted(set(re.findall(r"BGC\d{7}", " ".join(m for m in mem.split("; ") if m.upper().startswith("MIBIG")))))
    s = f"family {b['gcf_tightest']} (tightest cutoff {cut}); members by layer: " + \
        (", ".join(f"{k} {v}" for k, v in sorted(layers.items())) or "none")
    if own:
        s += f"; cohort strains: {', '.join(own[:10])}" + (f" and {len(own) - 10} more" if len(own) > 10 else "")
    if mib:
        s += f"; MIBiG: {', '.join(mib[:4])}" + (f" and {len(mib) - 4} more" if len(mib) > 4 else "")
    return s


# Steps a reader looks for in each class, matched case-insensitively against each gene's antiSMASH domains, Pfam/TIGRFAM
# names and descriptions, smCOG and rule annotations. A step not matched is "not seen in the annotations", never absent.
CLASS_STEPS = {
    "T2PKS": [("KS-alpha (minimal PKS)", ["t2ks"]), ("chain-length factor (KS-beta)", ["t2clf", "chain length factor", "chain-length factor"]),
              ("acyl carrier protein", ["t2acp", "pp-binding", "acyl carrier", "acp"]), ("ketoreductase", ["t2kr", "ketoreductase", "adh_short"]),
              ("cyclase / aromatase", ["polyketide_cyc", "dimerisation", "aromatase", "cyclase"]),
              ("oxygenase (ring tailoring)", ["monooxygenase", "abm", "fad_binding", "oxygenase"])],
    "T1PKS": [("ketosynthase (KS)", ["pks_ks", "ketoacyl-synt", "ks domain"]), ("acyltransferase (AT)", ["pks_at", "acyl_transf_1", "acyltransferase"]),
              ("acyl carrier (ACP)", ["acp", "pp-binding"]), ("reductive loop (KR/DH/ER)", ["pks_kr", "pks_dh", "pks_er", "ketoreductase"]),
              ("release (thioesterase)", ["thioesterase", "pks_te"])],
    "NRPS": [("condensation (C)", ["condensation"]), ("adenylation (A)", ["amp-binding"]), ("carrier (T/PCP)", ["pp-binding", "pcp"]),
             ("release (TE / reductase)", ["thioesterase", "nad_binding_4", "td domain", "ntd", "thioester reductase"]), ("MbtH-like partner", ["mbth"])],
    "terpene": [("terpene synthase / cyclase", ["terpene_synth", "terpene_cyclase", "terpene synthase", "sqs_psy", "polyprenyl_synt", "terpene"]),
                ("isoprenoid precursor supply", ["idi", "isopentenyl", "polyprenyl", "geranyl"])],
    "lanthipeptide": [("precursor peptide", ["precursor", "lant_", "lanthipeptide"]), ("dehydratase (LanB/LanM)", ["lant_dehydr", "lanb", "lanm", "dehydratase"]),
                      ("cyclase (LanC)", ["lanc", "lanc_like"])],
    "lassopeptide": [("precursor peptide", ["precursor", "lasso"]), ("B protease / RRE", ["transglut_core", "pqqd", "rre", "b protein"]),
                     ("C lactam synthetase", ["asn_synthase", "asparagine synth"])],
    "NRP-metallophore": [("siderophore synthetase / NRPS core", ["iuca", "amp-binding", "condensation"]),
                         ("catecholate or hydroxamate tailoring", ["isochorismat", "chorismate_bind", "2,3-dihydro", "ornithine", "monooxygenase"]),
                         ("iron uptake / export", ["siderophore", "fecb", "abc transporter", "iron"])],
    "NI-siderophore": [("NIS synthetase", ["iuca", "iucc", "nis synthetase"]), ("hydroxamate tailoring", ["ornithine", "lysine 6-mono", "acetyltrans"]),
                       ("iron uptake", ["siderophore", "fecb", "iron"])],
    "halogenated": [("flavin-dependent halogenase", ["trp_halogenase", "halogenase"]), ("flavin reductase partner", ["flavin_reduct", "flavin reductase"])],
    "saccharide": [("glycosyltransferase", ["glycos_transf", "glycosyltransferase", "udpgt"]),
                   ("sugar activation / NDP-sugar enzymes", ["epimerase", "ntp_transf", "dtdp", "ndp-hexose", "rmld", "nad_dependent"])],
}
TAILOR = [("methyltransferase", ["methyltransf", "o-methyl", "c-methyl"]), ("halogenase", ["halogenase"]),
          ("glycosyltransferase", ["glycos_transf", "glycosyltransferase", "udpgt"]), ("cytochrome P450", ["p450"]),
          ("monooxygenase / oxidoreductase", ["monooxygenase", "fad_binding", "oxidoreductase", "adh_short"]),
          ("aminotransferase", ["aminotran", "aminotransferase"]), ("acyltransferase", ["acetyltransf", "acyltransferase"])]
REG = ["tetr", "marr", "luxr", "sarp", "gntr", "arsr", "lysr", "laci", "response_reg", "sigma", "hth"]
TRANSPORT = ["abc_tran", "mfs", "abc2_membrane", "efflux", "transporter", "permease"]


def gene_text(D, g):
    t = g["locus_tag"]
    parts = [g.get("sec_met_domains", ""), g.get("product_qualifier", ""), g.get("gene_function_inference", "")]
    parts += [f"{d[1]} {d[2]}" for d in D["dom"].get(t, [])]
    c = D["cds"].get(t, {})
    parts += [c.get("product", ""), c.get("gene_functions", ""), c.get("sec_met_domains", "")]
    return " ".join(parts).lower()


def biosynthetic_logic(D, bgc):
    """Paragraphs (runs) describing the region's biosynthetic logic from gene annotations only. Class level."""
    gs = D["genes"].get(bgc, [])
    inv = D["inv"][bgc]
    txt = {g["locus_tag"]: gene_text(D, g) for g in gs}
    prods = [p.strip() for p in inv["Products"].split(";") if p.strip()]
    out, seen_classes = [], []
    for p in prods:
        key = next((k for k in CLASS_STEPS if k.lower() == p.lower() or (k == "T1PKS" and p.lower() in ("t1pks", "transat-pks"))
                    or (k == "NRPS" and p.lower() in ("nrps", "nrps-like")) or (k == "terpene" and "terpene" in p.lower())), None)
        if not key or key in seen_classes:
            continue
        seen_classes.append(key)
        found, missing = [], []
        for label, al in CLASS_STEPS[key]:
            hit = [t for t, x in txt.items() if any(a in x for a in al)]
            (found if hit else missing).append(f"{label} ({', '.join(hit[:3])}{'...' if len(hit) > 3 else ''})" if hit else label)
        out.append([(f"{key}: ", True, INK), ("found " + "; ".join(found) if found else "no step matched", False, INK)]
                   + ([(f". Not seen in the annotations: {', '.join(missing)}.", False, MUTED)] if missing else [(".", False, INK)]))
    # A-domain substrate calls and predicted polymer
    calls = []
    for g in gs:
        for a in D["adom"].get(g["locus_tag"], []):
            sub = a.get("consensus_substrate") or a.get("substrate") or ""
            if sub and sub not in ("X", "-"):
                calls.append(f"{g['locus_tag']} {sub}")
            elif a.get("class_or_alternatives"):
                calls.append(f"{g['locus_tag']} {a['class_or_alternatives'].split(' | ')[0]}")
    if calls:
        out.append([("A-domain substrate calls: ", True, INK), ("; ".join(calls[:6]) + ("; ..." if len(calls) > 6 else "")
                                                               + " (prediction, similarity level)", False, INK)])
    pol = D["polymer"].get((node_key(inv["Contig"]), str(int(re.sub(r"\D", "", inv["antiSMASH_Region"]) or 0))), [])
    for r in pol[:2]:
        if r.get("predicted_polymer"):
            out.append([("antiSMASH predicted polymer: ", True, INK), (r["predicted_polymer"] + (
                "; region may hold two or more clusters, split before reading a product" if r.get("over_merge_flag", "").startswith("YES") else ""), False, INK)])
    # tailoring, transport, regulation
    tail = []
    for label, al in TAILOR:
        hit = [t for t, x in txt.items() if any(a in x for a in al)]
        if hit:
            tail.append(f"{label} x{len(hit)}")
    reg = sum(1 for x in txt.values() if any(a in x for a in REG))
    tr = sum(1 for x in txt.values() if any(a in x for a in TRANSPORT))
    if tail or reg or tr:
        out.append([("Tailoring: ", True, INK), (", ".join(tail) or "none matched", False, INK),
                    (f". Transport/resistance genes {tr}; regulators {reg}.", False, INK)])
    # one class-level reading
    if seen_classes:
        core_ok = []
        for k in seen_classes:
            first = CLASS_STEPS[k][:2]
            if all(any(any(a in x for a in al) for x in txt.values()) for _, al in first):
                core_ok.append(k)
        if core_ok:
            read = (f"The core genes for {' + '.join(core_ok)} biosynthesis are present at the domain level"
                    + (f", with {tail[0].split(' x')[0]}" + (f" and {tail[1].split(' x')[0]}" if len(tail) > 1 else "") + " tailoring" if tail else "")
                    + ": capacity consistent with that class, not a product call.")
        else:
            read = "The class's first core steps are not both seen in the annotations: treat the label as unconfirmed until the genes are read."
        out.append([("Reading: ", True, ACCENT), (read, False, INK)])
    return out


def gene_domains(D, tag, n=2):
    """Up to n Pfam / antiSMASH domain names of one gene, best E-value first: the package's domains table, else GECCO's."""
    names = [x[1] for x in D["dom"].get(tag, []) if x[3] in ("PFAM_domain", "aSDomain") and x[1]]
    return list(dict.fromkeys(names or pfam_names(D, tag)))[:n]


def typed_paras(D, bgc):
    """Family-evidence lines from the typed class rules (T3PKS, fatty_acid, RiPP-like, quinone, butyrolactone,
    nucleoside): exact domain tokens, one count per gene, optional steps silent when absent."""
    if not D.get("typed_rules"):
        return []
    import typed_class_rules
    return [[(f"{cls}: ", True, INK), ("; ".join(lines) + ".", False, INK)]
            for cls, lines in typed_class_rules.builder_lines(D, bgc, D["typed_rules"]) if lines]


def domain_paras(D, bgc):
    """Every gene's domain families, for a region whose product has no class rule (nucleoside, fatty_acid, RiPP-like
    and others). Family names only: what each protein's domains are, not what the cluster makes."""
    gs = D["genes"].get(bgc, [])
    named = [(g["locus_tag"], gene_domains(D, g["locus_tag"])) for g in gs]
    with_d = [f"{tag} {' + '.join(d)}" for tag, d in named if d]
    without = [tag for tag, d in named if not d]
    if not with_d:
        return []
    out = [[("Domains per gene (Pfam): ", True, INK), ("; ".join(with_d) + ".", False, INK)]]
    if without:
        out.append([(f"No domain annotated: {', '.join(without)}.", False, MUTED)])
    out.append([("Reading: ", True, ACCENT), ("no class rule covers this product yet, so the domain families are listed as "
                 "annotated. A domain family is similarity, not a measured reaction or a product call.", False, INK)])
    return out


def top_domain(D, tag):
    d = [x for x in D["dom"].get(tag, []) if x[3] in ("PFAM_domain", "aSDomain")]
    return d[0][1] if d else ""


def short_domains(g):
    doms = [d.strip() for d in g.get("sec_met_domains", "").split(";") if d.strip()]
    return ", ".join(dict.fromkeys(doms[:2]))


LABEL_ROLE_COL = {"core": "#B91C1C", "tailoring": "#F59E0B", "regulation": "#15803D", "transport_resistance": "#2563EB"}


def gene_labels(path, D):
    """Admit labels only for current contig/tag/AA bindings; never infer identity from length."""
    import hashlib
    import io
    import zipfile
    from collections import defaultdict
    try:
        from Bio import SeqIO
    except ImportError:  # biopython absent: the bundle's GenBank shim
        from mamey._gbk_shim import SeqIO

    if not path or not Path(path).exists():
        return {}
    rs = [r for r in rows(path) if r.get('strain') == D['strain']]
    archive = Path(D.get('src', {}).get('antismash_zip') or '')
    if not archive.is_file():
        emit(f"{D['strain']}: gene labels refused (canonical sequence archive missing)", file=sys.stderr)
        return {}
    norm = lambda aa: ''.join(str(aa).split()).upper().rstrip('*')
    digest = lambda aa: hashlib.sha256(norm(aa).encode()).hexdigest()
    sequence_index = defaultdict(set)
    with zipfile.ZipFile(archive) as z:
        members = [name for name in z.namelist() if name.lower().endswith(('.gbk', '.gb', '.genbank'))
                   and '.region' not in name.lower() and '__MACOSX/' not in name
                   and not Path(name).name.startswith('._')]
        for member in members:
            with z.open(member) as raw:
                for rec in SeqIO.parse(io.TextIOWrapper(raw), 'genbank'):
                    for feature in rec.features:
                        if feature.type != 'CDS':
                            continue
                        tag = feature.qualifiers.get('locus_tag', [''])[0]
                        aa = feature.qualifiers.get('translation', [''])[0]
                        if tag and aa:
                            sequence_index[(node_key(rec.id), tag)].add((digest(aa), len(norm(aa))))
    # Package-region proteins provide an additional independent check where present.
    package_hashes = defaultdict(set)
    fasta = Path(D['pkg']) / (D['strain'] + '_proteins.faa')
    if fasta.is_file():
        for rec in SeqIO.parse(fasta, 'fasta'):
            package_hashes[rec.id].add(digest(rec.seq))
    accepted = defaultdict(set)
    rejected = 0
    for r in rs:
        tag = r.get('locus_tag', '')
        contig = node_key(r.get('full_contig', ''))
        try:
            expected = (r.get('aa_sha256', ''), int(r.get('protein_length_aa') or 0))
        except ValueError:
            rejected += 1
            continue
        if not contig or not tag or sequence_index.get((contig, tag)) != {expected}:
            rejected += 1
            continue
        c = D.get('cds', {}).get(tag)
        if c and (node_key(c.get('contig', '')) != contig or str(c.get('length_aa')) != str(expected[1])
                  or package_hashes.get(tag) != {expected[0]}):
            rejected += 1
            continue
        accepted[tag].add((contig, expected[0], r.get('short_label') or '', r.get('role') or ''))
    result = {tag: (next(iter(values))[2], next(iter(values))[3])
              for tag, values in accepted.items() if len(values) == 1}
    ambiguous = sum(len(values) != 1 for values in accepted.values())
    if rejected or ambiguous:
        emit(f"{D['strain']}: gene labels withheld ({rejected} sequence/contig/package mismatches; "
              f"{ambiguous} ambiguous locus tags)", file=sys.stderr)
    return result



def labelled(D, label_fn, colour_fn):
    """The strip's own label and colour, falling back to the gene-label table for a gene that would be blank or grey."""
    gl = D.get("glabel") or {}
    lf = lambda g: label_fn(g) or gl.get(g.get("tag"), ("", ""))[0]
    cf = lambda g: (lambda c: LABEL_ROLE_COL.get(gl.get(g.get("tag"), ("", ""))[1], c) if c == "#C7CBD1" else c)(colour_fn(g))
    return lf, cf


# ---------------------------------------------------------------- figures
def strip(genes, ggenes, start, end, out, label_fn, colour_fn, height=2.15, matched=(), flip=False):
    """Gene arrows over [start, end] with GECCO's per-gene probability as bars above them."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyArrow
    plt.rcParams.update({"font.family": "sans-serif", "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"]})
    span = max(end - start, 1)
    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(7.9, height), dpi=220, gridspec_kw=dict(height_ratios=[1, 2.1], hspace=0.28), sharex=True)
    for g in ggenes:
        s, e = int(g["start"]), int(g["end"])
        if e < start or s > end:
            continue
        p = float(g["average_p"])
        ax1.bar((s + e) / 2 / 1000, p, width=(e - s) / 1000 * 0.92, color="#1F4E79" if p >= 0.8 else "#9DB4CC", lw=0)
    ax1.axhline(0.8, color="#6B7280", lw=.6, ls="--")
    pin = [float(g["average_p"]) for g in ggenes if int(g["end"]) >= start and int(g["start"]) <= end]
    if pin and max(pin) < 0.1:
        ax1.text(0.5, 0.45, f"GECCO p below 0.1 for every gene (max {max(pin):.2f})", transform=ax1.transAxes,
                 fontsize=6, color="#6B7280", ha="center", va="center")
    ax1.set_ylim(0, 1)
    ax1.set_yticks([0, .8])
    ax1.set_ylabel("GECCO p", fontsize=6.5, color="#4B5563")
    ax1.tick_params(labelsize=6, colors="#6B7280", length=1.5)
    for s_ in ("top", "right"):
        ax1.spines[s_].set_visible(False)
    ax2.set_ylim(-1, 2.6)  # headroom so angled labels stay below the GECCO bars
    ax2.axhline(0, color="#9CA3AF", lw=.6, zorder=0)
    mg = [g for g in genes if g.get("tag") in matched]
    if mg:
        lo, hi = min(g["s"] for g in mg) / 1000, max(g["e"] for g in mg) / 1000
        for ax in (ax1, ax2):
            ax.axvspan(lo, hi, color="#FDE68A", alpha=.45, lw=0, zorder=0)
        ax2.text((lo + hi) / 2, -0.95, "genes matched to the MIBiG reference (the map above)", fontsize=5.6,
                 ha="center", va="bottom", color="#92400E")
    hl = span / 1000 * 0.012
    for g in genes:
        s, e, strand = g["s"], g["e"], g["strand"]
        L = (e - s) / 1000
        x0 = (s if strand > 0 else e) / 1000
        on = g.get("tag") in matched
        ax2.add_patch(FancyArrow(x0, 0, L * strand, 0, width=0.55, head_width=0.55, head_length=min(hl, L * .9),
                                 length_includes_head=True, color=colour_fn(g), ec="#92400E" if on else "none",
                                 lw=.9 if on else 0))
    labs = [((g["s"] + g["e"]) / 2 / 1000, label_fn(g), 2 if g.get("tag") in matched else 1) for g in genes]
    for x, lab, pri in thin_labels(labs, start / 1000, end / 1000, 7.9 * 0.9):
        ax2.text(x, 0.42, lab, rotation=35, fontsize=5.4, ha="left", va="bottom", color="#111827")
    ax2.set_yticks([])
    for s_ in ("top", "right", "left"):
        ax2.spines[s_].set_visible(False)
    ax2.set_xlim((end / 1000, start / 1000) if flip else (start / 1000, end / 1000))
    ax2.set_xlabel("position on the contig (kb)" + (", drawn right to left to match the map" if flip else ""),
                   fontsize=6.5, color="#4B5563")
    ax2.tick_params(labelsize=6, colors="#6B7280", length=1.5)
    fig.subplots_adjust(left=0.08, right=0.98, top=0.95, bottom=0.2)
    fig.savefig(out, facecolor="white", bbox_inches="tight", pad_inches=0.04)  # never clip the axis label
    plt.close(fig)
    return out


def region_strip(D, bgc, out, height, matched, flip):
    gs = D["genes"].get(bgc, [])
    if not gs:
        return None
    inv = D["inv"][bgc]
    start, end = int(gs[0]["bgc_start"]), int(gs[0]["bgc_end"])
    gg = D["ggenes"].get(node_key(inv["Contig"]), [])
    gp = {x["protein_id"]: float(x["average_p"]) for x in gg}
    genes = [dict(s=int(g["cds_start"]), e=int(g["cds_end"]), strand=1 if g["strand"] in ("+", "1") else -1,
                  role=g["gene_function_inference"], tag=g["locus_tag"],
                  dom=short_domains(g) if g["gene_function_inference"] == "core biosynthetic" else top_domain(D, g["locus_tag"]),
                  p=gp.get(g["locus_tag"], 0.0)) for g in gs]
    lf, cf = labelled(D, lambda g: g["dom"] if (g["role"] == "core biosynthetic" or g["tag"] in matched or g["p"] >= 0.8) else "",
                      lambda g: ROLE_COL.get(g["role"], "#C7CBD1"))
    return strip(genes, gg, start, end, out, label_fn=lf, colour_fn=cf, height=height, matched=matched, flip=flip)


def one_row_per_locus(rows):
    """{found gene: its rescue row}. When one gene is the best hit of several reference genes (paralogs such as two
    2OG oxygenases), the reciprocal best wins, then the higher identity, as on the map (the owner, 3 Oct: a validamycin-like BGC's
    strip read valJ where the map read valE, because the later row overwrote the earlier)."""
    out = {}
    for r in rows:
        k = r.get("best_locus")
        if not k:
            continue
        rank = (str(r.get("reciprocal_best", "")).strip().lower() in ("true", "1", "yes"), float(r.get("best_identity_pct") or 0))
        if k not in out or rank > out[k][0]:
            out[k] = (rank, r)
    return {k: r for k, (_, r) in out.items()}


def rescue_panels(D, bgc, st, lay, window=2500):
    """Partner contigs the map drew, in the map's order, each cut to its found genes +/- window bp."""
    inv = D["inv"][bgc]
    core = node_key(inv["Contig"])
    by = defaultdict(list)
    for r in st["table"]:
        if r["status"] in FOUND[1:] and r.get("best_locus") and r.get("best_contig"):
            by[node_key(r["best_contig"])].append(r)
    anchored = {c for c, rs in by.items() if any(r.get("partner_verdict") == "SUPPORTED" for r in rs)}
    drawn = [c for c in (lay or {}).get("contigs_drawn", []) if c != core and c in by and c in anchored]
    panels = []
    for c in drawn:
        gg = D["ggenes"].get(c, [])
        pos = {g["protein_id"]: (int(g["start"]), int(g["end"])) for g in gg}
        tags = {k: r for k, r in one_row_per_locus(by[c]).items() if k in pos}
        if not tags:
            continue
        lo = max(0, min(pos[t][0] for t in tags) - window)
        hi = max(pos[t][1] for t in tags) + window
        regs = [x for x in D["region_of"].get(c, []) if x[0] <= hi and x[1] >= lo]
        panels.append(dict(contig=c, start=lo, end=hi, matched=tags, flip=bool((lay or {}).get("flipped", {}).get(c, False)),
                           label=f"{c.split('_length')[0]} · " + (", ".join(x[2] for x in regs) if regs else "no antiSMASH region"),
                           regions=[x[2] for x in regs],
                           genes=[dict(s=int(g["start"]), e=int(g["end"]), strand=1 if g["strand"] == "+" else -1,
                                       tag=g["protein_id"], p=float(g["average_p"]))
                                  for g in gg if int(g["end"]) >= lo and int(g["start"]) <= hi], ggenes=gg))
    return panels


ACCESSION = r"[A-Z]{1,3}_?\d{5,}\.\d+"


def ref_label(D, r, n=16):
    """Strip label of a found reference gene: its reference name, or the found gene's first Pfam name when the
    reference names it only by a protein accession. Never a bare accession."""
    nm = r.get("name", "")
    if re.fullmatch(ACCESSION, nm):
        pf = pfam_names(D, r.get("best_locus", ""))
        nm = (pf[0][:max(n - 5, 11)].rstrip("_-,. ") + "-like") if pf else ("unnamed" if n > 16 else r.get("best_locus", ""))
    return nm[:n]


def thin_labels(items, lo, hi, inches, gap_in=0.17):
    """Labels that fit side by side: (x, text, priority) placed greedily, priority first, keeping gap_in inches
    between neighbours on an axis inches wide spanning lo..hi."""
    gap = (hi - lo) * gap_in / max(inches, 0.1)
    kept = []
    for x, lab, pri in sorted(items, key=lambda i: (-i[2], i[0])):
        if lab and all(abs(x - k[0]) >= gap for k in kept):
            kept.append((x, lab, pri))
    return kept


def partner_rank(P):
    """Partner contigs ranked by supported finds, then by all finds."""
    return (sum(r.get("partner_verdict") == "SUPPORTED" for r in P["matched"].values()), len(P["matched"]))


def main_locus(D, bgc, st, lay):
    """A partner region holding more of the reference's genes than this region: (partner panel, n there, n here)."""
    if not st:
        return None
    here = sum(r["status"] == "PRESENT_IN_CORE" for r in st["table"])
    best = max(rescue_panels(D, bgc, st, lay), key=lambda P: len(P["matched"]), default=None)
    if best and best.get("regions") and len(best["matched"]) >= 3 and len(best["matched"]) > here:
        return best, len(best["matched"]), here
    return None


def multi_strip(D, core_panel, partners, out, height, more=0, core_at=0):
    """The core region and its partner contigs side by side ("//" between them), each with GECCO bars, its own
    orientation and labels. Partner genes found for the reference carry its gene name; others their first Pfam name."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyArrow
    plt.rcParams.update({"font.family": "sans-serif", "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"]})
    panels = partners[:core_at] + [core_panel] + partners[core_at:]  # the map's left-to-right order
    spans = [max(p["end"] - p["start"], 1) for p in panels]
    tot = sum(spans)
    ratios = [max(s / tot, 0.16) for s in spans]
    fig, axes = plt.subplots(2, len(panels), figsize=(7.9, height), dpi=220, squeeze=False,
                             gridspec_kw=dict(height_ratios=[1, 2.2], width_ratios=ratios, wspace=0.12, hspace=0.28))
    for j, P in enumerate(panels):
        a1, a2 = axes[0][j], axes[1][j]
        for g in P["ggenes"]:
            s, e = int(g["start"]), int(g["end"])
            if e < P["start"] or s > P["end"]:
                continue
            p = float(g["average_p"])
            a1.bar((s + e) / 2 / 1000, p, width=(e - s) / 1000 * 0.92, color="#1F4E79" if p >= 0.8 else "#9DB4CC", lw=0)
        a1.axhline(0.8, color="#6B7280", lw=.6, ls="--")
        a1.set_ylim(0, 1)
        a1.set_yticks([0, .8] if j == 0 else [])
        if j == 0:  # the first panel carries the axis, whichever contig it is
            a1.set_ylabel("GECCO p", fontsize=6.5, color="#4B5563")
        inches = 7.9 * 0.87 * ratios[j] / sum(ratios)
        head = P["label"]
        if _width_pt(head, False, 6.6) > inches * 72:  # a narrow panel: one part per line, so headings never meet
            head = head.replace(" · ", "\n", 1)
        a1.set_title(head, fontsize=6.6 if "\n" not in head else 6.0, color="#111827", loc="left", pad=2)
        pin = [float(g["average_p"]) for g in P["ggenes"] if int(g["end"]) >= P["start"] and int(g["start"]) <= P["end"]]
        if pin and max(pin) < 0.1:
            msg = f"GECCO p below 0.1 for every gene (max {max(pin):.2f})"
            if _width_pt(msg, False, 6) > inches * 72 * 0.95:
                msg = f"GECCO p < 0.1\n(max {max(pin):.2f})"
            a1.text(0.5, 0.45, msg, transform=a1.transAxes, fontsize=6 if "\n" not in msg else 5.5, color="#6B7280",
                    ha="center", va="center", clip_on=True)
        a2.set_ylim(-1, 2.7)  # headroom so angled labels stay below the GECCO bars
        a2.axhline(0, color="#9CA3AF", lw=.6, zorder=0)
        span = (P["end"] - P["start"]) / 1000
        for g in P["genes"]:
            s, e, strand = g["s"], g["e"], g["strand"]
            L = (e - s) / 1000
            on = g["tag"] in P["matched"]
            a2.add_patch(FancyArrow((s if strand > 0 else e) / 1000, 0, L * strand, 0, width=0.55, head_width=0.55,
                                    head_length=min(span * 0.04, L * .9), length_includes_head=True,
                                    color=P["colour"](g), ec="#92400E" if on else "none", lw=.9 if on else 0))
        labs = []
        for g in P["genes"]:
            lab = P["label_fn"](g)
            pri = 1 if g["tag"] not in P["matched"] else 2 if lab.startswith("ctg") else 3  # a name before a bare locus tag
            labs.append(((g["s"] + g["e"]) / 2 / 1000, lab, pri))
        lo_, hi_ = P["start"] / 1000, P["end"] / 1000
        for x, lab, pri in thin_labels(labs, lo_, hi_, inches):
            a2.text(x, 0.42, lab, rotation=40, fontsize=5.2, ha="left", va="bottom",
                    color="#92400E" if pri > 1 else "#111827", fontweight="bold" if pri > 1 else "normal")
        for a in (a1, a2):
            a.set_xlim((P["end"] / 1000, P["start"] / 1000) if P["flip"] else (P["start"] / 1000, P["end"] / 1000))
            a.tick_params(labelsize=5.5, colors="#6B7280", length=1.5)
            for s_ in ("top", "right") + (("left",) if j else ()):
                a.spines[s_].set_visible(False)
        a1.set_xticks([])
        a2.set_yticks([])
        a2.spines["left"].set_visible(False)
        if j:
            a2.text(-0.06, 0, "//", transform=a2.get_yaxis_transform(), fontsize=9, color="#6B7280", ha="center", va="center")
    fig.text(0.5, 0.01, "kb on each contig" + ("; reversed contigs are drawn right to left, as on the map"
                                              if any(p["flip"] for p in panels) else "")
             + (f"; {more} more partner contig{'s' if more > 1 else ''} on the map above" if more else ""),
             fontsize=6.3, color="#4B5563", ha="center")
    fig.subplots_adjust(left=0.07, right=0.94, top=0.9, bottom=0.12, hspace=0.3)  # room for the last labels and for angled labels below the GECCO bars
    fig.savefig(out, facecolor="white")
    plt.close(fig)
    return out


def rescue_strip(D, bgc, st, lay, out, height, matched, flip):
    """Multi-contig strip for a gap rescue with partner finds; the single-contig strip otherwise."""
    partners = rescue_panels(D, bgc, st, lay) if st else []
    if not partners:
        return region_strip(D, bgc, out, height, matched, flip)
    gs = D["genes"].get(bgc, [])
    inv = D["inv"][bgc]
    k = node_key(inv["Contig"])
    gp = {x["protein_id"]: float(x["average_p"]) for x in D["ggenes"].get(k, [])}
    ref_core = {k: ref_label(D, r) for k, r in one_row_per_locus(
        [r for r in st["table"] if r["status"] == "PRESENT_IN_CORE"]).items()}
    core = dict(contig=k, start=int(gs[0]["bgc_start"]), end=int(gs[0]["bgc_end"]), flip=flip, matched=set(ref_core),
                label=f"{k.split('_length')[0]} · core, {bgc}", ggenes=D["ggenes"].get(k, []),
                genes=[dict(s=int(g["cds_start"]), e=int(g["cds_end"]), strand=1 if g["strand"] in ("+", "1") else -1,
                            tag=g["locus_tag"], role=g["gene_function_inference"], p=gp.get(g["locus_tag"], 0.0)) for g in gs])
    core["colour"] = lambda g: ROLE_COL.get(g.get("role", ""), "#C7CBD1")
    core["label_fn"] = lambda g: (ref_core[g["tag"]] if g["tag"] in ref_core else
                                  (short_domains(next(x for x in gs if x["locus_tag"] == g["tag"])) if g.get("role") == "core biosynthetic" else ""))
    inv_ = main_locus(D, bgc, st, lay)
    if inv_:
        core["label"] = f"{k.split('_length')[0]} · this region, {bgc}"
    main = inv_[0]["contig"] if inv_ else None
    keep = [P["contig"] for P in sorted(partners, key=partner_rank, reverse=True)][:2]
    if main and main not in keep:
        keep = [main] + keep[:1]
    more = len(partners) - len(keep)
    partners = [P for P in partners if P["contig"] in keep]  # the map's order
    for P in partners:
        if P["contig"] == main:
            P["label"] += " · more reference genes"
        P["colour"] = (lambda m: lambda g: STEP_COL.get(step_of(m[g["tag"]], D), "#C7CBD1") if g["tag"] in m else
                       ("#15803D" if any(a in " ".join(pfam_names(D, g["tag"])).lower() for a in REG) else
                        "#2563EB" if any(a in " ".join(pfam_names(D, g["tag"])).lower() for a in TRANSPORT) else "#C7CBD1"))(P["matched"])
        P["label_fn"] = (lambda m: lambda g: ref_label(D, m[g["tag"]]) if g["tag"] in m else "")(P["matched"])
        P["label_fn"], P["colour"] = labelled(D, P["label_fn"], P["colour"])
    core["label_fn"], core["colour"] = labelled(D, core["label_fn"], core["colour"])
    seq = [c for c in (lay or {}).get("contigs_drawn", []) if c == k or c in keep]
    return multi_strip(D, core, partners, out, height, more, seq.index(k) if k in seq else 0)


def map_for(D, st, out_maps, genome_box):
    """Gap-rescue map, with the renderer's own record of which contigs it drew reversed."""
    outd = out_maps / st["folder"].name
    side = outd / "map_layout.json"
    if (outd / "gap_rescue.png").exists() and side.exists():
        return outd / "gap_rescue.png", json.load(open(side))
    src = D["src"]
    if not (src.get("antismash_zip") and src.get("mibig_gbk_dir")):
        return None, None
    try:
        sys.path.insert(0, str(HERE))
        import gap_directed_rescue as g
        if genome_box.get("genome") is None:
            genome_box["genome"] = g.load_genome(Path(src["antismash_zip"]), D["strain"])
        prots, regions = genome_box["genome"]
        bad = [r["best_protein"] for r in st["table"] if r.get("best_protein") and
               (r["best_protein"] not in prots or prots[r["best_protein"]]["tag"] != r["best_locus"])]
        if bad:
            return None, None  # the table and the ZIP disagree: no map rather than a wrong one
        res = g.redraw(st["folder"], prots, regions, D["strain"], Path(src["mibig_gbk_dir"]), outd) or {}
        if not (outd / "gap_rescue.png").exists():
            return None, None
        lay = {"flipped": {node_key(k): v for k, v in (res.get("flipped") or {}).items()},
               "contigs_drawn": [node_key(c) for c in res.get("contigs_drawn", [])], "anchor": res.get("anchor", "")}
        side.write_text(json.dumps(lay, indent=1))
        return outd / "gap_rescue.png", lay
    except (Exception, SystemExit) as e:  # a map that cannot be drawn leaves the slide with its strip only
        emit(D["strain"], st["folder"].name, "map not drawn:", str(e)[:160], file=sys.stderr)
        return None, None


def product_chart(counts, out_png):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    top = counts.most_common(12)
    rest = counts.most_common()[12:]
    if rest:
        top.append((f"{len(rest)} other types", sum(n for _, n in rest)))
    labels, vals = [t for t, _ in top][::-1], [n for _, n in top][::-1]
    fig, ax = plt.subplots(figsize=(6.2, 4.6), dpi=200)
    bars = ax.barh(labels, vals, color="#" + ACCENT, height=0.62)
    for b, v in zip(bars, vals):
        ax.text(b.get_width() + 0.4, b.get_y() + b.get_height() / 2, str(v), va="center", fontsize=9, color="#374151")
    ax.set_xlabel("antiSMASH regions carrying the type", fontsize=9, color="#374151")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_xlim(0, max(vals or [1]) * 1.15)
    fig.tight_layout()
    fig.savefig(out_png, facecolor="white")
    plt.close(fig)


# ---------------------------------------------------------------- slide helpers
def text(slide, x, y, w, h, paras, size=14, space=6):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        for t, bold, col in para:
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = "Calibri", Pt(size), bold
            r.font.color.rgb = RGBColor.from_string(col)
    return tb


FONT_DIRS = [Path("/Applications/LibreOffice.app/Contents/Resources/fonts/truetype"), Path("/usr/share/fonts/truetype/crosextra")]
_FONTS = {}


def _width_pt(s, bold, size):
    """Width of s in points in Calibri (measured with Carlito, its metric twin); 0.5 em per character without it."""
    if bold not in _FONTS:
        f = next((d / ("Carlito-Bold.ttf" if bold else "Carlito-Regular.ttf") for d in FONT_DIRS
                  if (d / "Carlito-Regular.ttf").exists()), None)
        try:
            from PIL import ImageFont
            _FONTS[bold] = ImageFont.truetype(str(f), 100) if f else None
        except Exception:
            _FONTS[bold] = None
    f = _FONTS[bold]
    return f.getlength(s) * size / 100 if f else len(s) * size * 0.5


def text_height(paras, w, size, space):
    """Inches a list of paragraphs takes in a w-inch box: words wrapped by measured width, 1.2 line spacing."""
    tot = -space if paras else 0.0  # spacing falls between paragraphs only
    for para in paras:
        lines, x = 1, 0.0
        for t, bold, _ in para:
            for k, word in enumerate(re.split(r"(\s+)", t)):
                if not word:
                    continue
                ww = _width_pt(word, bold, size)
                if x + ww > w * 72 and not word.isspace() and x > 0:
                    lines, x = lines + 1, ww
                else:
                    x += ww
        tot += lines * size * 1.2 + space
    return tot / 72


def fit_paras(paras, w, h, sizes=(9, 8.5, 8, 7.5), space=2):
    """(size, paragraphs that fit, overflow): the largest size that fits everything, else the smallest size with the
    longest fitting run of paragraphs; a heading is never left last on the slide."""
    for size in sizes:
        if text_height(paras, w, size, space) <= h:
            return size, paras, []
    size = sizes[-1]
    n = len(paras)
    while n > 1 and text_height(paras[:n], w, size, space) > h:
        n -= 1
    while n > 1 and len(paras[n - 1]) == 1 and paras[n - 1][0][2] == ACCENT:  # a section heading alone at the bottom
        n -= 1
    return size, paras[:n], paras[n:]


def title(slide, strain, heading):
    text(slide, M, 0.25, SW - 2 * M, 0.6, [[(strain + "  ", True, INK), (heading, False, MUTED)]], size=24)


def picture(slide, path, x, y, w, h, alt):
    from PIL import Image
    from pptx.util import Inches
    iw, ih = Image.open(path).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    pic = slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y), Inches(pw), Inches(ph))
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return ph


def table(slide, header, widths, body, y=0.95, font=8.5, row_h=0.3):
    from pptx.util import Inches, Pt
    tb = slide.shapes.add_table(len(body) + 1, len(header), Inches(M), Inches(y), Inches(sum(widths)),
                                Inches(row_h * (len(body) + 1))).table
    for j, w in enumerate(widths):
        tb.columns[j].width = Inches(w)
    for j, h in enumerate(header):
        tb.cell(0, j).text = h
    for i, vals in enumerate(body, 1):
        for j, v in enumerate(vals):
            tb.cell(i, j).text = str(v)
    for i in range(len(body) + 1):
        for j in range(len(header)):
            c = tb.cell(i, j)
            c.margin_top = c.margin_bottom = Inches(0.01)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.name, r.font.bold = Pt(font if i else font + 0.5), "Calibri", i == 0
    return tb


def figure_slide(prs, strain, heading, img, caption):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, strain, heading)
    ph = picture(s, img, M, 0.9, SW - 2 * M, 5.85, heading)
    text(s, M, min(0.9 + ph + 0.08, SH - 0.55), SW - 2 * M, 0.5, [[(caption, False, MUTED)]], size=11)
    s.notes_slide.notes_text_frame.text = f"Figure: {img}"
    return s


# ---------------------------------------------------------------- slides
def overview(prs, D):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    md = D["src"].get("metadata") or {}
    org = md.get("organism", "")
    title(s, D["strain"], f"{org}: overview" if org else "overview (organism not recorded)")
    a, bc = D["manifest"].get("assembly", {}), D["manifest"].get("bgc_counts", {})
    left = [[("Isolation (as recorded)", True, ACCENT)]]
    for key, lab in (("host", "Host"), ("location", "Location"), ("collected", "Collected"),
                     ("accession_16s", "16S accession"), ("bioassay", "Bioassay (crude extract, as recorded)")):
        left.append([(f"{lab}: ", True, INK), (md.get(key) or "not recorded", False, INK)])
    left.append([("Genome", True, ACCENT)])
    if a:
        left.append([(f"{a['genome_bp'] / 1e6:.2f} Mb, {a['contigs']} contigs, N50 {a['n50']:,} bp, GC {a['gc_pct']:.1f}%", False, INK)])
    if bc:
        left.append([(f"{bc['raw']} antiSMASH regions: {bc['interior']} interior, {bc['edge']} at a contig edge, "
                      f"{bc['full_contig']} whole-contig", False, INK)])
    left.append([("A fragmented draft: region counts are a floor, not a measurement.", False, MUTED)])
    text(s, M, 1.1, 6.0, 5.9, left)
    near = [[("Nearest public genomes", True, ACCENT)]] + (
        [[(n["name"], True, INK), (f"  ANI {n['ani']}%, aligned {n['af']}%  ", False, INK), (n.get("method", ""), False, MUTED)]
         for n in D["src"].get("nearest") or [] if n.get("name")] or [[("Not given.", False, MUTED)]])
    near.append([("ANI is genome similarity. 95% or more with half the genome aligned places it in a known species.", False, MUTED)])
    text(s, 6.9, 1.1, SW - 6.9 - M, 5.9, near, size=13)
    s.notes_slide.notes_text_frame.text = "Metadata as recorded in the sources file; assembly and counts from manifest.json."


def trees(prs, D):
    panels = [t for t in D["src"].get("trees") or [] if t.get("image") and Path(t["image"]).exists()]
    if not panels:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        title(s, D["strain"], "phylogeny")
        text(s, M, 1.2, SW - 2 * M, 2, [[("No approved tree for this strain yet.", True, INK)],
             [("Only approved tree panels are shown. This slide is the place for that tree.", False, MUTED)]], size=16)
        return
    for t in panels:
        figure_slide(prs, D["strain"], "phylogeny", Path(t["image"]), f"{t.get('label', '')}. {t.get('note', '')}".strip())


def landscape(prs, D, assets):
    inv = list(D["inv"])
    counts = Counter(p for b in inv for p in {t.strip() for t in D["inv"][b]["Products"].split(";") if t.strip()})
    chart = assets / "product_types.png"
    product_chart(counts, chart)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, D["strain"], "biosynthetic gene cluster landscape")
    picture(s, chart, M, 1.0, 6.4, 5.6, "antiSMASH product types")
    runs = [b for b in inv if gap_stats(D, b)[1]]
    close = [b for b in runs if (gap_stats(D, b)[1]["med"] or 0) >= BAR and gap_stats(D, b)[1]["frac"] >= 0.5]
    high = [b for b in inv if rgg_for(D, b)[0]]
    fam = [b for b in inv if D["big"].get(b, {}).get("gcf_tightest", "unplaced") not in ("", "unplaced")]
    gec = [b for b in inv if gecco_max(D, b) is not None]
    sacc = sum(1 for b in inv if {t.strip() for t in D["inv"][b]["Products"].split(";")} == {"saccharide"})
    nos = [b for b in inv if left_out(D, b)]
    P = [[("What the region slides hold", True, ACCENT)],
         [(f"{len(inv) - len(nos)} of {len(inv)} antiSMASH regions, strongest evidence first, then {len(D['gonly'])} "
           "GECCO-only candidates.", False, INK)],
         [(f"{len(nos)} ectoine, NAPAA or geosmin regions have no slide of their own; each keeps its row in the region "
           "table. One with a split-link find keeps its slide.", False, MUTED)],
         [(f"{len(runs)} have a gap-rescue run against a MIBiG reference; {len(close)} of those find half or more of the "
           f"reference's genes at a median identity of {BAR:.0f}% or more.", False, INK)],
         [(f"{len(high)} have a HIGH RG-GMCI link that the two-proof check does not reject.", False, INK)],
         [(f"{len(fam)} share a BiG-SCAPE family with another region.", False, INK)],
         [(f"{len(gec)} are overlapped by a GECCO cluster ({len(D['gclusters'])} GECCO clusters in all).", False, INK)],
         [(f"{sacc} of {len(inv)} regions are labelled saccharide only. A label is not a call: each has its own slide.", False, MUTED)],
         [("Counts are per type: a hybrid region counts once under each type. Similarity to a characterized cluster is "
           "not proof of the same product.", False, MUTED)]]
    text(s, 7.2, 1.1, SW - 7.2 - M, 5.8, P, size=12.5, space=4)


def region_table(prs, D, order, slide_of):
    hdr = ["Slide", "Region (contig / region / BGC)", "antiSMASH type", "Edge", "KCB best match", "Gap rescue",
           "RG-GMCI", "BiG-SCAPE", "GECCO p"]
    widths = [0.5, 2.6, 2.0, 0.75, 1.9, 1.55, 0.95, 1.35, 0.8]
    for k in range(0, len(order), 15):
        chunk = order[k:k + 15]
        s = prs.slides.add_slide(prs.slide_layouts[6])
        title(s, D["strain"], f"every antiSMASH region ({k + 1}-{k + len(chunk)} of {len(order)})")
        body = []
        for b in chunk:
            inv = D["inv"][b]
            _, st = gap_stats(D, b)
            high, mod = rgg_for(D, b)
            big = D["big"].get(b, {})
            kc = inv.get("KCB_top", "").split(" | ")
            gm = gecco_max(D, b)
            body.append([slide_of.get(b, ""), f"{node_key(inv['Contig'])} / {inv['antiSMASH_Region']} / {b}", inv["Products"],
                         {"Interior": "no", "Edge": "edge", "Full-contig": "whole"}.get(inv["Boundary"], inv["Boundary"]),
                         kc[1][:28] if len(kc) > 1 else "",
                         (f"{st['core'] + st['clear']}/{st['n']}" + (f", {st['med']:.0f}%" if st["med"] is not None else "")) if st else "",
                         ("HIGH" if high else "") + (f"{' + ' if high else ''}{len(mod)} mod" if mod else ""),
                         big.get("gcf_tightest", "") or ("singleton" if big else ""), f"{gm:.2f}" if gm is not None else ""])
        table(s, hdr, widths, body)
        s.notes_slide.notes_text_frame.text = (
            "One row per antiSMASH region, in the order of the region slides. Slide 'none' = an ectoine, NAPAA or geosmin "
            "region with no slide of its own (no split-link find). Gap rescue = reference genes found (in the region "
            "+ clearly elsewhere) / reference genes, median identity. RG-GMCI = HIGH links the two-proof check does not reject, "
            "and MODERATE candidates. Blank = no evidence in that channel, not a negative test.")


def region_slide(prs, D, bgc, n_order, assets, genome_box):
    strain = D["strain"]
    inv = D["inv"][bgc]
    ident = f"{strain} / {inv['Contig']} / {inv['antiSMASH_Region']} / {bgc}"
    gs = D["genes"].get(bgc, [])
    prod = inv["Products"]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, strain, f"{bgc} · {prod if len(prod) < 58 else prod[:55] + '...'}")
    gr, st = gap_stats(D, bgc)
    mp, lay = map_for(D, st, assets / "maps", genome_box) if st else (None, None)
    y = 0.95
    if mp:
        from PIL import Image
        iw, ih = Image.open(mp).size
        natural_w = min(7.9, iw / ih * 3.55)  # the map's width when its height fills the box
        rs = region_structure(D, assets, bgc, inv["Contig"], gr.get("reference", ""))
        if rs:  # the region's reference structure sits right of the map; a wide map gives up width to make room
            map_w = min(natural_w, 7.9 - STRUCT_MIN_W - 0.1)
            sw = min(STRUCT_MAX_W, 7.9 - map_w - 0.1)
            mh = picture(s, mp, M, y, map_w, 3.55, f"{ident} against {gr['reference']}")
            e_, b_, png_ = rs
            structure_box(s, D, e_, b_, trimmed(png_, assets / "structure_crops"), M + 7.9 - sw, y, sw, min(3.55, max(mh, 2.2)))
            D.setdefault("_region_structures", []).append(bgc)
            y += mh + 0.05
        else:
            y += picture(s, mp, M, y, 7.9, 3.55, f"{ident} against {gr['reference']}") + 0.05
    matched = {r["best_locus"] for r in st["table"] if r["status"] == "PRESENT_IN_CORE"} if (mp and st) else set()
    flip = bool((lay or {}).get("flipped", {}).get(node_key(inv["Contig"]), False))
    sp = rescue_strip(D, bgc, st, lay, assets / f"{bgc}_strip.png", 2.15 if mp else 3.6, matched, flip) if mp else \
        region_strip(D, bgc, assets / f"{bgc}_strip.png", 3.6, matched, flip)
    if sp:
        picture(s, sp, M, y, 7.9, min(2.15, SH - y - 0.1) if mp else 3.6, f"{ident} genes and GECCO")
        if not mp:
            why = (gr or {}).get("check", "no gap-rescue row for this region")
            text(s, M, y + 3.75, 7.9, 0.8, [[("No gap-rescue map: " + (why if why != "run" else
                                              "the map could not be drawn from the saved tables"), False, MUTED)]], size=10)
    band = (" Yellow band: the stretch from the first to the last gene matched to the MIBiG reference in the map; not "
            "part of GECCO." if mp else "")
    text(s, M, SH - 0.5, 7.9, 0.45, [[(
        "GECCO bars: the probability (0-1) that each gene belongs to a biosynthetic cluster, from its protein domains; dark at "
        "0.8 or more, GECCO's cluster threshold (dashed line). Arrows: red core, amber tailoring, green regulation, blue "
        "transport/resistance, grey other." + band, False, MUTED)]], size=8)
    open_n, link_n = unsettled_links(D, bgc, st)
    if not D["src"].get("rescue_verdicts_tsv"):
        text(s, 8.6, 0.74, SW - 8.6 - 0.3, 0.2, [[("REVIEW DRAFT: links not yet adjudicated", True, WARN)]], size=9)
    elif open_n:
        text(s, 8.6, 0.74, SW - 8.6 - 0.3, 0.2, [[(f"REVIEW DRAFT: {open_n} of {link_n} links not settled", True, WARN)]], size=9)
    core = [g for g in gs if g["gene_function_inference"] == "core biosynthetic"]
    kc = inv.get("KCB_top", "").split(" | ")
    mb = D["mibig"].get(bgc, {})
    P = [[(ident, True, INK)],
         [(f"{prod}; {inv['Boundary'].lower()}; {inv['Length_kb']} kb; {len(gs)} genes, {len(core)} core", False, MUTED)],
         [("What the genes show", True, ACCENT)]]
    ml = main_locus(D, bgc, st, lay) if (st and mp) else None
    if ml:
        P.insert(2, [("More of the reference lies elsewhere: ", True, WARN), (
            f"{ml[0]['label'].split(' · ')[0]} ({', '.join(ml[0]['regions'])}) holds {ml[1]} of the "
            f"{gr['reference_name'].split('/')[0][:28]} genes, this region {ml[2]}. Read this region as a possible piece "
            f"of that locus; see {', '.join(ml[0]['regions'])}'s slide.", False, INK)])
    logic = biosynthetic_logic(D, bgc)
    logic = [p for k, p in enumerate(logic) if p not in logic[:k]]  # a line the annotations give twice is shown once
    path = pathway_paras(D, gr, st) if (st and mp) else []
    if path:  # the pathway block lists tailoring across contigs, so the core-only tailoring line is dropped
        logic = [p for p in logic if not p[0][0].startswith(("Tailoring", "Reading"))]
    logic = logic + typed_paras(D, bgc)  # typed rules for classes the legacy steps do not cover
    if not logic:  # no class steps for this product: list each gene's domain families instead
        logic = domain_paras(D, bgc)
    if logic:
        P += logic
    else:
        P.pop()  # no class steps and no annotated domains: drop the empty heading
    sp_ = split_paras(D, bgc, st)
    if sp_:
        P.append([("Possible pieces on other contigs", True, ACCENT)])
        P += sp_
    if path:
        P.append([(f"Genes of the {gr['reference_name'].split('/')[0][:30]} cluster", True, ACCENT)])
        P += path
    P.append([("Other evidence", True, ACCENT)])
    ev = []
    if len(kc) > 1:
        ev.append(f"best KnownClusterBlast match {kc[1]} ({kc[0].split('.')[0]})")
    if mb:
        med = mb.get("median_pct_identity")
        ev.append(f"{mb['recognizable_gene_count']} of {mb['query_gene_count']} genes match MIBiG proteins"
                  + (f", median {float(med):.0f}%" if med not in (None, "") else ""))
    P.append([("MIBiG 4.0: ", True, INK), (("; ".join(ev) + ".") if ev else "no match.", False, INK)])
    if st and not path:
        adj = (", partner genes " + ", ".join(f"{k.lower().replace('_', ' ')} {v}" for k, v in st["adj"].most_common(3))) if st.get("adj") else ""
        P.append([("Gap rescue: ", True, INK),
                  (f"{st['core']} of {st['n']} {gr['reference_name'].split('/')[0][:28]} genes in this region, {st['clear']} more "
                   "found elsewhere" + (f", median {st['med']:.0f}% identity" if st["med"] is not None else "") + adj + ".", False, INK)])
    allh = [r for r in D["rgg"] if bgc in (r["bgc_a"], r["bgc_b"]) and r["rggmci_confidence"] == "HIGH_RG_GMCI_RESCUE"]
    _, mod = rgg_for(D, bgc)
    if (allh or mod) and not sp_:
        tp = Counter(two_proof(D, r) or "not checked" for r in allh)
        P.append([("RG-GMCI: ", True, INK),
                  ((f"{len(allh)} high-confidence link{'s' if len(allh) > 1 else ''} by score ("
                    + ", ".join(r["bgc_b"] if r["bgc_a"] == bgc else r["bgc_a"] for r in allh[:3]) + "), not drawn"
                    if allh else "no high-confidence link")
                   + (f"; the two-proof check rates {tp['WEAK']} as weak (two similar loci, not one split cluster)"
                      if tp.get("WEAK") else "")
                   + f"; {len(mod)} moderate.", False, INK)])
    told = {x["other"] for x in split_partners(D, bgc, st)[0][:3]}  # rulings already given with their link above
    for v in sorted(D["verdicts"].get(bgc, []), key=lambda v: v.get("verdict") not in SETTLED):
        if (v.get("partner region") or "").split(" / ")[-1] in told:
            continue
        P.append([(f"Link to {node_key(v.get('partner contig') or v.get('partner_contig', '')).split('_length')[0]}: ", True, INK),
                  (verdict_sentence(v), False, INK)])
        break
    if D["big"]:
        b = D["big"].get(bgc, {})
        fam = b.get("gcf_tightest", "")
        P.append([("BiG-SCAPE: ", True, INK), ((f"family {fam.replace(' c', ' (cutoff ', 1) + ')' if ' c' in fam else fam}." if fam and fam != "unplaced"
                                                 else "no family at any cutoff."), False, INK)])
    hits = {g["locus_tag"]: D["blastp"][g["locus_tag"]] for g in gs if g["locus_tag"] in D["blastp"]}
    if hits:
        c1 = next(((t, h) for t, h in hits.items() if t in {g["locus_tag"] for g in core}), next(iter(hits.items())))
        P.append([("BLASTp: ", True, INK), (f"{len(hits)} of {len(gs)} genes searched; {c1[0]}'s best hit is "
                                           f"{(c1[1]['org'] or c1[1]['acc'])[:40]}, {float(c1[1]['pid'] or 0):.0f}%.", False, INK)])
    if D["gclusters"]:
        gm = gecco_max(D, bgc)
        sac = inv["Products"].strip().lower() == "saccharide"
        P.append([("GECCO: ", True, INK), ((f"cluster probability {gm:.2f}" if gm is not None else "no cluster here")
                                         + ("; GECCO rarely calls saccharide-only regions (32 of 731 in the 43 decks, and "
                                            "no better at looser domain filters), so a low score here is not evidence "
                                            "against the region" if sac and (gm is None or gm < 0.8) else "") + ".", False, INK)])
    # the BGC in the protein PCoA sits under the text, on the same slide (3 Oct); the text above it flows on to a
    # "continued" slide when it does not fit. A BGC in more than two PCoA sets also gets the full PCoA slide.
    for pch in (PC_H, PC_H_SHORT):  # a shorter PCoA when that saves a "continued" slide
        pc_img, pc_sum = pcoa_inline(D, bgc, st, assets, SW - 8.6 - 0.3, pch)
        th = SH - 0.95 - 0.4 - ((pch + 0.25) if pc_img else 0)
        size, P_, rest = fit_paras(P, SW - 8.6 - 0.3, th)
        if not (pc_img and rest):
            break
    P = P_
    text(s, 8.6, 0.95, SW - 8.6 - 0.3, th + 0.2, P, size=size, space=2)
    if pc_img:
        y_pc = SH - 0.38 - pch
        text(s, 8.6, y_pc - 0.22, SW - 8.6 - 0.3, 0.2, [[("Protein PCoA", True, ACCENT), (
            f" ({min(2, len(pc_sum))} of {len(pc_sum)} sets; stars: this BGC, fill: identity to its best reference/MIBiG match)",
            False, MUTED)]], size=7.5)
        picture(s, pc_img, 8.6, y_pc, SW - 8.6 - 0.3, pch, f"{ident} proteins in the protein PCoA")
    if rest:
        text(s, 8.6, SH - 0.32, SW - 8.6 - 0.3, 0.2, [[("Continued on the next slide.", True, MUTED)]], size=8)
    s.notes_slide.notes_text_frame.text = (
        ("PCoA sets and starred genes: " + "; ".join(f"{T}: {', '.join(tags) or f'{n} points (not shown here)'}"
                                                      for T, n, _, tags in pc_sum) + ". " if pc_img else "") +
        f"Region {n_order} in evidence order. Identity, products, boundary, length: {strain}_2_inventory.csv. Genes and roles: "
        f"{strain}_gene_by_gene_all_bgcs.csv. Domains: {strain}_domains.csv. Gap rescue: "
        f"{st['folder'] if st else 'SUMMARY.tsv'}. RG-GMCI: {strain}_4A_RGGMCI_ranked_pairs.csv and {strain}_4D_two_proof_rescue.csv. "
        "Map: rows lined up on the majority of matched genes, contigs in reference order, ribbons by protein identity; a partner contig is a candidate missing piece, nothing "
        "is joined. Similarity is not identity; capacity is not production.")
    if rest:
        c = prs.slides.add_slide(prs.slide_layouts[6])
        title(c, strain, f"{bgc} · continued")
        text(c, M, 0.95, SW - 2 * M, SH - 1.2, [[(ident, True, INK)]] + rest, size=11, space=4)
    if not pc_img or len(pc_sum) > 2:
        pcoa_slide(prs, D, bgc, st, assets, ident)
    for f in D["famfigs"].get(bgc, []):
        figure_slide(prs, strain, f"{bgc}: {Path(f['image']).stem.replace('_', ' ')[:70]}", Path(f["image"]), f.get("caption", ""))


def pcoa_tags(D, bgc, st=None):
    own = {g["locus_tag"] for g in D["genes"].get(bgc, [])}
    part = {r["best_locus"] for r in (st or {}).get("table", []) if r["status"] in FOUND[1:]
            and r.get("partner_verdict") == "SUPPORTED" and r.get("best_locus")}
    return own, part


def pcoa_tags_in_kit(D, bgc):
    """True when any protein of this BGC (or a supported partner gene) has a point in the PCoA kit."""
    kit = D["src"].get("pcoa_kit")
    if not (kit and Path(kit).is_dir()):
        return False
    sys.path.insert(0, str(HERE))
    import strain_slides_pcoa as spc
    import numpy as np
    own, part = pcoa_tags(D, bgc, gap_stats(D, bgc)[1])
    tags = list(own | part)
    for T in spc.kit_sets(kit):
        if T not in D["pcoa_cache"]:
            D["pcoa_cache"][T] = spc.load(Path(kit), T, D["pcoa_drop"])
            with open(Path(kit) / f"out_{T}" / f"PCOA_{T}.tsv", newline="") as fh:
                D["pcoa_cache"][T]["tag"] = np.array([r.get("locus_tag", "") for r in csv.DictReader(fh, delimiter="\t")])
        C = D["pcoa_cache"][T]
        if (C["iso"] & (C["strain"] == D["strain"]) & np.isin(C["tag"], tags)).any():
            return True
    return False


PC_H, PC_H_SHORT = 2.25, 1.75  # inches: the PCoA panel on the region slide, and its shorter form


def pcoa_inline(D, bgc, st, assets, w, h):
    """This BGC's proteins in its two largest PCoA sets, small enough for the region slide. Returns (path, summary of
    every set that holds the BGC) or (None, [])."""
    kit = D["src"].get("pcoa_kit")
    if not (kit and Path(kit).is_dir()):
        return None, []
    sys.path.insert(0, str(HERE))
    import strain_slides_pcoa as spc
    own, part = pcoa_tags(D, bgc, st)
    return spc.bgc_figure(kit, D["strain"], own | part, assets / f"{bgc}_pcoa_inline.png", D["pcoa_cache"], D["pcoa_drop"],
                          max_panels=2, width=w, height=h, compact=True)


def pcoa_slide(prs, D, bgc, st, assets, ident):
    """This BGC's proteins (and its SUPPORTED partner-contig genes) as stars in every protein PCoA set that holds them."""
    kit = D["src"].get("pcoa_kit")
    if not (kit and Path(kit).is_dir()):
        return
    sys.path.insert(0, str(HERE))
    import strain_slides_pcoa as spc
    own, part = pcoa_tags(D, bgc, st)
    img, summ = spc.bgc_figure(kit, D["strain"], own | part, assets / f"{bgc}_pcoa.png", D["pcoa_cache"], D["pcoa_drop"])
    if not img:
        return
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, D["strain"], f"{bgc} in the protein PCoA, every set")
    text(s, M, 0.72, SW - 2 * M, 0.3, [[(ident, False, MUTED)]], size=9)
    picture(s, img, M, 1.0, SW - 2 * M, SH - 1.75, f"{ident} proteins in the protein PCoA")
    starred = sorted({t for _, _, _, tags in summ for t in tags})
    pp = sorted(set(starred) & part)
    text(s, M, SH - 0.62, SW - 2 * M, 0.5, [[(
        f"Stars: this BGC's proteins in each set where they occur ({len(starred)} of its {len(own)} genes"
        + (f", plus {len(pp)} supported partner-contig gene{'s' if len(pp) > 1 else ''}: {', '.join(pp)}" if pp else "")
        + "). Each point is one protein or domain (DIAMOND all-against-all, classical PCoA); a star's fill is its identity to "
          "its best reference/MIBiG match. A position near other proteins is similarity, not function or product.", False, MUTED)]], size=8)
    s.notes_slide.notes_text_frame.text = "Sets and starred genes: " + "; ".join(f"{T}: {', '.join(tags)}" for T, _, _, tags in summ)


def gecco_only_slide(prs, D, c, assets):
    strain = D["strain"]
    k = node_key(c["sequence_id"])
    start, end = int(c["start"]), int(c["end"])
    gg = D["ggenes"].get(k, [])
    feats = defaultdict(set)
    for g in gg:
        for f in D["gfeat"].get(g["protein_id"], []):
            if f["domain"] in CORE_PFAM:
                feats[g["protein_id"]].add(CORE_PFAM[f["domain"]])
    genes = [dict(s=int(g["start"]), e=int(g["end"]), strand=1 if g["strand"] == "+" else -1, tag=g["protein_id"],
                  p=float(g["average_p"])) for g in gg if int(g["end"]) >= start and int(g["start"]) <= end]
    core = sorted({d for g in genes for d in feats.get(g["tag"], [])})
    doms = Counter(f["domain"] for g in genes for f in D["gfeat"].get(g["tag"], []))
    out = assets / f"GECCO_ONLY_{k}_{start}.png"
    lf, cf = labelled(D, lambda g: ", ".join(sorted(feats.get(g["tag"], []))),
                      lambda g: "#B91C1C" if g["tag"] in feats else ("#1F4E79" if g["p"] >= 0.8 else "#C7CBD1"))
    strip(genes, gg, start, end, out, label_fn=lf, colour_fn=cf, height=3.6)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, strain, f"GECCO-only candidate on {k}")
    picture(s, out, M, 0.95, 7.9, 3.6, f"GECCO cluster {c.get('cluster_id', '')}")
    P = [[(f"{strain} / {k} / {start:,}-{end:,} / GECCO {c.get('cluster_id', '').split('_')[-2:] and '_'.join(c.get('cluster_id', '').split('_')[-2:])}", True, INK)],
         [("GECCO: ", True, ACCENT), (f"type {c.get('type', '') or 'unclassified'}; average p {float(c['average_p']):.3f}; "
                                     f"{len(genes)} genes over {(end - start) / 1000:.1f} kb.", False, INK)],
         [("Core enzyme domains: ", True, ACCENT), (", ".join(core) or "none of the listed core families", False, INK)],
         [("Most frequent domains: ", True, ACCENT), ("; ".join(f"{d} x{n}" for d, n in doms.most_common(8)), False, INK)],
         [("antiSMASH: ", True, ACCENT), ("no region overlaps this stretch. A candidate to check, not a call: a protein-family "
                                         "model scored it, not a biosynthetic rule.", False, MUTED)]]
    text(s, 8.6, 0.95, SW - 8.6 - 0.3, SH - 1.1, P, size=10, space=3)
    s.notes_slide.notes_text_frame.text = (f"GECCO cluster {c.get('cluster_id', '')}, matched on {k}. Red arrows carry a core "
                                           "enzyme domain (CORE_PFAM); blue arrows have per-gene p of 0.8 or more.")


GALLERY_COLS, GALLERY_ROWS = 4, 3
STRUCT_MIN_W, STRUCT_MAX_W = 2.2, 3.2  # inches beside the map for the region's reference structure


def structure_assets(D, assets):
    """This strain's STRUCTURE_ASSETS.json entries, read once; [] when absent or made for another strain."""
    if "_structures" not in D:
        path = Path(D["src"].get("chemical_structures_json") or assets / "chemical_structures" / "STRUCTURE_ASSETS.json")
        data = json.loads(path.read_text()) if path.exists() else {}
        D["_structures"] = (path.parent, data.get("entries") or []) if data.get("strain") in (None, D["strain"]) else (path.parent, [])
    return D["_structures"]


def region_structure(D, assets, bgc, contig, map_reference=""):
    """(entry, binding, png) for the reference structure drawn beside this BGC's locus map. The structure must be the
    compound of the MIBiG cluster the map is drawn against (map_reference, e.g. BGC0001875): a map of one cluster beside
    the structure of another misleads (one region showed a nargenicin A1 map with gargantulide B, the region's top
    KnownClusterBlast hit). The entry bound to this BGC and contig wins, else the same accession bound to another BGC of
    the strain (the drawing is reference chemistry, the same for every binding; its KCB numbers are then left off). None when the map's reference has no
    drawing. Without a map reference, the earlier rule: the best-ranked KCB reference bound to this BGC and contig, then
    the larger share of matched query proteins."""
    base, entries = structure_assets(D, assets)
    want = re.sub(r"\.\d+$", "", (map_reference or "").strip())
    if want:
        hit = None
        for e in entries:
            png = base / e["representative_png"] if e.get("drawing_available") and e.get("representative_png") else None
            if not (png and png.exists()) or re.sub(r"\.\d+$", "", e.get("mibig_accession") or "") != want:
                continue
            bs = e.get("locus_bindings") or []
            own = [b for b in bs if b.get("bgc_alias") == bgc and b.get("full_contig") == contig]
            b = own[0] if own else {}   # another BGC's KCB numbers must not print under this BGC's structure
            if hit is None or (own and not hit[0]):
                hit = (bool(own), e, b, png)
        return hit[1:] if hit else None
    best = None
    for e in entries:
        png = base / e["representative_png"] if e.get("drawing_available") and e.get("representative_png") else None
        if not (png and png.exists()):
            continue
        for b in e.get("locus_bindings") or []:
            if b.get("bgc_alias") != bgc or b.get("full_contig") != contig:
                continue
            m = b.get("reference_match_metrics") or {}
            try:
                rank = int(b.get("kcb_rank") or 99)
            except ValueError:
                rank = 99
            key = (rank, -float(m.get("query_genes_matched_fraction") or 0))
            if best is None or key < best[0]:
                best = (key, e, b, png)
    return best[1:] if best else None


def trimmed(png, out_dir):
    """A copy of a structure drawing with its white margin cut to 4%, so a small molecule fills its box."""
    from PIL import Image, ImageOps
    out = Path(out_dir) / png.name
    if not out.exists():
        im = Image.open(png).convert("RGB")
        box = ImageOps.invert(im).getbbox()
        if box:
            pad = int(0.04 * max(box[2] - box[0], box[3] - box[1]))
            box = (max(0, box[0] - pad), max(0, box[1] - pad), min(im.width, box[2] + pad), min(im.height, box[3] + pad))
            im = im.crop(box)
        out.parent.mkdir(parents=True, exist_ok=True)
        im.save(out)
    return out


def structure_box(s, D, e, b, png, x, y, w, h):
    """The reference drawing with its name, accession, own KCB numbers and the claim boundary, as one group."""
    g = s.shapes.add_group_shape()
    g.name = f"reference structure {e.get('representative_name', '')} {e.get('mibig_accession', '')}"[:250]
    box = _InGroup(g)
    text_h = 0.62
    ph = picture(box, png, x, y, w, h - text_h - 0.03, f"{e.get('representative_name', '')} reference structure")
    m = b.get("reference_match_metrics") or {}
    nums = (f" \u00b7 KCB {m.get('matched_query_genes')}/{m.get('query_gene_denominator')}"
            + (f" \u00b7 {float(m['median_best_hit_identity_percent']):.1f}%" if m.get("median_best_hit_identity_percent") is not None else "")
            if m.get("state") == "REFERENCE_MATCH_METRICS_BOUND" else "")
    paras = [[(e.get("representative_name") or e.get("mibig_accession", ""), True, INK)],
             [(f"{e.get('mibig_accession', '')}{nums}", False, ACCENT)],
             [(f"Reference chemistry; production by {D['strain']} is unconfirmed.", False, MUTED)]]
    size, fit, _ = fit_paras(paras, w, text_h, sizes=(8, 7.5, 7))
    text(box, x, y + ph + 0.05, w, text_h, fit, size=size, space=0.5)  # directly under the drawing


class _InGroup:
    """Lets text() and picture() draw into a group shape, so each reference is one removable object."""

    def __init__(self, group):
        self.shapes = group.shapes


def _binding_lines(b):
    """One line per bound locus: its alias, contig and region, then its own reference-specific KCB metrics."""
    node = "_".join(str(b.get("full_contig", "")).split("_")[:2]) or "?"
    head = f"{b.get('bgc_alias', '?')} \u00b7 {node} \u00b7 {b.get('region', '?')}"
    m = b.get("reference_match_metrics") or {}
    if m.get("state") == "REFERENCE_MATCH_METRICS_BOUND":
        ident = m.get("median_best_hit_identity_percent")
        line = (f"KCB {m.get('matched_query_genes')}/{m.get('query_gene_denominator')}"
                + (f" \u00b7 {float(ident):.1f}%" if ident is not None else ""))
    else:
        line = "match metrics not bound"
    return [(head, True, INK), ("  " + line, False, INK)]


def structure_gallery(prs, D, assets):
    """MIBiG reference structures for this strain's KnownClusterBlast references, from
    <assets>/chemical_structures/STRUCTURE_ASSETS.json (staged per strain; merged into the assets folder before a build).

    Every bound reference is shown at any percentage, with its own locus line and metrics: no identity threshold, filter
    or mark (4 Oct). Drawn entries first, undrawn last with a short placeholder and no MW. Each reference is one
    native group. Reference chemistry only; production by the strain is unconfirmed. Returns (slides, entries, drawn)."""
    from pptx.util import Inches
    path = Path(D["src"].get("chemical_structures_json") or assets / "chemical_structures" / "STRUCTURE_ASSETS.json")
    if not path.exists():
        return 0, 0, 0
    data = json.loads(path.read_text())
    strain = D["strain"]
    if data.get("strain") not in (None, strain):
        emit(f"{strain}: structure assets are for {data.get('strain')}; gallery skipped", file=sys.stderr)
        return 0, 0, 0
    entries = data.get("entries") or []
    entries = [e for e in entries if e.get("drawing_available")] + [e for e in entries if not e.get("drawing_available")]
    per = GALLERY_COLS * GALLERY_ROWS
    top, foot = 0.95, 0.55
    cw = (SW - 2 * M) / GALLERY_COLS
    rh = (SH - top - foot) / GALLERY_ROWS
    draw_h, gap = 0.95, 0.04
    pages = -(-len(entries) // per)
    each = -(-len(entries) // pages) if pages else 0  # spread evenly: no near-empty last page
    for k in range(pages):
        page = entries[k * each:(k + 1) * each]
        s = prs.slides.add_slide(prs.slide_layouts[6])
        title(s, strain, f"MIBiG reference structures ({k + 1}/{pages})" if pages > 1 else "MIBiG reference structures")
        notes = []
        for i, e in enumerate(page):
            x = M + (i % GALLERY_COLS) * cw
            y = top + (i // GALLERY_COLS) * rh
            g = s.shapes.add_group_shape()
            g.name = (e.get("recommended_powerpoint_group_name") or e.get("representative_name") or e.get("mibig_accession", ""))[:250]
            box = _InGroup(g)
            img = path.parent / e["representative_png"] if e.get("drawing_available") and e.get("representative_png") else None
            if img and img.exists():
                picture(box, img, x + 0.05, y, cw - 0.25, draw_h, f"{e.get('representative_name', '')} reference structure")
            else:
                text(box, x + 0.05, y + draw_h / 2 - 0.1, cw - 0.25, 0.25, [[("No single structure bound", False, WARN)]], size=9)
            cls = ", ".join(e.get("reference_biosynthesis_classes") or []) or e.get("reference_class_label") or "class not given"
            mw = e.get("representative_average_mw_g_mol")
            paras = [[(e.get("representative_name") or e.get("mibig_accession", ""), True, INK)],
                     [(f"{cls} \u00b7 {e.get('mibig_accession', '')}", False, ACCENT)],
                     [((f"MW {float(mw):.2f} g/mol \u00b7 " if mw else "") + (e.get("reference_organism") or ""), False, MUTED)]]
            loci = [_binding_lines(b) for b in e.get("locus_bindings") or []]
            room = rh - draw_h - gap - 0.05
            shown = len(loci)
            while True:  # whole locus lines only: a locus never loses its numbers to the box edge
                more = len(loci) - shown
                fit = paras + loci[:shown] + ([[(f"+{more} more locus line(s) in the notes", False, MUTED)]] if more else [])
                size, fit2, over = fit_paras(fit, cw - 0.25, room, sizes=(8.5, 8, 7.5))
                if not over or shown == 0:
                    break
                shown -= 1
            fit = fit2
            text(box, x + 0.05, y + draw_h + gap, cw - 0.25, rh - draw_h - gap - 0.05, fit, size=size, space=0.5)
            notes += [f"{e.get('representative_name', '')} ({e.get('mibig_accession', '')})"] + list(e.get("notes_lines") or []) + [""]
        text(s, M, SH - foot + 0.08, SW - 2 * M, 0.4, [[(
            f"Reference chemistry. Production by {strain} is unconfirmed. Per locus: KCB a/b = this region's proteins with a hit "
            "to the reference / all its proteins; % = median identity of each protein's best hit to that reference. Full "
            "contig identities, other compound names and identifiers are in the speaker notes.", False, MUTED)]], size=8)
        s.notes_slide.notes_text_frame.text = "\n".join(notes) + f"\nSource: {path}"
    return pages, len(entries), sum(1 for e in entries if e.get("drawing_available"))


def locus_comparison_slide(prs, D, bgc):
    """A full-width locus comparison of this region with its gap-rescue MIBiG reference (and partner contigs), drawn by
    mamey.figures.locus_comparison from <locus_maps_dir>/<BGC>/render/comparison.png. Added after the region's slides."""
    d = D["src"].get("locus_maps_dir")
    png = Path(d) / bgc / "render" / "comparison.png" if d else None
    if not (png and png.exists()):
        return False
    inv = D["inv"][bgc]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, D["strain"], f"{bgc} · locus comparison with its MIBiG reference")
    h = picture(s, png, M, 0.85, SW - 2 * M, SH - 1.55, f"{bgc} locus comparison")
    text(s, M, min(0.9 + h, SH - 0.62), SW - 2 * M, 0.55, [[(
        f"{D['strain']} / {inv['Contig']} / {inv['antiSMASH_Region']} / {bgc}. Ribbons join each gene to its best protein hit "
        "in the gap-rescue search (DIAMOND, ultra-sensitive); the percent after an AS gene's label is that hit's identity. "
        "Other contigs are separate tracks, never joined." + same_locus_pointer(D, bgc),
        False, MUTED)]], size=8)
    D.setdefault("_locus_slides", []).append(bgc)
    return True


def same_locus_pointer(D, bgc):
    """'Best MIBiG matches for this locus: see BGCnnn.' when this BGC's multi-contig locus was drawn once, under an
    earlier BGC whose figure exists ('' otherwise)."""
    d = D["src"].get("multi_ref_dir")
    f = Path(d) / bgc / "SAME_LOCUS_AS.txt" if d else None
    if not (f and f.exists()):
        return ""
    c = f.read_text().strip()
    return f" Best MIBiG matches for this locus: see the {c} slide." if (Path(d) / c / "slide.png").exists() else ""


def multi_reference_slides(prs, D, bgc):
    """The region against its best MIBiG clusters, and against reference-genome antiSMASH regions where that set exists:
    tools/multi_reference_comparison.py output, <multi_ref_dir>/<BGC>/ and <ref_genome_dir>/<BGC>/ (slide.png and
    caption.txt). One slide each, after the locus comparison."""
    made = 0
    for key, what in (("multi_ref_dir", "best MIBiG matches"), ("ref_genome_dir", "best reference-genome matches")):
        d = D["src"].get(key)
        png = Path(d) / bgc / "slide.png" if d else None
        if not (png and png.exists()):
            continue
        cap = (png.parent / "caption.txt").read_text().strip() if (png.parent / "caption.txt").exists() else ""
        also = (png.parent / "ALSO_BGCS.txt").read_text().split() if (png.parent / "ALSO_BGCS.txt").exists() else []
        s = prs.slides.add_slide(prs.slide_layouts[6])
        title(s, D["strain"], f"{bgc} · {what}" + (f" (one locus with {', '.join(also)})" if also else ""))
        h = picture(s, png, M, 0.85, SW - 2 * M, SH - 2.05, f"{bgc} {what}")
        if cap:
            text(s, M, min(0.95 + h, SH - 1.05), SW - 2 * M, 1.0, [[(cap, False, INK)]], size=8)
        D.setdefault("_multi_ref_slides", []).append(f"{bgc}:{key}")
        made += 1
        rib = png.parent / "ribbon" / "render" / "comparison.png"
        if key == "multi_ref_dir" and rib.exists():  # the top match with ribbons, when the earlier ribbon slide shows another
            s2 = prs.slides.add_slide(prs.slide_layouts[6])
            title(s2, D["strain"], f"{bgc} · ribbons with its best MIBiG match")
            h2 = picture(s2, rib, M, 0.85, SW - 2 * M, SH - 1.55, f"{bgc} ribbons with its best MIBiG match")
            text(s2, M, min(0.9 + h2, SH - 0.62), SW - 2 * M, 0.55, [[(
                "Ribbons join each gene to its best protein hit in the multi-reference search (DIAMOND blastp, ultra-sensitive; one "
                "pair per gene); the percent after an AS gene's label is that hit's identity. Other contigs are separate tracks, "
                "never joined.", False, MUTED)]], size=8)
            D.setdefault("_multi_ref_slides", []).append(f"{bgc}:ribbon")
    return made


def pointer_slide(prs, D):
    items = [[(p.get("label", "") + ": ", True, INK), (p.get("path", ""), False, INK)] + ([("  " + p["note"], False, MUTED)] if p.get("note") else [])
             for p in D["src"].get("pointers") or [] if p.get("path")]
    if not items:
        return
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, D["strain"], "where the rest lives")
    text(s, M, 1.0, SW - 2 * M, 5.8, items + [[("Package: ", True, ACCENT), (str(D["pkg"]), False, INK)]], size=13)


def gene_tables(D, order, path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    hdr = ["Gene", "aa", "antiSMASH role", "Domains (Pfam / TIGRFAM / rules)", "Description / smCOG", "GECCO p",
           "MIBiG gene (gap rescue)", "BLASTp best (stored)"]
    widths = [0.85, 0.45, 1.25, 2.4, 3.05, 0.6, 1.9, 1.9]
    for bgc in order:
        inv = D["inv"][bgc]
        gs = D["genes"].get(bgc, [])
        gp = {x["protein_id"]: x["average_p"] for x in D["ggenes"].get(node_key(inv["Contig"]), [])}
        _, st = gap_stats(D, bgc)
        mib = {k: f"{r['name']} {r['reference_product'][:22]} ({float(r['best_identity_pct']):.0f}%)"
               for k, r in one_row_per_locus([r for r in (st["table"] if st else []) if r["status"] == "PRESENT_IN_CORE"]).items()}
        for k in range(0, len(gs), 20):
            chunk = gs[k:k + 20]
            s = prs.slides.add_slide(prs.slide_layouts[6])
            title(s, D["strain"], f"{bgc} genes ({k + 1}-{k + len(chunk)} of {len(gs)})")
            text(s, M, 0.72, SW - 2 * M, 0.25, [[(f"{D['strain']} / {inv['Contig']} / {inv['antiSMASH_Region']} / {bgc}", False, MUTED)]], size=9)
            body = []
            for g in chunk:
                t = g["locus_tag"]
                doms = list(dict.fromkeys(x[1] for x in D["dom"].get(t, [])))[:3]
                desc = next((x[2] for x in D["dom"].get(t, []) if x[2]), "")
                cds = D["cds"].get(t, {})
                sm = re.search(r"SMCOG\d+:[^()]*", cds.get("gene_functions", "") or g.get("product_qualifier", ""))
                d2 = (desc[:48] + ("; " + sm.group(0).strip()[:40] if sm else "")) or (cds.get("product", "") or "")[:60]
                h = D["blastp"].get(t)
                body.append([t, g.get("aa_length", ""), g["gene_function_inference"], ", ".join(doms), d2,
                             f"{float(gp[t]):.2f}" if t in gp else "", mib.get(t, ""),
                             f"{h['org'] or h['acc']} {float(h['pid'] or 0):.0f}%" if h else ""])
            table(s, hdr, widths, body, y=1.0, font=7, row_h=0.24)
            s.notes_slide.notes_text_frame.text = ("Blank = no record in that source, not a negative result. Domains sorted by "
                                                   "E-value, best first.")
    prs.save(path)
    return len(prs.slides)


def build(src, out, tag=TAG_DEFAULT):
    need_libs()
    from pptx import Presentation
    from pptx.util import Inches
    D = load(src)
    strain = D["strain"]
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    stamp = datetime.date.today().isoformat()
    deck = out / f"{strain}_strain_slides_{tag}_{stamp}.pptx"
    if deck.exists():
        raise SystemExit(f"{deck} exists; a delivered deck is never overwritten (use a new --tag)")
    assets = out / f"{strain}_{tag}_assets"
    (assets / "maps").mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    overview(prs, D)
    trees(prs, D)
    landscape(prs, D, assets)
    order = sorted(D["inv"], key=lambda b: (-score(D, b), b))
    at = len(prs.slides._sldIdLst)  # the region table goes here once the slide numbers are known
    n_table = -(-len(order) // 15)
    if D["src"].get("bigscape_figure") and Path(D["src"]["bigscape_figure"]).exists():
        figure_slide(prs, strain, "BiG-SCAPE families shared with reference genomes", Path(D["src"]["bigscape_figure"]),
                     "A filled cell means that genome has a region in the same BiG-SCAPE family. Same family = similar domain "
                     "architecture, not compound or species identity.")
    genome_box = {}
    slide_of = {}
    dropped = {b: left_out(D, b) for b in order}
    dropped = {b: w for b, w in dropped.items() if w}
    for n, b in enumerate([b for b in order if b not in dropped], 1):
        slide_of[b] = len(prs.slides._sldIdLst) + 1 + n_table  # counted after the build: continued slides shift them
        region_slide(prs, D, b, n, assets, genome_box)
        locus_comparison_slide(prs, D, b)
        multi_reference_slides(prs, D, b)
    for b in dropped:
        slide_of[b] = "none"
    for c in sorted(D["gonly"], key=lambda c: -float(c.get("average_p") or 0)):
        gecco_only_slide(prs, D, c, assets)
    gallery = structure_gallery(prs, D, assets)
    for f in D["src"].get("extra_figures") or []:
        if f.get("image") and Path(f["image"]).exists():
            figure_slide(prs, strain, f.get("heading", ""), Path(f["image"]), f.get("caption", ""))
    pointer_slide(prs, D)
    before = len(prs.slides._sldIdLst)
    region_table(prs, D, order, slide_of)
    ids = prs.slides._sldIdLst
    for k, el in enumerate(list(ids)[before:]):  # move the table slides to their place after the landscape
        ids.remove(el)
        ids.insert(at + k, el)
    prs.core_properties.title = f"{strain} strain slides {tag}"
    prs.save(deck)
    ng = gene_tables(D, order, out / f"{strain}_gene_tables_{tag}_{stamp}.pptx")
    receipt = {"strain": strain, "tag": tag, "deck": deck.name, "slides": len(prs.slides), "regions": len(order),
               "gecco_only": len(D["gonly"]), "gene_table_slides": ng, "no_region_slide": dropped,
               "structure_gallery": dict(zip(("slides", "entries", "drawn"), gallery)),
               "region_slide_structures": D.get("_region_structures", []),
               "locus_comparison_slides": D.get("_locus_slides", []), "multi_reference_slides": D.get("_multi_ref_slides", []), "adjudicated": bool(src.get("rescue_verdicts_tsv")),
               "sources": {k: v for k, v in src.items() if isinstance(v, str)}, "region_order": order}
    (out / f"{strain}_strain_slides_{tag}_{stamp}_RECEIPT.json").write_text(json.dumps(receipt, indent=1))
    return receipt


def audit_fit(deck):
    """Text-fit problems of a built deck: text past its box, past the slide bottom, or into a text box below it."""
    from pptx import Presentation
    EMU = 914400
    prs = Presentation(deck)
    sh_in = prs.slide_height / EMU
    bad = []
    for i, s in enumerate(prs.slides, 1):
        boxes = []
        for sh in s.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            paras, size, space = [], 9, 2
            for p in sh.text_frame.paragraphs:
                runs = [(r.text, bool(r.font.bold), "") for r in p.runs]
                if runs:
                    paras.append(runs)
                    size = max([r.font.size.pt for r in p.runs if r.font.size] + [0]) or size
                space = p.space_after.pt if p.space_after is not None else space
            x, y, w, h = (v / EMU for v in (sh.left, sh.top, sh.width, sh.height))
            need = text_height(paras, w, size, space)
            lab = sh.text_frame.text[:40].replace("\n", " ")
            boxes.append((x, y, w, need, lab))
            if y + need > sh_in + 0.02:
                bad.append(f"slide {i}: text runs {y + need - sh_in:.2f} in past the slide bottom: {lab!r}")
            elif need > h + 0.15:
                bad.append(f"slide {i}: text needs {need:.2f} in, box {h:.2f} in: {lab!r}")
        for b1 in boxes:
            for b2 in boxes:
                if b2[1] > b1[1] and min(b1[0] + b1[2], b2[0] + b2[2]) - max(b1[0], b2[0]) > 0.05 \
                        and b1[1] + b1[3] > b2[1] + 0.03:
                    bad.append(f"slide {i}: {b1[4]!r} runs into {b2[4]!r}")
    return bad


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("--sources", required=True, nargs="+", help="one sources JSON per strain")
    b.add_argument("--out", required=True)
    b.add_argument("--tag", default=TAG_DEFAULT)
    sub.add_parser("template")
    f = sub.add_parser("audit-fit", help="report text past its box, past the slide, or into another box")
    f.add_argument("decks", nargs="+")
    a = ap.parse_args(argv)
    if a.cmd == "audit-fit":
        n = 0
        for d in a.decks:
            bad = audit_fit(d)
            n += len(bad)
            emit(f"{Path(d).name}: {len(bad)} text-fit problems")
            for b_ in bad[:12]:
                emit("  ", b_)
        return 1 if n else 0
    if a.cmd == "template":
        emit(json.dumps(TEMPLATE, indent=2))
        return 0
    for sp in a.sources:
        p = Path(sp)
        r = build(resolve(json.load(open(p)), p.parent), a.out, a.tag)
        emit(f"{r['deck']}: {r['slides']} slides; {r['regions']} regions; {r['gecco_only']} GECCO-only; "
              f"{r['gene_table_slides']} gene-table slides")
    return 0


if __name__ == "__main__":
    sys.exit(main())
