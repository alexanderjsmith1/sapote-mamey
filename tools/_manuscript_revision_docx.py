"""Word (.docx) table helpers for tools/manuscript_revision.py; python-docx is imported only when a Word file is made or read.

python-docx is the optional `documents` extra (pyproject.toml). When it is absent, `require_docx()` raises
`DocxUnavailable` with the install hint, and the calling subcommand exits with a typed refusal instead of a traceback.
Importing this module performs no work.
"""
from __future__ import annotations

from typing import Iterable, Sequence

INSTALL_HINT = "python-docx is not installed; install the documents extra: pip install 'mamey[documents]' (README.md, Optional extras)"


class DocxUnavailable(RuntimeError):
    """python-docx could not be imported."""


def require_docx():
    """Return the python-docx module, or raise DocxUnavailable with the install hint."""
    try:
        import docx  # noqa: PLC0415 - optional dependency, imported on use
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise DocxUnavailable(INSTALL_HINT) from exc
    return docx


def new_landscape_document(font: str = "Arial", size_pt: float = 9, margin_in: float = 0.5):
    """A blank landscape Word document with one body font and even margins."""
    docx = require_docx()
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt

    d = docx.Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Inches(margin_in))
    d.styles["Normal"].font.name = font
    d.styles["Normal"].font.size = Pt(size_pt)
    return d


def add_table(d, head: Sequence[str], widths_in: Sequence[float], rows: Iterable[Sequence[str]]):
    """Append a gridded table: bold header row, one row per item, fixed column widths in inches."""
    from docx.shared import Inches

    t = d.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, head):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for vals in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, vals):
            c.text = "" if v is None else str(v)
    t.autofit = False
    for row in t.rows:
        for c, w in zip(row.cells, widths_in):
            c.width = Inches(w)
    return t


def read_tables(path) -> list[list[list[str]]]:
    """Every table in a .docx as rows of cell text (header row included), in document order."""
    docx = require_docx()
    out = []
    for t in docx.Document(str(path)).tables:
        out.append([[c.text for c in r.cells] for r in t.rows])
    return out
