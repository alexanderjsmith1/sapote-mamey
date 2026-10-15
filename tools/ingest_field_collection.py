#!/usr/bin/env python3
"""ingest_field_collection.py -- normalize the owner's field-collection workbook into a per-strain
collection-metadata TSV the rest of the engine can bind (placement host/origin tables, bioassay
cohort grouping, strain-metadata reconciliation).

FIRST-PASS / PRELIMINARY (peer-lane, first staged at v9.7.427). Reads the `Isolate Index` sheet of
`Field_Collection_Master_Data.xlsx` -- one row per AS code with the collector's own fields
(Experiment, Sample, Description, Simple Group, Category, Location, Date) -- and emits a normalized
TSV keyed by canonical AS id.

GOVERNANCE (governing, do not soften):
  * This is DESCRIPTIVE COLLECTION-RECORD provenance -- the collector's field note of what the
    isolate was picked off and where. It is NOT a tested host association and NOT authority over
    deposited GenBank /host metadata; it is a source to RECONCILE against those, per the standing
    rulings ([[deposited-genbank-is-a-provenance-source]], [[strain-summary-is-not-an-ecology-source]]).
  * Raw values are copied VERBATIM. This tool does not categorize, does not force a value into a
    fixed palette, and does not invent a mapping ([[report-isolation-source-do-not-categorize]]).
    `category_raw` / `location_raw` are exactly the cells; any bucketing is a downstream decision.
  * A duplicated AS code with conflicting non-empty values is REFUSED (typed), never silently
    merged -- an identity conflict is a hold, not a pick.

No workspace path is baked in: the workbook is an explicit --input. Emission is via sys.stdout/
sys.stderr only (no print()/emit()), keeping the strict-health print_calls ratchet flat.

EXPERIMENT-SHEETS LAYOUT (--layout experiment-sheets, added v9.7.452). A lab workbook kept as one
sheet per experiment (or per run of experiments) has no `Isolate Index`. This mode walks every
sheet (minus any --exclude-sheet pattern) and reads each row that names an AS/SID code:
  * AS codes are found with a CASE-SENSITIVE pattern (`AS-1`, `AS 1`, `AS1`, `SID-3`), so prose
    such as "same tube as 9" never becomes a strain. A cell naming several codes
    ("AS-1, AS-2", "AS-3 + AS-4") gives one record per code; `as_cell_raw` keeps the cell.
  * If one of a sheet's first 3 rows is a header (Date / Experiment number / Sample Number /
    Description / Location), those columns map to date / experiment / sample / description_raw /
    location_raw. Otherwise the experiment is the cell reading "Exp N" / "Experiment N" (else the
    sheet title), the sample is the "#N" cell, and description_raw is every remaining text cell of
    the row joined with " | ", verbatim; location_raw stays blank because no cell is named as one.
  * Rows with no AS code and no sample (headings, colony or site lines) are kept, verbatim, as
    `context_raw` for the rows of the same experiment block below them; a date on such a row is
    the block date. Nothing in them is parsed into a category or a place.
  * Every value carries its origin in `source_refs` (`field=Sheet!A1;...`).
  * A code seen on several rows is merged only when the rows agree (blank fills blank). If any of
    experiment, sample, date, description_raw or location_raw disagree, the strain is NOT written
    to the main TSV; it goes to the --holds TSV with code FIELD_COLLECTION_IDENTITY_CONFLICT and
    every source value, for a person to resolve. Experiment and sample compare by their number,
    so "Exp 7" and "Experiment 7" agree.

Usage:
  ingest_field_collection.py --input <Field_Collection_Master_Data.xlsx> --out <collection.tsv>
      [--sheet "Isolate Index"]
  ingest_field_collection.py --layout experiment-sheets --input <lab_workbook.xlsx>
      --out <collection.tsv> [--holds <holds.tsv>] [--exclude-sheet REGEX ...]
"""
from __future__ import annotations

import argparse
import collections
import datetime
import re
import sys

SOURCE_COLUMNS = ("AS code", "Experiment", "Sample", "Description", "Simple Group",
                  "Category", "Location", "Date")
# output column -> source column (raw is preserved; only the AS id is canonicalized)
OUT_FIELDS = (
    ("as_id", "AS code"),
    ("category_raw", "Category"),
    ("simple_group_raw", "Simple Group"),
    ("location_raw", "Location"),
    ("date", "Date"),
    ("experiment", "Experiment"),
    ("sample", "Sample"),
    ("description_raw", "Description"),
)
_TYPED_CODES = frozenset({"FIELD_COLLECTION_SCHEMA", "FIELD_COLLECTION_EMPTY",
                          "FIELD_COLLECTION_IDENTITY_CONFLICT"})


def _refuse(code: str, detail: str) -> "NoReturn":  # type: ignore[name-defined]
    sys.stderr.write(f"{code}: {detail}\n")
    raise SystemExit(2)


def _canon_as(raw) -> str:
    """Canonicalize a strain cell to AS-<n> / SID-<n>; '' if it is not a strain code."""
    s = ("" if raw is None else str(raw)).strip()
    if not s:
        return ""
    m = re.search(r"(AS|SID)[\s\-_]*0*([0-9]+)", s.upper())
    return f"{m.group(1)}-{m.group(2)}" if m else ""


def _clean(v) -> str:
    return "" if v is None else str(v).strip()


def load_isolate_rows(path: str, sheet: str):
    """Return list[dict] of the sheet's rows keyed by the source header, or raise a typed code."""
    try:
        import openpyxl  # declared bundle dep (>=3.1.2)
    except ImportError:
        _refuse("FIELD_COLLECTION_SCHEMA", "openpyxl is required to read the .xlsx workbook")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 -- surface the real reason, do not swallow
        _refuse("FIELD_COLLECTION_SCHEMA", f"cannot open workbook {path}: {exc}")
    if sheet not in wb.sheetnames:
        _refuse("FIELD_COLLECTION_SCHEMA",
                f"sheet {sheet!r} not found; available: {wb.sheetnames}")
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("FIELD_COLLECTION_EMPTY")
    header = [_clean(c) for c in rows[0]]
    missing = [c for c in SOURCE_COLUMNS if c not in header]
    if missing:
        raise ValueError("FIELD_COLLECTION_SCHEMA")
    idx = {name: header.index(name) for name in SOURCE_COLUMNS}
    out = []
    for r in rows[1:]:
        rec = {name: (r[i] if i < len(r) else None) for name, i in idx.items()}
        out.append(rec)
    return out


def normalize(rows):
    """Collapse to one record per canonical AS id; refuse on a genuine value conflict."""
    by_id = {}
    dropped_no_code = 0
    for rec in rows:
        aid = _canon_as(rec.get("AS code"))
        if not aid:
            dropped_no_code += 1
            continue
        norm = {out: _clean(rec.get(src)) for out, src in OUT_FIELDS}
        norm["as_id"] = aid
        if aid in by_id:
            prev = by_id[aid]
            for k in norm:
                if k == "as_id":
                    continue
                if norm[k] and prev[k] and norm[k] != prev[k]:
                    raise ValueError("FIELD_COLLECTION_IDENTITY_CONFLICT")
                if norm[k] and not prev[k]:
                    prev[k] = norm[k]
        else:
            by_id[aid] = norm
    if not by_id:
        raise ValueError("FIELD_COLLECTION_EMPTY")
    return by_id, dropped_no_code


# ---- experiment-sheets layout -------------------------------------------------------------------
# Case-sensitive on purpose: a lower-case "as 33" in prose is English, not a strain code.
_AS_IN_TEXT = re.compile(r"\b(AS|SID)[\-_ ]?0*([0-9]+)\b")
_EXP_CELL = re.compile(r"^\s*exp(?:eriment)?\.?\s*0*([0-9]+)\s*$", re.IGNORECASE)
_SAMPLE_CELL = re.compile(r"^\s*#\s*0*([0-9]+)\s*$")
_HEADER_MAP = {"date": "date", "experiment number": "experiment", "experiment": "experiment",
               "sample number": "sample", "sample": "sample", "description": "description_raw",
               "location": "location_raw"}
EXTRA_FIELDS = ("context_raw", "as_cell_raw", "source_sheet", "source_refs", "n_source_rows")
CONFLICT_FIELDS = ("experiment", "sample", "date", "description_raw", "location_raw")
HOLD_FIELDS = ("as_id", "hold_code", "conflict_fields", "n_source_rows", "source_values")


def _cell_text(v) -> str:
    """Render one cell verbatim: dates as ISO, integer-valued floats as integers, text stripped."""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _col_letter(i: int) -> str:
    out = ""
    i += 1
    while i:
        i, rem = divmod(i - 1, 26)
        out = chr(65 + rem) + out
    return out


def _ref(sheet: str, row_i: int, col_i: int) -> str:
    return f"{sheet}!{_col_letter(col_i)}{row_i + 1}"


def _as_codes(text: str):
    seen = []
    for m in _AS_IN_TEXT.finditer(text):
        aid = f"{m.group(1)}-{m.group(2)}"
        if aid not in seen:
            seen.append(aid)
    return seen


def _num_key(v: str) -> str:
    """Comparison key for experiment / sample: the number when the cell is one, else the text."""
    m = _EXP_CELL.match(v) or _SAMPLE_CELL.match(v) or re.match(r"^\s*0*([0-9]+)\s*$", v)
    return m.group(1) if m else v


def _find_header(rows):
    for hi, r in enumerate(rows[:3]):
        cols = {}
        for ci, c in enumerate(r):
            key = _clean(c).lower()
            if key in _HEADER_MAP and _HEADER_MAP[key] not in cols.values():
                cols[ci] = _HEADER_MAP[key]
        if len(cols) >= 3 and "experiment" in cols.values():
            return hi, cols
    return None, {}


def read_experiment_sheets(path: str, exclude_patterns=()):
    """Return (records, sheets_read): one record per (AS code, source row), values verbatim."""
    try:
        import openpyxl  # declared bundle dep (>=3.1.2)
    except ImportError:
        _refuse("FIELD_COLLECTION_SCHEMA", "openpyxl is required to read the .xlsx workbook")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 -- surface the real reason, do not swallow
        _refuse("FIELD_COLLECTION_SCHEMA", f"cannot open workbook {path}: {exc}")
    excl = [re.compile(p) for p in exclude_patterns]
    records, sheets_read = [], []
    for sheet in wb.sheetnames:
        if any(p.search(sheet) for p in excl):
            continue
        sheets_read.append(sheet)
        rows = [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
        hi, hcols = _find_header(rows)
        title_exp = _EXP_CELL.match(sheet)
        block_exp, context, ctx_refs, block_date = "", [], [], ("", "")
        for ri, r in enumerate(rows):
            if hi is not None and ri <= hi:
                continue
            texts = [(ci, _cell_text(c)) for ci, c in enumerate(r)]
            texts = [(ci, t) for ci, t in texts if t]
            if not texts:
                continue
            vals = {}  # field -> (value, ref)
            for ci, field in hcols.items():
                t = _cell_text(r[ci]) if ci < len(r) else ""
                if t:
                    vals[field] = (t, _ref(sheet, ri, ci))
            as_cells = [(ci, t) for ci, t in texts if _AS_IN_TEXT.search(t)]
            as_cols = {ci for ci, _ in as_cells}
            used = set(as_cols) | set(hcols)
            if "experiment" not in vals:
                for ci, t in texts:
                    if ci not in used and _EXP_CELL.match(t):
                        vals["experiment"] = (t, _ref(sheet, ri, ci))
                        used.add(ci)
                        break
            if "sample" not in vals:
                exp_ci = None
                for ci, t in texts:
                    if "experiment" in vals and vals["experiment"][1] == _ref(sheet, ri, ci):
                        exp_ci = ci
                for ci, t in texts:
                    if ci in used:
                        continue
                    bare_int = (exp_ci is not None and ci == exp_ci + 1
                                and isinstance(r[ci], (int, float))
                                and float(r[ci]).is_integer())
                    if _SAMPLE_CELL.match(t) or bare_int:
                        vals["sample"] = (t, _ref(sheet, ri, ci))
                        used.add(ci)
                        break
            if "date" not in vals:
                for ci, c in enumerate(r):
                    if ci not in used and isinstance(c, (datetime.datetime, datetime.date)):
                        vals["date"] = (_cell_text(c), _ref(sheet, ri, ci))
                        used.add(ci)
                        break
            exp_now = vals.get("experiment", ("", ""))[0]
            if exp_now and block_exp and _num_key(exp_now) != _num_key(block_exp):
                context, ctx_refs, block_date = [], [], ("", "")
            if exp_now:
                block_exp = exp_now
            is_sample_row = "sample" in vals and _num_key(vals["sample"][0]).isdigit()
            if not as_cells and not is_sample_row:
                # heading / colony / site line: verbatim context for the block below it
                for ci, t in texts:
                    if ci in hcols and hcols[ci] == "experiment":
                        continue
                    if "date" in vals and vals["date"][1] == _ref(sheet, ri, ci):
                        block_date = vals["date"]
                        continue
                    if _EXP_CELL.match(t):
                        continue
                    context.append(t)
                    ctx_refs.append(_ref(sheet, ri, ci))
                continue
            if not as_cells:
                continue
            if "description_raw" not in vals:
                rest = [(ci, t) for ci, t in texts if ci not in used]
                if rest:
                    vals["description_raw"] = (" | ".join(t for _, t in rest),
                                               ",".join(_ref(sheet, ri, ci) for ci, _ in rest))
            if "experiment" not in vals:
                if block_exp:
                    vals["experiment"] = (block_exp, f"{sheet}!block")
                elif title_exp:
                    vals["experiment"] = (sheet.strip(), f"{sheet}!title")
            if "date" not in vals and block_date[0]:
                vals["date"] = block_date
            if context:
                vals["context_raw"] = (" | ".join(context), ",".join(ctx_refs))
            as_cell_raw = " ; ".join(t for _, t in as_cells)
            as_refs = ",".join(_ref(sheet, ri, ci) for ci, _ in as_cells)
            for ci, t in as_cells:
                for aid in _as_codes(t):
                    rec = {f: "" for f, _ in OUT_FIELDS}
                    rec["as_id"] = aid
                    for f, (v, _ref_) in vals.items():
                        rec[f] = v
                    rec["context_raw"] = vals.get("context_raw", ("", ""))[0]
                    rec["as_cell_raw"] = as_cell_raw
                    rec["source_sheet"] = sheet
                    refs = [f"as_id={as_refs}"] + [f"{f}={rf}" for f, (_v, rf) in sorted(vals.items())]
                    rec["source_refs"] = ";".join(refs)
                    records.append(rec)
    if not sheets_read:
        raise ValueError("FIELD_COLLECTION_EMPTY")
    return records, sheets_read


def merge_experiment_records(records):
    """One row per AS id when its source rows agree; otherwise a typed hold with every value."""
    groups = collections.OrderedDict()
    for rec in records:
        groups.setdefault(rec["as_id"], []).append(rec)
    merged, holds = {}, {}
    for aid, recs in groups.items():
        # the same code twice in one row (two AS cells) is one source row, not two
        uniq = list({r["source_refs"]: r for r in recs}.values())
        conflicts = []
        for f in CONFLICT_FIELDS:
            keyf = _num_key if f in ("experiment", "sample") else (lambda v: v)
            vs = {keyf(r[f]) for r in uniq if r[f]}
            if len(vs) > 1:
                conflicts.append(f)
        if conflicts:
            holds[aid] = {
                "as_id": aid, "hold_code": "FIELD_COLLECTION_IDENTITY_CONFLICT",
                "conflict_fields": ",".join(conflicts), "n_source_rows": str(len(uniq)),
                "source_values": " || ".join(
                    "; ".join([r["source_sheet"]] + [f"{f}={r[f]}" for f in CONFLICT_FIELDS if r[f]])
                    for r in uniq)}
            continue
        out = dict(uniq[0])
        for r in uniq[1:]:
            for f in [f for f, _ in OUT_FIELDS] + ["context_raw"]:
                if r[f] and not out[f]:
                    out[f] = r[f]
            if r["context_raw"] and r["context_raw"] not in out["context_raw"].split(" || "):
                out["context_raw"] = " || ".join(x for x in (out["context_raw"], r["context_raw"]) if x)
        out["source_sheet"] = " | ".join(dict.fromkeys(r["source_sheet"] for r in uniq))
        out["as_cell_raw"] = " || ".join(dict.fromkeys(r["as_cell_raw"] for r in uniq))
        out["source_refs"] = " || ".join(r["source_refs"] for r in uniq)
        out["n_source_rows"] = str(len(uniq))
        merged[aid] = out
    if not merged and not holds:
        raise ValueError("FIELD_COLLECTION_EMPTY")
    return merged, holds


def _write_tsv(path, fields, rows):
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write("\t".join(fields) + "\n")
        for row in rows:
            fh.write("\t".join(row.get(f, "").replace("\t", " ").replace("\n", " ")
                               for f in fields) + "\n")


def _aid_sort(aid: str):
    pre, _, n = aid.partition("-")
    return (pre, int(n) if n.isdigit() else 0, aid)


def main_experiment_sheets(a):
    records, sheets = read_experiment_sheets(a.input, a.exclude_sheet)
    merged, holds = merge_experiment_records(records)
    fields = [out for out, _ in OUT_FIELDS] + list(EXTRA_FIELDS)
    _write_tsv(a.out, fields, [merged[k] for k in sorted(merged, key=_aid_sort)])
    holds_path = a.holds or (a.out[:-4] if a.out.endswith(".tsv") else a.out) + ".holds.tsv"
    _write_tsv(holds_path, list(HOLD_FIELDS), [holds[k] for k in sorted(holds, key=_aid_sort)])
    sys.stdout.write(
        f"[field-collection] experiment-sheets: {len(sheets)} sheets read, {len(records)} "
        f"strain-row records; {len(merged)} strains -> {a.out}; {len(holds)} identity "
        f"conflicts held -> {holds_path} (FIELD_COLLECTION_IDENTITY_CONFLICT)\n")
    sys.stdout.write("[field-collection] NOTE: collection-record provenance, verbatim; reconcile "
                     "against deposited GenBank /host -- not authority over it, not a tested "
                     "host association.\n")
    return 0


def main():
    ap = argparse.ArgumentParser(allow_abbrev=False, description=__doc__)
    ap.add_argument("--input", required=True, help="Field_Collection_Master_Data.xlsx")
    ap.add_argument("--out", required=True, help="output normalized collection-metadata TSV")
    ap.add_argument("--sheet", default="Isolate Index", help="worksheet name (default 'Isolate Index')")
    ap.add_argument("--layout", choices=("isolate-index", "experiment-sheets"),
                    default="isolate-index",
                    help="isolate-index (default): one header sheet, one row per AS code; "
                         "experiment-sheets: every sheet, rows that name an AS code")
    ap.add_argument("--holds", default=None,
                    help="experiment-sheets: TSV of identity-conflict holds "
                         "(default <out>.holds.tsv)")
    ap.add_argument("--exclude-sheet", action="append", default=[], metavar="REGEX",
                    help="experiment-sheets: skip sheets whose name matches (repeatable)")
    a = ap.parse_args()
    if a.layout == "experiment-sheets":
        return main_experiment_sheets(a)

    rows = load_isolate_rows(a.input, a.sheet)
    by_id, dropped = normalize(rows)

    fields = [out for out, _ in OUT_FIELDS]
    with open(a.out, "w", encoding="utf-8", newline="") as fh:
        fh.write("\t".join(fields) + "\n")
        for aid in sorted(by_id):
            fh.write("\t".join(by_id[aid][f] for f in fields) + "\n")

    cats = collections.Counter(v["category_raw"] for v in by_id.values() if v["category_raw"])
    sys.stdout.write(
        f"[field-collection] {len(by_id)} strains -> {a.out}"
        f" ({dropped} non-strain rows skipped);"
        f" top raw categories: {', '.join(f'{k}={n}' for k, n in cats.most_common(6))}\n")
    sys.stdout.write("[field-collection] NOTE: collection-record provenance, verbatim; reconcile "
                     "against deposited GenBank /host -- not authority over it, not a tested "
                     "host association.\n")
    return 0


def _run_or_refuse():
    try:
        return main()
    except ValueError as exc:
        code = str(exc)
        if code not in _TYPED_CODES:
            raise
        _refuse(code, "input contract violated; no output was written. The code above names the "
                      "contract that failed.")


if __name__ == "__main__":
    sys.exit(_run_or_refuse())
