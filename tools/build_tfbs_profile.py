#!/usr/bin/env python3
"""build_tfbs_profile.py — per-strain transcription-factor binding-site (TFBS) regulator profile.

Banks genome-wide TFBS motif counts (from gene_data.json['tfbs']) into a sortable TFBS_Profile sheet: one
row per strain, one column per regulator family, plus total and per-Mbp density (TFBS counts scale with
genome size, so density is the cross-strain-comparable metric). Sorted by SARP_BTAD_like — the pathway-
specific antibiotic activators, the most directly actionable regulatory signal (a SARP-rich strain is more
likely to carry actively-regulated antibiotic pathways). GBL/AdpA (master secondary-metabolism switch) and
DasR (GlcNAc/chitin-responsive antibiotic regulator — the regulatory counterpart to the chitinolytic
ecology) are the dominant families; both are near-uniform baselines, so density highlights the outliers.

Usage: python tools/build_tfbs_profile.py --workbook <xlsx> --banked-dir <dir>
"""
import os as _os, sys as _sys  # v9.7.407: resolve the tools-local emitter from any cwd
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from _console import emit  # noqa: E402
import argparse, json, os

from _wbio import atomic_save


def _read_json(_path, *, encoding="utf-8"):
    """P3b: context-managed JSON read; closes the handle a bare open() leaked."""
    import json as _json
    with open(_path, encoding=encoding) as _fh:
        return _json.load(_fh)

# display order: actionable/specific first, then master regulators, then stress/metal/other
REG_ORDER=['SARP_BTAD_like','GBL_AdpA_like','DasR_like_palindrome','BldD_like','PhoP_box_like',
           'DmdR_iron_box_like','IolR_like','LexA_SOS_like','ANR_FNR_like','PAS_LuxR_like']


_PLACEHOLDER={'not','unknown','unclassified','uncultured','unidentified','na','n/a','none',''}
def _genus(org):
    """First token of the organism string, but robust to placeholder taxonomies like 'not verified'
    (which naive split()[0] turns into 'not'). Returns 'unclassified' for placeholders."""
    org=(org or '').strip()
    if not org: return 'unclassified'
    first=org.split()[0]
    if first.lower() in _PLACEHOLDER: return 'unclassified'
    if first.lower()=='candidatus' and len(org.split())>1: return org.split()[1]
    return first

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--workbook', required=True); ap.add_argument('--banked-dir', default='cohort')
    a=ap.parse_args()
    gene=_read_json(os.path.join(a.banked_dir,'gene_data.json'))
    tfbs=gene.get('tfbs',{})
    if not tfbs:
        emit('  [tfbs] no tfbs data in gene_data.json — skipping'); return
    bgc=_read_json(os.path.join(a.banked_dir,'bgc_data.json')); strains=bgc['strains']
    # discover any regulator families beyond the known order
    seen=set()
    for d in tfbs.values(): seen|=set(d)
    cols=REG_ORDER+[r for r in sorted(seen) if r not in REG_ORDER]
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb=openpyxl.load_workbook(a.workbook)
    if 'TFBS_Profile' in wb.sheetnames: del wb['TFBS_Profile']
    ws=wb.create_sheet('TFBS_Profile')
    hdr=['strain','genus','genome_Mbp']+cols+['total_TFBS','TFBS_per_Mbp','SARP_per_Mbp']
    ws.append(hdr)
    for c in range(1,len(hdr)+1):
        cell=ws.cell(1,c); cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='2E5E7E')
    rows=[]
    for sid,d in tfbs.items():
        if sid not in strains: continue
        s=strains[sid]; mbp=(s.get('genome_bp',0) or 0)/1e6 or 0.01
        genus=_genus(s.get('organism',''))
        vals=[d.get(c,0) for c in cols]; tot=sum(d.values())
        sarp=d.get('SARP_BTAD_like',0)
        rows.append([sid,genus,round(mbp,2)]+vals+[tot,round(tot/mbp,1),round(sarp/mbp,2)])
    sarp_i=hdr.index('SARP_per_Mbp')
    rows.sort(key=lambda r:-r[sarp_i])  # actionable: SARP density first
    HI=PatternFill('solid',fgColor='E6F4EA')
    import statistics
    sarps=[r[sarp_i] for r in rows]; mu=statistics.mean(sarps); sd=statistics.pstdev(sarps) or 1
    for r in rows:
        ws.append(r)
        if (r[sarp_i]-mu)/sd > 1.3:  # SARP-rich = pathway-activator-rich, flag green
            for c in range(1,len(hdr)+1): ws.cell(ws.max_row,c).fill=HI
    ws.freeze_panes='D2'
    widths=[9,13,10]+[12]*len(cols)+[11,13,12]
    for i,w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
    ws.append([])
    ws.append(['NOTE','genome-wide motif counts; sort by any regulator column. SARP_BTAD = pathway-specific '
               'activators (most actionable); GBL/AdpA + DasR are near-uniform baselines.']+['']*(len(hdr)-2))
    atomic_save(wb, a.workbook)
    emit(f"  TFBS_Profile sheet written: {len(rows)} strains x {len(cols)} regulator families "
          f"(sorted by SARP density; {sum(1 for r in rows if (r[sarp_i]-mu)/sd>1.3)} SARP-rich strains flagged)")

if __name__=='__main__': main()
