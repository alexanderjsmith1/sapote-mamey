"""Real local writers and readers, with invented policy/sequence fixtures."""
from pathlib import Path
from types import SimpleNamespace
from contextlib import closing
import csv
import hashlib
import importlib.util
import json
import sys
import pytest
from openpyxl import Workbook, load_workbook
from mamey.release_policy import resolve_run_release
from mamey.report_release_policy import policy_for_package, policy_for_master_rows, policy_for_run, row_policy
from mamey.package_inspector import explain_command
from mamey.compile_report import _read_manifests, _cover

def export_module():
    root = Path(__file__).resolve().parents[1]
    bundle = root
    sys.path.insert(0, str(bundle / "tools"))
    spec = importlib.util.spec_from_file_location("candidate_figure_ready_export", root / "tools/export_figure_ready.py")
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    return module

def tree(path):
    return {str(p.relative_to(path)): hashlib.sha256(p.read_bytes()).hexdigest() for p in path.rglob("*") if p.is_file()}

@pytest.fixture
def synthetic_single_contig_full_locus_zip(tmp_path):
    from tools.fixture_inputs import prepare_full_locus_single_contig_fixture
    bundle = Path(__file__).resolve().parents[1]
    destination = tmp_path / "prepared_input.zip"
    recipe = prepare_full_locus_single_contig_fixture(bundle / "tests/fixtures/synthetic_single_contig_antismash.zip", destination)
    (tmp_path / "INPUT_TRANSFORMATION.json").write_text(json.dumps(recipe, indent=2))
    return destination

def profile(path, strain="reference"):
    value = {"schema_version": "sapote_privacy_profile_v1", "profile_id": "synthetic_report_policy",
             "tiers": [{"tier_id": "OPEN", "public_export": True},
                       {"tier_id": "LOCAL", "public_export": False}],
             "default_tier": "LOCAL", "strain_assignments": [{"strain_id": strain, "tier_id": "OPEN"}]}
    path.write_text(json.dumps(value)); return path

@pytest.mark.parametrize("case,expected,state", [
    ("private", "PRIVATE", "RECORDED"), ("public", "PUBLIC", "EXPLICIT"),
    ("legacy_public", "PRIVATE", "HOLD"), ("missing", "PRIVATE", "HOLD"),
    ("conflict", "PRIVATE", "HOLD"), ("unreadable", "PRIVATE", "HOLD"),
    ("private_decision_conflict", "PRIVATE", "HOLD"), ("refused", "PRIVATE", "REFUSED"),
    ("row_conflict", "PRIVATE", "HOLD"),
])
def test_saved_metadata_to_explain_and_report(tmp_path, capsys, case, expected, state):
    pkg = tmp_path / "package"; pkg.mkdir()
    strain = "synthetic_reference"
    decision = resolve_run_release(strain, "PRIVATE" if case == "private" else "PUBLIC").to_dict()
    man = {"strain_id": strain}
    if case != "missing": man["release"] = "PRIVATE" if case in ("private", "private_decision_conflict", "refused") else "PUBLIC"
    if case in ("private", "public", "private_decision_conflict", "refused", "row_conflict"):
        man["release_decision"] = decision
        if case == "private_decision_conflict": decision["strain_id"] = "different_reference"
        if case == "refused": decision.update(release="PRIVATE", refused=True, reason="public_override_refused_by_profile")
    if case == "row_conflict": man["bgcs"] = [{"bgc_id": "BGC001", "release": "PRIVATE"}]
    (pkg / "manifest.json").write_text(json.dumps(man))
    short = dict(man)
    if case == "conflict": short["release"] = "PRIVATE"
    (pkg / "manifest_short.json").write_text("{" if case == "unreadable" else json.dumps(short))
    before = tree(pkg)
    policy = policy_for_package(pkg, strain)
    assert (policy["effective_release"], policy["privacy_assignment_state"]) == (expected, state)
    assert explain_command(SimpleNamespace(package_dir=str(pkg))) == 0
    explain = capsys.readouterr().out
    report = _cover(_read_manifests(pkg))
    for text in (explain, report):
        assert "Release policy: " + expected in text
        assert "state: " + state in text
        assert "not transfer authorization" in text
        assert "upload manifest" not in text.lower()
    assert "Review the package locally" in explain
    assert tree(pkg) == before
    (tmp_path / "explain.txt").write_text(explain)
    (tmp_path / "cover.md").write_text(report)
    (tmp_path / "POLICY.json").write_text(json.dumps(policy, indent=2))

def test_direct_run_without_decision_is_a_visible_hold():
    run = SimpleNamespace(context=SimpleNamespace(strain_id="legacy", release_decision=None), bgcs=[])
    assert policy_for_run(run)["privacy_assignment_state"] == "HOLD"

@pytest.mark.parametrize("policy_mix,expected", [("public", "PUBLIC"), ("mixed", "PRIVATE"), ("missing", "PRIVATE"), ("a3_conflict", "PRIVATE")])
def test_export_all_csvs_carry_conservative_scope_and_source_hash(tmp_path, policy_mix, expected):
    wb = Workbook(); a2 = wb.active; a2.title = "A2_Strain_Registry"
    keys = list(row_policy(policy_for_package(tmp_path, "unknown")))
    a2.append(["strain", "taxonomy", "bgc_count"] + keys)
    b1 = wb.create_sheet("B1_BGC_Master")
    b1.append(["strain", "BGC_ID", "contig", "region", "products"] + keys)
    for number in (1, 2):
        strain = "ref" + str(number)
        release = "PRIVATE" if policy_mix == "mixed" and number == 2 else "PUBLIC"
        decision = resolve_run_release(strain, release)
        run = SimpleNamespace(context=SimpleNamespace(strain_id=strain, release_decision=decision), bgcs=[])
        values = row_policy(policy_for_run(run)) if policy_mix != "missing" else {}
        a2.append([strain, "Syntheticus sp.", 1] + [values.get(k) for k in keys])
        b1.append([strain, "BGC001", "FULL_CONTIG_" + str(number), "region001", "NRPS"] + [values.get(k) for k in keys])
    source = tmp_path / "master.xlsx"; wb.save(source)
    if policy_mix == "a3_conflict":
        a3 = wb.create_sheet("A3_Run_Manifest")
        a3.append(["strain", "raw_bgcs", "corrected_bgcs"] + keys)
        run = SimpleNamespace(context=SimpleNamespace(strain_id="ref1", release_decision=resolve_run_release("ref1", "PRIVATE")), bgcs=[])
        values = row_policy(policy_for_run(run))
        a3.append(["ref1", 1, 1] + [values.get(k) for k in keys])
        wb.save(source)
    before = source.read_bytes(); out = tmp_path / "export"
    export_module().main(str(source), str(out))
    policy = json.loads((out / "release_policy.json").read_text())
    assert policy["policy"]["effective_release"] == expected
    policy_sha = hashlib.sha256((out / "release_policy.json").read_bytes()).hexdigest()
    for path in out.glob("*.csv"):
        with path.open() as handle: rows = list(csv.DictReader(handle))
        assert rows
        assert all(r["release"] == expected and r["release_policy_sha256"] == policy_sha
                   and r["source_master_sha256"] == hashlib.sha256(before).hexdigest() for r in rows)
    assert source.read_bytes() == before

@pytest.mark.parametrize("override,profile_case,expected", [
    (None, False, "PRIVATE"), ("PUBLIC", False, "PUBLIC"),
    (None, True, "PUBLIC"), ("PRIVATE", True, "PRIVATE"),
    ("PUBLIC", "unassigned", "PRIVATE"),
])
def test_actual_gold_writer_policy_chain(tmp_path, synthetic_single_contig_full_locus_zip, capsys,
                                        override, profile_case, expected):
    from mamey.cli import run_one_strain
    from mamey.validate import validate_package
    strain = "synthetic_policy_reference"
    profile_path = profile(tmp_path / "privacy.json", "different_reference" if profile_case == "unassigned" else strain) if profile_case else None
    master = tmp_path / "master.xlsx"
    summary = run_one_strain(strain, "Invented privacy fixture", str(synthetic_single_contig_full_locus_zip),
        str(tmp_path / "runs"), "gold", "Streptomyces sp.", "synthetic software fixture",
        master_path=str(master), json_mode="off", brief="none", release=override,
        privacy_profile=str(profile_path) if profile_path else None,
        locus_maps="off", reference_completion="off", rggmci_threads=4)
    capsys.readouterr()
    assert summary["status"] in ("MAMEY_COMPLETE", "MAMEY_COMPLETE_WITH_ISSUES")
    pkg = tmp_path / "runs" / strain / "package"
    man = json.loads((pkg / "manifest.json").read_text())
    assert man["release"] == expected
    state = "REFUSED" if profile_case == "unassigned" else None
    with closing(load_workbook(pkg / (strain + "_5_workbook.xlsx"), read_only=True, data_only=True)) as wb:
        cover = {r[0]: r[1] for r in wb["README"].iter_rows(values_only=True) if len(r) >= 2 and r[0]}
        assert cover["Release"] == expected and cover["Release reason"] == man["release_decision"]["reason"]
        assert not any("upload manifest" in str(v).lower() for v in cover.values())
        for name in ("BGC_Crosswalk", "BGC_Inventory"):
            rows = list(wb[name].iter_rows(values_only=True)); headers = rows[0]
            assert all(row[headers.index("Strain")] == strain and row[headers.index("release")] == expected for row in rows[1:])
    with closing(load_workbook(master, read_only=True, data_only=True)) as wb:
        for name in ("A2_Strain_Registry", "A3_Run_Manifest", "B1_BGC_Master"):
            rows = list(wb[name].iter_rows(values_only=True)); headers = rows[0]
            assert rows[1:]
            assert all(row[headers.index("release")] == expected for row in rows[1:])
            assert all(json.loads(row[headers.index("release_decision_json")]) == man["release_decision"] for row in rows[1:])
    before = tree(pkg)
    assert explain_command(SimpleNamespace(package_dir=str(pkg))) == 0
    text = capsys.readouterr().out
    assert "Release policy: " + expected in text and "upload manifest" not in text.lower()
    if state: assert "state: " + state in text
    (tmp_path / "explain.txt").write_text(text)
    report = (pkg / (strain + "_compiled_report.md")).read_text()
    assert "Release policy: " + expected in report and "upload manifest" not in report.lower()
    assert policy_for_package(pkg, strain)["policy_source"] in report
    validation = validate_package(pkg, enrichment_check=True, write_status_receipt=False)
    assert validation["status"] == "MAMEY_COMPLETE" and tree(pkg) == before
    export_module().main(str(master), str(tmp_path / "figure_ready"))
    (tmp_path / "RESULT.json").write_text(json.dumps({"release": expected, "summary": summary,
        "policy": policy_for_package(pkg, strain), "package_bindings": before, "master_sha256": hashlib.sha256(master.read_bytes()).hexdigest()}, indent=2))
