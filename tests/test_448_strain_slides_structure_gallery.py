"""The strain deck's MIBiG reference-structure gallery (4 Oct): every bound reference at any percentage, each with
its own locus line and metrics, one removable group per reference, undrawn entries last, and the claim-boundary footer."""
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
import strain_slides as ss  # noqa: E402

pptx = pytest.importorskip("pptx")
PIL = pytest.importorskip("PIL.Image")


def _binding(alias, node, matched, den, ident):
    return {"bgc_alias": alias, "full_contig": f"{node}_length_9000_cov_10.5", "region": "region001",
            "reference_match_metrics": {"state": "REFERENCE_MATCH_METRICS_BOUND", "matched_query_genes": matched,
                                        "query_gene_denominator": den, "median_best_hit_identity_percent": ident}}


def _assets(tmp_path, strain="AS-XXX"):
    cs = tmp_path / f"{strain}_v2n_assets" / "chemical_structures"
    (cs / "structures").mkdir(parents=True)
    PIL.new("RGB", (300, 120), "white").save(cs / "structures" / "a.png")
    entries = [
        {"mibig_accession": "BGC0000001", "representative_name": "undrawnmycin", "drawing_available": False,
         "reference_biosynthesis_classes": ["PKS"], "reference_organism": "Streptomyces sp.",
         "locus_bindings": [_binding("BGC003", "NODE_10", 8, 21, 62.0)], "notes_lines": ["undrawn"]},
        {"mibig_accession": "BGC0000002", "representative_name": "drawnamide", "drawing_available": True,
         "representative_png": "structures/a.png", "representative_average_mw_g_mol": 468.46,
         "reference_biosynthesis_classes": ["NRPS"], "reference_organism": "Streptomyces sp. LZ35",
         "recommended_powerpoint_group_name": "drawnamide BGC0000002",
         "locus_bindings": [_binding("BGC021", "NODE_20", 4, 23, 69.0), _binding("BGC054", "NODE_80", 6, 23, 85.5)],
         "notes_lines": ["drawn"]},
    ]
    (cs / "STRUCTURE_ASSETS.json").write_text(json.dumps({"strain": strain, "entries": entries}))
    return tmp_path / f"{strain}_v2n_assets"


def _deck():
    from pptx.util import Inches
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(ss.SW), Inches(ss.SH)
    return prs


def _texts(slide):
    out = []
    for sh in slide.shapes:
        for s in (sh.shapes if sh.shape_type == 6 else [sh]):
            if s.has_text_frame:
                out.append(s.text_frame.text)
    return "\n".join(out)


def test_every_reference_is_shown_with_its_own_locus_metrics(tmp_path):
    prs = _deck()
    pages, n, drawn = ss.structure_gallery(prs, {"strain": "AS-XXX", "src": {}}, _assets(tmp_path))
    assert (pages, n, drawn) == (1, 2, 1)
    txt = _texts(prs.slides[0])
    assert "BGC021 · NODE_20 · region001" in txt and "KCB 4/23 · 69.0%" in txt
    assert "BGC054 · NODE_80 · region001" in txt and "KCB 6/23 · 85.5%" in txt  # both loci, no threshold
    assert "Production by AS-XXX is unconfirmed" in txt


def test_each_reference_is_one_group_and_undrawn_comes_last_without_mw(tmp_path):
    prs = _deck()
    ss.structure_gallery(prs, {"strain": "AS-XXX", "src": {}}, _assets(tmp_path))
    groups = [sh for sh in prs.slides[0].shapes if sh.shape_type == 6]
    assert [g.name for g in groups] == ["drawnamide BGC0000002", "undrawnmycin"]
    last = "\n".join(s.text_frame.text for s in groups[-1].shapes if s.has_text_frame)
    assert "No single structure bound" in last and "MW" not in last


def test_no_assets_means_no_gallery_and_a_foreign_strain_is_refused(tmp_path):
    prs = _deck()
    assert ss.structure_gallery(prs, {"strain": "AS-XXX", "src": {}}, tmp_path / "missing") == (0, 0, 0)
    assert ss.structure_gallery(prs, {"strain": "AS-YYY", "src": {}}, _assets(tmp_path)) == (0, 0, 0)
    assert len(prs.slides) == 0
