#!/usr/bin/env python3
"""Fail-before / pass-after tests for CLAUDE_409_concurrency_atomic.

Covers the two hazards the patch closes (DEEP_AUDIT2_concurrency.md):

  * C2/C3/C4 — concurrent writers to ONE target collide on a fixed ``<name>.tmp`` sibling and the
    loser crashes in ``os.replace`` with ``FileNotFoundError``. The patch swaps every fixed sibling
    temp in packaging.py (`_atomic_write_text`, `atomic_open`, `zip_package`) and
    judgment_store.py (`_atomic_write_text`) for a UNIQUE ``tempfile.mkstemp`` temp, mirroring
    tools/_wbio._temp_path.

  * C1 — two runs sharing one ``--master`` workbook do a read-modify-write PUBLISHED with a
    non-atomic ``shutil.copy2`` and no lock, so they silently drop each other's strain rows (lost
    update). The patch (cli.py) takes a BLOCKING advisory ``flock`` on ``<master>.lock`` around the
    whole read-append-publish and publishes via unique-temp + ``os.replace``.

This file is SELF-CONTAINED (no bundle import needed): each test runs an OLD implementation that
mirrors the shipping v408 code (asserted to FAIL) and a NEW implementation that mirrors the patched
code verbatim (asserted to PASS). The NEW blocks are byte-for-byte the logic added by
CLAUDE_409_concurrency_atomic.patch; the mapping is noted at each block.

Run:  python3 TESTS_409_concurrency_atomic.py
Exit 0 = every fail-before failed as expected AND every pass-after passed.
"""
import os
import sys
import shutil
import tempfile
import threading
import traceback
from pathlib import Path

try:
    import fcntl
except ImportError:  # pragma: no cover - POSIX-only test host expected
    fcntl = None

N_WRITERS = 8
TRIALS = 6


# ----------------------------------------------------------------------------------------------
# OLD (v408 shipping) — fixed sibling temp. Mirrors the pre-patch bodies of
# packaging._atomic_write_text / judgment_store._atomic_write_text.
# ----------------------------------------------------------------------------------------------
def old_atomic_write_text(path: Path, text: str) -> None:
    tmp = path.with_name(path.name + ".tmp")          # <-- fixed, SHARED across concurrent writers
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, path)


# ----------------------------------------------------------------------------------------------
# NEW (patched) — unique mkstemp temp. Verbatim logic from the patched
# packaging._unique_tmp + _atomic_write_text.
# ----------------------------------------------------------------------------------------------
def _unique_tmp(path: Path) -> Path:
    path = Path(path)
    fd, tmp = tempfile.mkstemp(prefix="." + path.name + ".", suffix=".tmp", dir=str(path.parent))
    os.close(fd)
    return Path(tmp)


def new_atomic_write_text(path: Path, text: str) -> None:
    path = Path(path)
    tmp = _unique_tmp(path)
    try:
        tmp.write_text(text, encoding="utf-8")
        os.replace(tmp, path)
    except BaseException:
        try:
            if tmp.exists():
                tmp.unlink()
        except OSError:
            pass
        raise


def _hammer_writers(writer, target: Path):
    """Fire N_WRITERS threads all writing `target` at once; return list of exceptions raised."""
    errors = []
    barrier = threading.Barrier(N_WRITERS)

    def work(i):
        try:
            barrier.wait()                       # release all writers simultaneously -> max contention
            for _ in range(4):
                writer(target, f"payload-from-writer-{i}\n")
        except BaseException as exc:             # noqa: BLE001 - we are collecting them on purpose
            errors.append(exc)

    threads = [threading.Thread(target=work, args=(i,)) for i in range(N_WRITERS)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    return errors


def test_c234_fixed_tmp_collision():
    """C2/C3/C4: fixed .tmp crashes under concurrency (fail-before); mkstemp does not (pass-after)."""
    print("\n[C2/C3/C4] concurrent writers to one target")
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)

        # --- FAIL-BEFORE: the old fixed-.tmp writer must crash at least once ---
        old_crashes = 0
        old_fnf = 0
        for _ in range(TRIALS):
            errs = _hammer_writers(old_atomic_write_text, d / "manifest.json")
            old_crashes += len(errs)
            old_fnf += sum(1 for e in errs if isinstance(e, FileNotFoundError))
        print(f"  OLD fixed '<name>.tmp' : {old_crashes} crashes across {TRIALS} trials "
              f"({old_fnf} FileNotFoundError in os.replace)")
        assert old_crashes > 0, "FAIL-BEFORE did not reproduce: old fixed-.tmp writer never crashed"

        # --- PASS-AFTER: the new mkstemp writer must never crash, and leave a valid final file ---
        new_crashes = 0
        for _ in range(TRIALS):
            errs = _hammer_writers(new_atomic_write_text, d / "manifest_new.json")
            new_crashes += len(errs)
        final = (d / "manifest_new.json").read_text(encoding="utf-8")
        # no stray leftover temps
        leftover = [p.name for p in d.iterdir() if p.name.startswith(".manifest_new.json.")]
        print(f"  NEW mkstemp temp       : {new_crashes} crashes across {TRIALS} trials; "
              f"final file valid={bool(final.strip())}; leftover temps={len(leftover)}")
        assert new_crashes == 0, f"PASS-AFTER regressed: mkstemp writer crashed {new_crashes} times"
        assert final.strip().startswith("payload-from-writer-"), "final file is not a clean whole write"
        assert not leftover, f"mkstemp writer leaked temp files: {leftover}"
    print("  -> PASS (fixed-.tmp crash reproduced; mkstemp is crash-free)")


# ----------------------------------------------------------------------------------------------
# C1 — shared --master workbook lost update.
# A real openpyxl workbook stands in for project_master.xlsx; each worker appends one strain row.
# OLD: read-modify-write + copy2 publish, no lock (mirrors cli.py ~2053/2065 pre-patch).
# NEW: blocking flock on <master>.lock + mkstemp/os.replace publish (mirrors the patched helpers
#      _acquire_master_workbook_lock / _atomic_publish_master and the finally-release).
# ----------------------------------------------------------------------------------------------
try:
    from openpyxl import Workbook, load_workbook
    HAVE_OPENPYXL = True
except Exception:  # pragma: no cover
    HAVE_OPENPYXL = False


def _make_master(master: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["strain_id"])
    wb.save(master)


def _read_strains(master: Path):
    wb = load_workbook(master, read_only=True)
    ws = wb["Master"]
    rows = [c[0] for c in ws.iter_rows(min_row=2, values_only=True) if c and c[0]]
    wb.close()
    return rows


def _append_and_stage(master: Path, out: Path, strain: str):
    """Read master, append `strain`, write the per-run staged copy atomically (this half is already
    atomic+unique in v408 via _wbio; the corruption is the PUBLISH, below)."""
    wb = load_workbook(master)
    ws = wb["Master"]
    ws.append([strain])
    wb.save(out)
    wb.close()


def _publish_copy2(out: Path, master: Path):
    shutil.copy2(out, master)                       # OLD non-atomic shared publish


def _publish_atomic(out: Path, master: Path):
    # verbatim logic from patched cli._atomic_publish_master
    master = Path(master)
    fd, tmpname = tempfile.mkstemp(prefix="." + master.name + ".", suffix=".tmp",
                                   dir=str(master.parent))
    os.close(fd)
    try:
        shutil.copy2(out, tmpname)
        os.replace(tmpname, master)
    except BaseException:
        try:
            if os.path.exists(tmpname):
                os.remove(tmpname)
        except OSError:
            pass
        raise


def _run_master_workers(master: Path, staging: Path, use_lock: bool):
    errors = []
    # For the NO-LOCK case a barrier forces every worker to read the same base before any publish,
    # deterministically reproducing the lost update. The lock case must NOT use that barrier (the
    # lock serializes reads); it relies on flock for correctness under real contention.
    barrier = threading.Barrier(N_WRITERS) if not use_lock else None

    def work(i):
        strain = f"STRAIN_{i:02d}"
        out = staging / f"staged_{i}.xlsx"
        try:
            if use_lock:
                lock_fh = open(str(master) + ".lock", "a+")
                fcntl.flock(lock_fh.fileno(), fcntl.LOCK_EX)     # blocking: serialize writers
                try:
                    _append_and_stage(master, out, strain)
                    _publish_atomic(out, master)
                finally:
                    lock_fh.close()                              # drops the flock
            else:
                _append_and_stage(master, out, strain)
                barrier.wait()                                   # all have read; now all publish
                _publish_copy2(out, master)
        except BaseException as exc:  # noqa: BLE001
            errors.append((strain, repr(exc)))

    threads = [threading.Thread(target=work, args=(i,)) for i in range(N_WRITERS)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    return errors


def test_c1_master_lost_update():
    """C1: concurrent master appends lose rows without a lock (fail-before); all survive with the
    flock + atomic publish (pass-after)."""
    print("\n[C1] concurrent shared --master appends")
    if not HAVE_OPENPYXL:
        print("  SKIP: openpyxl not installed")
        return
    if fcntl is None:
        print("  SKIP: fcntl (POSIX advisory locks) unavailable")
        return

    # --- FAIL-BEFORE: no lock + copy2 publish -> rows lost ---
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        master = d / "project_master.xlsx"
        staging = d / "staging_old"
        staging.mkdir()
        _make_master(master)
        _run_master_workers(master, staging, use_lock=False)
        survivors_old = _read_strains(master)
        print(f"  OLD copy2 + no lock : {len(survivors_old)}/{N_WRITERS} strains survived "
              f"-> {sorted(survivors_old)}")
        assert len(survivors_old) < N_WRITERS, (
            "FAIL-BEFORE did not reproduce: expected lost updates, all rows survived")

    # --- PASS-AFTER: blocking flock + atomic publish -> every row survives ---
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        master = d / "project_master.xlsx"
        staging = d / "staging_new"
        staging.mkdir()
        _make_master(master)
        errs = _run_master_workers(master, staging, use_lock=True)
        survivors_new = _read_strains(master)
        # the .lock sidecar must sit beside the master (never inside a package tree)
        lock_beside = (Path(str(master) + ".lock")).exists()
        print(f"  NEW flock + atomic  : {len(survivors_new)}/{N_WRITERS} strains survived; "
              f"errors={errs}; .lock beside master={lock_beside}")
        assert not errs, f"PASS-AFTER raised errors: {errs}"
        assert len(survivors_new) == N_WRITERS, (
            f"PASS-AFTER lost rows: {len(survivors_new)}/{N_WRITERS} survived: {sorted(survivors_new)}")
        assert sorted(survivors_new) == [f"STRAIN_{i:02d}" for i in range(N_WRITERS)]
    print("  -> PASS (lost update reproduced without lock; every row retained with lock)")


def main():
    tests = [test_c234_fixed_tmp_collision, test_c1_master_lost_update]
    failed = 0
    for t in tests:
        try:
            t()
        except AssertionError as exc:
            failed += 1
            print(f"  !! ASSERTION FAILED in {t.__name__}: {exc}")
        except Exception:  # noqa: BLE001
            failed += 1
            print(f"  !! ERROR in {t.__name__}:")
            traceback.print_exc()
    print("\n" + ("=" * 60))
    print("RESULT:", "ALL TESTS PASSED" if failed == 0 else f"{failed} TEST(S) FAILED")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
