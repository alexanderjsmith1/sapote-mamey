"""v9.7.390 regression coverage for the cross-strain sheet builder."""
from __future__ import annotations

import importlib
import pathlib
import subprocess
import sys

import openpyxl


ROOT = pathlib.Path(__file__).resolve().parents[1]
TOOLS = ROOT / "tools"


def test_module_is_import_safe_under_foreign_argv(monkeypatch):
    monkeypatch.syspath_prepend(str(TOOLS))
    monkeypatch.setattr(sys, "argv", ["pytest", "--foreign-option"])
    module = importlib.import_module("add_xstrain_sheets")
    assert callable(module.main)


def test_generated_sheet_matcher_does_not_delete_legitimate_prefix_sheets(monkeypatch):
    monkeypatch.syspath_prepend(str(TOOLS))
    module = importlib.import_module("add_xstrain_sheets")
    canonical = "Cross_Strain_Findings"
    assert module._generated_sheet_name(canonical, canonical)
    assert module._generated_sheet_name(canonical + "1", canonical)
    assert module._generated_sheet_name(canonical + "27", canonical)
    assert not module._generated_sheet_name(canonical + "_Notes", canonical)
    assert not module._generated_sheet_name(canonical + "Archive", canonical)


def test_empty_cohort_fails_before_mutating_workbook(tmp_path):
    bank = tmp_path / "bank"
    bank.mkdir()
    (bank / "bgc_data.json").write_text('{"strains": {}, "bgcs": []}')
    for name in ("gene_data.json", "deep_data.json", "rggmci_full.json", "tigrfam.json"):
        (bank / name).write_text("{}")
    workbook_path = tmp_path / "master.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "original"
    workbook.save(workbook_path)
    before = workbook_path.read_bytes()

    result = subprocess.run(
        [sys.executable, str(TOOLS / "add_xstrain_sheets.py"), "--banked-dir", str(bank), "--out", str(workbook_path)],
        text=True,
        capture_output=True,
        timeout=30,
    )

    assert result.returncode != 0
    assert "zero strains" in result.stderr
    assert workbook_path.read_bytes() == before


def test_exact_locator_uses_permanent_four_part_order(tmp_path):
    bank = tmp_path / "bank"
    bank.mkdir()
    (bank / "bgc_data.json").write_text(
        '{"strains":{"AS-FIX":{"n50":1000000,"corrected_bgcs":1,"raw_bgcs":1}},'
        '"bgcs":[{"sid":"AS-FIX","products":"T1PKS","kcb_top":"hit","kcb_cumulative":10,'
        '"bgc_id":"BGC001","contig":"NODE_1_length_10000_cov_10.0","region":"region001",'
        '"edge_status":"Interior"}]}'
    )
    (bank / "tigrfam.json").write_text('{"AS-FIX":{"present":{"TIGR03828":{}}}}')
    for name in ("gene_data.json", "deep_data.json", "rggmci_full.json"):
        (bank / name).write_text("{}")
    workbook_path = tmp_path / "master.xlsx"
    workbook = openpyxl.Workbook()
    workbook.save(workbook_path)

    result = subprocess.run(
        [sys.executable, str(TOOLS / "add_xstrain_sheets.py"), "--banked-dir", str(bank), "--out", str(workbook_path)],
        text=True,
        capture_output=True,
        timeout=30,
    )
    assert result.returncode == 0, result.stderr
    workbook = openpyxl.load_workbook(workbook_path, read_only=True)
    try:
        values = [cell for row in workbook["Cross_Strain_Findings"].iter_rows(values_only=True) for cell in row if isinstance(cell, str)]
    finally:
        workbook.close()
    locator = "AS-FIX / NODE_1_length_10000_cov_10.0 / region001 / BGC001"
    assert any(locator in value for value in values)
    assert not any("BGC001 (NODE_1" in value for value in values)
