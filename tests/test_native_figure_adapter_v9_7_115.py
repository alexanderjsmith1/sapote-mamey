"""v9.7.115: native-figure engine reads the canonical CODED-schema workbook.

The engine (mamey_native_figures) was written against pre-v1.1 legacy sheet/column names
(Strain_Registry, DAPR_Antibacterial, 'AB score', …). The canonical v1.1 workbook emits coded
names (A2_Strain_Registry, C1_DAPR_Antibacterial, snake_case columns). Before the adapter the
engine raised "missing required sheets" or returned NO_FIGURES on a current workbook — the figure
subsystem was effectively dead. These tests pin that a coded workbook renders, and that legacy
workbooks still render (backward compatibility).
"""
import os
import sys
import tempfile
import pathlib

import openpyxl
import pytest

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from mamey.mamey_native_figures import render_mamey_native_figure_set  # noqa: E402


def _coded_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A2_Strain_Registry"
    ws.append(["strain", "taxonomy", "contigs", "n50", "bgc_count",
               "package_status", "ab_score", "af_score"])
    for i in range(1, 6):
        ws.append([f"AS-{900+i}", "Streptomyces sp.", 5, 120000, 30 + i,
                   "MAMEY_COMPLETE", 70 + i, 60 + i])
    for sheet in ["C1_DAPR_Antibacterial", "C2_DAPR_Antifungal"]:
        w = wb.create_sheet(sheet)
        w.append(["strain", "rank", "BGC_ID", "contig", "products", "score", "lead_tier"])
        for i in range(1, 6):
            w.append([f"AS-{900+i}", 1, f"BGC00{i}", f"ctg{i}", "NRPS", 70 + i, "High"])
    w = wb.create_sheet("D1_RGGMCI_All_Strains")
    w.append(["strain", "bgc_id", "rggmci"])
    for i in range(1, 6):
        w.append([f"AS-{900+i}", f"BGC00{i}", 0.5])
    wb.save(path)


def _legacy_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Strain_Registry"
    ws.append(["Strain", "Taxonomy", "Contigs", "N50", "BGC regions",
               "Package status", "AB score", "AF score"])
    for i in range(1, 6):
        ws.append([f"AS-{900+i}", "Streptomyces sp.", 5, 120000, 30 + i, "OK", 70 + i, 60 + i])
    for sheet in ["DAPR_Antibacterial", "DAPR_Antifungal"]:
        w = wb.create_sheet(sheet)
        w.append(["Strain", "Score", "BGC_ID"])
        for i in range(1, 6):
            w.append([f"AS-{900+i}", 70 + i, f"BGC00{i}"])
    w = wb.create_sheet("RG-GMCI_All_Strains")
    w.append(["Strain", "bgc_id"])
    for i in range(1, 6):
        w.append([f"AS-{900+i}", f"BGC00{i}"])
    wb.save(path)


def test_coded_schema_workbook_renders():
    d = tempfile.mkdtemp()
    wbp = os.path.join(d, "coded.xlsx")
    _coded_workbook(wbp)
    result = render_mamey_native_figure_set(wbp, outdir=os.path.join(d, "figs"))
    assert result["status"] == "PASS"
    assert result["figure_count"] > 0   # not NO_FIGURES


def test_legacy_schema_workbook_still_renders():
    d = tempfile.mkdtemp()
    wbp = os.path.join(d, "legacy.xlsx")
    _legacy_workbook(wbp)
    result = render_mamey_native_figure_set(wbp, outdir=os.path.join(d, "figs"))
    assert result["status"] == "PASS"
    assert result["figure_count"] > 0   # backward compatibility preserved


def _unscored_coded_workbook(path):
    """Coded sheets present, but DAPR leads empty (Sapote scaffold) and no registry scores."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A2_Strain_Registry"
    # registry populated with genome metadata but NO ab/af score columns
    ws.append(["strain", "taxonomy", "Contigs", "N50", "Genome_bp", "Lead_Tier"])
    for i in range(1, 6):
        ws.append([f"AS-{900+i}", "not verified", 5, 120000, 6_000_000, "—"])
    for sheet in ["C1_DAPR_Antibacterial", "C2_DAPR_Antifungal"]:
        w = wb.create_sheet(sheet)
        w.append(["Rank", "strain", "BGC_ID", "Product_Class", "AN_Score", "WL_Score"])
        # no data rows — empty scaffold
    w = wb.create_sheet("D1_RGGMCI_All_Strains")
    w.append(["strain", "Status", "Pairs_Total"])
    for i in range(1, 6):
        w.append([f"AS-{900+i}", "REVIEW", 0])
    wb.save(path)


def test_unscored_cohort_reports_NO_DATA_not_NO_FIGURES():
    """v9.7.115: an unscored cohort (sheets resolve, leads empty) must report NO_DATA, distinct
    from NO_FIGURES — so the benign unscored case can't mask the next genuine schema drift."""
    d = tempfile.mkdtemp()
    wbp = os.path.join(d, "unscored.xlsx")
    _unscored_coded_workbook(wbp)
    result = render_mamey_native_figure_set(wbp, outdir=os.path.join(d, "figs"))
    assert result["status"] == "NO_DATA"
    assert result["figure_count"] == 0
    assert "scoring" in result.get("detail", "").lower()


def test_renders_when_contigs_column_absent():
    """v9.7.115 self-audit: the AB-vs-AF portfolio figure crashed (df.get('Contigs',100).fillna)
    when the registry lacked a Contigs column — a path the adapter newly makes reachable. Guard it."""
    d = tempfile.mkdtemp()
    wbp = os.path.join(d, "no_contigs.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A2_Strain_Registry"
    ws.append(["strain", "taxonomy", "n50", "ab_score", "af_score"])   # no contigs/bgc_count
    for i in range(1, 6):
        ws.append([f"AS-{900+i}", "Streptomyces sp.", 120000, 70 + i, 60 + i])
    for sheet in ["C1_DAPR_Antibacterial", "C2_DAPR_Antifungal"]:
        w = wb.create_sheet(sheet)
        w.append(["strain", "score", "BGC_ID"])
        for i in range(1, 6):
            w.append([f"AS-{900+i}", 70 + i, f"BGC00{i}"])
    w = wb.create_sheet("D1_RGGMCI_All_Strains")
    w.append(["strain", "bgc_id"])
    for i in range(1, 6):
        w.append([f"AS-{900+i}", f"BGC00{i}"])
    wb.save(wbp)
    result = render_mamey_native_figure_set(wbp, outdir=os.path.join(d, "figs"))
    assert result["status"] == "PASS"   # no crash, figures render
    assert result["figure_count"] > 0
