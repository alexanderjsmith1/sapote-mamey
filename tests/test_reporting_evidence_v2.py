from __future__ import annotations

import csv
import json
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from mamey.antismash_evidence import extract_rrefinder_hits
from mamey.antismash_tables import (
    MODULE_COLUMNS,
    MOTIF_COLUMNS,
    RIPP_MOTIF_COLUMNS,
    RREFINDER_COLUMNS,
    build_structured_tables,
)
from mamey.models import AssemblyMetrics, BGCRecord, RunContext
from mamey.validate import validate_reporting_v2_outputs
from mamey.workbook import write_per_strain_workbook


def _rre_hit(location, domain, identifier, score=50.0):
    return {
        "location": location,
        "label": "ctg1_10",
        "locus_tag": "ctg1_10",
        "domain": domain,
        "description": f"{domain} description",
        "identifier": identifier,
        "evalue": 1e-20,
        "score": score,
        "protein_start": 2,
        "protein_end": 82,
    }


def _rre_zip(tmp_path: Path, records: list[dict]) -> Path:
    path = tmp_path / "rrefinder.zip"
    payload = {
        "version": "8.0.4",
        "schema": 4,
        "records": records,
    }
    with ZipFile(path, "w") as archive:
        archive.writestr("synthetic.json", json.dumps(payload))
    return path


def _record(record_id: str, hits_by_cds: dict, hits_by_protocluster=None):
    return {
        "id": record_id,
        "modules": {
            "antismash.modules.rrefinder": {
                "schema_version": 1,
                "bitscore_cutoff": 25.0,
                "hits_by_protocluster": hits_by_protocluster or {},
                "hits_by_cds": hits_by_cds,
                "min_length": 50,
                "record_id": record_id,
            }
        },
    }


def test_rrefinder_multiple_domains_are_rows_not_index_duplicates(tmp_path):
    first = _rre_hit("[120:150](+)", "PqqD_RRE", "RREFam006.1")
    second = _rre_hit("[125:155](+)", "NHLP_CD_RRE", "RREFam015.1")
    path = _rre_zip(tmp_path, [
        _record(
            "NODE_1",
            {"ctg1_10": [first, second, dict(first)]},
            {"2": ["ctg1_10"]},
        )
    ])
    rows = extract_rrefinder_hits(path)
    assert len(rows) == 2
    assert {row["rrefam_identifier"] for row in rows} == {
        "RREFam006.1", "RREFam015.1"
    }
    assert all(row["protoclusters"] == ["2"] for row in rows)


def test_rrefinder_coordinate_mapping_retains_ambiguous_and_unmapped(tmp_path):
    path = _rre_zip(tmp_path, [
        _record(
            "NODE_1",
            {
                "mapped": [_rre_hit("[110:120](+)", "A", "RREFam001.1")],
                "ambiguous": [_rre_hit("[170:180](+)", "B", "RREFam002.1")],
                "outside": [_rre_hit("[300:310](-)", "C", "RREFam003.1")],
            },
        )
    ])
    bgcs = [
        SimpleNamespace(
            bgc_id="BGC001", contig="NODE_1", start=100, end=200,
            products=["RiPP"], source_gbk="",
        ),
        SimpleNamespace(
            bgc_id="BGC002", contig="NODE_1", start=150, end=250,
            products=["RiPP"], source_gbk="",
        ),
    ]
    evidence = {"rrefinder_hits": extract_rrefinder_hits(path)}
    out = build_structured_tables(path, evidence, bgcs)
    by_family = {row["rre_family"]: row for row in out["rrefinder"]}
    assert by_family["A"]["bgc_id"] == "BGC001"
    assert by_family["A"]["mapping_status"] == "MAPPED"
    assert by_family["B"]["mapping_status"] == "AMBIGUOUS_MULTIPLE_BGC_OVERLAPS"
    assert by_family["B"]["bgc_id"] == "UNMAPPED"
    assert by_family["C"]["mapping_status"] == "UNMAPPED_NO_BGC_OVERLAP"
    assert out["counts"]["rrefinder"] == 3
    assert out["counts"]["rrefinder_mapped"] == 1
    assert out["report_only_contract"] == "REPORT_ONLY_NO_SCORING"


def test_empty_rrefinder_is_valid_null_table(tmp_path):
    path = _rre_zip(tmp_path, [_record("NODE_1", {})])
    evidence = {"rrefinder_hits": extract_rrefinder_hits(path)}
    out = build_structured_tables(path, evidence, [])
    assert out["rrefinder"] == []
    assert out["rrefinder_status"] == "NULL_NO_RREFINDER_HITS"
    assert out["table_columns"]["rrefinder"] == RREFINDER_COLUMNS


class _FakeRun:
    def __init__(self, path: Path, structured):
        self.context = RunContext(
            strain_id="AS-TEST",
            display_name="AS-TEST",
            version="1.9.test",
            analysis_mode="gold",
            input_zip=str(path),
            outdir=str(path.parent),
        )
        self.assembly = AssemblyMetrics(
            genome_bp=2_000, contigs=1, n50=2_000, gc_pct=70.0,
            largest_contig=2_000,
        )
        self.bgcs = [
            BGCRecord(
                bgc_id="BGC001", contig="NODE_1", region_number=1,
                start=100, end=1_000, contig_length=2_000,
                products=["RiPP"], edge_status="Interior",
                antismash_region="region001", source_gbk="NODE_1.region001.gbk",
                node_id="NODE_1",
            )
        ]
        mpg_row = {
            "bgc_id": "BGC001", "query_gene": "q1", "subject_gene": "s1",
            "mibig_accession": "BGC0000001", "mibig_compound": "reference",
            "reference_type": "RiPP", "pct_identity": 50.0,
            "pct_coverage": 90.0, "blast_score": 100.0, "evalue": "1e-20",
            "reference_rank": 1, "reference": "BGC0000001.1",
            "source_file": "knownclusterblast/NODE_1_c1.txt",
        }
        convergence = {
            "bgc_id": "BGC001", "products": "RiPP", "boundary": "Interior",
            "mibig_accession": "BGC0000001", "mibig_compound": "reference",
            "reference_type": "RiPP", "distinct_query_genes": 1,
            "distinct_subject_genes": 1, "query_gene_count_total": 1,
            "recognizable_query_gene_count": 1, "query_gene_share": 1.0,
            "recognizable_gene_share": 1.0, "median_pct_identity": 50.0,
            "median_pct_coverage": 90.0, "minimum_pct_identity": 50.0,
            "best_reference_rank": 1, "source_row_count": 1,
            "class_concordance": "CONCORDANT",
            "convergence_tier": "H5_SINGLE_OR_WEAK",
            "convergence_tier_basis": "recognizable_gene_share",
            "claim_safety": "report only",
        }
        profile = {
            "bgc_id": "BGC001", "query_gene_count": 1,
            "recognizable_gene_count": 1, "recognizable_gene_fraction": 1.0,
            "recognizable_min_pct_identity": 30.0, "median_pct_identity": 50.0,
            "distinct_mibig_refs": 1, "interpretation_class": "INTERPRETABLE_DARK_MATTER",
            "dominant_mibig_accession": "BGC0000001",
            "dominant_mibig_compound": "reference",
            "dominant_distinct_query_genes": 1,
            "dominant_convergence_tier": "H5_SINGLE_OR_WEAK",
            "report_only_contract": "REPORT_ONLY_NO_SCORING",
        }
        scan_names = [
            "flbr", "rggmci", "chitinase", "tfbs", "blda_tta",
            "regulators", "transporters", "resistance", "cctt", "cassettes",
            "umed", "efls", "domain_architecture", "resistance_tiers",
            "wetlab_rows", "qs_signals", "glycosylation_arms", "per_bgc_dss",
        ]
        self.source_scans = SimpleNamespace(**{name: {} for name in scan_names})
        self.source_scans.mibig_per_gene = {
            "per_gene_mibig": {"BGC001": [mpg_row]},
            "mibig_convergence": [convergence],
            "bgc_mibig_profile": {"BGC001": profile},
        }
        self.source_scans.antismash_structured = structured
        self.scan_status = {"scans": []}
        self.triage = []

    def resistance_gene_summary(self):
        return {}


def test_workbook_has_dedicated_reporting_sheets_and_row_parity(tmp_path):
    path = _rre_zip(tmp_path, [
        _record(
            "NODE_1",
            {"ctg1_10": [_rre_hit("[120:150](+)", "PqqD_RRE", "RREFam006.1")]},
        )
    ])
    bgc = SimpleNamespace(
        bgc_id="BGC001", contig="NODE_1", start=100, end=1000,
        products=["RiPP"], source_gbk="",
    )
    structured = build_structured_tables(
        path, {"rrefinder_hits": extract_rrefinder_hits(path)}, [bgc]
    )
    output = tmp_path / "reporting.xlsx"
    write_per_strain_workbook(_FakeRun(path, structured), output)
    workbook = load_workbook(output, read_only=True, data_only=True)
    assert "MIBiG_Per_Gene" in workbook.sheetnames
    assert "antiSMASH_RREfinder" in workbook.sheetnames
    assert "P_LWC_Trial" in workbook.sheetnames
    assert workbook["MIBiG_Per_Gene"].max_row == 2
    assert workbook["antiSMASH_RREfinder"].max_row == 2
    assert workbook["P_LWC_Trial"].max_row == 2
    workbook.close()


def _write_csv(path: Path, headers: list[str], rows: list[dict]):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def test_reporting_validator_checks_json_csv_and_workbook_parity(tmp_path):
    per_gene_row = {
        "bgc_id": "BGC001", "query_gene": "q1", "subject_gene": "s1",
        "mibig_accession": "BGC0000001", "pct_identity": 50,
        "pct_coverage": 90, "pct_coverage_interpretation": 90,
        "coverage_qc_flag": "WITHIN_EXPECTED_RANGE", "reference_rank": 1,
    }
    (tmp_path / "AS_3_mibig_per_gene.json").write_text(json.dumps({
        "schema_version": "mibig_per_gene_v3",
        "per_gene_mibig": {"BGC001": [per_gene_row]},
    }), encoding="utf-8")
    _write_csv(
        tmp_path / "AS_3_mibig_per_gene.csv",
        list(per_gene_row),
        [per_gene_row],
    )
    convergence_row = {
        "bgc_id": "BGC001", "mibig_accession": "BGC0000001",
        "distinct_query_genes": 1, "query_gene_share": 1.0,
        "recognizable_gene_share": 1.0, "class_concordance": "CONCORDANT",
        "convergence_tier": "H5_SINGLE_OR_WEAK",
        "median_pct_coverage_interpretation": 90,
        "coverage_qc_flag": "WITHIN_EXPECTED_RANGE",
        "convergence_rank": 1, "dominant_reference": True,
        "dominance_status": "UNIQUE_DOMINANT",
    }
    (tmp_path / "AS_3_mibig_convergence.json").write_text(json.dumps({
        "schema": "mibig_pathway_convergence_v2", "rows": [convergence_row],
    }), encoding="utf-8")
    _write_csv(
        tmp_path / "AS_3_mibig_convergence.csv",
        list(convergence_row),
        [convergence_row],
    )
    profile_row = {
        "bgc_id": "BGC001", "recognizable_gene_fraction": 1.0,
    }
    _write_csv(
        tmp_path / "AS_3_mibig_profile.csv",
        list(profile_row),
        [profile_row],
    )
    rre_row = {column: "" for column in RREFINDER_COLUMNS}
    rre_row.update({"row_id": "RRE000001", "bgc_id": "BGC001"})
    structured = {
        "schema_version": "antismash_structured_tables_v2",
        "modules": [], "ripp_motifs": [], "motifs": [],
        "rrefinder": [rre_row],
    }
    (tmp_path / "AS_3_antismash_structured.json").write_text(
        json.dumps(structured), encoding="utf-8"
    )
    for suffix, headers, rows in (
        ("modules", MODULE_COLUMNS, []),
        ("ripp_motifs", RIPP_MOTIF_COLUMNS, []),
        ("motifs", MOTIF_COLUMNS, []),
        ("rrefinder", RREFINDER_COLUMNS, [rre_row]),
    ):
        _write_csv(tmp_path / f"AS_3_antismash_{suffix}.csv", headers, rows)
    lwc_row = {"bgc_id": "BGC001"}
    _write_csv(
        tmp_path / "AS_3_length_weighted_capacity.csv",
        list(lwc_row),
        [lwc_row],
    )
    (tmp_path / "AS_3_length_weighted_summary.json").write_text(json.dumps({
        "status": "TRIAL_ONLY_CAPACITY_METRIC", "raw_regions": 1,
    }), encoding="utf-8")

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, headers, rows in (
        ("MIBiG_Per_Gene", list(per_gene_row), [per_gene_row]),
        ("MIBiG_Convergence", list(convergence_row), [convergence_row]),
        ("MIBiG_Profile", list(profile_row), [profile_row]),
        ("antiSMASH_Modules", MODULE_COLUMNS, []),
        ("antiSMASH_RiPP_Motifs", RIPP_MOTIF_COLUMNS, []),
        ("antiSMASH_Motifs", MOTIF_COLUMNS, []),
        ("antiSMASH_RREfinder", RREFINDER_COLUMNS, [rre_row]),
        ("P_LWC_Trial", list(lwc_row), [lwc_row]),
        ("P_LWC_Summary", ["Field", "Value"], [{"Field": "status", "Value": "trial"}]),
    ):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
    workbook.save(tmp_path / "AS_5_workbook.xlsx")

    assert validate_reporting_v2_outputs(tmp_path)["status"] == "PASS"
    _write_csv(
        tmp_path / "AS_3_antismash_rrefinder.csv",
        RREFINDER_COLUMNS,
        [],
    )
    result = validate_reporting_v2_outputs(tmp_path)
    assert result["status"] == "FAIL"
    assert any("rrefinder JSON/CSV row mismatch" in error for error in result["errors"])


def test_reporting_evidence_is_not_consumed_by_scoring():
    from mamey import scoring

    source = Path(scoring.__file__).read_text(encoding="utf-8")
    assert "mibig_per_gene" not in source
    assert "antismash_structured" not in source
    assert "length_weighted" not in source
