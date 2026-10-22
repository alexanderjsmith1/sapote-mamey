"""v9.7.456: tools/manuscript_qa.py, a read-only check of a Word (.docx) deliverable.

Every document here is built in the test with python-docx from synthetic text. Author names, years, titles and
species epithets are invented; any match with a real publication is accidental. Genus names are real, because the
genus check needs them.
Each test names the check it covers; the checks are listed in the tool's docstring.
"""
from __future__ import annotations

import ast
import csv
import hashlib
import importlib.util
import sys
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
PIL_Image = pytest.importorskip("PIL.Image")
from docx.enum.text import WD_BREAK  # noqa: E402,F401  (python-docx present)
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches  # noqa: E402
from PIL.PngImagePlugin import PngInfo  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
TOOL = ROOT / "tools" / "manuscript_qa.py"


def _tool():
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    spec = importlib.util.spec_from_file_location("manuscript_qa_under_test", TOOL)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _runs(paragraph, parts):
    """parts: list of (text, italic)."""
    for text, italic in parts:
        paragraph.add_run(text).italic = italic or None
    return paragraph


def _hyperlink(paragraph, r_id, text, italic=False):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    if italic:
        rpr = OxmlElement("w:rPr")
        rpr.append(OxmlElement("w:i"))
        run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _png(path, px, dpi=None, text_pt=None):
    img = PIL_Image.new("RGB", px, (250, 250, 250))
    img.putpixel((0, 0), (len(path.name), 0, 0))     # distinct bytes per file
    info = PngInfo()
    if text_pt is not None:
        info.add_text("min_text_pt", str(text_pt))
    kw = {"pnginfo": info}
    if dpi:
        kw["dpi"] = (dpi, dpi)
    img.save(path, **kw)
    return path


def _findings(mod, path, **kw):
    return mod.run_checks(path, **kw)


def _by(findings, check):
    return [f for f in findings if f["check"] == check]


# ------------------------------------------------------------------------------------------------- check 1

@pytest.fixture
def cited_doc(tmp_path):
    d = docx.Document()
    d.add_paragraph("Earlier work (Smith et al., 2010; Jones, 2015) agrees with Brown (2001).")
    d.add_paragraph("A further claim rests on (Green, 2019) and on (Obdam, 2003; O'Bryn, 2004).")
    d.add_heading("References", level=2)
    for ref in ["Brown AB. 2001. A synthetic title. J Synth 1:1-2.",
                "Jones CD. 2014. Another synthetic title. J Synth 2:3-4.",
                "Smith EF, Lee GH. 2010. Third synthetic title. J Synth 3:5-6.",
                "Adams IJ. 2000. Never cited title. J Synth 4:7-8.",
                "Obdam KL. 2003. Ordered title. J Synth 5:9-10.",
                "O'Bryn MN. 2004. Apostrophe title. J Synth 6:11-12."]:
        d.add_paragraph(ref)
    d.add_heading("Supplementary Information", level=1)
    p = tmp_path / "cited.docx"
    d.save(p)
    return p


def test_check1_unmatched_citation(cited_doc):
    rows = _by(_findings(_tool(), cited_doc), "citation_unmatched")
    assert [r["para"] for r in rows] == [1]
    assert "green" in rows[0]["detail"].lower() and "2019" in rows[0]["detail"]


def test_check1_year_mismatch(cited_doc):
    rows = _by(_findings(_tool(), cited_doc), "citation_year_mismatch")
    assert len(rows) == 1 and rows[0]["para"] == 0
    assert "2015" in rows[0]["detail"] and "2014" in rows[0]["detail"]


def test_check1_uncited_reference(cited_doc):
    rows = _by(_findings(_tool(), cited_doc), "reference_uncited")
    texts = " ".join(r["quote"] for r in rows)
    assert "Adams" in texts and "Jones" in texts     # Jones 2014 is cited only with the wrong year
    assert "Brown" not in texts and "Smith" not in texts


def test_check1_alphabetical_order_ignores_apostrophes(cited_doc):
    rows = _by(_findings(_tool(), cited_doc), "reference_order")
    # Adams after Smith is out of order; Obdam before O'Bryn is in order once the apostrophe is ignored.
    assert len(rows) == 1 and "Adams" in rows[0]["quote"]
    assert rows[0]["para"] == 6


def test_check1_mixed_groups_brackets_particles_and_corporate_names(tmp_path):
    d = docx.Document()
    d.add_paragraph("Ordination (PCoA; Gowan 1971) used trees (TreeTool, Kozik et al. 2017; 100 replicates) "
                    "and markers (Lee, Kim, et al. 2020).")
    d.add_paragraph("Domains [Henrot 1993] were embedded as in van der Molen and Hint (2004), "
                    "searched (database of 14 July 2026; Camaro et al. 2011; Example Consortium 2025).")
    d.add_heading("References", level=2)
    for ref in ["Camaro C. 2011. Synthetic title one.", "Gowan JC. 1971. Synthetic title two.",
                "Henrot B. 1993. Synthetic title three.", "Kozik AM. 2017. Synthetic title four.",
                "Lee AB, Kim CD, Park EF. 2020. Synthetic title seven.",
                "Molen, L. van der, and G. Hint. 2004. Synthetic title five.",
                "The Example Consortium. 2025. Synthetic title six."]:
        d.add_paragraph(ref)
    p = tmp_path / "mixed.docx"
    d.save(p)
    found = _findings(_tool(), p)
    assert _by(found, "citation_unmatched") == [] and _by(found, "citation_year_mismatch") == []
    assert _by(found, "reference_uncited") == []      # a month and day are not an author


# ------------------------------------------------------------------------------------------------- check 2

def test_check2_skips_italic_words_and_species_epithets(tmp_path):
    d = docx.Document()
    _runs(d.add_paragraph(), [("Strain X1 matched S. sulphurinus and Streptomyces sulphurinus; the ", False),
                              ("colour", True), (" term is quoted, but the colour of plates is not.", False)])
    p = tmp_path / "epithet.docx"
    d.save(p)
    rows = _by(_findings(_tool(), p), "british_spelling")
    assert len(rows) == 1 and rows[0]["detail"].startswith('"colour"') and "plates" in rows[0]["quote"]


def test_check2_british_spellings_skip_reference_titles(tmp_path):
    d = docx.Document()
    d.add_paragraph("The colour was labelled grey, and we characterised the analyses in one modelling programme.")
    d.add_heading("References", level=2)
    d.add_paragraph("Smith EF. 2010. Colour and behaviour of synthetic things. J Synth 3:5-6.")
    p = tmp_path / "spell.docx"
    d.save(p)
    rows = _by(_findings(_tool(), p), "british_spelling")
    found = sorted(r["detail"].split(" ")[0].strip('"').lower() for r in rows)
    assert found == ["characterised", "colour", "grey", "labelled", "modelling", "programme"]
    assert all(r["para"] == 0 for r in rows)          # "analyses" (US plural noun) and the reference are not flagged


def test_check2_spelling_list_is_configurable(tmp_path):
    d = docx.Document()
    d.add_paragraph("The flavour and colour of the sample were noted.")
    p = tmp_path / "spell2.docx"
    d.save(p)
    words = tmp_path / "words.tsv"
    words.write_text("# british\tamerican\nflavour\tflavor\n")
    rows = _by(_findings(_tool(), p, spelling_file=words), "british_spelling")
    assert len(rows) == 1 and "flavor" in rows[0]["detail"]


# ------------------------------------------------------------------------------------------------- check 3

def test_check3_upright_genus_italic_family_and_hyperlink_text(tmp_path):
    d = docx.Document()
    para = _runs(d.add_paragraph(), [("Strains of ", False), ("Streptomyces", False), (" and ", False),
                                     ("Bacillus", True), (" in the family ", False), ("Streptomycetaceae", True),
                                     (" were kept.", False)])
    rel = d.part.relate_to("https://example.org/synthetic", RT.HYPERLINK, is_external=True)
    _hyperlink(para, rel, " See Micromonospora here.")
    p = tmp_path / "genus.docx"
    d.save(p)
    found = _findings(_tool(), p)
    upright = sorted(r["detail"] for r in _by(found, "genus_upright"))
    assert len(upright) == 2 and "Streptomyces" in upright[1] and "Micromonospora" in upright[0]
    fam = _by(found, "family_italic")
    assert len(fam) == 1 and "Streptomycetaceae" in fam[0]["detail"]


def test_check3_genus_file_and_caption_panel_labels(tmp_path):
    d = docx.Document()
    _runs(d.add_paragraph(), [("Fooella", False), (" grew; Streptomyces is not in the supplied list.", False)])
    _runs(d.add_paragraph(), [("Fig. 1. Two panels. (", False), ("A", True), (") Left side. (B) Right side.", False)])
    _runs(d.add_paragraph(), [("Body text with an italic (", False), ("C", True), (") outside a caption.", False)])
    p = tmp_path / "panel.docx"
    d.save(p)
    genera = tmp_path / "genera.txt"
    genera.write_text("Fooella\n")
    found = _findings(_tool(), p, genus_file=genera)
    assert [r["para"] for r in _by(found, "genus_upright")] == [0]
    panels = _by(found, "panel_label_italic")
    assert [r["para"] for r in panels] == [1] and "(A)" in panels[0]["detail"]


# ------------------------------------------------------------------------------------------------- check 4

def test_check4_short_sentences_in_body_only(tmp_path):
    d = docx.Document()
    d.add_heading("Short heading", level=2)
    d.add_paragraph("This is short. Smith et al. reported that this sentence is long enough to pass the check.")
    d.add_paragraph("Fig. 2. Short caption title. The caption sentence after it is also long enough here.")
    d.add_paragraph("A second paragraph holds a sentence that is long enough. Too brief here.")
    p = tmp_path / "short.docx"
    d.save(p)
    rows = _by(_findings(_tool(), p), "short_sentence")
    assert [(r["para"], r["quote"]) for r in rows] == [(1, "This is short."), (3, "Too brief here.")]


def test_check4_run_in_labels_are_not_sentences(tmp_path):
    d = docx.Document()
    para = d.add_paragraph()
    para.add_run("Assembly.").bold = True
    para.add_run(" Reads were assembled and the contigs were kept above a fixed length.")
    d.add_paragraph("Methods. Counts come from the deposited records of every isolate in the set.")
    d.add_paragraph("A long enough opening sentence sits here in the paragraph. It was short.")
    p = tmp_path / "runin.docx"
    d.save(p)
    rows = _by(_findings(_tool(), p), "short_sentence")
    assert [(r["para"], r["quote"]) for r in rows] == [(2, "It was short.")]


# ------------------------------------------------------------------------------------------------- check 5

def test_check5_internal_names_default_list(tmp_path):
    d = docx.Document()
    d.add_paragraph("Values came from results.tsv, from tools/run_all.py and from Tier A in run 7.")
    d.add_paragraph("Doses were 5 mg/L and/or 2 µg/mL; see https://doi.org/10.1000/xyz and run 70.")
    p = tmp_path / "internal.docx"
    d.save(p)
    rows = _by(_findings(_tool(), p), "internal_name")
    hits = sorted(r["detail"].split('"')[1] for r in rows)
    assert hits == ["Tier A", "results.tsv", "run 7", "tools/run_all.py"]
    assert all(r["para"] == 0 for r in rows)


def test_check5_internal_names_configurable(tmp_path):
    d = docx.Document()
    d.add_paragraph("The SecretTool output was checked against results.tsv.")
    p = tmp_path / "internal2.docx"
    d.save(p)
    names = tmp_path / "names.txt"
    names.write_text("SecretTool\n")
    rows = _by(_findings(_tool(), p, internal_names_file=names), "internal_name")
    assert len(rows) == 1 and "SecretTool" in rows[0]["detail"]


# ------------------------------------------------------------------------------------------------- check 6

def test_check6_internal_hyperlink_target_and_undefined_style(tmp_path):
    d = docx.Document()
    img = _png(tmp_path / "a.png", (300, 150), dpi=100)
    d.add_picture(str(img), width=Inches(3))
    image_rid = d.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
    para = d.add_paragraph("A reference with ")
    _hyperlink(para, image_rid, "doi:10.1000/synthetic")
    good = d.add_paragraph("A good link ")
    _hyperlink(good, d.part.relate_to("https://doi.org/10.1000/ok", RT.HYPERLINK, is_external=True), "here")
    ghost = d.add_paragraph("A paragraph in a style the document does not define.")
    ghost._p.get_or_add_pPr().style = "GhostStyle"
    p = tmp_path / "links.docx"
    d.save(p)
    found = _findings(_tool(), p)
    links = _by(found, "hyperlink_internal_target")
    assert [r["para"] for r in links] == [1] and "media" in links[0]["detail"]
    styles = _by(found, "style_undefined")
    assert [r["para"] for r in styles] == [3] and "GhostStyle" in styles[0]["detail"]


def test_check6_undefined_run_style_is_one_row_with_a_count(tmp_path):
    d = docx.Document()
    for i in range(3):
        run = d.add_paragraph(f"Paragraph {i} carries a comment anchor run.").runs[0]
        run._r.get_or_add_rPr().style = "GhostRunStyle"
    p = tmp_path / "runstyle.docx"
    d.save(p)
    rows = _by(_findings(_tool(), p), "style_undefined")
    assert len(rows) == 1 and rows[0]["para"] == 0 and "3 paragraphs" in rows[0]["detail"]


# ------------------------------------------------------------------------------------------------- checks 7 and 8

@pytest.fixture
def image_doc(tmp_path):
    d = docx.Document()                                        # default template: 8.5 x 11 in, 1 in margins
    small = _png(tmp_path / "shrunk.png", (1000, 500), dpi=100, text_pt=8)      # drawn 10.0 in
    exact = _png(tmp_path / "exact.png", (600, 300), dpi=100)                    # drawn 6.0 in
    tall = _png(tmp_path / "tall.png", (100, 1000), dpi=100)                     # drawn 1 x 10 in
    nodpi = _png(tmp_path / "nodpi.png", (400, 200))
    d.add_picture(str(small), width=Inches(6.0))
    d.add_paragraph("Fig. 1. A figure placed at 60 percent of its drawn width.")
    d.add_picture(str(exact), width=Inches(6.0))
    d.add_paragraph("Fig. 2. A figure placed at its drawn width.")
    d.add_picture(str(tall), width=Inches(1.0))
    d.add_paragraph("Fig. 3. A figure taller than the text block.")
    d.add_picture(str(nodpi), width=Inches(2.0))
    p = tmp_path / "images.docx"
    d.save(p)
    old = tmp_path / "superseded"
    old.mkdir()
    (old / "old_render.png").write_bytes(small.read_bytes())
    return p, old


def test_check7_print_scale_and_declared_text_size(image_doc):
    path, _ = image_doc
    found = _findings(_tool(), path)
    scale = _by(found, "print_scale")
    assert [r["para"] for r in scale] == [0]
    assert "60.0%" in scale[0]["detail"] and "4.8 pt" in scale[0]["detail"]
    assert "Fig. 1." in scale[0]["quote"]
    assert [r["para"] for r in _by(found, "image_taller_than_text_block")] == [4]
    assert [r["para"] for r in _by(found, "image_no_dpi")] == [6]


def test_check7_threshold_is_configurable(image_doc):
    path, _ = image_doc
    assert _by(_findings(_tool(), path, min_scale=0.5), "print_scale") == []


def test_check8_superseded_render_by_bytes(image_doc):
    path, old = image_doc
    mod = _tool()
    assert _by(_findings(mod, path), "superseded_render") == []      # optional: off without a folder
    rows = _by(_findings(mod, path, superseded_dir=old), "superseded_render")
    assert [r["para"] for r in rows] == [0] and "old_render.png" in rows[0]["detail"]


# ------------------------------------------------------------------------------------------------- reports and guards

def test_reports_read_only_short_quotes_and_tsv_columns(cited_doc, tmp_path):
    mod = _tool()
    before = hashlib.sha256(cited_doc.read_bytes()).hexdigest()
    out = tmp_path / "out"
    assert mod.main([str(cited_doc), "--out", str(out)]) == 0
    assert hashlib.sha256(cited_doc.read_bytes()).hexdigest() == before
    tsv = next(out.glob("*_manuscript_qa.tsv"))
    md = next(out.glob("*_manuscript_qa.md"))
    with tsv.open(encoding="utf-8", newline="") as fh:
        rows = list(csv.reader(fh, delimiter="\t"))
    assert rows[0] == ["check", "¶", "quote", "detail"]
    assert len(rows) > 1 and all(len(r[2].split()) <= 15 for r in rows[1:])
    text = md.read_text(encoding="utf-8")
    assert before in text and "citation_unmatched" in text


def test_long_paragraph_quote_is_at_most_15_words(tmp_path):
    d = docx.Document()
    d.add_paragraph(" ".join(["word"] * 40) + " colour " + " ".join(["word"] * 40) + ".")
    p = tmp_path / "long.docx"
    d.save(p)
    rows = _by(_findings(_tool(), p), "british_spelling")
    assert len(rows) == 1 and len(rows[0]["quote"].split()) <= 15 and "colour" in rows[0]["quote"]


def test_refuses_non_docx(tmp_path):
    bad = tmp_path / "notes.txt"
    bad.write_text("plain text")
    assert _tool().main([str(bad), "--out", str(tmp_path / "o")]) == 2


def test_tool_emits_through_logging_not_print():
    tree = ast.parse(TOOL.read_text(encoding="utf-8"))
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)]
    assert not [n for n in calls if isinstance(n.func, ast.Name) and n.func.id == "print"]
    assert not [n for n in calls if isinstance(n.func, ast.Attribute) and n.func.attr == "write"
                and isinstance(n.func.value, ast.Attribute) and n.func.value.attr in {"stdout", "stderr"}]


# ------------------------------------------------------------------------------------------------- check 1, review round

def _ref_doc(tmp_path, name, body, refs, heading="References"):
    d = docx.Document()
    for text in body:
        d.add_paragraph(text)
    d.add_heading(heading, level=2)
    for ref in refs:
        d.add_paragraph(ref)
    p = tmp_path / name
    d.save(p)
    return p


def test_check1_year_suffix_list_cites_every_suffix(tmp_path):
    p = _ref_doc(tmp_path, "suffix.docx", ["Both reports agree (Fakeman et al. 2019a,b) and (Toller 2020a, b)."],
                 ["Fakeman A. 2019a. Synthetic one.", "Fakeman A. 2019b. Synthetic two.",
                  "Toller B. 2020a. Synthetic three.", "Toller B. 2020b. Synthetic four."])
    found = _findings(_tool(), p)
    assert _by(found, "reference_uncited") == [] and _by(found, "citation_unmatched") == []


def test_check1_three_authors_with_and_is_the_first_author(tmp_path):
    p = _ref_doc(tmp_path, "three.docx", ["One view (Quill, Varno, and Brisk 2019) and another (Quill, Varno and Brisk 2019)."],
                 ["Brisk D. 2019. A different synthetic paper.", "Quill, A., B. Varno, and C. Brisk. 2019. Synthetic."])
    found = _findings(_tool(), p)
    assert _by(found, "citation_unmatched") == []
    uncited = _by(found, "reference_uncited")
    assert len(uncited) == 1 and "Brisk" in uncited[0]["detail"]   # Brisk's own 2019 paper is not cited


def test_check1_order_compares_given_names_within_a_surname(tmp_path):
    body = ["Cited (Quill 2019; Quill 2020; Quill 2021; Rauvik 2018; R Core Team 2022)."]
    p = _ref_doc(tmp_path, "given.docx", body,
                 ["Quill, Adam. 2019. Synthetic one.", "Quill, John. 2020. Synthetic two.",
                  "Quill, Bert. 2021. Synthetic three.", "R Core Team. 2022. Synthetic manual.",
                  "Rauvik, C. 2018. Synthetic four."])
    rows = _by(_findings(_tool(), p), "reference_order")
    # Quill, Bert after Quill, John is out of order; letter by letter "rauvik" sorts before "rcoreteam".
    assert sorted(r["para"] for r in rows) == [4, 6]


def test_check1_order_same_initials_are_one_person(tmp_path):
    body = ["Cited (Quill 2019; Quill 2020)."]
    p = _ref_doc(tmp_path, "initials.docx", body,
                 ["Quill, C. R. 2019. Synthetic one.", "Quill, Carl R. 2020. Synthetic two."])
    assert _by(_findings(_tool(), p), "reference_order") == []


def test_check1_other_reference_headings_are_recognized(tmp_path):
    for heading in ("Works Cited", "Literature Cited", "Bibliography", "References Cited", "Reference List"):
        p = _ref_doc(tmp_path, "works.docx", ["A claim (Nobodi 2019) and (Fakeman 2018)."],
                     ["Fakeman A. 2018. Synthetic."], heading=heading)
        rows = _by(_findings(_tool(), p), "citation_unmatched")
        assert len(rows) == 1 and "Nobodi" in rows[0]["detail"], heading


def test_check1_missing_reference_list_is_a_hold_not_a_clean_report(tmp_path):
    p = _ref_doc(tmp_path, "nolist.docx", ["A claim (Nobodi 2019) rests here."], ["Fakeman A. 2018. Synthetic."],
                 heading="Sources and further notes")
    found = _findings(_tool(), p)
    holds = _by(found, "citation_hold")
    assert len(holds) == 1 and "no reference list" in holds[0]["detail"] and "1 citation" in holds[0]["detail"]
