"""Round-2 caller-oriented regressions for swallowed evidence failures."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from mamey import cli
from mamey import judgment_store
from mamey import mode_b_receipt as receipt
from tests._modeb_card_fixtures import valid_modeb_card_stub


STRAIN = "FIXTURE-02"
BGC = "BGC001"
NODE = "NODE_1_length_12000_cov_20.000000"
REGION = "region001"


def _package(tmp_path: Path) -> Path:
    package = tmp_path / "package"
    package.mkdir(parents=True)
    (package / "manifest.json").write_text(
        json.dumps({"strain_id": STRAIN}), encoding="utf-8"
    )
    judgment_store.init_register(package, strain_id=STRAIN, bgc_ids=[BGC])
    (package / f"{STRAIN}_2_inventory.csv").write_text(
        "BGC_ID,Contig,Node_ID,antiSMASH_Region\n"
        f"{BGC},{NODE},{NODE},{REGION}\n",
        encoding="utf-8",
    )
    (package / f"{STRAIN}_4_triage_board.csv").write_text(
        "BGC_ID,Products,Corrected_rank,Boundary,Lead_tier_auto\n"
        f"{BGC},NRPS,1,Interior,Inventory\n",
        encoding="utf-8",
    )
    judgment = package / "judgment"
    judgment.mkdir()
    (judgment / f"{STRAIN}_{BGC}_mode_b.md").write_text(
        valid_modeb_card_stub(BGC, node=NODE, strain_id=STRAIN), encoding="utf-8"
    )
    return package


def _card() -> str:
    return valid_modeb_card_stub(BGC, node=NODE, strain_id=STRAIN)


def _receipt_file(tmp_path: Path) -> Path:
    path = tmp_path / "mode_b_receipt.json"
    path.write_text(json.dumps({
        "schema_version": receipt.RECEIPT_SCHEMA_VERSION,
        "strain_id": STRAIN,
        "cards": [{"bgc_id": BGC, "mode_b_md": _card()}],
    }), encoding="utf-8")
    return path


def test_package_context_merge_failure_is_explicit_and_success_is_clean(
    tmp_path, monkeypatch
):
    package = _package(tmp_path)
    from mamey import authored_verify

    original = authored_verify._bgc_context_from_package
    monkeypatch.setattr(
        authored_verify,
        "_bgc_context_from_package",
        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("fixture merge failure")),
    )
    failed = receipt.lint_card_in_package(package, BGC, card_md=_card())
    assert any(row.get("code") == "VERIFICATION_CONTEXT_MERGE_FAILED" for row in failed)

    monkeypatch.setattr(authored_verify, "_bgc_context_from_package", original)
    clean = receipt.lint_card_in_package(package, BGC, card_md=_card())
    assert not any(row.get("code") == "VERIFICATION_CONTEXT_MERGE_FAILED" for row in clean)


def test_context_findings_survive_context_lint_failure_and_bare_fallback(
    tmp_path, monkeypatch
):
    package = _package(tmp_path)
    (package / f"{STRAIN}_4_triage_board.csv").write_text(
        "wrong,header\nvalue,row\n", encoding="utf-8"
    )
    from mamey import modeb_structure_gate

    def fail_context(card_md, **kwargs):
        if "bgc_context" in kwargs:
            raise RuntimeError("fixture contextual lint failure")
        return [{"severity": "WARN", "code": "BARE_FALLBACK", "section": None,
                 "message": "context-free fallback exercised"}]

    monkeypatch.setattr(modeb_structure_gate, "lint_card", fail_context)
    findings = receipt.lint_card_in_package(package, BGC, card_md=_card())
    codes = {row.get("code") for row in findings}
    assert {"VERIFICATION_CONTEXT_MALFORMED", "CONTEXT_LINT_FAILED", "BARE_FALLBACK"} <= codes


def test_context_findings_survive_dual_lint_failure(tmp_path, monkeypatch):
    package = _package(tmp_path)
    (package / f"{STRAIN}_4_triage_board.csv").write_text(
        "wrong,header\nvalue,row\n", encoding="utf-8"
    )
    from mamey import modeb_structure_gate

    monkeypatch.setattr(
        modeb_structure_gate,
        "lint_card",
        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("fixture dual failure")),
    )
    findings = receipt.lint_card_in_package(package, BGC, card_md=_card())
    codes = {row.get("code") for row in findings}
    assert {"VERIFICATION_CONTEXT_MALFORMED", "CONTEXT_LINT_FAILED", "GATE_UNAVAILABLE"} <= codes


def test_invalid_utf8_triage_refused_by_batch_ingest_but_absence_is_accepted(tmp_path):
    malformed = _package(tmp_path / "malformed")
    (malformed / f"{STRAIN}_4_triage_board.csv").write_bytes(
        b"BGC_ID,Products\nBGC001,\xff\n"
    )
    malformed_findings = receipt.lint_card_in_package(
        malformed, BGC, card_md=_card()
    )
    assert any(row.get("code") == "VERIFICATION_CONTEXT_MALFORMED"
               for row in malformed_findings)
    refused = receipt.ingest_receipt(malformed, _receipt_file(tmp_path / "malformed"))
    assert refused["recorded"] == []
    assert refused["skipped_structure_invalid"] == [(BGC, 1)]

    absent = _package(tmp_path / "absent")
    (absent / f"{STRAIN}_4_triage_board.csv").unlink()
    accepted = receipt.ingest_receipt(absent, _receipt_file(tmp_path / "absent"))
    assert accepted["recorded"] == [BGC]


def test_valid_non_object_gene_context_is_explicit_at_lint_caller(tmp_path):
    package = _package(tmp_path)
    ledger = package / f"{STRAIN}_gene_context.jsonl"
    ledger.write_text("[]\n", encoding="utf-8")
    malformed = receipt.lint_card_in_package(package, BGC, card_md=_card())
    assert any(row.get("code") == "VERIFICATION_CONTEXT_MALFORMED" for row in malformed)

    ledger.unlink()
    absent = receipt.lint_card_in_package(package, BGC, card_md=_card())
    assert not any(row.get("code") == "VERIFICATION_CONTEXT_MALFORMED" for row in absent)


def test_partial_satellite_csv_refused_by_auto_detect_but_absence_is_accepted(tmp_path):
    malformed = _package(tmp_path / "malformed")
    (malformed / f"{STRAIN}_3_antismash_modules.csv").write_text(
        "bgc_id,feature_type\n"
        f"{BGC},aSModule,unexpected\n",
        encoding="utf-8",
    )
    malformed_findings = receipt.lint_card_in_package(
        malformed, BGC, card_md=_card()
    )
    assert any(row.get("code") == "VERIFICATION_CONTEXT_MALFORMED"
               for row in malformed_findings)
    refused = receipt.auto_detect_ingest(malformed)
    assert refused["recorded"] == []
    assert refused["skipped_structure_invalid"]

    absent = _package(tmp_path / "absent")
    accepted = receipt.auto_detect_ingest(absent)
    assert accepted["recorded"] == [BGC]


def test_auto_detect_recovery_diagnostics_use_one_renderer_call(
    tmp_path, monkeypatch, capsys
):
    package = _package(tmp_path)
    summary = receipt.auto_detect_ingest(package)
    summary.update({
        "recorded": [],
        "skipped_unreadable": [{"identity": "UNBOUND", "error_type": "OSError"}],
        "skipped_no_content": ["UNBOUND"],
        "record_failures": [{"identity": "UNBOUND", "error_type": "OSError"}],
    })
    calls = []
    original_emit = receipt.emit

    def counted_emit(*args, **kwargs):
        if kwargs.get("file") is not None:
            calls.append(args)
        return original_emit(*args, **kwargs)

    monkeypatch.setattr(receipt, "auto_detect_ingest", lambda *a, **k: summary)
    monkeypatch.setattr(receipt, "emit", counted_emit)
    args = argparse.Namespace(
        package=str(package), auto_detect=True, receipt=None, card=None,
        force_structure=False, master=None,
    )
    assert receipt.ingest_receipts_command(args) == 1
    err = capsys.readouterr().err
    assert "FAILED (unreadable): 1 card(s)" in err
    assert "SKIPPED (empty): 1 card(s)" in err
    assert "FAILED (persistence): 1 card(s)" in err
    assert len(calls) == 1


def test_terminal_status_upgrades_normalized_completion_when_late_issues_exist():
    assert cli._terminal_mamey_status("MAMEY_COMPLETE", []) == "MAMEY_COMPLETE"
    assert cli._terminal_mamey_status(
        "MAMEY_COMPLETE", ["fixture late issue"]
    ) == "MAMEY_COMPLETE_WITH_ISSUES"


def _run_fixture(
    input_zip: Path,
    outdir: Path,
    *,
    brief: str = "none",
    master_path=None,
    locus_maps: str = "auto",
):
    return cli.run_one_strain(
        strain_id=STRAIN,
        display_name="Generic fixture strain",
        input_zip=str(input_zip),
        outdir=str(outdir),
        mode="gold",
        taxonomy="not verified",
        source="not supplied",
        bioactivity="",
        master_path=str(master_path) if master_path is not None else None,
        json_mode="off",
        brief=brief,
        locus_maps=locus_maps,
    )


def test_bgc_bank_failure_is_structured_and_success_remains_clean(tmp_path):
    issues = []
    malformed = {
        "strain_id": STRAIN,
        "assembly": {},
        "bgc_counts": {},
        "bgcs": [{"bgc_id": BGC, "region_number": "not-an-integer"}],
    }
    assert cli._emit_bgc_bank(tmp_path / "bad", malformed, issues=issues) is False
    assert any("BGC_BANK_WRITE_FAILED" in issue for issue in issues)

    good_dir = tmp_path / "good"
    good_dir.mkdir()
    clean_issues = []
    valid = {
        "strain_id": STRAIN,
        "assembly": {},
        "bgc_counts": {},
        "bgcs": [{"bgc_id": BGC, "region_number": 1, "contig": NODE}],
    }
    assert cli._emit_bgc_bank(good_dir, valid, issues=clean_issues) is True
    assert (good_dir / "bgc_data.json").is_file()
    assert clean_issues == []


def test_run_records_entrypoint_write_and_record_limit_probe_failures(
    tmp_path, monkeypatch, synthetic_single_contig_full_locus_zip
):
    original_write_text = Path.write_text
    original_atomic_text = cli._atomic_write_manifest_text

    def fail_entrypoints(self, *args, **kwargs):
        if self.name in {
            f"{STRAIN}_BATCH_PLAN.md", "START_HERE.md", "FIGURES_SUPPLEMENTARY.md"
        }:
            raise OSError(f"fixture refusal for {self.name}")
        return original_write_text(self, *args, **kwargs)

    def fail_node_map(path, content, *args, **kwargs):
        if Path(path).name == "node_citation_map.json":
            raise OSError("fixture node map refusal")
        return original_atomic_text(path, content, *args, **kwargs)

    from mamey import parsers
    from mamey import assembly
    monkeypatch.setattr(Path, "write_text", fail_entrypoints)
    monkeypatch.setattr(cli, "_atomic_write_manifest_text", fail_node_map)
    monkeypatch.setattr(cli, "_input_zip_sha256", lambda *a, **k: None)
    monkeypatch.setattr(
        cli,
        "_render_brief_nonblocking",
        lambda *a, **k: {"status": "COMPLETE", "tier": "standard", "files": []},
    )
    monkeypatch.setattr(
        parsers,
        "antismash_record_limit_truncation",
        lambda *a, **k: (_ for _ in ()).throw(ValueError("fixture limit probe failure")),
    )
    monkeypatch.setattr(
        assembly,
        "assembly_sanity_check",
        lambda *a, **k: (_ for _ in ()).throw(ValueError("fixture sanity probe failure")),
    )
    from mamey import rescue_two_proof
    monkeypatch.setattr(
        rescue_two_proof,
        "two_proof_join",
        lambda *a, **k: (_ for _ in ()).throw(ValueError("fixture two-proof failure")),
    )
    result = _run_fixture(
        synthetic_single_contig_full_locus_zip, tmp_path, brief="standard"
    )
    issues = result["issues"]
    expected_issue_codes = {
        "BATCH_PLAN_WRITE_FAILED",
        "START_HERE_WRITE_FAILED",
        "RECORD_LIMIT_PROBE_FAILED",
        "ASSEMBLY_SANITY_PROBE_FAILED",
        "FIGURES_SUPPLEMENTARY_MARKER_WRITE_FAILED",
        "TWO_PROOF_RESCUE_WRITE_FAILED",
        "NODE_CITATION_MAP_WRITE_FAILED",
        "INPUT_ZIP_SHA256_UNAVAILABLE",
    }
    observed_issue_codes = {
        code for code in expected_issue_codes
        if any(code in issue for issue in issues)
    }
    assert observed_issue_codes == expected_issue_codes
    assert result["status"] == "MAMEY_COMPLETE_WITH_ISSUES"
    package = tmp_path / STRAIN / "package"
    phase_rows = [json.loads(line) for line in
                  (package / "run_phase_receipts.jsonl").read_text().splitlines()]
    assert any(row["phase"] == "batch_plan" and row["status"] == "ERROR"
               for row in phase_rows)
    assert any(row["phase"] == "start_here" and row["status"] == "ERROR"
               for row in phase_rows)
    assert any(row["phase"] == "record_limit_probe" and row["status"] == "ERROR"
               for row in phase_rows)
    assert any(row["phase"] == "assembly_sanity_probe" and row["status"] == "ERROR"
               for row in phase_rows)
    assert any(row["phase"] == "figures_supplementary_marker" and row["status"] == "ERROR"
               for row in phase_rows)
    assert any(row["phase"] == "two_proof_rescue" and row["status"] == "ERROR"
               for row in phase_rows)
    assert any(row["phase"] == "node_citation_map" and row["status"] == "ERROR"
               for row in phase_rows)
    assert any(row["phase"] == "input_zip_sha256" and row["status"] == "ERROR"
               for row in phase_rows)


def test_run_entrypoint_and_record_limit_negative_control(
    tmp_path, monkeypatch, synthetic_single_contig_full_locus_zip
):
    from mamey import parsers
    monkeypatch.setattr(
        parsers, "antismash_record_limit_truncation", lambda *a, **k: {"truncated": False}
    )
    result = _run_fixture(synthetic_single_contig_full_locus_zip, tmp_path)
    package = tmp_path / STRAIN / "package"
    assert (package / f"{STRAIN}_BATCH_PLAN.md").is_file()
    assert (package / "START_HERE.md").is_file()
    assert not any("_WRITE_FAILED" in issue for issue in result["issues"])
    assert not any("RECORD_LIMIT_PROBE_FAILED" in issue for issue in result["issues"])
    assert not any("INPUT_ZIP_SHA256_UNAVAILABLE" in issue for issue in result["issues"])
    manifest = json.loads((package / "manifest.json").read_text())
    assert manifest["input_zip_sha256"]
    forbidden = re.compile(r"/Users/|Codex Alex|Claude_Alex|\bAS-\d{1,6}\b")
    for path in package.rglob("*"):
        if path.is_file() and path.suffix.lower() in {
            ".md", ".txt", ".json", ".jsonl", ".csv", ".html"
        }:
            assert not forbidden.search(path.read_text(encoding="utf-8", errors="replace")), path


def test_nonstandard_gold_figure_note_is_path_generic(tmp_path, monkeypatch):
    from mamey import cohort_figures

    monkeypatch.setattr(
        cohort_figures, "generate", lambda **kwargs: {"figures": 0}
    )
    package = tmp_path / "operator-private-root" / "odd-package-name"
    package.mkdir(parents=True)
    cli._emit_gold_figures(package, STRAIN, "gold")
    note = (package / "gold_figures" / "RUNS_DIR_NOTE.md").read_text()
    assert str(tmp_path) not in note
    assert "operator-private-root" not in note
    assert "<runs_dir>/<strain_id>/package/" in note


def test_outdir_resolution_failure_refuses_instead_of_allowing(tmp_path, monkeypatch, capsys):
    original_resolve = Path.resolve
    target = tmp_path / "unresolvable-output"
    target.mkdir()

    def fail_selected(self, *args, **kwargs):
        if self.name == "unresolvable-output":
            raise OSError("fixture resolution failure")
        return original_resolve(self, *args, **kwargs)

    monkeypatch.setattr(Path, "resolve", fail_selected)
    assert cli._refuse_outdir_inside_bundle(target) is True
    assert "refusing output" in capsys.readouterr().err


def test_master_lock_failure_preserves_master_and_records_late_issue(
    tmp_path, monkeypatch, synthetic_single_contig_full_locus_zip
):
    import openpyxl

    master = tmp_path / "generic-master.xlsx"
    workbook = openpyxl.Workbook()
    workbook.save(master)
    workbook.close()
    before = master.read_bytes()
    monkeypatch.setattr(cli, "_acquire_master_workbook_lock", lambda *a, **k: None)
    result = _run_fixture(
        synthetic_single_contig_full_locus_zip, tmp_path / "run", master_path=master
    )
    assert master.read_bytes() == before
    assert any("MASTER_WORKBOOK_LOCK_UNAVAILABLE" in issue for issue in result["issues"])
    assert result["status"] == "MAMEY_COMPLETE_WITH_ISSUES"


def test_master_schema_probe_failure_preserves_master_and_records_receipt(
    tmp_path, monkeypatch, synthetic_single_contig_full_locus_zip
):
    import openpyxl

    master = tmp_path / "generic-master.xlsx"
    master.write_bytes(b"fixture master bytes")
    before = master.read_bytes()
    original_load = openpyxl.load_workbook

    def fail_selected(path, *args, **kwargs):
        if Path(path) == master:
            raise OSError("fixture schema probe failure")
        return original_load(path, *args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", fail_selected)
    result = _run_fixture(
        synthetic_single_contig_full_locus_zip, tmp_path / "run", master_path=master
    )
    assert master.read_bytes() == before
    assert any("MASTER_WORKBOOK_SCHEMA_PROBE_FAILED" in issue for issue in result["issues"])
    assert result["status"] == "MAMEY_COMPLETE_WITH_ISSUES"
    rows = [json.loads(line) for line in (
        tmp_path / "run" / STRAIN / "package" / "run_phase_receipts.jsonl"
    ).read_text().splitlines()]
    assert any(row["phase"] == "master_workbook_schema_probe" and row["status"] == "ERROR"
               for row in rows)


def test_no_figures_marker_failure_is_structured(
    tmp_path, monkeypatch, synthetic_single_contig_full_locus_zip
):
    original_write_text = Path.write_text

    def fail_marker(self, *args, **kwargs):
        if self.name == "NO_FIGURES_RENDERED.md":
            raise OSError("fixture marker refusal")
        return original_write_text(self, *args, **kwargs)

    monkeypatch.setattr(cli, "_emit_gold_figures", lambda *a, **k: None)
    monkeypatch.setattr(Path, "write_text", fail_marker)
    result = _run_fixture(
        synthetic_single_contig_full_locus_zip, tmp_path, locus_maps="off"
    )
    assert any("NO_FIGURES_MARKER_WRITE_FAILED" in issue for issue in result["issues"])
    assert result["status"] == "MAMEY_COMPLETE_WITH_ISSUES"


def test_locus_map_catalog_failure_is_structured(
    tmp_path, monkeypatch, synthetic_single_contig_full_locus_zip
):
    import builtins

    original_open = builtins.open

    def fail_catalog(path, *args, **kwargs):
        if Path(path).name == "locus_map_manifest.csv":
            raise OSError("fixture catalog refusal")
        return original_open(path, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", fail_catalog)
    result = _run_fixture(
        synthetic_single_contig_full_locus_zip, tmp_path, locus_maps="on"
    )
    assert any("LOCUS_MAP_CATALOG_WRITE_FAILED" in issue for issue in result["issues"])
    assert result["status"] == "MAMEY_COMPLETE_WITH_ISSUES"
