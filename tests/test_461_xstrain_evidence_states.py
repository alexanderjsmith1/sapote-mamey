"""Reporting regressions for scoped policy and missing KCB evidence.

All records below are synthetic software fixtures, not biological results.
"""
from __future__ import annotations

import json
import pathlib
import subprocess
import sys

import openpyxl
import pytest

ROOT = pathlib.Path(__file__).resolve().parents[1]


def _run_overlay(tmp_path, values):
    bank = tmp_path / 'bank'
    bank.mkdir()
    strains = {'SYNTHETIC_REFERENCE': {'n50': 1000000, 'raw_bgcs': len(values),
                                       'corrected_bgcs': len(values)}}
    bgcs = []
    for index, value in enumerate(values, 1):
        row = {'sid': 'SYNTHETIC_REFERENCE', 'products': 'NAPAA',
               'bgc_id': f'BGC{index:03d}', 'contig': 'SYNTHETIC_CONTIG.1',
               'region': f'region{index:03d}', 'edge_status': 'Interior',
               'kcb_cumulative': None,
               'source_kcb_file': 'knownclusterblast/synthetic_c1.txt',
               'kcb_provenance': 'knownclusterblast'}
        if value != '__ABSENT_FIXTURE_FIELD__':
            row['kcb_top'] = value
        bgcs.append(row)
    (bank / 'bgc_data.json').write_text(json.dumps({'strains': strains, 'bgcs': bgcs}))
    for name in ['gene_data.json', 'tigrfam.json', 'deep_data.json', 'rggmci_full.json']:
        (bank / name).write_text('{}')
    original = {p.name: p.read_bytes() for p in bank.glob('*.json')}
    path = tmp_path / 'fixture.xlsx'
    wb = openpyxl.Workbook()
    wb.active.title = 'UserNotes'
    wb.active['A1'] = 'Preserve this user cell'
    wb.save(path)
    wb.close()
    result = subprocess.run([sys.executable, str(ROOT / 'tools/add_xstrain_sheets.py'),
                             '--banked-dir', str(bank), '--out', str(path)],
                            cwd=ROOT, text=True, capture_output=True, timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    for name, data in original.items():
        assert (bank / name).read_bytes() == data
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        assert wb['UserNotes']['A1'].value == 'Preserve this user cell'
        findings = {r[0]: r for r in wb['Cross_Strain_Findings'].iter_rows(values_only=True)
                    if isinstance(r[0], int)}
        context_rows = list(wb['Strain_Cohort_Context'].iter_rows(values_only=True))
        context = dict(zip(context_rows[0], context_rows[1]))
    finally:
        wb.close()
    return findings, context, json.loads((bank / 'cohort_context.json').read_text())


def test_napaa_note_states_family_deck_scope_and_neutrality(tmp_path):
    findings, _, _ = _run_overlay(tmp_path, ['synthetic comparison label'])
    row = findings[6]
    assert row[1] == 'NAPAA family/deck scope'
    assert row[4] == 'SCOPED POLICY'
    assert 'family analyses and strain decks' in row[5]
    assert 'neutral elsewhere' in row[5]
    assert 'does not apply an engine exclusion' in row[5]


@pytest.mark.parametrize('missing', [None, '', '   ', 'UNRESOLVED', '__ABSENT_FIXTURE_FIELD__'])
def test_missing_top_hit_is_unknown_search_state_not_novelty(tmp_path, missing):
    findings, context, derived = _run_overlay(tmp_path, [missing, 'synthetic comparison label'])
    row = findings[7]
    assert row[1] == 'No KnownClusterBlast hit recorded'
    assert row[3] == '1/2 (50%)'
    assert row[4] == 'FIELD SUMMARY'
    assert 'Search state UNKNOWN' in row[5]
    assert 'not searched versus searched with no hit' in row[5]
    assert 'not novelty results' in row[5]
    assert context['no_KCB_hit_recorded_regions'] == 1
    assert context['no_KCB_hit_recorded_share'] == '1 of 1'
    assert context['missing_KCB_search_state'] == 'UNKNOWN'
    assert 'zero_KCB_regions' not in context
    # Compatibility keys retain a count of missing labels, never a filled score.
    assert derived['SYNTHETIC_REFERENCE']['zero_kcb'] == 1
    assert derived['SYNTHETIC_REFERENCE']['zk_share'] == '1 of 1'
    assert derived['SYNTHETIC_REFERENCE']['kcb_missing_hit_search_state'] == 'UNKNOWN'
    assert derived['SYNTHETIC_REFERENCE']['rank'] == 1


def test_source_filename_alone_does_not_prove_a_completed_search(tmp_path):
    # The normal bank only records a source filename/provenance, not a search receipt.
    findings, _, _ = _run_overlay(tmp_path, [None])
    assert findings[7][4] == 'FIELD SUMMARY'
    assert 'Search state UNKNOWN' in findings[7][5]
