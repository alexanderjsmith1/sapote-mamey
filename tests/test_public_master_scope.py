"""Public saved-master projections: synthetic identities, sealed JSON, no source scans/jobs."""
import csv
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from mamey.class_architecture import with_family_labels
from mamey.cohort_deliverable import run_cohort_deliverable
from mamey.cohort_synthesis import (
    PublicCohortScopeError, load_public_master, render_public_synthesis,
    recheck_public_scope, write_synthesis,
)
from mamey.crosswalk import assembly_locator
from mamey.master_workbook import (
    CANONICAL_V1_HEADERS, _b2_product_class_counts, _b2_product_class_rows,
    _gene_backed_carotenoid_bgcs,
)


def _json(path, value):
    path.write_text(json.dumps(value), encoding="utf-8")


def _seal(package):
    rows = [f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.name}"
            for p in sorted(package.iterdir()) if p.is_file() and p.name != "checksums_sha256.txt"]
    (package / "checksums_sha256.txt").write_text("\n".join(rows) + "\n")


def _bgc(index, products, *, backed=False):
    return {"bgc_id": f"BGC{index:03d}", "contig": f"ctg{index}", "node_id": f"ctg{index}",
            "antismash_region": f"region{index:03d}", "region_number": index,
            "start": 0, "end": 100, "products": products, "kcb_top": None,
            "_backed": backed}


def _manifest(sid, bgcs, *, release="PUBLIC"):
    for b in bgcs:
        b["release"] = release
    per = {b["bgc_id"]: {"domain_counts": {"PKS_KS": i + 1, "NRPS_A": i % 2},
            "domains": [{"domain": "phytoene synthase"}] if b["_backed"] else []}
           for i, b in enumerate(bgcs)}
    return {"strain_id": sid, "release": release, "workflow_version": "Mamey v1.9.174",
            "taxonomy": "Referenceus sp.", "source": "public reference record", "bgcs": bgcs,
            "bgc_counts": {"raw": len(bgcs), "corrected": len(bgcs), "assembly_tier": "GOOD"},
            "source_scans": {"domain_architecture": {"status": "SOURCE_DERIVED_STRUCTURED_DOMAINS", "per_bgc": per},
             "primary_metabolism": {"status": "SOURCE_DERIVED", "hits": {"pigment": []},
               "per_bgc": {b["bgc_id"]: {"families": []} for b in bgcs}}}}


def _package(runs, sid, man, *, folder=None):
    pkg = runs / (folder or sid) / "package"
    pkg.mkdir(parents=True)
    _json(pkg / "manifest.json", man)
    _json(pkg / "manifest_short.json", {"strain_id": sid, "release": man["release"],
          "taxonomy": man["taxonomy"], "raw_bgcs": len(man["bgcs"]),
          "corrected_bgcs": len(man["bgcs"]), "assembly_tier": "GOOD"})
    da = man["source_scans"]["domain_architecture"]["per_bgc"]
    _json(pkg / "deep_data.json", {"bgc_profile": [{"sid": sid, "bgc_id": bid,
          **rec["domain_counts"], "total_domains": sum(rec["domain_counts"].values())}
          for bid, rec in da.items()], "active_sites": [], "class_pred": [], "domain_hits": []})
    _json(pkg / "gene_data.json", {"domain_arch": [{"sid": sid, "bgc_id": bid,
          "architecture": repr(rec)} for bid, rec in da.items()], "substrates": [], "ripp": [], "tfbs": {}})
    headers = ["BGC_ID", "Strain", "Contig", "Node_ID", "antiSMASH_Region", "Products", "KCB_top",
               "Boundary", "Length_kb", "KCB_score", "Arch", "TTA_tier", "CCTT_triggers", "Resistance_tier"]
    with (pkg / f"{sid}_2_inventory.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        for b in man["bgcs"]:
            writer.writerow({"BGC_ID": b["bgc_id"], "Strain": sid, "Contig": b["contig"],
                "Node_ID": b["node_id"], "antiSMASH_Region": b["antismash_region"],
                "Products": "; ".join(b["products"]), "Boundary": "Interior", "Length_kb": .1, "Arch": "A"})
    from ._gold_identity_fixture import bind_synthetic_gold_package
    bind_synthetic_gold_package(pkg, sid, manifest=man)
    _seal(pkg)
    return pkg


def _master(path, manifests):
    wb = Workbook(); wb.remove(wb.active)
    for name in ("A2_Strain_Registry", "A3_Run_Manifest", "B1_BGC_Master", "B2_Product_Class_Matrix"):
        ws = wb.create_sheet(name); ws.append(CANONICAL_V1_HEADERS[name])
    def row(sheet, values):
        wb[sheet].append([values.get(k) for k in CANONICAL_V1_HEADERS[sheet]])
    for man in manifests:
        sid = man["strain_id"]
        row("A2_Strain_Registry", {"strain": sid, "taxonomy": man["taxonomy"], "ecology_source": man["source"],
            "workflow_version": man["workflow_version"], "bgc_count": len(man["bgcs"])})
        row("A3_Run_Manifest", {"strain": sid, "version": man["workflow_version"]})
        for b in man["bgcs"]:
            row("B1_BGC_Master", {"strain": sid, "contig": b["node_id"], "region": b["antismash_region"],
                "BGC_ID": b["bgc_id"], "assembly_locator": assembly_locator(b),
                "products": "; ".join(b["products"]), "kcb_top": b["kcb_top"], "engine_version": "1.9.174"})
        raw = _b2_product_class_counts([c for b in man["bgcs"] for c in with_family_labels(b["products"])])
        for view in ("RAW_ANTISMASH", "GENE_BACKED"):
            counts = dict(raw)
            if view == "GENE_BACKED":
                counts["carotenoid"] = sum(b["_backed"] for b in man["bgcs"])
            row("B2_Product_Class_Matrix", {**counts, "strain": sid, "label_provenance": view})
    # Full-cohort aggregates must never supply public numerators or prose.
    ws = wb.create_sheet("Cross_Strain_Class_Prevalence")
    ws.append(["product_class", "n_strains", "band", "informative_for_comparison"])
    ws.append(["held_scientific_sentinel", 999999, "core", "yes"])
    wb.save(path); wb.close()


@pytest.fixture
def cohort(tmp_path):
    runs = tmp_path / "runs"
    alpha = _manifest("public_alpha", [_bgc(1, ["carotenoid", "phosphonate"], backed=True),
                                       _bgc(2, ["carotenoid"])])
    beta = _manifest("public_beta", [_bgc(1, ["ectoine"], backed=True)])
    held = _manifest("held_scientific_sentinel", [_bgc(1, ["phosphonate"])], release="PRIVATE")
    packages = {m["strain_id"]: _package(runs, m["strain_id"], m,
                 folder="display_alias" if m is alpha else None) for m in (alpha, beta, held)}
    # Exclusion must occur before any scientific payload validation/read.
    (packages[held["strain_id"]] / "deep_data.json").write_text("not JSON")
    profile = tmp_path / "privacy.json"
    _json(profile, {"schema_version": "sapote_privacy_profile_v1", "profile_id": "synthetic",
        "default_tier": "INTERNAL", "tiers": [{"tier_id": "INTERNAL", "public_export": False},
        {"tier_id": "OPEN", "public_export": True}], "strain_assignments": [
            {"strain_id": x, "tier_id": "OPEN"} for x in ("public_alpha", "public_beta")]})
    master = tmp_path / "master.xlsx"
    _master(master, [alpha, beta, held])
    return SimpleNamespace(runs=runs, master=master, profile=profile, packages=packages,
                           manifests=[alpha, beta, held], tmp=tmp_path)


def _load(c, view="RAW_ANTISMASH", **kwargs):
    return load_public_master(c.master, runs_dir=c.runs, privacy_profile=c.profile, class_view=view, **kwargs)


def _change_cell(c, sheet, header, value, row=2):
    wb = load_workbook(c.master)
    ws = wb[sheet]; col = [x.value for x in ws[1]].index(header) + 1
    ws.cell(row, col, value)
    wb.save(c.master); wb.close()


@pytest.mark.parametrize("view,carotenoid", [("RAW_ANTISMASH", 1), ("GENE_BACKED", 2)])
def test_independent_projection_both_views(cohort, view, carotenoid):
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in cohort.runs.rglob("*") if p.is_file()}
    scoped = _load(cohort, view)
    assert scoped["N"] == 2 and len(scoped["bgcs"]) == 3
    assert scoped["prevalence"]["carotenoid"]["n_strains"] == carotenoid
    assert scoped["prevalence"]["phosphonate"]["n_strains"] == 1
    text = render_public_synthesis(scoped, "fully_dark")
    assert "held_scientific_sentinel" not in text and "999999" not in text
    assert "phosphonate → public_alpha" in text and "display_alias" not in text
    assert all(hashlib.sha256(p.read_bytes()).hexdigest() == digest for p, digest in before.items())


@pytest.mark.parametrize("view", ["RAW_ANTISMASH", "GENE_BACKED"])
def test_direct_writer_is_independently_scoped(cohort, view):
    text = write_synthesis(cohort.master, public_only=True, runs_dir=cohort.runs,
                          privacy_profile=cohort.profile, strains=["public_alpha"], class_view=view)
    assert "**Cohort:** 1 strains · 2 BGCs" in text and "public_beta" not in text
    with pytest.raises(PublicCohortScopeError, match="PUBLIC_METADATA_OVERRIDE_UNBOUND"):
        write_synthesis(cohort.master, public_only=True, runs_dir=cohort.runs,
                        class_view=view, host_overrides={"public_alpha": "unbound"})


def test_held_master_changes_do_not_change_public_science(cohort):
    first = render_public_synthesis(_load(cohort), "fully_dark")
    wb = load_workbook(cohort.master)
    wb["Cross_Strain_Class_Prevalence"].append(["another_held_sentinel", 10**9, "core", "yes"])
    ws = wb["A3_Run_Manifest"]; ws.cell(4, 3, "Mamey v0.0.1")
    wb.save(cohort.master); wb.close()
    assert render_public_synthesis(_load(cohort), "fully_dark") == first


@pytest.mark.parametrize("sheet,header,value,code", [
    ("A2_Strain_Registry", "taxonomy", "conflicting", "REGISTRY_CONFLICT"),
    ("B1_BGC_Master", "products", "conflicting", "BGC_VALUE_CONFLICT"),
    ("B1_BGC_Master", "BGC_ID", "BGC999", "BGC_MEMBERSHIP_CONFLICT"),
    ("B1_BGC_Master", "region", "region999", "BGC_VALUE_CONFLICT"),
    ("B1_BGC_Master", "engine_version", "0.0.1", "ENGINE_PROVENANCE_CONFLICT"),
    ("B2_Product_Class_Matrix", "carotenoid", -1, "COUNT_INVALID"),
    ("B2_Product_Class_Matrix", "carotenoid", 999, "CLASS_VIEW_CONFLICT"),
    ("B2_Product_Class_Matrix", "carotenoid", "=2", "FORMULA_UNSUPPORTED"),
])
def test_admitted_conflicts_refuse_without_output(cohort, sheet, header, value, code):
    _change_cell(cohort, sheet, header, value)
    result = run_cohort_deliverable(cohort.runs, cohort.tmp / "out", master_path=cohort.master,
        public_only=True, privacy_profile=cohort.profile, class_view="RAW_ANTISMASH", with_figures=False)
    assert not result.ok and code in result.warnings[0]
    assert not result.out_dir.exists()
    assert "display_alias" not in result.warnings[0]


@pytest.mark.parametrize("sheet", ["A2_Strain_Registry", "A3_Run_Manifest", "B1_BGC_Master", "B2_Product_Class_Matrix"])
def test_duplicate_admitted_rows_never_overwrite(cohort, sheet):
    wb = load_workbook(cohort.master); ws = wb[sheet]
    ws.append([x.value for x in ws[2]]); wb.save(cohort.master); wb.close()
    with pytest.raises(PublicCohortScopeError, match="CARDINALITY|MEMBERSHIP"):
        _load(cohort)


@pytest.mark.parametrize("damage", ["missing", "status", "coverage", "domains", "pigment"])
def test_gene_backed_missing_evidence_is_not_zero_or_raw_fallback(cohort, damage):
    pkg = cohort.packages["public_alpha"]
    man = json.loads((pkg / "manifest.json").read_text())
    scans = man["source_scans"]
    if damage == "missing": man["source_scans"] = None
    if damage == "status": scans["domain_architecture"]["status"] = "UNAVAILABLE"
    if damage == "coverage": scans["domain_architecture"]["per_bgc"].pop("BGC002")
    if damage == "domains": scans["domain_architecture"]["per_bgc"]["BGC002"]["domains"] = None
    if damage == "pigment": scans["primary_metabolism"]["hits"].pop("pigment")
    _json(pkg / "manifest.json", man); _seal(pkg)
    with pytest.raises(PublicCohortScopeError, match="GENE_BACKED"):
        _load(cohort, "GENE_BACKED")


def test_complete_gene_negative_and_existing_selector_parity(cohort):
    scoped_views = {view: _load(cohort, view) for view in ("RAW_ANTISMASH", "GENE_BACKED")}
    classes = [c for c in CANONICAL_V1_HEADERS["B2_Product_Class_Matrix"]
               if c not in ("strain", "label_provenance", "counts_reliability")]
    for man in cohort.manifests[:2]:
        sid = man["strain_id"]
        saved = SimpleNamespace(bgcs=[SimpleNamespace(**b) for b in man["bgcs"]],
                                source_scans=SimpleNamespace(**man["source_scans"]),
                                context=SimpleNamespace(strain_id=sid))
        assert _gene_backed_carotenoid_bgcs(saved) == {b["bgc_id"] for b in man["bgcs"] if b["_backed"]}
        producer_rows = {r["label_provenance"]: r for r in _b2_product_class_rows(saved, "GOOD")}
        assert producer_rows["RAW_ANTISMASH"]["carotenoid"] == (2 if sid == "public_alpha" else 0)
        assert producer_rows["GENE_BACKED"]["carotenoid"] == 1
        for c in classes:
            if c != "carotenoid":
                assert producer_rows["GENE_BACKED"][c] == producer_rows["RAW_ANTISMASH"][c]
        for view, scoped in scoped_views.items():
            assert scoped["registry"][sid]["classes"] == {c for c in classes if producer_rows[view][c] > 0}
    # The existing family helper places the saved "carotenoid" token in OTHER;
    # the gene-backed builder changes only the carotenoid column, preserving OTHER.
    assert with_family_labels(["carotenoid"]) == ["carotenoid", "other"]
    raw_bgc = next(b for b in scoped_views["RAW_ANTISMASH"]["bgcs"]
                   if b["strain"] == "public_alpha" and b["bgc_id"] == "BGC002")
    gene_bgc = next(b for b in scoped_views["GENE_BACKED"]["bgcs"]
                    if b["strain"] == "public_alpha" and b["bgc_id"] == "BGC002")
    assert raw_bgc["classes"] == ["carotenoid", "other"]
    assert gene_bgc["classes"] == ["other"]
    # Existing selector can back a saved observed BGC without a raw carotenoid product.
    assert "carotenoid" in next(b for b in scoped_views["GENE_BACKED"]["bgcs"]
                               if b["strain"] == "public_beta")["classes"]


@pytest.mark.parametrize("view", ["RAW_ANTISMASH", "GENE_BACKED"])
def test_canonical_figure_loader_uses_same_membership_and_selected_classes(cohort, view):
    from mamey.cohort_figures import load, products_per_strain
    scoped = _load(cohort, view)
    rows = load(cohort.runs, public_only=True, privacy_profile=cohort.profile, canonical_scope=scoped)
    assert list(rows) == ["public_alpha", "public_beta"]
    counts = products_per_strain(rows, scoped["strains"])
    assert counts["carotenoid"]["public_alpha"] == (2 if view == "RAW_ANTISMASH" else 1)
    assert counts["carotenoid"].get("public_beta", 0) == (0 if view == "RAW_ANTISMASH" else 1)


@pytest.mark.parametrize("view", ["RAW_ANTISMASH", "GENE_BACKED"])
def test_real_combined_public_figures_and_no_figures(cohort, view):
    for with_figures in (False, True):
        out = cohort.tmp / f"{view}_{with_figures}"
        result = run_cohort_deliverable(cohort.runs, out, strains=["public_alpha"],
            master_path=cohort.master, public_only=True, privacy_profile=cohort.profile,
            class_view=view, with_figures=with_figures, figure_series="F")
        assert result.ok, result.warnings
        assert result.strains == ["public_alpha"]
        assert (result.figure_count > 0) == with_figures
        receipt = json.loads((out / "cohort_scope_receipt.json").read_text())
        assert receipt["class_view"] == view and receipt["strain_denominator"] == 1
        for path in out.rglob("*"):
            if path.is_file() and path.suffix not in (".png", ".pdf"):
                assert "held_scientific_sentinel" not in path.read_text()
                assert "display_alias" not in path.read_text()


def test_input_mutation_and_nonempty_output_refuse(cohort):
    scoped = _load(cohort)
    (cohort.packages["public_alpha"] / "gene_data.json").write_text("{}")
    with pytest.raises(PublicCohortScopeError, match="INPUT_CHANGED"):
        recheck_public_scope(scoped)
    out = cohort.tmp / "out"; out.mkdir(); sentinel = out / "existing.txt"; sentinel.write_text("retain")
    result = run_cohort_deliverable(cohort.runs, out, public_only=True)
    assert not result.ok and result.warnings == ["PUBLIC_COHORT_OUTPUT_NOT_EMPTY"]
    assert sentinel.read_text() == "retain"


def test_explicit_class_view_and_master_required(cohort):
    with pytest.raises(PublicCohortScopeError, match="CLASS_VIEW_REQUIRED"):
        load_public_master(cohort.master, runs_dir=cohort.runs, privacy_profile=cohort.profile)
    with pytest.raises(PublicCohortScopeError, match="MASTER_PATH_REQUIRED"):
        load_public_master(None, runs_dir=cohort.runs, class_view="RAW_ANTISMASH")
    with pytest.raises(PublicCohortScopeError, match="REQUESTED_IDENTITY_UNAVAILABLE"):
        _load(cohort, strains=["held_scientific_sentinel"])


def test_conflicting_bgc_release_and_duplicate_package_hold(cohort):
    import shutil
    pkg = cohort.packages["public_alpha"]
    man = json.loads((pkg / "manifest.json").read_text()); man["bgcs"][0]["release"] = "PRIVATE"
    _json(pkg / "manifest.json", man); _seal(pkg)
    with pytest.raises(PublicCohortScopeError, match="ROW_RELEASE_CONFLICT"):
        _load(cohort)
    man["bgcs"][0]["release"] = "PUBLIC"; _json(pkg / "manifest.json", man); _seal(pkg)
    shutil.copytree(pkg, cohort.runs / "another_alias" / "package")
    with pytest.raises(PublicCohortScopeError, match="PACKAGE_IDENTITY_DUPLICATE"):
        _load(cohort)


def test_missing_public_figure_member_rolls_back_every_output(cohort):
    pkg = cohort.packages["public_alpha"]
    (pkg / "deep_data.json").unlink(); _seal(pkg)
    out = cohort.tmp / "out"
    result = run_cohort_deliverable(cohort.runs, out, strains=["public_alpha"],
        master_path=cohort.master, public_only=True, privacy_profile=cohort.profile,
        class_view="GENE_BACKED", with_figures=True)
    assert not result.ok and result.warnings == ["PUBLIC_FIGURE_PAYLOAD_UNAVAILABLE"]
    assert not out.exists() and not list(cohort.tmp.glob(".sapote-public-cohort-*"))


def test_configured_private_literal_is_content_hold(cohort):
    policy = json.loads(cohort.profile.read_text())
    policy["public_export_policy"] = {"private_identifiers": [{"identifier_id": "held_text", "literal": "Referenceus"}]}
    _json(cohort.profile, policy)
    with pytest.raises(PublicCohortScopeError, match="PRIVATE_LITERAL_IN_CONTENT"):
        render_public_synthesis(_load(cohort), "fully_dark")


@pytest.mark.parametrize("view", ["RAW_ANTISMASH", "GENE_BACKED"])
def test_zero_bgc_member_counts_in_registry_denominator(cohort, view):
    alpha, _, held = cohort.manifests
    zero = _manifest("public_beta", [])
    pkg = cohort.packages["public_beta"]
    # Replace only this synthetic package's saved data; fixture mutation is deliberate and resealed.
    import shutil
    shutil.rmtree(pkg)
    cohort.packages["public_beta"] = _package(cohort.runs, "public_beta", zero)
    _master(cohort.master, [alpha, zero, held])
    scoped = _load(cohort, view)
    assert scoped["N"] == 2 and len(scoped["bgcs"]) == 2
    assert scoped["prevalence"]["phosphonate"]["n_strains"] == 1
    with pytest.raises(PublicCohortScopeError, match="NO_BGCS"):
        _load(cohort, view, strains=["public_beta"])


@pytest.mark.parametrize("view", ["RAW_ANTISMASH", "GENE_BACKED"])
def test_cohort_cli_wires_class_view_profile_and_subset(cohort, view, capsys):
    from mamey.cli import build_parser
    out = cohort.tmp / "cli_output"
    args = build_parser().parse_args(["cohort", "--runs-dir", str(cohort.runs),
        "--master-path", str(cohort.master), "--out", str(out), "--public-only",
        "--privacy-profile", str(cohort.profile), "--class-view", view,
        "--strains", "public_alpha", "--no-figures"])
    assert args.func(args) == 0
    assert json.loads((out / "cohort_scope_receipt.json").read_text())["strains"] == ["public_alpha"]
    assert "held_scientific_sentinel" not in capsys.readouterr().out
