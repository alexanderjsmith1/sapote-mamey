from pathlib import Path
from types import SimpleNamespace
import json
from zipfile import ZipFile
import pytest
from openpyxl import load_workbook
from mamey import antismash_evidence as ae
from mamey.antismash_tables import build_structured_tables
from mamey.workbook import write_per_strain_workbook
from mamey.package_inspector import explain_command
from mamey.compile_report import build_report
from mamey.rrefinder_status import derive_rrefinder_state, read_package_state, state_from_structured
from tests.test_reporting_evidence_v2 import _FakeRun, _record, _rre_hit


def archive(tmp_path, records=None, members=None):
    path = tmp_path / "source.zip"
    with ZipFile(path, "w") as zf:
        if members is None:
            members = {"source.json": json.dumps({"version": "8.0.4", "schema": 4, "records": records})}
        for name, body in members.items():
            zf.writestr(name, body)
    return path


def hit_record(record_id="NODE_1"):
    return _record(record_id, {"gene": [_rre_hit("[120:150](+)", "PqqD_RRE", "RREFam006.1")]})


@pytest.mark.parametrize("case,expected", [
    ("not_requested", "NOT_REQUESTED"),
    ("missing_json", "SOURCE_UNAVAILABLE"),
    ("missing_module", "SOURCE_UNAVAILABLE"),
    ("invalid_json", "PARSE_FAILED"),
    ("invalid_hits", "PARSE_FAILED"),
    ("empty_source", "SOURCE_OBSERVED_EMPTY"),
    ("source_hits", "SOURCE_OBSERVED_HITS"),
    ("partial_file", "SOURCE_PARTIAL"),
    ("missing_module_partial", "SOURCE_PARTIAL"),
    ("identity_unmapped", "SOURCE_OBSERVED_HITS"),
    ("legacy_empty", "UNKNOWN_COMPLETENESS"),
    ("request_conflict", "SOURCE_CONFLICT"),
])
def test_source_to_report_states(tmp_path, capsys, case, expected):
    records = [hit_record()]
    members = None
    if case == "missing_json": members = {"readme.txt": "No source JSON supplied"}
    if case == "missing_module": records = [{"id": "NODE_1", "modules": {}}]
    if case == "invalid_json": members = {"source.json": "{"}
    if case == "invalid_hits": records = [_record("NODE_1", {"gene": [None]})]
    if case == "empty_source": records = [_record("NODE_1", {})]
    if case == "partial_file":
        members = {"source.json": json.dumps({"records": records}), "bad.json": "{"}
    if case == "missing_module_partial": records += [{"id": "NODE_2", "modules": {}}]
    if case == "identity_unmapped": records = [hit_record("ANOTHER_CONTIG")]
    path = archive(tmp_path, records, members)
    # Main walker is deliberately OFF; include_structured isolates the actual record path.
    # TIGRFAM-only requests must never make RREfinder look requested.
    evidence = ae.parse_antismash_evidence(path, json_mode="off", want_tigrfam=True,
                                           include_structured=case != "not_requested")
    if case == "legacy_empty": evidence = {"rrefinder_hits": []}
    if case == "request_conflict": evidence["record_extras_requested"] = False
    run = _FakeRun(path, {})
    structured = build_structured_tables(path, evidence, run.bgcs)
    run.source_scans.antismash_structured = structured
    assert structured["rrefinder_status"] == expected
    state = structured["rrefinder_evidence_state"]
    assert state["upstream_search"] == "UNKNOWN"
    assert state["biological_absence"] == "NOT_ESTABLISHED"
    if case == "identity_unmapped":
        assert len(structured["rrefinder"]) == 1
        assert structured["rrefinder"][0]["bgc_id"] == "UNMAPPED"
    pkg = tmp_path / "package"
    pkg.mkdir()
    (pkg / "manifest.json").write_text(json.dumps({"strain_id": "AS-TEST", "mode": "gold",
                                                    "release_policy": "PRIVATE"}))
    (pkg / "AS-TEST_3_antismash_structured.json").write_text(json.dumps(structured))
    # Exercise actual engine workbook and reader implementations, not hand-built renderings.
    output = pkg / "AS-TEST.xlsx"
    write_per_strain_workbook(run, output)
    wb = load_workbook(output, read_only=True, data_only=True)
    status_rows = list(wb["antiSMASH_Evidence_Status"].values)
    assert status_rows[1][1] == expected
    assert wb["antiSMASH_RREfinder"].max_row == len(structured["rrefinder"]) + 1
    wb.close()
    assert explain_command(SimpleNamespace(package_dir=str(pkg))) == 0
    explanation = capsys.readouterr().out
    report = build_report(pkg, generate_figures=False)
    for text in (explanation, report):
        assert "RREfinder: " + expected in text
        assert "Biological absence is not established" in text
        assert "NULL_NO_RREFINDER_HITS" not in text
    (tmp_path / "case_summary.json").write_text(json.dumps({"case": case, "state": state}, indent=2))


def test_size_hold_is_not_empty(tmp_path, monkeypatch):
    path = archive(tmp_path, [_record("NODE_1", {})])
    monkeypatch.setattr(ae, "_should_stream", lambda *a: False)
    monkeypatch.setattr(ae, "FULL_MAX_JSON_BYTES", 1)
    ev = ae.parse_antismash_evidence(path, json_mode="off", include_structured=True)
    assert derive_rrefinder_state(ev, 0)["state"] == "SOURCE_HELD"


def test_lazy_stream_failure_retains_rows(tmp_path, monkeypatch):
    path = archive(tmp_path, [hit_record()])
    def broken(*args):
        yield hit_record()
        raise ValueError("truncated stream")
    monkeypatch.setattr(ae, "_iter_records", broken)
    monkeypatch.setattr(ae, "_should_stream", lambda *a: True)
    ev = ae.parse_antismash_evidence(path, json_mode="off", include_structured=True)
    assert len(ev["rrefinder_hits"]) == 1
    assert derive_rrefinder_state(ev, 1)["state"] == "SOURCE_PARTIAL"
    assert ev["record_extraction"]["files"][0]["status"] == "PARSE_FAILED"


def test_other_handler_failure_does_not_invalidate_rrefinder(tmp_path, monkeypatch):
    path = archive(tmp_path, [hit_record()])
    def broken(*args): raise ValueError("different evidence channel")
    broken.__name__ = "_nrps_pks_from_rec"
    monkeypatch.setattr(ae, "_nrps_pks_from_rec", broken)
    ev = ae.parse_antismash_evidence(path, json_mode="off", include_structured=True)
    assert ev["record_extraction"]["files"][0]["errors"]
    assert derive_rrefinder_state(ev, 1)["state"] == "SOURCE_OBSERVED_HITS"


def test_not_requested_performs_no_rre_extraction(tmp_path, monkeypatch):
    path = archive(tmp_path, [hit_record()])
    def forbidden(*args): raise AssertionError("record pass must not run")
    monkeypatch.setattr(ae, "_run_record_extractors", forbidden)
    ev = ae.parse_antismash_evidence(path, json_mode="off")
    assert derive_rrefinder_state(ev, 0)["state"] == "NOT_REQUESTED"


def test_legacy_label_does_not_become_negative(tmp_path):
    structured = {"rrefinder": [], "rrefinder_status": "NULL_NO_RREFINDER_HITS"}
    assert state_from_structured(structured)["state"] == "UNKNOWN_COMPLETENESS"
    (tmp_path / "OTHER_3_antismash_structured.json").write_text(json.dumps(structured))
    assert read_package_state(tmp_path, "AS-TEST")["state"] == "SOURCE_UNAVAILABLE"


def test_reader_rejects_corrupt_and_conflicting_sidecar(tmp_path):
    path = tmp_path / "AS-TEST_3_antismash_structured.json"
    path.write_text("{")
    assert read_package_state(tmp_path, "AS-TEST")["state"] == "PARSE_FAILED"
    state = derive_rrefinder_state({"record_extras_requested": False}, 0)
    path.write_text(json.dumps({"rrefinder": [{}], "rrefinder_evidence_state": state}))
    assert read_package_state(tmp_path, "AS-TEST")["state"] == "SOURCE_CONFLICT"


def test_capped_error_detail_preserves_channel_failure(tmp_path, monkeypatch):
    # Other-channel errors fill the detail cap before the RREfinder failure.
    records = [_record("NODE_1", {}) for _ in range(105)]
    records.append(_record("NODE_1", {"gene": [None]}))
    path = archive(tmp_path, records)
    def broken(*args): raise ValueError("other channel")
    broken.__name__ = "_nrps_pks_from_rec"
    monkeypatch.setattr(ae, "_nrps_pks_from_rec", broken)
    ev = ae.parse_antismash_evidence(path, json_mode="off", include_structured=True)
    receipt = ev["record_extraction"]["files"][0]
    assert len(receipt["errors"]) == 100
    assert receipt["errors_omitted"] == 7
    assert receipt["error_counts"]["_rrefinder_from_rec"] == 1
    assert derive_rrefinder_state(ev, 0)["state"] == "SOURCE_PARTIAL"
