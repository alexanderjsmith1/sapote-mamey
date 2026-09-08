import csv
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def _write_json(path: Path, obj):
    path.write_text(json.dumps(obj, indent=2), encoding="utf-8")


def _minimal_banked(tmp_path: Path) -> Path:
    bank = tmp_path / "cohort"
    bank.mkdir()
    _write_json(bank / "bgc_data.json", {
        "strains": {"PUBLIC-FIXTURE-001": {"n50": 200000, "corrected_bgcs": 1.0, "raw_bgcs": 1}},
        "bgcs": [{
            "sid": "PUBLIC-FIXTURE-001",
            "bgc_id": "BGC001",
            "products": "NRPS",
            "kcb_top": "",
            "length_kb": 12.0,
            "contig": "ctg1",
            "region": "region001",
            "kcb_cumulative": 0,
            "edge_status": "Interior",
        }],
    })
    _write_json(bank / "deep_data.json", {"bgc_profile": [], "active_sites": [], "class_pred": []})
    _write_json(bank / "gene_data.json", {"tfbs": {}, "substrates": []})
    _write_json(bank / "rggmci_full.json", [])
    _write_json(bank / "tigrfam.json", {"PUBLIC-FIXTURE-001": {"present": {}}})
    _write_json(bank / "tfbs_coupling.json", {})
    return bank


def test_build_size_profile_help_imports_from_outside_repo(tmp_path):
    script = ROOT / "tools" / "build_size_profile.py"
    res = subprocess.run(
        [sys.executable, str(script), "--help"],
        cwd=tmp_path,
        text=True,
        capture_output=True,
        timeout=30,
    )
    assert res.returncode == 0, res.stderr
    assert "build_size_profile.py" in res.stdout or "usage:" in res.stdout.lower()


def test_build_thesis_vignettes_refuses_default_sid_placeholders(tmp_path):
    bank = _minimal_banked(tmp_path)
    out = tmp_path / "thesis_vignettes.md"
    res = subprocess.run(
        [sys.executable, str(ROOT / "tools" / "build_thesis_vignettes.py"), "--banked-dir", str(bank), "--out", str(out)],
        cwd=ROOT,
        text=True,
        capture_output=True,
        timeout=30,
    )
    assert res.returncode != 0
    combined = res.stdout + res.stderr
    assert "default SID placeholder targets are absent" in combined
    assert not out.exists() or "SID-XXX" not in out.read_text(encoding="utf-8")


def test_build_thesis_vignettes_fails_loudly_when_modeb_targets_do_not_match(tmp_path):
    bank = _minimal_banked(tmp_path)
    with (bank / "modeb_verdicts.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["strain", "bgc", "status", "modeb_class", "note"])
        w.writeheader()
        w.writerow({"strain": "PUBLIC-FIXTURE-001", "bgc": "BGC999", "status": "CONFIRM", "modeb_class": "NRPS", "note": "synthetic mismatch"})
    out = tmp_path / "thesis_vignettes.md"
    res = subprocess.run(
        [sys.executable, str(ROOT / "tools" / "build_thesis_vignettes.py"), "--banked-dir", str(bank), "--out", str(out)],
        cwd=ROOT,
        text=True,
        capture_output=True,
        timeout=30,
    )
    assert res.returncode != 0
    combined = res.stdout + res.stderr
    assert "no vignette target matches banked bgc_data.json" in combined
    assert not out.exists() or "SID-XXX" not in out.read_text(encoding="utf-8")


def test_add_xstrain_sheets_single_strain_workbook_prints_dapr_skip(tmp_path):
    bank = _minimal_banked(tmp_path)
    wb_path = tmp_path / "public_fixture.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["strain", "note"])
    ws.append(["PUBLIC-FIXTURE-001", "single-strain public fixture"])
    wb.save(wb_path)

    res = subprocess.run(
        [sys.executable, str(ROOT / "tools" / "add_xstrain_sheets.py"), "--banked-dir", str(bank), "--out", str(wb_path)],
        cwd=ROOT,
        text=True,
        capture_output=True,
        timeout=30,
    )
    assert res.returncode == 0, res.stderr
    assert "DAPR rescue refresh SKIPPED" in res.stdout
    wb2 = openpyxl.load_workbook(wb_path, read_only=True)
    try:
        assert "Cross_Strain_Findings" in wb2.sheetnames
        assert "Strain_Cohort_Context" in wb2.sheetnames
    finally:
        wb2.close()


def test_causemap_overrides_have_distinct_real_keys():
    """Guards the FIXED v9.7.308 state of CAUSEMAP (sibling of the master_figure_atlas AS-XXX fix).

    History: CAUSEMAP was a 3-entry dict literal whose keys were all the placeholder 'SID-XXX'.
    Python keeps only the last duplicate key, so 2 of the 3 authored routes were dead and none
    ever fired. It is now an empty, real-strain-ID-keyed map (behavior-preserving: every real
    strain still resolves to DEFAULT_MAP via .get) with the three authored routes preserved in
    _AUTHORED_CAUSEMAP_ROUTES_PENDING_MAPPING.

    This test guards the INVARIANT: no duplicate or placeholder keys may reappear at the source
    level, values must be valid (image, description) pairs, real strains still fall through to
    DEFAULT_MAP, and the authored-route reference survives. Empty (current) or real-ID-filled
    (a future real fix) both pass; only a regression toward the old collision fails.
    """
    import ast
    import importlib
    import sys as _sys
    sys_path_added = str(ROOT / "tools") not in _sys.path
    if sys_path_added:
        _sys.path.insert(0, str(ROOT / "tools"))
    try:
        bv = importlib.import_module("build_thesis_vignettes")
        importlib.reload(bv)
    finally:
        if sys_path_added:
            _sys.path.remove(str(ROOT / "tools"))

    assert isinstance(bv.CAUSEMAP, dict)
    for sid, val in bv.CAUSEMAP.items():
        assert isinstance(sid, str) and "XXX" not in sid, f"placeholder key in CAUSEMAP: {sid!r}"
        assert isinstance(val, tuple) and len(val) == 2, f"bad CAUSEMAP route for {sid!r}: {val!r}"

    # Source-level duplicate/placeholder key check (dict runtime would hide duplicate literals).
    src = ast.parse(open(bv.__file__).read())
    dict_node = None
    for node in ast.walk(src):
        if isinstance(node, ast.Assign) and any(
            isinstance(t, ast.Name) and t.id == "CAUSEMAP" for t in node.targets
        ):
            dict_node = node.value
            break
    assert isinstance(dict_node, ast.Dict), "CAUSEMAP literal not found -- source moved?"
    literal_keys = [k.value for k in dict_node.keys if isinstance(k, ast.Constant)]
    assert len(literal_keys) == len(set(literal_keys)), (
        f"duplicate key literals in CAUSEMAP: {literal_keys} -- the old collapse bug is back."
    )
    assert not any("XXX" in str(k) for k in literal_keys), f"placeholder key literal: {literal_keys}"

    # Behavior preserved: real strains fall through to DEFAULT_MAP.
    for real_sid in ("AS-696", "AS-705", "SID-042"):
        assert bv.CAUSEMAP.get(real_sid, bv.DEFAULT_MAP) == bv.DEFAULT_MAP

    # Authored routes preserved as reference for the pending editorial mapping.
    assert len(bv._AUTHORED_CAUSEMAP_ROUTES_PENDING_MAPPING) == 3
