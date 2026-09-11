"""test_seal_integrity_f04_f05.py — v9.7.354 audit findings F04 + F05 (reworked).

F04: the final validation receipt (gate_validation.json) is finalized BEFORE the ZIP is sealed,
     so the receipt sealed inside the ZIP is byte-identical to the external on-disk one.
     _verify_sealed_receipt() proves that equality (and detects the F04 ordering regression).
     NOTE (v9.7.354 rework): the receipt is written AFTER write_manifest (so the package is
     complete when the enrichment-aware validation runs) and BEFORE zip_package (so the sealed
     copy equals the external). The reverted .353 attempt validated before write_manifest, which
     blocked the gold-path seal for smoke/minimal packages; that regression is not reintroduced.
F05: a blocking-gate FAIL exits NON-ZERO by default; advisory (report-only) mode is an explicit
     opt-in that exits 0 but always surfaces a loud banner.
"""
import json
from pathlib import Path

import openpyxl
import pytest

from mamey.packaging import zip_package
from mamey.cli import _verify_sealed_receipt
from mamey.seal_package import seal_package


# ── F04: sealed-internal receipt == external receipt ─────────────────────────────────
def test_verify_sealed_receipt_matches_when_finalized_before_seal(tmp_path):
    pkg = tmp_path / "AS-TEST_package"
    pkg.mkdir()
    receipt = {"status": "MAMEY_COMPLETE", "file_presence": "PASS"}
    (pkg / "gate_validation.json").write_text(json.dumps(receipt, indent=2), encoding="utf-8")
    (pkg / "manifest.json").write_text(json.dumps({"strain_id": "AS-TEST"}), encoding="utf-8")

    zip_path = tmp_path / "AS-TEST_package.zip"
    zip_package(pkg, zip_path)  # seals the receipt exactly as it stands on disk

    assert _verify_sealed_receipt(zip_path, pkg / "gate_validation.json") is True


def test_verify_sealed_receipt_detects_post_seal_divergence(tmp_path):
    """The F04 defect class: if the external receipt is rewritten AFTER sealing, internal!=external."""
    pkg = tmp_path / "AS-TEST_package"
    pkg.mkdir()
    (pkg / "gate_validation.json").write_text(json.dumps({"status": "MAMEY_COMPLETE"}), encoding="utf-8")
    zip_path = tmp_path / "AS-TEST_package.zip"
    zip_package(pkg, zip_path)
    # simulate the OLD ordering: rewrite the external receipt after the seal
    (pkg / "gate_validation.json").write_text(json.dumps({"status": "FAIL"}), encoding="utf-8")
    assert _verify_sealed_receipt(zip_path, pkg / "gate_validation.json") is False


def test_verify_sealed_receipt_missing_entry_is_false(tmp_path):
    pkg = tmp_path / "AS-TEST_package"
    pkg.mkdir()
    (pkg / "manifest.json").write_text("{}", encoding="utf-8")  # no gate_validation.json
    zip_path = tmp_path / "AS-TEST_package.zip"
    zip_package(pkg, zip_path)
    (pkg / "gate_validation.json").write_text("{}", encoding="utf-8")  # external appears later
    assert _verify_sealed_receipt(zip_path, pkg / "gate_validation.json") is False


# ── F05: blocking FAIL exits non-zero by default; advisory is an explicit opt-in ──────
@pytest.fixture
def blocking_pkg(tmp_path):
    """A package whose workbook has a header-only sheet -> blocking workbook_content FAIL."""
    pkg = tmp_path
    (pkg / "manifest.json").write_text(json.dumps({"strain_id": "AS-TEST", "mode": "gold"}))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("BGC_Inventory")
    ws.append(["BGC_ID", "Col"])  # header only -> FAIL_empty (blocking)
    for name in ("Triage_Board", "WetLab_Decision_Matrix", "Mode_B_Summary"):
        s = wb.create_sheet(name)
        s.append(["BGC_ID", "Col"]); s.append(["BGC001", "v"])
    wb.save(str(pkg / "AS-TEST_5_workbook.xlsx"))
    return pkg


def test_blocking_fail_exits_nonzero_by_default(blocking_pkg):
    r = seal_package(blocking_pkg)  # no flags -> enforce
    assert r["overall"] == "FAIL"
    assert r["exit_code"] == 1
    assert r["enforced"] is True
    assert r["non_strict_warning"] is None  # enforced -> not demoted, no banner


def test_advisory_opt_in_exits_zero_with_banner(blocking_pkg):
    r = seal_package(blocking_pkg, advisory=True)
    assert r["overall"] == "FAIL"
    assert r["exit_code"] == 0
    assert r["enforced"] is False
    assert r["non_strict_warning"] and "exit forced to 0" in r["non_strict_warning"]


def test_strict_overrides_advisory(blocking_pkg):
    r = seal_package(blocking_pkg, strict=True, advisory=True)
    assert r["exit_code"] == 1  # strict wins
    assert r["enforced"] is True
