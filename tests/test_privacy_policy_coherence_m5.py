"""Integrated release authority and read-only public figure admission.

Only synthetic software fixtures are used. Tests are scheduled by the campaign
coordinator after all source owners have stopped editing.
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import shlex
import re
import subprocess
import sys

import pytest

from mamey.release_policy import (
    load_bound_profile, package_release_policy, resolve_run_release,
)

ROOT = Path(__file__).resolve().parents[1]
SYNTHETIC_RESEARCH = "AS-" + "987654"


def _profile(tmp_path, public_id="public_reference"):
    path = tmp_path / "privacy.json"
    path.write_text(json.dumps({
        "schema_version": "sapote_privacy_profile_v1", "profile_id": "synthetic_policy",
        "default_tier": "INTERNAL",
        "tiers": [{"tier_id": "INTERNAL", "public_export": False},
                  {"tier_id": "OPEN", "public_export": True}],
        "strain_assignments": [{"strain_id": public_id, "tier_id": "OPEN"}],
    }))
    return path


def _tree(path):
    return {str(p.relative_to(path)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in path.rglob("*") if p.is_file()}


def _package(runs, folder, *, strain=None, release=None, decision=None, broken_payload=False):
    package = runs / folder / "package"
    package.mkdir(parents=True)
    metadata = {"strain_id": strain or folder, "taxonomy": "Streptomyces sp.",
                "raw_bgcs": 1, "corrected_bgcs": 1, "assembly_tier": "GOOD"}
    if release is not None:
        metadata["release"] = release
    if decision is not None:
        metadata["release_decision"] = decision.to_dict()
    (package / "manifest_short.json").write_text(json.dumps(metadata))
    (package / "deep_data.json").write_text("not JSON" if broken_payload else json.dumps({"bgc_profile": []}))
    (package / "gene_data.json").write_text("not JSON" if broken_payload else "{}")
    return package


@pytest.mark.parametrize("strain", [SYNTHETIC_RESEARCH, "SID987654", "SCLAV", "public_reference"])
def test_every_omitted_unbound_identifier_defaults_private(strain):
    decision = resolve_run_release(strain)
    assert decision.release == "PRIVATE" and decision.reason == "no_public_authority"


def test_exact_public_profile_preserves_omission_and_private_override(tmp_path):
    path = _profile(tmp_path)
    public = resolve_run_release("public_reference", privacy_profile=path)
    assert public.release == "PUBLIC" and public.assignment_state == "EXACT"
    assert public.authority_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert resolve_run_release("public_reference", "PRIVATE", path).release == "PRIVATE"
    refused = resolve_run_release("unassigned_synthetic", "PUBLIC", path)
    assert refused.release == "PRIVATE" and refused.refused


def test_malformed_profile_never_falls_back_to_explicit_public(tmp_path):
    path = _profile(tmp_path)
    raw = json.loads(path.read_text())
    raw["strain_assignments"].append(raw["strain_assignments"][0])
    path.write_text(json.dumps(raw))
    with pytest.raises(ValueError, match="duplicate exact"):
        resolve_run_release("public_reference", "PUBLIC", path)


@pytest.mark.parametrize("override,profile_public,expected", [
    (None, False, "PRIVATE"), ("PUBLIC", False, "PUBLIC"),
    (None, True, "PUBLIC"), ("PRIVATE", True, "PRIVATE"),
])
def test_actual_engine_all_release_surfaces_agree(tmp_path, synthetic_single_contig_full_locus_zip,
                                                override, profile_public, expected):
    # Real runner and package/workbook writers; no mocked privacy or execution.
    strain = SYNTHETIC_RESEARCH
    command = [sys.executable, str(ROOT / "mamey_run.py"), "run", "--strain", strain,
               "--input-zip", str(synthetic_single_contig_full_locus_zip),
               "--taxonomy", "Streptomyces sp.", "--source", "synthetic software fixture",
               "--outdir", str(tmp_path / "runs"), "--master", str(tmp_path / "master.xlsx"),
               "--mode", "gold", "--capped-session", "--locus-maps", "off",
               "--reference-completion", "off"]
    if override:
        command.extend(["--release", override])
    if profile_public:
        command.extend(["--privacy-profile", str(_profile(tmp_path, strain))])
    completed = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, timeout=180)
    assert completed.returncode == 0, completed.stdout[-4000:] + completed.stderr[-2000:]
    terminal = re.findall(r"MAMEY_STATUS: (\S+)", completed.stdout)
    assert terminal and terminal[-1] in {"MAMEY_COMPLETE", "MAMEY_COMPLETE_WITH_ISSUES"}, (
        completed.stdout[-4000:] + completed.stderr[-2000:])
    package = tmp_path / "runs" / strain / "package"
    documents = [json.loads((package / name).read_text()) for name in
                 ("manifest.json", "manifest_short.json", f"{strain}_1_intake.json", "commit_receipt.json")]
    assert all(d["release"] == expected for d in documents)
    assert all(d["release_decision"] == documents[0]["release_decision"] for d in documents)
    with (package / f"{strain}_2_inventory.csv").open(newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert rows and all(row["Release"] == expected for row in rows)
    import openpyxl
    workbook_paths = list(package.rglob("*.xlsx"))
    assert workbook_paths
    covers = []
    for path in workbook_paths:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "README" in book.sheetnames:
            covers.extend((row[0], row[1]) for row in book["README"].iter_rows(values_only=True)
                          if len(row) >= 2 and row[0] == "Release")
        book.close()
    assert ("Release", expected) in covers


@pytest.mark.parametrize("profile_public", [False, True])
def test_empty_inventory_writers_keep_the_original_public_decision(tmp_path, synthetic_single_contig_full_locus_zip,
                                                                  profile_public):
    from mamey import __version__, cli
    from mamey.models import AssemblyMetrics, MameyRun, RunContext, SourceScanBundle
    from mamey.workbook import write_per_strain_workbook
    import openpyxl
    # Serialization/writer coverage, not a successful zero-region analysis. The
    # extraction entrypoint deliberately refuses raw==0 before package writing.
    strain = "empty_reference"
    decision = resolve_run_release(
        strain, None if profile_public else "PUBLIC",
        _profile(tmp_path, strain) if profile_public else None)
    package = tmp_path / "writer_fixture" / "package"
    package.mkdir(parents=True)
    context = RunContext(strain, "Synthetic empty writer fixture", __version__, "gold",
                         str(synthetic_single_contig_full_locus_zip), str(package.parent),
                         taxonomy="Streptomyces sp.", source="synthetic software fixture",
                         release_decision=decision)
    # Supply all required scan channels explicitly. Empty mappings mean no scan
    # evidence was supplied to this writer fixture, not negative scan results.
    scans = SourceScanBundle(
        chitinase={}, tfbs={}, blda_tta={}, regulators={}, transporters={},
        resistance={}, cctt={}, flbr={}, cassettes={}, umed={}, efls={},
        domain_architecture={}, resistance_tiers={}, wetlab_rows={},
        qs_signals={}, glycosylation_arms={}, per_bgc_dss={}, rggmci={},
    )
    run = MameyRun(context, AssemblyMetrics(1000, 1, 1000, 50.0, 1000), [], {},
                   source_scans=scans)
    assert cli._write_package(run, package, {}, "unknown") is None
    # The actual writer's completion receipt must succeed before the remaining
    # handoff fields are inspected. This is not a sealed-analysis admission.
    commit = json.loads((package / "commit_receipt.json").read_text())
    assert commit["status"] in {"PASS", "PASS_WITH_ISSUES"}
    assert commit["raw_bgcs"] == 0
    documents = [commit, json.loads((package / "manifest.json").read_text()),
                 json.loads((package / f"{strain}_1_intake.json").read_text())]
    assert documents[1]["bgcs"] == []
    assert all(d["release"] == "PUBLIC" for d in documents)
    assert all(d["release_decision"] == decision.to_dict() for d in documents)
    workbook = package / "empty_workbook.xlsx"
    assert write_per_strain_workbook(run, workbook) is None
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        cover = {row[0]: row[1] for row in book["README"].iter_rows(values_only=True)
                 if len(row) >= 2 and isinstance(row[0], str)}
        assert cover["Release"] == "PUBLIC"
        assert cover["Release reason"] == decision.reason
    finally:
        book.close()


def test_raw_zero_pipeline_refuses_before_package_serialization(tmp_path, synthetic_single_contig_full_locus_zip,
                                                               monkeypatch):
    from mamey import cli
    monkeypatch.setattr(cli, "parse_bgcs_from_zip", lambda *_a, **_k: [])
    def must_not_run(*_args, **_kwargs):
        raise AssertionError("empty parse reached downstream analysis or serialization")
    for name in ("extract_gbk_pfam_hits", "extract_cds_features", "extract_contig_sequences",
                 "extract_domain_features", "run_source_scans", "_write_package"):
        monkeypatch.setattr(cli, name, must_not_run)
    from mamey import rggmci
    monkeypatch.setattr(rggmci, "run_archive", must_not_run)
    strain = "empty_reference"
    result = cli.run_one_strain(
        strain, "Synthetic zero-parse refusal", str(synthetic_single_contig_full_locus_zip),
        str(tmp_path / "runs"), "gold", "Streptomyces sp.", "synthetic software fixture",
        master_path=str(tmp_path / "master.xlsx"), json_mode="off", brief="none",
        release="PUBLIC", locus_maps="off", reference_completion="off")
    assert result["status"] == "MAMEY_FAILED" and result["raw_bgcs"] == 0
    assert result["package_zip"] is None and not result["master_updated"]
    package = tmp_path / "runs" / strain / "package"
    assert not (package / "manifest.json").exists()
    assert not (package / "commit_receipt.json").exists()
    assert not (tmp_path / "master.xlsx").exists()
    assert result["issues"][0].startswith("NO_REGIONS_PARSED: 3 antiSMASH region GBK(s)")
    phases = [json.loads(line) for line in (package / "run_phase_receipts.jsonl").read_text().splitlines()]
    assert any(row["phase"] == "inventory" and row["status"] == "ERROR" for row in phases)
    assert not any(row["phase"] in {"source_scans", "gene_context", "rggmci"} for row in phases)


def test_run_decision_rejects_conflicting_bgc_row_before_manifest(tmp_path):
    from mamey.models import AssemblyMetrics, BGCRecord, MameyRun, RunContext
    context = RunContext("public_reference", "Synthetic", "test", "gold", "synthetic.zip", str(tmp_path),
                         release_decision=resolve_run_release("public_reference", "PUBLIC"))
    row = BGCRecord("BGC001", "NODE_1", 1, 0, 100, 100, ["RiPP"], release="PRIVATE")
    run = MameyRun(context, AssemblyMetrics(100, 1, 100, 50, 100), [row], {})
    with pytest.raises(ValueError, match="RELEASE_DECISION_ROW_CONFLICT"):
        run.to_dict()


def test_generated_batch_command_is_private_and_preserves_arbitrary_source(tmp_path):
    from mamey.cli import _write_batch_plan
    source = 'synthetic source "quotes" $(literal) with spaces'
    assert _write_batch_plan(tmp_path, "synthetic_reference", "gold", source)
    text = (tmp_path / "synthetic_reference_BATCH_PLAN.md").read_text()
    command = text.split("```", 2)[1].replace("\\\n", "")
    argv = shlex.split(command)
    assert argv[argv.index("--release") + 1] == "PRIVATE"
    assert argv[argv.index("--source") + 1] == source
    assert "match their receipts" in text


def test_inspect_generated_command_and_profile_advice(tmp_path, capsys, synthetic_single_contig_full_locus_zip):
    from mamey.package_inspector import inspect_command
    assert inspect_command(argparse.Namespace(zip=str(synthetic_single_contig_full_locus_zip))) == 0
    text = capsys.readouterr().out
    block = text.split("Capped-session run command", 1)[1].split("Then validate:", 1)[0]
    argv = shlex.split(block[block.index("python "):].replace("\\\n", ""))
    assert argv[argv.index("--release") + 1] == "PRIVATE"
    assert "omit --release and provide --privacy-profile" in text


def test_public_filter_never_loads_excluded_payloads_or_relabels_old_packages(tmp_path):
    from mamey.cohort_figures import is_private, load
    runs = tmp_path / "runs"
    public = _package(runs, "public_reference", release="PUBLIC",
                      decision=resolve_run_release("public_reference", "PUBLIC"))
    for sid, release in (("private_synthetic_secret", "PRIVATE"),
                         ("legacy_public", "PUBLIC"), ("missing_release", None)):
        _package(runs, sid, release=release, broken_payload=True)
    conflict = _package(runs, "conflicting_release", release="PUBLIC", broken_payload=True)
    (conflict / "conflicting_release_1_intake.json").write_text(json.dumps({
        "strain_id": "conflicting_release", "release": "PRIVATE"}))
    before = _tree(runs)
    audit = {}
    rows = load(runs, public_only=True, privacy_audit=audit)
    assert list(rows) == ["public_reference"] and not is_private("public_reference", rows)
    assert audit["admitted"] == 1 and audit["excluded"] == 4
    assert "private_synthetic_secret" not in json.dumps(audit)
    assert _tree(runs) == before
    assert package_release_policy(public, "public_reference")["recorded_release"] == "PUBLIC"


def test_legacy_public_requires_exact_package_id_not_folder_alias(tmp_path):
    from mamey.cohort_figures import load
    runs = tmp_path / "runs"
    _package(runs, "display_alias", strain="public_reference", release="PUBLIC")
    _package(runs, "still_private", release="PRIVATE", broken_payload=True)
    before = _tree(runs)
    assert load(runs, public_only=True) == {}
    rows = load(runs, public_only=True, privacy_profile=_profile(tmp_path))
    assert list(rows) == ["display_alias"]
    assert rows["display_alias"]["privacy"]["strain_id"] == "public_reference"
    assert _tree(runs) == before


def test_profile_cannot_promote_stored_private_or_conflicting_identity(tmp_path):
    profile, digest = load_bound_profile(_profile(tmp_path))
    package = _package(tmp_path / "runs", "alias", strain="public_reference", release="PRIVATE")
    assert package_release_policy(package, "alias", profile=profile, profile_sha256=digest)["effective_release"] == "PRIVATE"
    (package / "other_1_intake.json").write_text(json.dumps({"strain_id": "other_synthetic", "release": "PUBLIC"}))
    policy = package_release_policy(package, "alias", profile=profile, profile_sha256=digest)
    assert policy["effective_release"] == "PRIVATE" and policy["privacy_reason"] == "identity_conflict"


def test_compact_manifest_fallback_retains_release_authority(tmp_path):
    from mamey.cohort_figures import _load_manifest_short
    package = _package(tmp_path / "runs", "public_reference", release="PUBLIC",
                       decision=resolve_run_release("public_reference", "PUBLIC"))
    short = package / "manifest_short.json"
    short.rename(package / "manifest.json")
    metadata = _load_manifest_short(str(package), "public_reference")
    assert metadata["release"] == "PUBLIC"
    assert metadata["release_decision"]["reason"] == "operator_public_override"


def test_public_figure_output_must_be_fresh(tmp_path):
    pytest.importorskip("matplotlib")
    from mamey.cohort_figures import generate
    out = tmp_path / "figures"
    out.mkdir()
    old = out / "private_previous_figure.svg"
    old.write_text("synthetic previous private figure")
    with pytest.raises(ValueError, match="PUBLIC_FIGURE_OUTPUT_NOT_EMPTY"):
        generate(tmp_path / "runs", out, public_only=True)
    assert list(out.iterdir()) == [old]


def test_public_figure_receipt_omits_excluded_identifiers(tmp_path, monkeypatch):
    pytest.importorskip("matplotlib")
    from mamey import cohort_figures as cf
    runs = tmp_path / "runs"
    _package(runs, "public_reference", release="PUBLIC",
             decision=resolve_run_release("public_reference", "PUBLIC"))
    _package(runs, "private_synthetic_secret", release="PRIVATE", broken_payload=True)
    # Bound the suite to a representative real renderer, retaining actual
    # public admission, plot/sidecar writes and privacy receipt publication.
    def representative(rows, sid, out):
        assert not cf.is_private(sid, rows)
        cf.heatmap([[1]], ["synthetic_count"], [sid], rows, "Synthetic correctness fixture", "privacy_fixture", out)
    monkeypatch.setattr(cf, "figs_single", representative)
    out = tmp_path / "figures"
    result = cf.generate(runs, out, public_only=True)
    assert result["figures"] == 1 and result["strains"] == ["public_reference"]
    receipt = json.loads((out / "cohort_privacy_receipt.json").read_text())
    assert receipt["excluded"] == 1
    for path in out.iterdir():
        if path.suffix in {".json", ".jsonl", ".csv", ".svg", ".md"}:
            assert "private_synthetic_secret" not in path.read_text()


@pytest.mark.parametrize("class_view,reason", [
    (None, "PUBLIC_CLASS_VIEW_REQUIRED"),
    ("RAW_ANTISMASH", "PUBLIC_MASTER_PATH_REQUIRED"),
    ("GENE_BACKED", "PUBLIC_MASTER_PATH_REQUIRED"),
])
def test_combined_public_cohort_refuses_missing_scope_before_write(tmp_path, class_view, reason):
    from mamey.cohort_deliverable import run_cohort_deliverable
    out = tmp_path / "combined"
    result = run_cohort_deliverable(tmp_path / "runs", out, public_only=True,
                                   class_view=class_view)
    assert result.warnings == [reason]
    assert not result.ok and not result.deliverables and result.strains == []
    assert not out.exists()



def test_extended_suite_receives_only_the_admitted_base_subset(tmp_path, monkeypatch):
    pytest.importorskip("matplotlib")
    from mamey import cli, cohort_figures, cohort_figures_extended
    seen = []
    monkeypatch.setattr(cohort_figures, "generate", lambda **_k: {
        "figures": 1, "strains": ["public_reference"], "out": str(tmp_path), "files": []})
    monkeypatch.setattr(cohort_figures_extended, "generate_extended", lambda **kw: (
        seen.append(kw["strains"]) or {"figures": 0, "errors": []}))
    args = argparse.Namespace(runs_dir=str(tmp_path / "runs"), out=str(tmp_path / "figures"),
                              strains=None, public_only=True, series="F", extended=True)
    assert cli.cohort_figures_command(args) == 0
    assert seen == [["public_reference"]]


@pytest.mark.parametrize("mutation", [
    {"refused": True}, {"strain_id": "other_synthetic"},
    {"authority_sha256": int("9" * 64), "reason": "exact_public_assignment"},
])
def test_malformed_public_decision_is_private_hold(tmp_path, mutation):
    package = _package(tmp_path / "runs", "public_reference", release="PUBLIC",
                       decision=resolve_run_release("public_reference", "PUBLIC"))
    path = package / "manifest_short.json"
    metadata = json.loads(path.read_text())
    metadata["release_decision"].update(mutation)
    path.write_text(json.dumps(metadata))
    assert package_release_policy(package, "public_reference")["effective_release"] == "PRIVATE"
