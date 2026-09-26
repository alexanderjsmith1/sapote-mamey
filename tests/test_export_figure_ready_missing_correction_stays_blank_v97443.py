"""A strain with no A3 correction exports a blank loss, not its whole raw count.

export_figure_ready.py read a missing A3 corrected count as 0, so fragmentation_loss became the
strain's entire raw BGC count, and a strain with no raw count became 0. Missing stays blank;
an observed 0 stays 0. plot_examples.py (the consumer) skips blank rows instead of crashing.
"""
import csv
import importlib.util
import os
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def _book(path):
    wb = openpyxl.Workbook()
    a2 = wb.active
    a2.title = "A2_Strain_Registry"
    a2.append(["strain", "taxonomy", "ecology_source", "assembly_tier", "contigs", "n50",
               "assembly_bp", "gc_pct", "interior_pct", "bgc_count"])
    a2.append(["REF-1", "Streptomyces sp.", "soil", "GOOD", 10, 250000, 8000000, 72, 90, 10])
    a2.append(["REF-2", "Streptomyces sp.", "soil", "POOR", 400, 12000, 7500000, 71, 40, 7])
    a2.append(["REF-3", "Streptomyces sp.", "soil", "GOOD", 12, 300000, 8100000, 72, 95, None])
    a2.append(["REF-4", "Streptomyces sp.", "soil", "GOOD", 9, 310000, 8200000, 72, 96, 5])
    a3 = wb.create_sheet("A3_Run_Manifest")
    a3.append(["strain", "raw_bgcs", "corrected_bgcs"])
    a3.append(["REF-1", 10, 8])
    a3.append(["REF-4", 5, 0])  # an observed zero
    b1 = wb.create_sheet("B1_BGC_Master")
    b1.append(["strain", "BGC_ID", "contig", "region", "products", "boundary", "length_kb",
               "kcb_top", "kcb_score", "safe_claim"])
    b1.append(["REF-1", "REF-1_001", "c1", 1, "NRPS", "interior", 40, "", "", ""])
    wb.save(path)


def _export(tmp_path):
    book = tmp_path / "master.xlsx"
    _book(book)
    out = tmp_path / "figure_ready"
    spec = importlib.util.spec_from_file_location("export_figure_ready_v97443", ROOT / "tools" / "export_figure_ready.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.main(str(book), str(out))
    with open(out / "strain_summary.csv", newline="", encoding="utf-8") as fh:
        return out, {r["strain"]: r for r in csv.DictReader(fh)}


def test_missing_correction_is_blank_not_a_fabricated_loss(tmp_path):
    _, rows = _export(tmp_path)
    assert rows["REF-1"]["corrected_bgcs"] in ("8", "8.0") and float(rows["REF-1"]["fragmentation_loss"]) == 2
    assert rows["REF-2"]["corrected_bgcs"] == "" and rows["REF-2"]["fragmentation_loss"] == ""
    assert rows["REF-3"]["raw_bgcs"] == "" and rows["REF-3"]["fragmentation_loss"] == ""
    assert float(rows["REF-4"]["corrected_bgcs"]) == 0 and float(rows["REF-4"]["fragmentation_loss"]) == 5


def test_plot_examples_skips_blank_loss(tmp_path):
    out, _ = _export(tmp_path)
    env = dict(os.environ, MPLBACKEND="Agg")
    proc = subprocess.run([sys.executable, str(ROOT / "tools" / "plot_examples.py"), str(out)],
                          capture_output=True, text=True, env=env, cwd=str(ROOT))
    assert proc.returncode == 0, proc.stderr[-600:]
    assert (out / "fig1_fragmentation_gradient.png").stat().st_size > 0
