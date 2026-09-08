"""M1 — merge_workbooks.py: normalize-then-append, PK uniqueness, fail-closed on collision (synthetic)."""
import os, subprocess, sys
import openpyxl

TOOL = os.path.join(os.path.dirname(__file__), "..", "tools", "merge_workbooks.py")


def _wb(path, sheet, header, rows, info="x"):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.create_sheet("_SCHEMA_INFO").append([info])
    wb.save(path)


def _mapping(path, rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Column_Mapping"
    ws.append(["canonical", "divergent", "action", "notes"])
    for r in rows:
        ws.append(r)
    wb.save(path)


def _setup(tmp):
    A = str(tmp / "A.xlsx"); B = str(tmp / "B.xlsx"); M = str(tmp / "map.xlsx")
    _wb(A, "B1_BGC_Master", ["strain", "BGC_ID", "region", "boundary", "products", "start"],
        [["AGLAU", "BGC001", "region001", "Interior", "NRPS;PKS", 100]])
    _wb(B, "B1_BGC_Master", ["strain", "BGC_ID", "Region", "Edge_Status", "Products"],
        [["MHUMI", "BGC001", 1, "full-contig", "PKS/T1PKS"]])
    _mapping(M, [
        ["strain", "strain", "direct", ""],
        ["BGC_ID", "BGC_ID", "direct", "PK"],
        ["region", "Region", "TRANSFORM", "int->regionNNN"],
        ["boundary", "Edge_Status", "RENAME+TRANSFORM", "case"],
        ["products", "Products", "TRANSFORM", "/ -> ;"],
        ["start", "(missing)", "BLANK", "no coords"],
    ])
    return A, B, M


def _run(A, B, M, out, extra=()):
    return subprocess.run([sys.executable, TOOL, "--canonical", A, "--sources", B,
                           "--mapping", M, "--out", out, *extra], capture_output=True, text=True)


def test_normalize_then_append(tmp_path):
    A, B, M = _setup(tmp_path)
    out = str(tmp_path / "merged.xlsx")
    r = _run(A, B, M, out)
    assert r.returncode == 0, r.stdout + r.stderr
    wb = openpyxl.load_workbook(out, data_only=True)
    rows = list(wb["B1_BGC_Master"].iter_rows(values_only=True))
    hdr = list(rows[0]); data = [dict(zip(hdr, x)) for x in rows[1:]]
    assert len(data) == 2  # 1 + 1 appended
    b = [d for d in data if d["strain"] == "MHUMI"][0]
    assert b["region"] == "region001"      # int -> regionNNN
    assert b["boundary"] == "Full-contig"  # case normalized
    assert b["products"] == "PKS;T1PKS"    # delimiter normalized
    assert b["start"] in (None, "")        # honest blank
    assert b["bgc_uid"] == "MHUMI:BGC001"


def test_pk_collision_fails_closed(tmp_path):
    A, B, M = _setup(tmp_path)
    # make B collide with A on strain:BGC_ID
    _wb(B, "B1_BGC_Master", ["strain", "BGC_ID", "Region", "Edge_Status", "Products"],
        [["AGLAU", "BGC001", 1, "edge", "PKS"]])
    out = str(tmp_path / "m.xlsx")
    r = _run(A, B, M, out)
    assert r.returncode == 2 and "PK COLLISION" in r.stdout


def test_named_transform_explicit_and_unknown(tmp_path):
    """Spec names a transform explicitly (transform=region_fmt) and an unknown name -> blank + conserved raw."""
    A = str(tmp_path / "A.xlsx"); B = str(tmp_path / "B.xlsx"); M = str(tmp_path / "map.xlsx")
    _wb(A, "B1_BGC_Master", ["strain", "BGC_ID", "region", "novelty_auto"],
        [["AGLAU", "BGC001", "region001", "Likely known"]])
    _wb(B, "B1_BGC_Master", ["strain", "BGC_ID", "Region", "RiQ"],
        [["MHUMI", "BGC001", 7, 0.94]])
    _mapping(M, [
        ["strain", "strain", "direct", ""],
        ["BGC_ID", "BGC_ID", "direct", ""],
        ["region", "Region", "TRANSFORM=region_fmt", "explicit name"],          # explicit
        ["novelty_auto", "RiQ", "TRANSFORM", "no threshold -> deferred/unknown"],  # unknown -> blank+conserve
    ])
    out = str(tmp_path / "m.xlsx")
    r = _run(A, B, M, out)
    assert r.returncode == 0, r.stdout + r.stderr
    wb = openpyxl.load_workbook(out, data_only=True)
    rows = list(wb["B1_BGC_Master"].iter_rows(values_only=True)); hdr = list(rows[0])
    b = [dict(zip(hdr, x)) for x in rows[1:] if dict(zip(hdr, x))["strain"] == "MHUMI"][0]
    assert b["region"] == "region007"          # explicit named transform applied
    assert b["novelty_auto"] in (None, "")     # unknown transform -> honest blank
    assert b.get("riq") == 0.94                # raw value conserved
