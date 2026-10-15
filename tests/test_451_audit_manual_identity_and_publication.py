"""Synthetic regressions through the real ingest-blastp CLI dispatch.

Identity used: SYN-A / NODE_1_length_900_cov_1 / region001 / BGC001.
The original workbook and raw CSV are synthetic and retained byte-for-byte on failure.
"""
import csv
import errno
import json
import os
from pathlib import Path

import openpyxl
import pytest
from mamey.cli import main
from mamey.master_workbook import CANONICAL_V1_HEADERS

def fixture(tmp_path, query="SYN-A|BGC001|gene=ctg1_1|aa=300", package=True):
    master = tmp_path / "master.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Retained_Evidence"
    wb.active.append(["original evidence", "must survive failed update"])
    ws = wb.create_sheet("B5_BLASTp_Hits")
    ws.append(CANONICAL_V1_HEADERS["B5_BLASTp_Hits"])
    wb.save(master)
    raw = tmp_path / "hits.csv"
    raw.write_text(f"{query},WP_SYN.1,70,300,90,0,1,300,1,300,1e-60,400,80\n")
    args = ["ingest-blastp", "--master", str(master), "--strain", "SYN-A", "--hit-table", str(raw)]
    pkg = tmp_path / "package"
    if package:
        pkg.mkdir()
        (pkg / "manifest.json").write_text(json.dumps({"strain_id":"SYN-A"}))
        context = {"bgc_id":"BGC001", "node":"NODE_1_length_900_cov_1", "region":"region001",
                   "cds":[{"locus_tag":"ctg1_1", "aa_length":300}]}
        (pkg / "SYN-A_gene_context.jsonl").write_text(json.dumps(context)+"\n")
        args += ["--package", str(pkg)]
    return master, raw, pkg, args

def b5_rows(master):
    wb = openpyxl.load_workbook(master, read_only=True)
    ws = wb["B5_BLASTp_Hits"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    rows = [dict(zip(header,r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    return rows

def test_foreign_query_cannot_be_relabelled_as_selected_strain(tmp_path):
    master, raw, pkg, args = fixture(tmp_path, "SYN-B|BGC001|gene=ctg1_1|aa=300")
    raw_before = raw.read_bytes()
    before = master.read_bytes()
    with pytest.raises(ValueError, match="QUERY_STRAIN_CONFLICT"):
        main(args)
    assert master.read_bytes() == before
    assert b5_rows(master) == [], "foreign query was attributed to SYN-A in B5"
    assert not (pkg / "blastp_online").exists()
    assert not (pkg / "blastp_quarantine").exists()
    assert raw.read_bytes() == raw_before

def test_matching_query_and_legacy_query_remain_idempotent(tmp_path):
    master, raw, pkg, args = fixture(tmp_path)
    assert main(args) == main(args) == 0
    rows = b5_rows(master)
    assert len(rows) == 1 and rows[0]["strain"] == "SYN-A"
    raw.write_text("BGC001_ctg1_1,WP_LEGACY.1,70,300,90,0,1,300,1,300,1e-60,400,80\n")
    assert main(args) == 0
    assert len(b5_rows(master)) == 2

def test_failed_save_preserves_existing_master_and_raw_input(tmp_path, monkeypatch):
    master, raw, pkg, args = fixture(tmp_path, package=False)
    before, raw_before = master.read_bytes(), raw.read_bytes()
    attempted = []
    def interrupted_save(self, destination):
        attempted.append(Path(destination))
        Path(destination).write_bytes(b"PK truncated synthetic disk-full write")
        raise OSError(errno.ENOSPC, "synthetic disk full after opening output")
    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, "save", interrupted_save)
    with pytest.raises(OSError, match="synthetic disk full"):
        main(args)
    assert attempted, "fixture did not reach the production workbook save"
    assert master.read_bytes() == before, "failed BLASTp update truncated retained master"
    assert raw.read_bytes() == raw_before
    assert sorted(p.name for p in tmp_path.iterdir()) == ["hits.csv", "master.xlsx"]

@pytest.mark.parametrize("alias", ["symlink", "hardlink"])
def test_selected_workbook_alias_cannot_modify_other_retained_evidence(tmp_path, alias):
    master, raw, pkg, args = fixture(tmp_path, package=False)
    retained = tmp_path / "retained_original.xlsx"
    master.rename(retained)
    if alias == "symlink": master.symlink_to(retained)
    else: os.link(retained, master)
    before = retained.read_bytes()
    if alias == "symlink":
        with pytest.raises(ValueError, match="symlink workbook destination"):
            main(args)
        assert master.is_symlink()
    else:
        assert main(args) == 0
        assert len(b5_rows(master)) == 1
        assert not os.path.samefile(master, retained)
    assert retained.read_bytes() == before, "workbook alias overwrote retained original"

@pytest.mark.parametrize("with_package", [True, False])
def test_foreign_xml_title_cannot_supply_selected_strain_metadata(tmp_path, with_package):
    master, raw, pkg, args = fixture(tmp_path, "BGC001_ctg1_1", package=with_package)
    xml = tmp_path / "alignment.xml"
    xml.write_text("<BlastOutput2><Report><Results><Search>"
        "<query-id>BGC001_ctg1_1</query-id><query-len>300</query-len>"
        "<query-title>SYN-B|BGC001|gene=ctg1_1|aa=300</query-title>"
        "<hits></hits></Search></Results></Report></BlastOutput2>")
    before, raw_before = master.read_bytes(), raw.read_bytes()
    with pytest.raises(ValueError, match="QUERY_STRAIN_CONFLICT"):
        main(args + ["--xml", str(xml)])
    assert master.read_bytes() == before and raw.read_bytes() == raw_before
    if with_package: assert not (pkg / "blastp_online").exists()

def test_same_strain_two_bgcs_keep_exact_query_subject_binding(tmp_path):
    # SYN-A / NODE_1_length_900_cov_1 / region001 / BGC001 and
    # SYN-A / NODE_2_length_900_cov_1 / region001 / BGC002.
    master, raw, pkg, args = fixture(tmp_path)
    context = {"bgc_id":"BGC002", "node":"NODE_2_length_900_cov_1", "region":"region001",
               "cds":[{"locus_tag":"ctg2_1", "aa_length":300}]}
    with (pkg / "SYN-A_gene_context.jsonl").open("a") as handle:
        handle.write(json.dumps(context)+"\n")
    raw.write_text(raw.read_text()+"SYN-A|BGC002|gene=ctg2_1|aa=300,WP_SECOND.1,60,300,120,0,1,300,1,300,1e-40,300,70\n")
    assert main(args) == 0
    assert {(r["query_locus"],r["BGC_ID"],r["subject_acc"],r["strain"]) for r in b5_rows(master)} == {
        ("SYN-A|BGC001|gene=ctg1_1|aa=300","BGC001","WP_SYN.1","SYN-A"),
        ("SYN-A|BGC002|gene=ctg2_1|aa=300","BGC002","WP_SECOND.1","SYN-A")}
    assert len(list((pkg / "blastp_online").glob("*_online_blastp.csv"))) == 2

def test_late_foreign_query_refuses_whole_input_before_any_publication(tmp_path):
    master, raw, pkg, args = fixture(tmp_path)
    raw.write_text(raw.read_text()+"SYN-B|BGC001|gene=ctg1_1|aa=300,WP_FOREIGN.1,80,300,60,0,1,300,1,300,1e-60,500,90\n")
    before, raw_before = master.read_bytes(), raw.read_bytes()
    package_before = {str(p.relative_to(pkg)):p.read_bytes() for p in pkg.rglob("*") if p.is_file()}
    with pytest.raises(ValueError,match="QUERY_STRAIN_CONFLICT"):
        main(args)
    assert master.read_bytes() == before and raw.read_bytes() == raw_before
    assert {str(p.relative_to(pkg)):p.read_bytes() for p in pkg.rglob("*") if p.is_file()} == package_before
