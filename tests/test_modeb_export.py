"""test_modeb_export.py — FA6 Mode B card exporter (markdown → .docx / .pdf).

A small synthetic §1-§30 Mode B card (with a §4 evidence grid table and the
verbatim claim-safety language) is exported. We assert:
  * a .pdf is written, non-empty, renders >0 pages, and carries the claim-safety
    footer text on the page;
  * a .docx is written, non-empty, and its section footer holds the footer text
    (SKIPPED cleanly when python-docx is unavailable);
  * the card's own claim-safety language survives verbatim into the PDF.

reportlab ships with the bundle, so the PDF assertions run unconditionally. The
DOCX assertions skip (not fail) when python-docx is absent.
"""
import pytest

from mamey import modeb_export
from mamey.modeb_export import CLAIM_SAFETY_FOOTER

CLAIM_SAFETY_LINE = "KCB comparisons are similarity signals, not identity."


def _synthetic_card() -> str:
    """A compact §1-§30 Mode B card with headings, a §4 evidence grid table,
    bullets, a blockquote, and the verbatim claim-safety block."""
    sections = [
        "# Mode B Card — BGC001 (TEST-STRAIN)",
        "",
        "> Claim-safety: All class assignments are bioinformatic predictions.",
        f"> {CLAIM_SAFETY_LINE}",
        "> Bioactivity links are mechanistic hypotheses requiring experimental validation.",
        "",
        "## §1 Identity",
        "BGC001 on ctg1_region001; class **NRPS**; length 42.1 kb.",
        "",
        "## §2 Boundary confidence",
        "Interior region; boundaries well supported.",
        "",
        "## §3 Architecture",
        "Type A architecture; class confidence *MODERATE*.",
        "",
        "## §4 Evidence grid",
        "",
        "| Evidence channel | Signal | Ceiling |",
        "|------------------|--------|---------|",
        "| KnownClusterBlast | BGC0000001 (score 512) | source-derived similarity |",
        "| Pfam domains | AMP-binding, Condensation | domain-level |",
        "| Resistance | T3 transporter only | routing |",
        "",
    ]
    # §5 … §30 to make it a full-depth card.
    for n in range(5, 31):
        sections.append(f"## §{n} Section {n}")
        sections.append(f"- Deterministic bullet for section {n}.")
        sections.append(f"Prose for section {n} with `inline code` and **bold** text.")
        sections.append("")
    return "\n".join(sections)


@pytest.fixture()
def card(tmp_path):
    p = tmp_path / "TEST-STRAIN_BGC001_ModeB.md"
    p.write_text(_synthetic_card(), encoding="utf-8")
    return p


def test_pdf_written_and_renders(card, tmp_path):
    out = tmp_path / "out"
    res = modeb_export.export_card_pdf(card, out / f"{card.stem}.pdf")
    assert res["status"] == "WRITTEN", res
    pdf = out / f"{card.stem}.pdf"
    assert pdf.exists() and pdf.stat().st_size > 0

    # page count > 0 via pypdf/PyPDF2 (whichever the bundle ships); fall back to a
    # byte-level %%EOF / /Page check if neither is importable.
    try:
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader
        reader = PdfReader(str(pdf))
        assert len(reader.pages) > 0
        text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
        assert "Class-level capacity hypothesis" in text  # footer on the page
        assert CLAIM_SAFETY_LINE in text                   # card's own language, verbatim
    except ImportError:
        raw = pdf.read_bytes()
        assert b"%%EOF" in raw and raw.count(b"/Type /Page") >= 1


def test_docx_written_or_skipped(card, tmp_path):
    out = tmp_path / "out"
    res = modeb_export.export_card_docx(card, out / f"{card.stem}.docx")
    if not modeb_export.docx_available():
        assert res["status"] == "SKIPPED_NO_DOCX"
        pytest.skip("python-docx not installed — docx path skipped gracefully")
    assert res["status"] == "WRITTEN", res
    docx_path = out / f"{card.stem}.docx"
    assert docx_path.exists() and docx_path.stat().st_size > 0

    from docx import Document

    doc = Document(str(docx_path))
    # footer carries the claim-safety line
    assert doc.sections[0].footer.paragraphs[0].text == CLAIM_SAFETY_FOOTER
    # the §4 evidence grid became a real Word table (>=1 table, header row intact)
    assert len(doc.tables) >= 1
    header = [c.text for c in doc.tables[0].rows[0].cells]
    assert "Evidence channel" in header


def test_export_card_both_formats(card, tmp_path):
    res = modeb_export.export_card(card, outdir=tmp_path / "both", formats=("docx", "pdf"))
    assert res["pdf"]["status"] == "WRITTEN"
    assert res["docx"]["status"] in ("WRITTEN", "SKIPPED_NO_DOCX")


# --- v9.7.347 NC-009: reportlab surfaced as a CORE render dependency, distinct from optional extras ---
def test_nc009_core_render_dependency_status():
    import importlib.util
    from mamey import modeb_export
    st = modeb_export.core_render_dependency_status()
    assert st["dependency"] == "reportlab"
    assert st["tier"] == "CORE"
    assert st["state"] in ("OK", "CORE_DEPENDENCY_MISSING")
    assert "optional" not in st["state"].lower()  # must NOT be treated as an optional figure extra
    have = importlib.util.find_spec("reportlab") is not None
    assert st["present"] is have
    if not have:
        assert st["state"] == "CORE_DEPENDENCY_MISSING"
        assert "reportlab" in st["detail"] and "PDF-ready" in st["detail"]


# --- v9.7.347 NC-010: cut-time Mode-B PDF/DOCX renderer smoke tests ---
def _tiny_modeb_card(tmp_path):
    md = tmp_path / "AS-1_BGC001_ModeB.md"
    md.write_text(
        "# AS-1 BGC001 — Mode B\n\n"
        "## §1 Overview\n\nClass-level capacity hypothesis; judgment deferred; similarity not identity.\n\n"
        "## §3 Gene table\n\n"
        "| locus | role |\n|---|---|\n| ctg1_1 | KS |\n| ctg1_2 | AT |\n",
        encoding="utf-8")
    return md


def test_nc010_pdf_render_smoke(tmp_path):
    """A real Mode-B card must render to a PDF that opens (in the cut environment)."""
    pytest.importorskip("reportlab")
    from mamey import modeb_export
    out = tmp_path / "card.pdf"
    modeb_export.export_card_pdf(_tiny_modeb_card(tmp_path), out)
    assert out.exists() and out.stat().st_size > 0
    assert out.read_bytes()[:5] == b"%PDF-"  # opens as a PDF


def test_nc010_docx_render_smoke(tmp_path):
    """A real Mode-B card must render to a valid .docx (in the cut environment)."""
    pytest.importorskip("docx")
    from mamey import modeb_export
    if not modeb_export.docx_available():
        pytest.skip("python-docx not available")
    out = tmp_path / "card.docx"
    modeb_export.export_card_docx(_tiny_modeb_card(tmp_path), out)
    assert out.exists() and out.stat().st_size > 0
    import zipfile
    with zipfile.ZipFile(out) as z:
        assert "word/document.xml" in z.namelist()  # structurally valid docx
