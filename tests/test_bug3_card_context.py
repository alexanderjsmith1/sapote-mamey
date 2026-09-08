"""Regression test (Bug 3): cross_strain_card_context.load_prevalence must derive from
B2_Product_Class_Matrix when Cross_Strain_Class_Prevalence is absent (bare gold master),
instead of silently returning {} and dropping the cross-strain card §-blocks."""
import openpyxl, tempfile, os
from mamey.cross_strain_card_context import load_prevalence

def _master(path, with_prevalence):
    wb = openpyxl.Workbook(); wb.active.title = "A2_Strain_Registry"
    wb.active.append(["Strain"]); [wb.active.append([s]) for s in ("AS-1","AS-2","AS-3")]
    b2 = wb.create_sheet("B2_Product_Class_Matrix")
    b2.append(["strain","NRPS","PKS","terpene","counts_reliability"])
    b2.append(["AS-1",2,0,1,"ok"]); b2.append(["AS-2",1,3,0,"ok"]); b2.append(["AS-3",0,1,1,"ok"])
    if with_prevalence:
        p = wb.create_sheet("Cross_Strain_Class_Prevalence")
        p.append(["product_class","n_strains","band","informative_for_comparison"])
        p.append(["NRPS",2,"common","yes"])
    wb.save(path)

def test_derives_from_b2_when_dedicated_absent():
    with tempfile.TemporaryDirectory() as t:
        m = os.path.join(t,"Mamey_Master.xlsx"); _master(m, with_prevalence=False)
        prev = load_prevalence(m)
        assert prev, "bare master must yield prevalence derived from B2, not {}"
        assert prev["NRPS"]["n_strains"] == 2   # AS-1, AS-2
        assert prev["PKS"]["n_strains"] == 2    # AS-2, AS-3
        assert prev["terpene"]["n_strains"] == 2  # AS-1, AS-3

def test_dedicated_sheet_still_read():
    with tempfile.TemporaryDirectory() as t:
        m = os.path.join(t,"Mamey_Master.xlsx"); _master(m, with_prevalence=True)
        assert load_prevalence(m)["NRPS"]["n_strains"] == 2
