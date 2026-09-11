"""SCHEMA-P03 / SCHEMA-P04 (v9.7.328): workbook_schema_check.validate must not crash — and must
not leak the read_only zip handle — on a master workbook whose worksheets omit the stored
<dimension> element. In openpyxl read_only mode that makes ws.max_row / ws.max_column return None,
so the old code did `range(1, None + 1)` / `None - 1` -> TypeError (past any caller guard), and
because the body was not under try/finally the open archive handle leaked.
"""
import os
import re
import gc
import sys
import zipfile
import warnings

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
openpyxl = pytest.importorskip("openpyxl")
from mamey.workbook_schema_check import validate, CANONICAL_REQUIRED


def _make_master(path, strip_dimension):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in sorted(CANONICAL_REQUIRED):
        ws = wb.create_sheet(sheet)
        ws.append(["col1", "col2"])
        ws.append(["AS-001", "x"])
    wb.save(path)
    if not strip_dimension:
        return path
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    for n in list(data):
        if re.match(r"xl/worksheets/sheet\d+\.xml$", n):
            data[n] = re.sub(rb"<dimension[^/]*/>", b"", data[n])
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    return path


def test_missing_dimension_does_not_crash(tmp_path):
    nodim = _make_master(str(tmp_path / "master_nodim.xlsx"), strip_dimension=True)
    good = _make_master(str(tmp_path / "master_good.xlsx"), strip_dimension=False)
    # Both must return a structured result (no TypeError), and the dimension scan must recover the
    # same picture the stored-dimension workbook gives.
    r_nodim = validate(nodim)
    r_good = validate(good)
    assert r_nodim["sheets_found"] == r_good["sheets_found"] == len(CANONICAL_REQUIRED)
    assert r_nodim["strain_count"] == r_good["strain_count"] == 1


def test_missing_dimension_does_not_leak_handle(tmp_path):
    nodim = _make_master(str(tmp_path / "master_nodim.xlsx"), strip_dimension=True)
    with warnings.catch_warnings():
        warnings.simplefilter("error", ResourceWarning)
        validate(nodim)
        gc.collect()  # would raise ResourceWarning-as-error if the read_only archive leaked


# Invalid schema inputs must be reported, not disappear through fallback handlers.
def _schema_master(path, n50=1_000_000, efls=0):
    from mamey.workbook_schema_check import REQUIRED_SHEETS, D5_COLUMNS
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, spec in REQUIRED_SHEETS.items():
        ws = wb.create_sheet(name)
        headers = list(spec.get("columns", spec.get("prefix", ["label"])))
        if name in {"C1_DAPR_Antibacterial", "C2_DAPR_Antifungal"} and "Activity_Ref" not in headers:
            headers.append("Activity_Ref")
        ws.append(headers)
    ws = wb.create_sheet("Fragment_Rescue_Tiers")
    ws.append(D5_COLUMNS)
    ws.append(["A", "fixture-strain", "complete", 1, n50, 0, efls, 0, 0, "review"])
    wb.save(path)
    wb.close()
    return path


@pytest.mark.parametrize("n50,efls", [("invalid", 0), (1_000_000, "invalid")])
def test_malformed_tier_values_fail_with_row_diagnostic(tmp_path, n50, efls):
    result = validate(_schema_master(tmp_path / "master.xlsx", n50, efls), mode="full", check_v12=True)
    assert result["status"] == "FAIL"
    assert any("Fragment_Rescue_Tiers row 2" in e and "cannot validate tier inputs" in e
               for e in result.get("v1_2_errors", []))


def test_valid_tier_and_explicit_fast_scope_remain_supported(tmp_path):
    good = _schema_master(tmp_path / "good.xlsx")
    assert validate(good, mode="full", check_v12=True)["status"] == "PASS"
    bad = _schema_master(tmp_path / "bad.xlsx", "invalid")
    result = validate(bad, mode="fast", check_v12=True)
    assert result["status"] == "PASS"
    assert result["validation_path"] == "fast-structural"
    assert any("per-row value re-validation" in e for e in result["warnings"])


def test_failed_dimension_scan_is_explicit_and_closes_workbook(tmp_path, monkeypatch):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from openpyxl.workbook.workbook import Workbook
    path = _schema_master(tmp_path / "master.xlsx")
    original_scan = ReadOnlyWorksheet.calculate_dimension
    original_load = openpyxl.load_workbook
    original_close = Workbook.close
    opened, closed = [], []

    def load(*args, **kwargs):
        wb = original_load(*args, **kwargs)
        wb["A1_Dashboard"].reset_dimensions()
        opened.append(wb)
        return wb

    def scan(self, force=False):
        if self.title == "A1_Dashboard" and force:
            raise OSError("fixture scan failure")
        return original_scan(self, force=force)

    def close(self):
        closed.append(self)
        return original_close(self)

    monkeypatch.setattr(openpyxl, "load_workbook", load)
    monkeypatch.setattr(ReadOnlyWorksheet, "calculate_dimension", scan)
    monkeypatch.setattr(Workbook, "close", close)
    result = validate(path, mode="full")
    assert result["status"] == "FAIL"
    assert any(e.get("sheet") == "A1_Dashboard" and e.get("code") == "DIMENSION_SCAN_FAILED"
               for e in result["column_errors"])
    assert opened and all(wb in closed for wb in opened)


def test_malformed_tier_cli_returns_json_and_nonzero(tmp_path):
    import json
    import subprocess
    from pathlib import Path
    path = _schema_master(tmp_path / "master.xlsx", "invalid")
    root = Path(__file__).resolve().parents[1]
    for command in ([sys.executable, str(root / "mamey/workbook_schema_check.py")],
                    [sys.executable, "-m", "mamey.workbook_schema_check"]):
        proc = subprocess.run(command + ["--full", "--v12", str(path)], cwd=root,
                              text=True, capture_output=True)
        result = json.loads(proc.stdout)
        assert proc.returncode == 1
        assert result["status"] == "FAIL"
        assert "cannot validate tier inputs" in " ".join(result["v1_2_errors"])
