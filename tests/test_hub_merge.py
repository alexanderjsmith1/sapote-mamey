"""hub_merge — composed ingest→schema-gate→normalize→merge→verify. No-drift auto-path + drift fail-closed."""
import os, subprocess, sys
import openpyxl

TOOL = os.path.join(os.path.dirname(__file__), "..", "tools", "hub_merge.py")


def _wb(path, header, rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "B1_BGC_Master"
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_no_drift_auto_direct_merge(tmp_path):
    # identical column signatures → NO_DRIFT → synthesized direct mapping, no --mapping needed
    cols = ["strain", "BGC_ID", "products"]
    A = str(tmp_path / "A.xlsx"); B = str(tmp_path / "B.xlsx"); out = str(tmp_path / "m.xlsx")
    _wb(A, cols, [["AGLAU", "BGC001", "NRPS"]])
    _wb(B, cols, [["MHUMI", "BGC001", "PKS"]])
    r = subprocess.run([sys.executable, TOOL, "--canonical", A, "--sources", B, "--out", out],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "NO_DRIFT" in r.stdout
    wb = openpyxl.load_workbook(out, data_only=True)
    assert len(list(wb["B1_BGC_Master"].iter_rows(values_only=True))) - 1 == 2


def test_drift_without_mapping_fails_closed(tmp_path):
    A = str(tmp_path / "A.xlsx"); B = str(tmp_path / "B.xlsx"); out = str(tmp_path / "m.xlsx")
    _wb(A, ["strain", "BGC_ID", "products"], [["AGLAU", "BGC001", "NRPS"]])
    _wb(B, ["strain", "BGC_ID", "Products"], [["MHUMI", "BGC001", "PKS"]])  # 'Products' renamed → drift
    r = subprocess.run([sys.executable, TOOL, "--canonical", A, "--sources", B, "--out", out],
                       capture_output=True, text=True)
    assert r.returncode == 2 and "SCHEMA DRIFT" in r.stdout
    assert not os.path.exists(out)
