"""Synthetic T15/T16/T17 controls; no real workbooks or scientific processing."""
import importlib.util
import os
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('merge_workbooks_safety', ROOT / 'tools/merge_workbooks.py')
merge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(merge)


def wb(path, header, rows, extra=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = merge.SHEET_DEFAULT
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    if extra:
        book.create_sheet('KEEP_FIXTURE_METADATA').append(['synthetic sentinel'])
    book.save(path)
    book.close()
    return path


def mapping(path, rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Column_Mapping'
    sheet.append(['canonical', 'divergent', 'action', 'notes'])
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()
    return path


def setup(tmp_path):
    a = wb(tmp_path / 'canonical.xlsx', ['note', 'strain', 'BGC_ID'],
           [['canonical', 'SYNTH_A', 'BGC001']], extra=True)
    b = wb(tmp_path / 'source.xlsx', ['note', 'strain', 'BGC_ID'],
           [['source', 'SYNTH_B', 'BGC002']], extra=True)
    m = mapping(tmp_path / 'mapping.xlsx', [[name, name, 'direct', ''] for name in ['note', 'strain', 'BGC_ID']])
    return a, b, m


def test_optional_first_blank_preserves_both_input_rows_and_counts(tmp_path):
    a, b, m = setup(tmp_path)
    wb(a, ['note', 'strain', 'BGC_ID'], [[None, 'SYNTH_A', 'BGC001'], [None, None, None]])
    wb(b, ['note', 'strain', 'BGC_ID'], [[None, 'SYNTH_B', 'BGC002']])
    result = merge.run_merge(a, [b], m, tmp_path / 'new.xlsx')
    assert result['status'] == 'OK' and result['rows'] == 2
    assert result['src_counts'] == {a.name: 1, b.name: 1}


@pytest.mark.parametrize('value', [0, False, ''])
def test_false_or_empty_first_value_does_not_drop_other_cells(tmp_path, value):
    path = wb(tmp_path / 'a.xlsx', ['optional', 'strain', 'BGC_ID'], [[value, 'SYNTH_A', 'BGC001']])
    assert len(merge._read_sheet(path, merge.SHEET_DEFAULT)[1]) == 1


@pytest.mark.parametrize('input_role', ['canonical', 'source', 'mapping'])
@pytest.mark.parametrize('destination_role', ['out', 'report'])
@pytest.mark.parametrize('alias', ['direct', 'symlink', 'hardlink'])
def test_artifact_input_aliases_are_refused_before_any_write(tmp_path, input_role, destination_role, alias):
    a, b, m = setup(tmp_path)
    paths = {'canonical': a, 'source': b, 'mapping': m}
    target = paths[input_role]
    destination = target
    if alias == 'symlink':
        destination = tmp_path / 'alias.xlsx'
        destination.symlink_to(target)
    elif alias == 'hardlink':
        destination = tmp_path / 'alias.xlsx'
        os.link(target, destination)
    before = {p: p.read_bytes() for p in (a, b, m)}
    out, report = tmp_path / 'new.xlsx', tmp_path / 'report.md'
    if destination_role == 'out':
        out = destination
    else:
        report = destination
    with pytest.raises(ValueError, match='alias'):
        merge.run_merge(a, [b], m, out, report=report)
    assert {p: p.read_bytes() for p in before} == before
    assert not (tmp_path / 'new.xlsx').exists()
    assert not (tmp_path / 'report.md').exists()


@pytest.mark.parametrize('alias', ['same_path', 'resolved_path', 'symlink', 'hardlink'])
def test_output_and_report_aliases_are_refused(tmp_path, alias):
    a, b, m = setup(tmp_path)
    out = tmp_path / 'new.xlsx'
    report = out
    if alias == 'resolved_path':
        nested = tmp_path / 'nested'
        nested.mkdir()
        report = nested / '..' / out.name
    elif alias in ('symlink', 'hardlink'):
        out.write_bytes(b'existing independent output sentinel')
        report = tmp_path / 'alias.md'
        if alias == 'symlink':
            report.symlink_to(out)
        else:
            os.link(out, report)
    previous = out.read_bytes() if out.exists() else None
    with pytest.raises(ValueError, match='alias'):
        merge.run_merge(a, [b], m, out, report=report)
    assert (out.read_bytes() if out.exists() else None) == previous


def test_new_outputs_work_and_leave_all_source_sheets_untouched(tmp_path):
    a, b, m = setup(tmp_path)
    before = {p: p.read_bytes() for p in (a, b, m)}
    result = merge.run_merge(a, [b], m, tmp_path / 'new.xlsx', report=tmp_path / 'report.md')
    assert result['status'] == 'OK' and result['rows'] == 2
    assert {p: p.read_bytes() for p in before} == before
    assert (tmp_path / 'report.md').read_text().startswith('# Merge report')


@pytest.mark.parametrize('action', ['BLANK', 'BLANK+FLAG'])
def test_real_source_blank_actions_leave_target_blank(tmp_path, action):
    m = mapping(tmp_path / 'mapping.xlsx', [['target', 'source', action, '']])
    rules, flags, conserved, _ = merge.parse_mapping(m)
    result, raw = merge.normalize_row({'source': 'must not copy'}, rules, conserved, ['target'])
    assert result['target'] is None
    assert flags == (['target'] if action == 'BLANK+FLAG' else [])


@pytest.mark.parametrize('action', ['ERASE', 'DIRECTISH', 'BLANK+TRANSFORM', 'NOT_TRANSFORM=identity'])
def test_unknown_or_mixed_actions_refused(tmp_path, action):
    m = mapping(tmp_path / 'mapping.xlsx', [['target', 'source', action, '']])
    with pytest.raises(ValueError, match='action'):
        merge.parse_mapping(m)


def test_missing_source_blank_flag_and_raw_map_conservation_still_work(tmp_path):
    m = mapping(tmp_path / 'mapping.xlsx', [
        ['target', 'source', 'BLANK+FLAG', ''],
        ['(conserved)', 'source', 'MAP', '-> source_raw'],
        ['missing', '(missing)', 'BLANK+FLAG', 'provenance gap'],
    ])
    rules, flags, conserved, _ = merge.parse_mapping(m)
    result, raw = merge.normalize_row({'source': 'raw evidence'}, rules, conserved, ['target', 'missing'])
    assert result == {'target': None, 'missing': None}
    assert flags == ['target', 'missing']
    assert raw == {'source_raw': 'raw evidence'}


def test_documented_transform_and_move_controls(tmp_path):
    m = mapping(tmp_path / 'mapping.xlsx', [
        ['renamed', 'a', 'RENAME', ''],
        ['region', 'b', 'RENAME+TRANSFORM:region_fmt', ''],
        ['products', 'c', 'TRANSFORM', 'legacy canonical fallback'],
        ['unknown_transform', 'd', 'TRANSFORM=unregistered', ''],
        ['(conserved)', 'e', 'MAP', '→ evidence_raw'],
        ['(move)', 'taxonomy', 'DROP/move', '→ A2_Strain_Registry'],
    ])
    rules, _, conserved, taxonomy = merge.parse_mapping(m)
    result, raw = merge.normalize_row({'a': 'rename control', 'b': 4, 'c': 'A/B', 'd': 'raw', 'e': 'evidence'},
                                      rules, conserved, ['renamed', 'region', 'products', 'unknown_transform'])
    assert result == {'renamed': 'rename control', 'region': 'region004', 'products': 'A;B', 'unknown_transform': None}
    assert raw == {'evidence_raw': 'evidence', 'd': 'raw'} and taxonomy == 'taxonomy'


@pytest.mark.parametrize('header', [['strain', 'strain'], ['strain', None], ['strain', ' ']])
def test_invalid_headers_are_explicitly_refused(tmp_path, header):
    path = wb(tmp_path / 'invalid.xlsx', header, [['SYNTH_A', 'value']])
    with pytest.raises(ValueError, match='header'):
        merge._read_sheet(path, merge.SHEET_DEFAULT)


def test_missing_required_key_reports_source_row_and_never_writes(tmp_path):
    a, b, m = setup(tmp_path)
    wb(b, ['note', 'strain', 'BGC_ID'], [[None, 'SYNTH_B', None]])
    with pytest.raises(ValueError, match=r'source.xlsx.*row 2.*BGC_ID'):
        merge.run_merge(a, [b], m, tmp_path / 'new.xlsx')
    assert not (tmp_path / 'new.xlsx').exists()


def test_conflicting_canonical_target_mappings_refused(tmp_path):
    m = mapping(tmp_path / 'mapping.xlsx', [['target', 'a', 'DIRECT', ''], ['target', 'b', 'RENAME', '']])
    with pytest.raises(ValueError, match='conflict'):
        merge.parse_mapping(m)


def test_missing_blank_target_cannot_conflict_with_real_column_rule(tmp_path):
    m = mapping(tmp_path / 'mapping.xlsx', [['target', '(missing)', 'BLANK', ''], ['target', 'source', 'DIRECT', '']])
    with pytest.raises(ValueError, match='conflict'):
        merge.parse_mapping(m)


def test_blank_flag_survives_end_to_end_ledger_without_copying_source_value(tmp_path):
    a = wb(tmp_path / 'a.xlsx', ['strain', 'BGC_ID', 'target'], [['SYNTH_A', 'BGC001', 'canonical']])
    b = wb(tmp_path / 'b.xlsx', ['strain', 'BGC_ID', 'raw'], [['SYNTH_B', 'BGC002', 'DO_NOT_COPY']])
    m = mapping(tmp_path / 'mapping.xlsx', [['strain', 'strain', 'DIRECT', ''],
                                           ['BGC_ID', 'BGC_ID', 'DIRECT', ''],
                                           ['target', 'raw', 'BLANK+FLAG', '']])
    out = tmp_path / 'out.xlsx'
    result = merge.run_merge(a, [b], m, out)
    assert result['prov_gap'] == ['target']
    book = openpyxl.load_workbook(out, data_only=True)
    try:
        cells = list(book[merge.SHEET_DEFAULT].iter_rows(values_only=True))
        header = list(cells[0])
        source_row = dict(zip(header, cells[2]))
        assert source_row['target'] is None
        ledger = {row[0]: row for row in book['_MERGE_LEDGER'].iter_rows(min_row=2, values_only=True)}
        assert ledger['target'][1] == 1 and 'PROVENANCE GAP' in ledger['target'][2]
    finally:
        book.close()


def test_unknown_action_cli_refuses_before_writing_any_artifact(tmp_path):
    import subprocess
    import sys
    a, b, m = setup(tmp_path)
    mapping(m, [['strain', 'strain', 'DIRECT', ''], ['BGC_ID', 'BGC_ID', 'DIRECT', ''], ['note', 'note', 'ERASE', '']])
    out = tmp_path / 'new.xlsx'
    report = tmp_path / 'report.md'
    p = subprocess.run([sys.executable, '-B', str(ROOT / 'tools/merge_workbooks.py'),
                        '--canonical', str(a), '--sources', str(b), '--mapping', str(m),
                        '--out', str(out), '--report', str(report)], capture_output=True, text=True)
    assert p.returncode == 2 and 'MERGE_REFUSED' in p.stderr
    assert not out.exists() and not report.exists()


@pytest.mark.parametrize('role', ['canonical', 'source'])
def test_styled_empty_trailing_columns_do_not_refuse_valid_table(tmp_path, role):
    a, b, m = setup(tmp_path)
    path = a if role == 'canonical' else b
    book = openpyxl.load_workbook(path)
    book[merge.SHEET_DEFAULT].cell(40, 40).font = openpyxl.styles.Font(bold=True)
    book.save(path)
    book.close()
    before = path.read_bytes()
    header, rows = merge._read_sheet(path, merge.SHEET_DEFAULT)
    assert header == ['note', 'strain', 'BGC_ID'] and len(rows) == 1
    result = merge.run_merge(a, [b], m, tmp_path / 'new.xlsx')
    assert result['status'] == 'OK' and result['rows'] == 2
    assert path.read_bytes() == before


@pytest.mark.parametrize('internal', [False, True])
def test_blank_headers_are_not_trimmed_when_internal_or_with_data(tmp_path, internal):
    header = ['strain', None, 'BGC_ID'] if internal else ['strain', 'BGC_ID', None]
    row = ['SYNTH_A', None, 'BGC001'] if internal else ['SYNTH_A', 'BGC001', 'retain this evidence']
    path = wb(tmp_path / 'invalid.xlsx', header, [row])
    with pytest.raises(ValueError, match='header'):
        merge._read_sheet(path, merge.SHEET_DEFAULT)


@pytest.mark.parametrize('reverse', [False, True])
@pytest.mark.parametrize('collision', ['map', 'canonical', 'case_sources'])
def test_raw_fallback_namespaces_refuse_collisions_before_writes(tmp_path, reverse, collision):
    canonical_columns = ['strain', 'BGC_ID', 'target', 'other_target']
    if collision == 'canonical':
        canonical_columns.append('raw')
    a = wb(tmp_path / 'canonical.xlsx', canonical_columns, [['SYNTH_A', 'BGC001']])
    b = wb(tmp_path / 'source.xlsx', ['strain', 'BGC_ID', 'Raw', 'Other', 'raw'],
           [['SYNTH_B', 'BGC002', 'transform raw', 'mapped raw', 'case raw']])
    extra = [['target', 'Raw', 'TRANSFORM=unregistered', '']]
    if collision == 'map':
        extra.append(['(conserved)', 'Other', 'MAP', '-> raw'])
    elif collision == 'case_sources':
        extra.append(['other_target', 'raw', 'TRANSFORM=unregistered', ''])
    if reverse:
        extra.reverse()
    m = mapping(tmp_path / 'mapping.xlsx', [['strain', 'strain', 'DIRECT', ''],
                                            ['BGC_ID', 'BGC_ID', 'DIRECT', '']] + extra)
    before = {p: p.read_bytes() for p in (a, b, m)}
    out, report = tmp_path / 'new.xlsx', tmp_path / 'report.md'
    with pytest.raises(ValueError, match=r'mapping.xlsx.*mapping row [45].*raw.*collision'):
        merge.run_merge(a, [b], m, out, report=report)
    assert not out.exists() and not report.exists()
    assert {p: p.read_bytes() for p in before} == before


@pytest.mark.parametrize('reverse', [False, True])
def test_same_source_identical_raw_destination_is_safe(tmp_path, reverse):
    a = wb(tmp_path / 'canonical.xlsx', ['strain', 'BGC_ID', 'target'], [['SYNTH_A', 'BGC001']])
    b = wb(tmp_path / 'source.xlsx', ['strain', 'BGC_ID', 'Raw'], [['SYNTH_B', 'BGC002', 'retain raw']])
    extra = [['target', 'Raw', 'TRANSFORM=unregistered', ''],
             ['(conserved)', 'Raw', 'MAP', '-> raw']]
    if reverse:
        extra.reverse()
    m = mapping(tmp_path / 'mapping.xlsx', [['strain', 'strain', 'DIRECT', ''],
                                            ['BGC_ID', 'BGC_ID', 'DIRECT', '']] + extra)
    out = tmp_path / 'new.xlsx'
    assert merge.run_merge(a, [b], m, out)['status'] == 'OK'
    book = openpyxl.load_workbook(out, data_only=True)
    try:
        rows = list(book[merge.SHEET_DEFAULT].iter_rows(values_only=True))
        header = list(rows[0])
        assert header.count('raw') == 1
        source = dict(zip(header, rows[2]))
        assert source['raw'] == 'retain raw' and source['target'] is None
    finally:
        book.close()


@pytest.mark.parametrize('raw_column', ['_SRC', 'BGC_UID'])
def test_raw_fallback_cannot_replace_generated_metadata(tmp_path, raw_column):
    a = wb(tmp_path / 'canonical.xlsx', ['strain', 'BGC_ID', 'target'], [['SYNTH_A', 'BGC001']])
    b = wb(tmp_path / 'source.xlsx', ['strain', 'BGC_ID', raw_column],
           [['SYNTH_B', 'BGC002', 'retain raw']])
    m = mapping(tmp_path / 'mapping.xlsx', [['strain', 'strain', 'DIRECT', ''],
                                           ['BGC_ID', 'BGC_ID', 'DIRECT', ''],
                                           ['target', raw_column, 'TRANSFORM=unregistered', '']])
    with pytest.raises(ValueError, match=r'mapping row 4.*collision.*generated column'):
        merge.run_merge(a, [b], m, tmp_path / 'new.xlsx')
    assert not (tmp_path / 'new.xlsx').exists()


def test_namespace_preflight_runs_even_without_source_rows(tmp_path):
    a = wb(tmp_path / 'canonical.xlsx', ['strain', 'BGC_ID', 'raw'], [['SYNTH_A', 'BGC001']])
    b = wb(tmp_path / 'source.xlsx', ['strain', 'BGC_ID', 'Raw'], [])
    m = mapping(tmp_path / 'mapping.xlsx', [['strain', 'strain', 'DIRECT', ''],
                                           ['BGC_ID', 'BGC_ID', 'DIRECT', ''],
                                           ['(conserved)', 'Raw', 'MAP', '-> raw']])
    with pytest.raises(ValueError, match=r'mapping row 4.*collision.*canonical header'):
        merge.run_merge(a, [b], m, tmp_path / 'new.xlsx')
    assert not (tmp_path / 'new.xlsx').exists()


def test_multiple_map_destinations_for_same_source_are_explicit_hold(tmp_path):
    m = mapping(tmp_path / 'mapping.xlsx', [['(conserved)', 'Raw', 'MAP', '-> raw_one'],
                                           ['(conserved)', 'Raw', 'MAP', '-> raw_two']])
    with pytest.raises(ValueError, match=r'mapping row 3.*conflicting'):
        merge.parse_mapping(m)



def test_identical_map_source_and_destination_rows_are_idempotent(tmp_path):
    row = ['(conserved)', 'Raw', 'MAP', '-> raw']
    m = mapping(tmp_path / 'mapping.xlsx', [row, row])
    rules, _, conserved, _ = merge.parse_mapping(m)
    assert dict(conserved) == {'Raw': 'raw'}
    assert merge.normalize_row({'Raw': 'retain raw'}, rules, conserved, []) == ({}, {'raw': 'retain raw'})



def test_merge_output_can_be_next_canonical_repeatedly(tmp_path):
    a, b, m = setup(tmp_path)
    current = a
    for turn in range(3):
        wb(b, ['note', 'strain', 'BGC_ID'], [['source', f'SYNTH_{turn}', f'BGC{turn + 10:03}']])
        out = tmp_path / f'roundtrip_{turn}.xlsx'
        assert merge.run_merge(current, [b], m, out)['status'] == 'OK'
        header, rows = merge._read_sheet(out, merge.SHEET_DEFAULT)
        assert len(header) == len(set(header)) and header.count('bgc_uid') == 1
        assert len(rows) == turn + 2
        assert all(row['bgc_uid'] == f"{row['strain']}:{row['BGC_ID']}" for row in rows)
        current = out


@pytest.mark.parametrize('uid', ['SYNTH_OTHER:BGC999', 0])
def test_nonempty_canonical_uid_conflicting_with_keys_refuses_before_writes(tmp_path, uid):
    a, b, m = setup(tmp_path)
    wb(a, ['bgc_uid', 'note', 'strain', 'BGC_ID'], [[uid, 'canonical', 'SYNTH_A', 'BGC001']])
    before = {path: path.read_bytes() for path in (a, b, m)}
    out, report = tmp_path / 'new.xlsx', tmp_path / 'report.md'
    with pytest.raises(ValueError, match=r'canonical.xlsx.*row 2.*bgc_uid.*SYNTH_A:BGC001'):
        merge.run_merge(a, [b], m, out, report=report)
    assert not out.exists() and not report.exists()
    assert {path: path.read_bytes() for path in before} == before


@pytest.mark.parametrize('uid', [None, '', ' ', 'SYNTH_A:BGC001'])
def test_optional_blank_or_matching_canonical_uid_is_derived_once(tmp_path, uid):
    a, b, m = setup(tmp_path)
    wb(a, ['bgc_uid', 'note', 'strain', 'BGC_ID'], [[uid, 'canonical', 'SYNTH_A', 'BGC001']])
    out = tmp_path / 'new.xlsx'
    assert merge.run_merge(a, [b], m, out)['status'] == 'OK'
    book = openpyxl.load_workbook(out, data_only=True)
    try:
        rows = list(book[merge.SHEET_DEFAULT].iter_rows(values_only=True))
        header = list(rows[0])
        assert header.count('bgc_uid') == 1
        assert dict(zip(header, rows[1]))['bgc_uid'] == 'SYNTH_A:BGC001'
    finally:
        book.close()
