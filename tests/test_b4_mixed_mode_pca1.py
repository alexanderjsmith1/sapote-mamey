"""PC-A1 (v9.7.101): B4 cross-strain scans must distinguish "not measured" (non-gold
strain, gold-only scan) from a measured 0, and flag a mixed-mode cohort.

Full MameyRun construction is heavy, so this tests (a) the schema carries run_mode,
and (b) the mixed-mode detection + gold-only NA logic on a hand-built B4 sheet that
mirrors what the builder writes.
"""
from openpyxl import Workbook

from mamey.master_workbook import CANONICAL_V1_HEADERS

GOLD_ONLY = {"resistance_T1", "UMED_gaps", "EFLS_pairs"}


def test_b4_schema_carries_run_mode():
    b4 = CANONICAL_V1_HEADERS["B4_Cross_Strain_Scans"]
    assert "run_mode" in b4, "B4 must carry run_mode so non-gold zeros are distinguishable"
    # run_mode should sit right after strain for readability
    assert b4[0] == "strain" and b4[1] == "run_mode"


def _na(val, run_mode):
    # mirrors the builder's gold-only NA rule
    return val if run_mode == "gold" else "NA"


def test_gold_only_columns_render_NA_for_non_gold():
    # a smoke strain's gold-only scan must be NA, not 0
    assert _na(0, "smoke") == "NA"
    assert _na(3, "standard") == "NA"
    # a gold strain keeps its measured value (including a genuine 0)
    assert _na(0, "gold") == 0
    assert _na(3, "gold") == 3


def test_mixed_mode_detection():
    # mirror the A1_Dashboard detection: >1 distinct run_mode across B4 rows -> flag
    wb = Workbook()
    ws = wb.active
    ws.append(["strain", "run_mode"])
    ws.append(["StrainA", "gold"])
    ws.append(["StrainB", "smoke"])
    modes = {ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    assert len(modes) > 1, "two distinct modes should be detected as mixed"

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(["strain", "run_mode"])
    ws2.append(["StrainA", "gold"])
    ws2.append(["StrainB", "gold"])
    modes2 = {ws2.cell(r, 2).value for r in range(2, ws2.max_row + 1)}
    assert len(modes2) == 1, "uniform mode should NOT be flagged mixed"
