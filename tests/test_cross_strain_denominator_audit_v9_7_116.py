"""v9.7.116: cross-strain denominator invariant (Bunny Hop / cohort-master verification).

A cohort-size denominator on a cross-strain sheet (e.g. 'ectoine 17/18') must equal the registry
length. The 24-strain master shipped with two Cross_Strain_Findings rows frozen at 18-strain
literals — a silent stale-number bug a synthesis writer would faithfully reproduce. These tests pin
the fail-closed audit that catches it.
"""
import os
import sys
import tempfile
import pathlib

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from cross_strain_denominator_audit import audit_denominators  # noqa: E402


def _make_master(path, findings_rows, n_strains=24, n_bgc=100):
    wb = openpyxl.Workbook()
    reg = wb.active
    reg.title = "A2_Strain_Registry"
    reg.append(["strain", "taxonomy"])
    for i in range(n_strains):
        reg.append([f"AS-{900 + i}", "sp."])
    b1 = wb.create_sheet("B1_BGC_Master")
    b1.append(["bgc_uid"])
    for i in range(n_bgc):
        b1.append([f"AS-900:BGC{i}"])
    f = wb.create_sheet("Cross_Strain_Findings")
    f.append(["#", "finding", "value"])
    for r in findings_rows:
        f.append(r)
    wb.save(path)


def test_clean_master_passes():
    d = tempfile.mkdtemp()
    p = os.path.join(d, "clean.xlsx")
    _make_master(p, [[3, "universal", "24/24"],
                     [4, "ectoine", "20/24 / 18/24"],
                     [7, "dark", "15/100 (15%)"]])   # BGC-total ratio, not a cohort count
    assert audit_denominators(p) == []


def test_stale_denominator_caught():
    d = tempfile.mkdtemp()
    p = os.path.join(d, "stale.xlsx")
    _make_master(p, [[3, "universal", "18/18"],
                     [4, "ectoine", "17/18 / 15/18"]])
    violations = audit_denominators(p)
    assert len(violations) == 3
    assert all("Stale cohort count" in v for v in violations)


def test_bgc_total_ratio_not_flagged():
    """A ratio over the BGC total (e.g. 164/1131) must not be flagged as a stale cohort count."""
    d = tempfile.mkdtemp()
    p = os.path.join(d, "bgc.xlsx")
    _make_master(p, [[7, "dark", "164/1131 (15%)"]], n_strains=24, n_bgc=1131)
    assert audit_denominators(p) == []


def test_empty_registry_fails_closed():
    d = tempfile.mkdtemp()
    p = os.path.join(d, "empty.xlsx")
    _make_master(p, [[3, "x", "1/1"]], n_strains=0)
    violations = audit_denominators(p)
    assert violations and "cannot establish cohort size" in violations[0]
