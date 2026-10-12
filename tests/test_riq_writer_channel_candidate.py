"""RiQ reaches BGC/report fields without changing KnownClusterBlast precedence."""
import csv
import json
from pathlib import Path
import zipfile

import pytest

from mamey.antismash_evidence import apply_evidence_to_bgcs, parse_antismash_evidence
from mamey.models import BGCRecord

CONTIG = 'SYNTHETIC_CONTIG_000001'

def modules(scores):
    return {'cluster_compare':{'by_region':{str(n):{
        'RegionToRegion_RiQ':{'scores_by_region':{'reference':v}}}
        for n,v in scores.items()}}}

def block(rank, accession, score, proteins):
    return (f'>>\n{rank}. {accession}\nSource: synthetic comparator\nType: PKS\n'
        f'Number of proteins with BLAST hits to this cluster: {proteins}\nCumulative BLAST score: {score}\n\n'
        'Table of Blast hits (query gene, subject gene, %identity, blast score, %coverage, e-value):\n'
        f'query\t{accession}_g1\t60\t100\t95.0\t0.0\n\n')

def bgc(contig=CONTIG):
    return BGCRecord(bgc_id='BGC001',contig=contig,region_number=1,start=0,end=20000,contig_length=90000,products=['T1PKS'])

@pytest.mark.parametrize('score',[0.0,0.73])
def test_kcb_priority_keeps_independent_riq(tmp_path,score):
    z=tmp_path/'source.zip'
    header=f'ClusterBlast scores for {CONTIG}\n\nSignificant hits: \n'
    with zipfile.ZipFile(z,'w') as out:
        out.writestr(f'knownclusterblast/{CONTIG}_c1.txt',header+
            '1. BGC0000002\tsynthetic rank1\n2. BGC0001234\tsynthetic rank2\n\n\nDetails:\n\n'+
            block(1,'BGC0000002',37992.0,5)+block(2,'BGC0001234',61753.0,7))
        out.writestr(f'clusterblast/{CONTIG}_c1.txt',header+
            '1. NZ_SYNTHETIC\tsynthetic genome comparator\n\n\nDetails:\n\n'+block(1,'NZ_SYNTHETIC',73038.0,20))
        out.writestr('source.json',json.dumps({'records':[{'id':CONTIG,'modules':modules({1:score})}]}))
    evidence=parse_antismash_evidence(z,json_mode='bounded')
    b=bgc();apply_evidence_to_bgcs([b],evidence)
    assert b.riq_score==score
    assert b.riq_label==('Potentially novel' if score==0.0 else 'Possibly novel / structural variant')
    assert b.kcb_cumulative==37992.0 and b.kcb_protein_hits==5
    assert b.closest_mibig_accession=='BGC0000002'
    assert b.kcb_evidence_state=='KNOWNCLUSTERBLAST_OBSERVED'
    # Parser and CLI each apply evidence; the second pass must retain both channels.
    apply_evidence_to_bgcs([b],evidence)
    assert b.riq_score==score and b.kcb_cumulative==37992.0

def test_record_region_binding_does_not_cross_contigs_with_same_alias():
    a=bgc();b=bgc('SYNTHETIC_CONTIG_000002')
    evidence={'by_region':{CONTIG+'_c1':[{'source_kind':'knownclusterblast','cumulative_score':10},
        {'source':'source.json','riq_score':0.0}]},'loose_hits':[{'riq_score':.99}]}
    apply_evidence_to_bgcs([a,b],evidence)
    assert a.riq_score==0.0 and b.riq_score is None
    assert evidence['riq_mapping_status']=='UNMAPPED_LOOSE_VALUES_NOT_ASSIGNED_TO_BGC'

def test_missing_riq_remains_null_with_kcb():
    b=bgc()
    apply_evidence_to_bgcs([b],{'by_region':{CONTIG+'_c1':[{'source_kind':'knownclusterblast','cumulative_score':10}]}})
    assert b.riq_score is None and b.riq_label is None

def test_gold_writer_preserves_zero_positive_and_missing_scores(tmp_path):
    from mamey.cli import run_one_strain
    from openpyxl import load_workbook
    from tools.fixture_inputs import prepare_full_locus_single_contig_fixture
    bundle=Path(__file__).resolve().parents[1]
    base=tmp_path/'base.zip'
    recipe=prepare_full_locus_single_contig_fixture(bundle/'tests/fixtures/synthetic_single_contig_antismash.zip',base)
    (tmp_path/'TRANSFORMATION.json').write_text(json.dumps(recipe,indent=2)+'\n')
    source=tmp_path/'scores.zip'
    with zipfile.ZipFile(base) as inp,zipfile.ZipFile(source,'x') as out:
        for info in inp.infolist():out.writestr(info,inp.read(info))
        out.writestr('syn/source.json',json.dumps({'version':'8.0.4','schema':4,'records':[
            {'id':CONTIG,'modules':modules({1:0.0,2:0.73})}]}))
    strain='SYNTHETIC_RIQ_WRITER'
    result = run_one_strain(strain_id=strain,display_name=strain,input_zip=str(source),outdir=str(tmp_path/'runs'),
        mode='gold',taxonomy='not verified',source='invented synthetic acceptance fixture',bioactivity='',
        master_path=None,json_mode='bounded',brief='none',release='PRIVATE',require_workbook=True,
        locus_maps='off',reference_completion='off',rggmci_threads=4)
    assert result['status'] in {'MAMEY_COMPLETE', 'MAMEY_COMPLETE_WITH_ISSUES'}
    assert result['raw_bgcs'] == 3
    assert result['workbook_status']['per_strain']['state'] == 'PRODUCED'
    assert result['package_zip'] and Path(result['package_zip']).is_file()
    package=tmp_path/'runs'/strain/'package'
    with (package/f'{strain}_2_inventory.csv').open(newline='') as f:
        rows={int(r['Region']):r for r in csv.DictReader(f)}
    assert set(rows)=={1,2,3}
    assert float(rows[1]['RiQ_score'])==0.0
    assert float(rows[2]['RiQ_score'])==0.73
    assert rows[3]['RiQ_score']==''
    assert all(r['KCB_evidence_state']=='KNOWNCLUSTERBLAST_OBSERVED' for r in rows.values())
    wb=load_workbook(package/f'{strain}_5_workbook.xlsx',read_only=True,data_only=True)
    try:
        for sheet in ('BGC_Inventory','KCB_RiQ_Sweep'):
            values=list(wb[sheet].iter_rows(values_only=True));headers=list(values[0])
            column='RiQ' if sheet=='BGC_Inventory' else 'RiQ score'
            data={r[headers.index('BGC')]:r[headers.index(column)] for r in values[1:]}
            assert data=={'BGC001':0,'BGC002':0.73,'BGC003':None}
    finally:wb.close()
    manifest=json.loads((package/'manifest.json').read_text())
    assert manifest['release']=='PRIVATE'
