"""Regression coverage for workbook staging and publication contracts."""
from __future__ import annotations

import os
from pathlib import Path

import pytest

from mamey import xlsx_determinism


def _directory_inventory(path: Path) -> set[str]:
    return {entry.name for entry in path.iterdir()}


def test_canonicalize_xlsx_failure_leaves_directory_unchanged(tmp_path, monkeypatch):
    """A failed final rename must leave neither old nor hidden canonical stages."""
    import openpyxl

    target = tmp_path / "report.xlsx"
    openpyxl.Workbook().save(target)
    before = _directory_inventory(tmp_path)
    original = target.read_bytes()

    def fail_replace(source, destination):
        raise OSError("injected final publication failure")

    monkeypatch.setattr(xlsx_determinism.os, "replace", fail_replace)
    with pytest.raises(OSError, match="injected"):
        xlsx_determinism.canonicalize_xlsx(target)

    assert _directory_inventory(tmp_path) == before
    assert target.read_bytes() == original


def test_canonicalize_xlsx_preserves_existing_destination_mode(tmp_path):
    import openpyxl

    target = tmp_path / "report.xlsx"
    openpyxl.Workbook().save(target)
    os.chmod(target, 0o640)

    xlsx_determinism.canonicalize_xlsx(target)

    assert target.stat().st_mode & 0o7777 == 0o640


def test_atomic_save_workbook_safely_cleans_stage_and_preserves_destination(tmp_path):
    """A save failure preserves the destination and leaves no unique stage behind."""
    target = tmp_path / "report.xlsx"
    original = b"known-good-destination"
    target.write_bytes(original)
    before = _directory_inventory(tmp_path)

    class BrokenWorkbook:
        worksheets = []

        def save(self, path):
            Path(path).write_bytes(b"partial-workbook")
            raise OSError("injected save failure")

    with pytest.raises(OSError, match="injected"):
        xlsx_determinism.atomic_save_workbook_safely(BrokenWorkbook(), target)

    assert _directory_inventory(tmp_path) == before
    assert target.read_bytes() == original


def test_atomic_save_workbook_safely_preserves_existing_mode_and_canonicalizes(tmp_path):
    import openpyxl

    target = tmp_path / "report.xlsx"
    original = openpyxl.Workbook()
    original.active["A1"] = "old"
    original.save(target)
    os.chmod(target, 0o640)
    replacement = openpyxl.Workbook()
    replacement.active["A1"] = "new"

    xlsx_determinism.atomic_save_workbook_safely(replacement, target, canonicalize=True)

    assert target.stat().st_mode & 0o7777 == 0o640
    assert openpyxl.load_workbook(target, read_only=True).active["A1"].value == "new"


def test_workbook_publication_refuses_symlink_destinations_without_side_effects(tmp_path):
    import openpyxl

    outside = tmp_path / "outside.xlsx"
    openpyxl.Workbook().save(outside)
    destination = tmp_path / "destination.xlsx"
    destination.symlink_to(outside)
    before = _directory_inventory(tmp_path)
    outside_before = outside.read_bytes()

    with pytest.raises(ValueError, match="symlink"):
        xlsx_determinism.canonicalize_xlsx(destination)
    with pytest.raises(ValueError, match="symlink"):
        xlsx_determinism.atomic_save_workbook_safely(openpyxl.Workbook(), destination)

    assert _directory_inventory(tmp_path) == before
    assert destination.is_symlink()
    assert outside.read_bytes() == outside_before


def test_atomic_save_workbook_safely_new_destination_is_private(tmp_path):
    import openpyxl

    target = tmp_path / "new.xlsx"
    xlsx_determinism.atomic_save_workbook_safely(openpyxl.Workbook(), target)
    assert target.stat().st_mode & 0o7777 == 0o600
