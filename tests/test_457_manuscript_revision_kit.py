"""tools/manuscript_revision.py: staged edits, change log, answerable ledger, asks-only round and the two prose checks.

Every document here is synthetic and built inside the test. The Word steps need python-docx (the `documents` extra);
those tests skip cleanly when it is absent.
"""
from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
TOOL = ROOT / "tools" / "manuscript_revision.py"


def _load():
    sys.path.insert(0, str(TOOL.parent))
    spec = importlib.util.spec_from_file_location("manuscript_revision_under_test", TOOL)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


mr = _load()
try:
    import docx  # noqa: F401
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False
needs_docx = pytest.mark.skipif(not HAVE_DOCX, reason="python-docx (documents extra) not installed")


def _work(tmp_path, text="The survey found 42 plots. The survey ended in spring.\n"):
    root = tmp_path / "work"
    root.mkdir()
    (root / "01_intro.md").write_text(text, encoding="utf-8")
    return root


def _edits(tmp_path, edits, name="edits.json"):
    p = tmp_path / name
    p.write_text(json.dumps(edits), encoding="utf-8")
    return p


# ------------------------------------------------------------------------------------------------ apply

def test_apply_refuses_a_non_unique_find_and_leaves_the_file_untouched(tmp_path):
    root = _work(tmp_path)
    before = (root / "01_intro.md").read_bytes()
    edits = _edits(tmp_path, [
        {"doc": "01_intro.md", "find": "42 plots", "replace": "forty-two plots", "row": "N1"},
        {"doc": "01_intro.md", "find": "The survey", "replace": "This survey", "row": "N2"},
    ])
    code = mr.main(["apply", str(edits), "--root", str(root)])
    assert code == mr.EXIT_FINDINGS
    assert (root / "01_intro.md").read_bytes() == before
    assert not (root / "CHANGES.jsonl").exists()


def test_apply_is_idempotent_and_reports_already_applied(tmp_path, capsys):
    root = _work(tmp_path)
    edits = _edits(tmp_path, [{"doc": "01_intro.md", "find": "ended in spring", "replace": "ended in late spring",
                               "ledger": "Notes", "row": "N3", "why": "author note"}])
    assert mr.main(["apply", str(edits), "--root", str(root)]) == mr.EXIT_OK
    after_first = (root / "01_intro.md").read_text(encoding="utf-8")
    assert "late spring" in after_first
    capsys.readouterr()
    assert mr.main(["apply", str(edits), "--root", str(root)]) == mr.EXIT_OK
    out = capsys.readouterr().out
    assert "applied 0, already applied 1, refused 0" in out
    assert (root / "01_intro.md").read_text(encoding="utf-8") == after_first
    assert len((root / "CHANGES.jsonl").read_text(encoding="utf-8").splitlines()) == 1


def test_apply_dry_run_writes_nothing(tmp_path):
    root = _work(tmp_path)
    before = (root / "01_intro.md").read_bytes()
    edits = _edits(tmp_path, [{"doc": "01_intro.md", "find": "42 plots", "replace": "42 field plots"}])
    assert mr.main(["apply", str(edits), "--root", str(root), "--dry-run"]) == mr.EXIT_OK
    assert (root / "01_intro.md").read_bytes() == before
    assert not (root / "CHANGES.jsonl").exists()


def test_apply_refuses_a_path_outside_the_working_copy(tmp_path):
    root = _work(tmp_path)
    edits = _edits(tmp_path, [{"doc": "../outside.md", "find": "x", "replace": "y"}])
    assert mr.main(["apply", str(edits), "--root", str(root)]) == mr.EXIT_INPUT


# ------------------------------------------------------------------------------------------------ changelog

@needs_docx
def test_changelog_rows_equal_record_lines(tmp_path):
    import docx
    root = _work(tmp_path)
    edits = _edits(tmp_path, [
        {"doc": "01_intro.md", "find": "42 plots", "replace": "42 field plots", "row": "N1", "why": "clarity"},
        {"doc": "01_intro.md", "find": "in spring", "replace": "in late spring", "row": "N2", "why": "author note"},
    ])
    assert mr.main(["apply", str(edits), "--root", str(root)]) == mr.EXIT_OK
    record = root / "CHANGES.jsonl"
    n_lines = len(record.read_text(encoding="utf-8").splitlines())
    labels = tmp_path / "labels.tsv"
    labels.write_text("prefix\tlabel\n01\tIntroduction\n", encoding="utf-8")
    out = tmp_path / "log.docx"
    assert mr.main(["changelog", str(record), "--out", str(out), "--doc-labels", str(labels)]) == mr.EXIT_OK
    table = docx.Document(str(out)).tables[0]
    assert len(table.rows) - 1 == n_lines == 2
    assert table.rows[1].cells[1].text == "Introduction"
    assert "field plots" in table.rows[1].cells[3].text


# ------------------------------------------------------------------------------------------------ ledger + answers

def _ledger_inputs(tmp_path):
    notes = tmp_path / "notes.md"
    notes.write_text(
        "| id | location | note | action | status |\n|---|---|---|---|---|\n"
        "| N001 | p3, paragraph 2 | 'too long' | Shorten the sentence | clear |\n"
        "| N002 | p4, Table 1 | 'which plots?' | Ask: name the plots | ASK |\n"
        "| N003 | p5 | 'faint mark' | Ask: is this a mark? | ASK |\n"
        "| N004 | p6 | 'word [?]' | Replace the word | clear |\n", encoding="utf-8")
    findings = tmp_path / "findings.tsv"
    findings.write_text(
        "id\tgroup\tlocation\tfinding\ttype\trecommendation\toverlap\n"
        "R-01\tIntroduction\tparagraph 2\tSentence repeats the aim.\tFIX\tDelete the repeat.\tN001\n"
        "R-02\tIntroduction\tTable 1\tPlot count not sourced.\tCHECK\tLook up the plot record.\t\n"
        "R-03\tMethods\tsection 2\tScope of the survey.\tRULE\tRecommend: keep the scope as written.\t\n", encoding="utf-8")
    packages = tmp_path / "packages.tsv"
    packages.write_text("id\ttitle\tcovers\tcodes\trecommendation\nP1\tTighten aims\tAims across documents\tR-01\tAccept.\n",
                        encoding="utf-8")
    return notes, findings, packages


@needs_docx
def test_ledger_cross_links_both_sides(tmp_path):
    import docx
    notes, findings, packages = _ledger_inputs(tmp_path)
    out = tmp_path / "ledger.docx"
    assert mr.main(["ledger", "--notes", str(notes), "--findings", str(findings), "--packages", str(packages),
                    "--out", str(out), "--author-label", "Writer", "--reviewer-label", "Panel"]) == mr.EXIT_OK
    d = docx.Document(str(out))
    part_a = d.tables[0]
    assert "Same passage: R-01" in part_a.rows[1].cells[3].text
    assert "Same passage: N001" in d.tables[1].rows[1].cells[3].text
    assert part_a.rows[0].cells[-1].text == "Answer"
    assert any(p.text == "Part B. Panel findings" for p in d.paragraphs)


def _answer(path, out, edits):
    import docx
    d = docx.Document(str(path))
    for ti, ri, ci, text in edits:
        d.tables[ti].rows[ri].cells[ci].text = text
    d.save(str(out))


@needs_docx
def test_read_answers_extracts_answers_and_refuses_a_changed_non_answer_cell(tmp_path):
    notes, findings, packages = _ledger_inputs(tmp_path)
    issued = tmp_path / "ledger.docx"
    mr.main(["ledger", "--notes", str(notes), "--findings", str(findings), "--packages", str(packages), "--out", str(issued)])
    good = tmp_path / "good.docx"
    _answer(issued, good, [(0, 2, 4, "Plots A and B")])
    out = tmp_path / "answers.tsv"
    assert mr.main(["read-answers", str(good), "--original", str(issued), "--out", str(out)]) == mr.EXIT_OK
    assert out.read_text(encoding="utf-8").splitlines() == ["id\tanswer", "N002\tPlots A and B"]
    bad = tmp_path / "bad.docx"
    _answer(issued, bad, [(0, 2, 4, "Plots A and B"), (0, 1, 3, "edited the proposal instead")])
    out2 = tmp_path / "answers2.tsv"
    assert mr.main(["read-answers", str(bad), "--original", str(issued), "--out", str(out2)]) == mr.EXIT_INPUT
    assert not out2.exists()


@needs_docx
def test_asks_drops_answered_and_clear_rows(tmp_path):
    import docx
    notes, findings, packages = _ledger_inputs(tmp_path)
    answers = tmp_path / "answers.tsv"
    answers.write_text("id\tanswer\nN002\tPlots A and B\nP1\tR\n", encoding="utf-8")
    out = tmp_path / "asks.docx"
    assert mr.main(["asks", "--notes", str(notes), "--findings", str(findings), "--packages", str(packages),
                    "--answers", str(answers), "--drop-regex", "faint mark", "--out", str(out)]) == mr.EXIT_OK
    ids = {r.cells[0].text for t in docx.Document(str(out)).tables for r in t.rows[1:]}
    assert ids == {"N004", "R-03"}  # N004 stays: its note has an unread word [?]
    assert not ids & {"N001", "N002", "N003", "R-01", "R-02", "P1"}


# ------------------------------------------------------------------------------------------------ prose checks

def test_check_numbers_flags_an_invented_number_but_not_one_already_present(tmp_path, capsys):
    root = _work(tmp_path, "Plots: 42 in total. Spring visits: 17.\n")
    edits = _edits(tmp_path, [
        {"doc": "01_intro.md", "find": "42 in total", "replace": "42 in total, 17 visited twice", "row": "OK1"},
        {"doc": "01_intro.md", "find": "Spring visits", "replace": "Spring visits (63 hours)", "row": "NEW1"},
    ])
    assert mr.main(["check-numbers", str(edits), "--root", str(root)]) == mr.EXIT_FINDINGS
    out = capsys.readouterr().out
    assert "NEW1\t01_intro.md\t63" in out
    assert "OK1" not in out
    src = tmp_path / "record.tsv"
    src.write_text("hours\n63\n", encoding="utf-8")
    assert mr.main(["check-numbers", str(edits), "--root", str(root), "--source", str(src)]) == mr.EXIT_OK


def test_scan_names_ignores_comments_and_author_notes(tmp_path, capsys):
    doc = tmp_path / "chapter.md"
    doc.write_text(
        "Plots were mapped as before.\n"
        "<!-- source: results/plot_counts_final.tsv, run 7 -->\n"
        "[AUTHOR: check tables/plot_list.csv before print]\n"
        "The count came from plot_counts_final.tsv in run 7.\n"
        "![](figures/map.png)\n", encoding="utf-8")
    pats = tmp_path / "names.txt"
    pats.write_text("# assistant and tool names\nHelperBot\n", encoding="utf-8")
    assert mr.main(["scan-names", str(doc), "--patterns", str(pats)]) == mr.EXIT_FINDINGS
    out = capsys.readouterr().out
    hit_lines = {ln.split("\t")[0].rsplit(":", 1)[1] for ln in out.splitlines() if "\t" in ln}
    assert hit_lines == {"4"}
    clean = tmp_path / "clean.md"
    clean.write_text("Plots were mapped.\n<!-- HelperBot drafted this -->\n", encoding="utf-8")
    assert mr.main(["scan-names", str(clean), "--patterns", str(pats)]) == mr.EXIT_OK


def test_cli_help_runs_from_a_foreign_directory(tmp_path):
    r = subprocess.run([sys.executable, str(TOOL), "--help"], cwd=tmp_path, capture_output=True, text=True, timeout=60)
    assert r.returncode == 0
    for sub in ("apply", "changelog", "ledger", "read-answers", "asks", "check-numbers", "scan-names"):
        assert sub in r.stdout
