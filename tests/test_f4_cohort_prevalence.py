"""Regression test for F4: load_master must derive Cross_Strain_Class_Prevalence from
B2_Product_Class_Matrix when the dedicated sheet is absent (bare gold master), and still
read the dedicated sheet when present (precomputed master)."""
import openpyxl, tempfile, os
from mamey.cohort_synthesis import load_master

def _base(tmp, with_prevalence):
    wb = openpyxl.Workbook()
    a2 = wb.active; a2.title = "A2_Strain_Registry"
    # COH-02 (v9.7.338): use the SHIPPED master_workbook headers, not fake capitalized ones, so
    # load_master's real column resolution is exercised (it now RAISES on a foreign schema).
    a2.append(["strain","taxonomy","ecology_source","habitat","assembly_bp","contigs","n50","gc_pct","bgc_count"])
    for s in ("AS-001","AS-002","AS-003"):
        a2.append([s,"Streptomyces sp.","Bombus terrestris","nest",8_000_000,5,20000,70,2])
    b1 = wb.create_sheet("B1_BGC_Master")
    b1.append(["strain","assembly_locator","contig","region","BGC_ID","start","end","length_kb","products","boundary","arch","kcb_top"])
    for s in ("AS-001","AS-002","AS-003"): b1.append([s]+[""]*11)  # all dark (empty kcb_top)
    b2 = wb.create_sheet("B2_Product_Class_Matrix")
    b2.append(["strain","NRPS","PKS","terpene","counts_reliability"])
    b2.append(["AS-001",2,0,1,"ok"])   # NRPS+terpene
    b2.append(["AS-002",1,3,0,"ok"])   # NRPS+PKS
    b2.append(["AS-003",0,1,1,"ok"])   # PKS+terpene
    if with_prevalence:
        p = wb.create_sheet("Cross_Strain_Class_Prevalence")
        p.append(["product_class","n_strains","band","informative_for_comparison"])
        p.append(["NRPS",2,"common","yes"]); p.append(["PKS",2,"common","yes"]); p.append(["terpene",2,"common","yes"])
    path = os.path.join(tmp,"m.xlsx"); wb.save(path); return path

def test_derives_prevalence_from_b2_when_dedicated_sheet_absent():
    with tempfile.TemporaryDirectory() as t:
        M = load_master(_base(t, with_prevalence=False))
        assert M, "bare master (A2+B1+B2, no dedicated sheet) must NOT be rejected"
        prev = M["prevalence"]
        # NRPS carried by AS-001,AS-002 = 2; PKS by AS-002,AS-003 = 2; terpene by AS-001,AS-003 = 2
        assert prev["NRPS"]["n_strains"] == 2
        assert prev["PKS"]["n_strains"] == 2
        assert prev["terpene"]["n_strains"] == 2

def test_still_reads_dedicated_sheet_when_present():
    with tempfile.TemporaryDirectory() as t:
        M = load_master(_base(t, with_prevalence=True))
        assert M and M["prevalence"]["NRPS"]["n_strains"] == 2

def test_rejects_master_missing_both():
    import openpyxl as ox
    with tempfile.TemporaryDirectory() as t:
        wb = ox.Workbook(); wb.active.title="A2_Strain_Registry"; wb.active.append(["Strain"])
        wb.create_sheet("B1_BGC_Master").append(["strain"])
        p=os.path.join(t,"m.xlsx"); wb.save(p)
        assert load_master(p) == {}, "master with neither prevalence sheet nor B2 must be rejected"
