"""v9.7.90 cohort class-capacity heatmap: matrix build, exclusions, render, skip card."""
from __future__ import annotations
import csv, os, warnings
import pytest

from mamey.cohort_class_heatmap import build_class_matrix, render_cohort_class_heatmap, _EXCLUDED_CLASSES


def _b2_rows():
    return [
        {"strain": "AS-901", "NRPS": 11, "PKS": 9, "RiPP": 10, "saccharide": 5,
         "NAPAA": 3, "terpene": 8, "other": 13},
        {"strain": "AS-902", "NRPS": 18, "PKS": 13, "RiPP": 8, "saccharide": 2,
         "NAPAA": 0, "terpene": 10, "other": 14},
    ]


def test_matrix_excludes_saccharide_but_keeps_napaa():
    # v9.7.101: NAPAA is registry-neutral and no longer dropped; saccharide is still excluded.
    strains, classes, matrix = build_class_matrix(_b2_rows())
    assert "saccharide" not in classes
    assert "NAPAA" in classes, "NAPAA is registry-neutral and should appear in the class matrix"
    assert "saccharide" in _EXCLUDED_CLASSES and "NAPAA" not in _EXCLUDED_CLASSES
    assert "NRPS" in classes and "RiPP" in classes


def test_matrix_drops_all_zero_columns():
    rows = [{"strain": "AS-1", "NRPS": 5, "PKS": 0, "other": 2},
            {"strain": "SYNTHETIC-CASE-002", "NRPS": 3, "PKS": 0, "other": 1}]
    strains, classes, matrix = build_class_matrix(rows)
    assert "PKS" not in classes  # all-zero column dropped
    assert "NRPS" in classes


def test_matrix_values_preserved():
    strains, classes, matrix = build_class_matrix(_b2_rows())
    assert strains == ["AS-901", "AS-902"]
    nrps_idx = classes.index("NRPS")
    assert matrix[0][nrps_idx] == 11
    assert matrix[1][nrps_idx] == 18


def test_matrix_preserves_missing_separately_from_observed_zero():
    rows = [
        {"strain": "AS-1", "NRPS": 5, "PKS": "", "RiPP": 0},
        {"strain": "SYNTHETIC-CASE-002", "NRPS": 3, "PKS": 2, "RiPP": 1},
    ]
    _, classes, matrix = build_class_matrix(rows)
    assert matrix[0][classes.index("PKS")] is None
    assert matrix[0][classes.index("RiPP")] == 0


def test_matrix_treats_malformed_count_as_missing_not_zero():
    rows = [
        {"strain": "AS-1", "NRPS": "not-a-count"},
        {"strain": "SYNTHETIC-CASE-002", "NRPS": 2},
    ]
    _, classes, matrix = build_class_matrix(rows)
    assert matrix[0][classes.index("NRPS")] is None


@pytest.mark.skipif(pytest.importorskip("openpyxl") is None, reason="openpyxl needed")
def test_render_from_workbook(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook(); wb.active.title = "A1"
    b2 = wb.create_sheet("B2_Product_Class_Matrix")
    cols = ["strain", "NRPS", "PKS", "RiPP", "saccharide", "NAPAA", "terpene", "other"]
    for j, c in enumerate(cols, 1):
        b2.cell(1, j).value = c
    for r, row in enumerate(_b2_rows(), 2):
        for j, c in enumerate(cols, 1):
            b2.cell(r, j).value = row.get(c, 0)
    mp = str(tmp_path / "cohort.xlsx"); wb.save(mp)

    png = tmp_path / "h.png"; csvp = tmp_path / "h.csv"
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        res = render_cohort_class_heatmap(mp, png, csvp, claim_prefix="PRIVATE")
    assert res["status"] == "OK"
    assert res["n_strains"] == 2
    assert png.exists() and png.stat().st_size > 0
    assert res["independent_gates"]["publication_approval"] == "NOT_ASSESSED"
    # companion CSV excludes saccharide; NAPAA (registry-neutral) is kept
    with open(csvp) as fh:
        header = next(csv.reader(fh))
    assert "saccharide" not in header and "NAPAA" in header
    deprecated_colormap_mutations = [
        str(item.message) for item in caught
        if "set_bad" in str(item.message) or "set_under" in str(item.message)
    ]
    assert not deprecated_colormap_mutations


def test_missing_b2_writes_skip_card(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook(); mp = str(tmp_path / "empty.xlsx"); wb.save(mp)
    png = tmp_path / "h.png"; csvp = tmp_path / "h.csv"
    res = render_cohort_class_heatmap(mp, png, csvp)
    assert res["status"] == "SKIPPED"
    assert png.with_suffix(".SKIPPED.md").exists()


# --- CODEX13B: browser-safe cohort raster (emit _SCREEN.png; keep visual QA REQUIRED) ---
from mamey.cohort_class_heatmap import (
    _browser_safe_raster_dpi, _BROWSER_SAFE_MAX_EDGE_PX, _PUBLICATION_DPI)


def test_small_figure_is_browser_safe_no_screen_raster(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook(); wb.active.title = "A1"
    b2 = wb.create_sheet("B2_Product_Class_Matrix")
    cols = ["strain", "NRPS", "PKS", "terpene", "other"]
    for j, c in enumerate(cols, 1):
        b2.cell(1, j).value = c
    for r, row in enumerate(_b2_rows(), 2):
        for j, c in enumerate(cols, 1):
            b2.cell(r, j).value = row.get(c, 0)
    mp = str(tmp_path / "small.xlsx"); wb.save(mp)
    res = render_cohort_class_heatmap(mp, tmp_path / "s.png", tmp_path / "s.csv")
    assert res["browser_raster"] is None                 # publication PNG already browser-safe
    assert res["raster_eligibility"] == "PASS"
    assert res["independent_gates"]["browser_visual_review"] == "REQUIRED"  # never auto-PASS
    assert not (tmp_path / "s_SCREEN.png").exists()


def test_281_row_geometry_requires_a_browser_raster():
    # the maintained cohort: 281 rows -> height 0.42*281+2.5 = 120.52 in -> 36156 px at 300 DPI
    height_in = 0.42 * 281 + 2.5
    est_pub_edge = int(height_in * _PUBLICATION_DPI)
    assert est_pub_edge > _BROWSER_SAFE_MAX_EDGE_PX          # too big for the browser
    screen_dpi = _browser_safe_raster_dpi(20.0, height_in)
    assert 72 <= screen_dpi < _PUBLICATION_DPI               # a genuinely lower screen DPI
    assert int(height_in * screen_dpi) <= _BROWSER_SAFE_MAX_EDGE_PX  # screen raster fits the browser


def _merge_child_page_text_observations(request, observations):
    parent_guard = getattr(request.config, "_page_text_ratchet", None)
    if parent_guard is not None:
        parent_guard.seen_dirty.update(observations["seen_dirty"])
        parent_guard.seen_clean.update(observations["seen_clean"])


@pytest.mark.slow
def test_large_cohort_emits_screen_raster(tmp_path, request):
    import openpyxl
    wb = openpyxl.Workbook(); wb.active.title = "A1"
    b2 = wb.create_sheet("B2_Product_Class_Matrix")
    cols = ["strain", "NRPS", "PKS", "terpene"]
    for j, c in enumerate(cols, 1):
        b2.cell(1, j).value = c
    for r in range(2, 2 + 110):                              # 110 strains -> just over the threshold
        b2.cell(r, 1).value = f"AS-{r:03d}"
        b2.cell(r, 2).value = (r % 4); b2.cell(r, 3).value = (r % 3); b2.cell(r, 4).value = (r % 2)
    mp = str(tmp_path / "big.xlsx"); wb.save(mp)
    # This fixture's native raster workspace is large even after its Figure dies.
    # Give that workspace a process lifetime; the parent still checks real outputs.
    import inspect, json, pickle, subprocess, sys
    from pathlib import Path
    import matplotlib
    matplotlib.use("Agg")
    settings = tmp_path / "renderer_rc.pickle"
    with settings.open("wb") as handle:
        pickle.dump(dict(matplotlib.rcParams), handle)
    receipt = tmp_path / "renderer_result.json"
    root = Path(inspect.getfile(render_cohort_class_heatmap)).resolve().parents[1]
    worker = r'''
import importlib.util, json, pickle, sys
from pathlib import Path
root, workbook, png, csv, settings, receipt = map(Path, sys.argv[1:])
sys.path.insert(0, str(root))
import matplotlib
matplotlib.use("Agg")
with settings.open("rb") as handle:
    matplotlib.rcParams.update(pickle.load(handle))
# Retain the same source-bound figure-text guard as the parent pytest session.
spec = importlib.util.spec_from_file_location("figure_page_text_ratchet", root / "tests/figure_page_text_ratchet.py")
ratchet = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ratchet)
guard = ratchet.PageTextRatchet(ratchet.load_known_sites())
ratchet.install(guard)
from mamey.cohort_class_heatmap import render_cohort_class_heatmap
result = render_cohort_class_heatmap(str(workbook), png, csv)
receipt.write_text(json.dumps(result), encoding="utf-8")
receipt.with_suffix(".guard.json").write_text(json.dumps({
    "seen_dirty": sorted(guard.seen_dirty), "seen_clean": sorted(guard.seen_clean),
}), encoding="utf-8")
'''
    rendered = subprocess.run(
        [sys.executable, "-B", "-c", worker, str(root), mp,
         str(tmp_path / "b.png"), str(tmp_path / "b.csv"), str(settings), str(receipt)],
        cwd=root, capture_output=True, text=True, timeout=120,
    )
    assert rendered.returncode == 0, rendered.stdout + rendered.stderr
    res = json.loads(receipt.read_text(encoding="utf-8"))
    child_guard = json.loads(receipt.with_suffix(".guard.json").read_text(encoding="utf-8"))
    _merge_child_page_text_observations(request, child_guard)
    assert res["browser_raster"] is not None
    assert res["browser_raster"]["dpi"] < _PUBLICATION_DPI
    assert (tmp_path / "b_SCREEN.png").exists()
    assert res["independent_gates"]["browser_visual_review"] == "REQUIRED"


def test_child_page_text_observations_preserve_dirty_veto_and_clean_reporting(request, monkeypatch):
    from figure_page_text_ratchet import PageTextRatchet
    mixed = "mamey/cohort_class_heatmap.py::render_cohort_class_heatmap"
    child_clean = "mamey/render_brief.py::fig_landscape"
    existing_dirty = "mamey/locus_map.py::render_locus_map"
    parent = PageTextRatchet({mixed, child_clean, existing_dirty})
    parent.seen_clean.add(mixed)
    parent.seen_dirty.add(existing_dirty)
    monkeypatch.setattr(request.config, "_page_text_ratchet", parent, raising=False)
    _merge_child_page_text_observations(request, {
        "seen_dirty": [mixed], "seen_clean": [child_clean],
    })
    assert isinstance(parent.seen_dirty, set) and isinstance(parent.seen_clean, set)
    assert parent.seen_dirty == {existing_dirty, mixed}
    assert parent.seen_clean == {mixed, child_clean}
    assert mixed not in parent.cleared_sites()
    assert parent.cleared_sites() == [child_clean]
