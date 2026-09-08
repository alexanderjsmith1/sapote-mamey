"""v9.7.409 — second hostile pass on sealed .408 (Black Cherry, 2026-09-04 evening): H8 formula injection,
H9 phantom-locus card ingest, H10 outdir inside the engine tree."""
from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SMOKE = ROOT / "examples" / "test_data" / "smoke_antismash_small.zip"


def _run(args, timeout=600, env=None):
    e = dict(os.environ); e.update(env or {})
    return subprocess.run([sys.executable, str(ROOT / "mamey_run.py")] + args, cwd=str(ROOT), capture_output=True, text=True, timeout=timeout, env=e)


# H8 — spreadsheet formula injection from annotation text --------------------------------------------
def test_workbooks_never_carry_live_formulas_from_annotation_text(tmp_path):
    import openpyxl
    inj = tmp_path / "inject.zip"
    with zipfile.ZipFile(SMOKE) as src, zipfile.ZipFile(inj, "w") as dst:
        for n in src.namelist():
            data = src.read(n)
            if n.endswith(".gbk"):
                t = data.decode("utf-8", "replace")
                t = re.sub(r'/product="([^"]*)"', lambda m: '/product="=HYPERLINK(\\"http://evil.example\\",\\"x\\") ' + m.group(1) + '"', t, count=3)
                data = t.encode()
            dst.writestr(n, data)
    out = tmp_path / "out"; master = tmp_path / "master.xlsx"
    r = _run(["run", "--strain", "INJ1", "--input-zip", str(inj), "--outdir", str(out), "--mode", "gold",
              "--capped-session", "--json-evidence", "off", "--master", str(master)])
    assert r.returncode == 0, r.stdout[-600:]
    books = list((out / "INJ1" / "package").glob("*.xlsx")) + list(tmp_path.glob("*.xlsx"))
    assert books
    live = []
    for b in books:
        for ws in openpyxl.load_workbook(b).worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value[:1] in "=+-@" and c.data_type == "f":
                        live.append((b.name, ws.title, c.coordinate))
    assert not live, f"live formula cells (on .408 there were 21): {live[:5]}"
    # the text itself is preserved — only the cell TYPE changed
    inv = next((out / "INJ1" / "package").glob("*_5_workbook.xlsx"))
    found = any(isinstance(c.value, str) and "HYPERLINK" in c.value for ws in openpyxl.load_workbook(inv).worksheets for row in ws.iter_rows() for c in row)
    assert found, "annotation text must be kept verbatim as a string"


def test_neutralise_formula_strings_unit():
    import openpyxl
    from mamey.xlsx_determinism import neutralise_formula_strings
    wb = openpyxl.Workbook(); wb.active.append(['=1+1', '=HYPERLINK("http://x","y")', '+x', 'plain', 12])
    assert neutralise_formula_strings(wb) == 2   # only the two openpyxl typed as formulas
    assert all(c.data_type == "s" for c in wb.active[1] if isinstance(c.value, str))


# H9 — a card naming a locus the package does not contain must not be recorded ---------------------
def _package(tmp_path):
    out = tmp_path / "run"
    r = _run(["run", "--strain", "TEST-01", "--input-zip", str(SMOKE), "--outdir", str(out), "--mode", "gold", "--capped-session", "--json-evidence", "off"])
    assert r.returncode == 0, r.stdout[-600:]
    return out / "TEST-01" / "package"


def test_ingest_refuses_card_with_foreign_node(tmp_path):
    from mamey.mode_b_receipt import ingest_one_card
    pkg = _package(tmp_path)
    card = tmp_path / "card.md"
    r = _run(["emit-modeb-template", "--package", str(pkg), "--bgc", "BGC001", "--out", str(card)])
    assert r.returncode == 0 and card.is_file(), r.stdout[-300:] + r.stderr[-300:]
    text = card.read_text(encoding="utf-8")
    node = re.search(r"\bnode:\s*([^|>\s]+)", text).group(1)
    bad = tmp_path / "bad.md"; bad.write_text(text.replace(node, "NODE_9_length_99999_cov_1"), encoding="utf-8")
    s = ingest_one_card(pkg, bad)
    assert s["status"] == "SKIPPED_IDENTITY_MISMATCH", s["status"]
    assert s["structure_findings"][0]["code"] == "CARD_IDENTITY_MISMATCH"
    reg = json.loads(next(pkg.glob("*judgment_register.json")).read_text(encoding="utf-8")) if list(pkg.glob("*judgment_register.json")) else {}
    assert (reg.get("bgcs", {}).get("BGC001", {}).get("status") or "") != "COMPLETE", "the phantom-locus card must not complete the register"


def test_ingest_refuses_card_with_foreign_strain(tmp_path):
    from mamey.mode_b_receipt import ingest_one_card
    pkg = _package(tmp_path)
    card = tmp_path / "card.md"
    _run(["emit-modeb-template", "--package", str(pkg), "--bgc", "BGC001", "--out", str(card)])
    bad = tmp_path / "bad.md"; bad.write_text(card.read_text(encoding="utf-8").replace("strain: TEST-01", "strain: AS-999"), encoding="utf-8")
    s = ingest_one_card(pkg, bad)
    assert s["status"] == "SKIPPED_IDENTITY_MISMATCH"


def test_ingest_still_accepts_the_untampered_template(tmp_path):
    from mamey.mode_b_receipt import ingest_one_card
    pkg = _package(tmp_path)
    card = tmp_path / "card.md"
    _run(["emit-modeb-template", "--package", str(pkg), "--bgc", "BGC001", "--out", str(card)])
    s = ingest_one_card(pkg, card)
    assert s["status"] != "SKIPPED_IDENTITY_MISMATCH", s


# H10 — an outdir that is the engine root itself, or inside an engine-owned directory, is refused;
# a NEW top-level directory under the bundle (the documented `out/` / `runs/` first-run layout) is allowed ------
def test_run_refuses_outdir_that_is_the_bundle_root():
    r = _run(["run", "--strain", "SYM2", "--input-zip", str(SMOKE), "--outdir", str(ROOT), "--mode", "gold", "--capped-session", "--json-evidence", "off"])
    assert r.returncode == 2 and "[outdir]" in r.stderr, r.stderr[-400:]
    assert not (ROOT / "SYM2").exists()


def test_run_refuses_outdir_inside_engine_owned_dir_via_symlink(tmp_path):
    link = tmp_path / "link"; link.symlink_to(ROOT)
    r = _run(["run", "--strain", "SYM1", "--input-zip", str(SMOKE), "--outdir", str(link / "mamey" / "runs_sym_test"), "--mode", "gold", "--capped-session", "--json-evidence", "off"])
    assert r.returncode == 2 and "[outdir]" in r.stderr and "mamey/" in r.stderr, r.stderr[-400:]
    assert not (ROOT / "mamey" / "runs_sym_test").exists()


def test_new_top_level_outdir_under_bundle_is_allowed():
    from mamey.cli import _refuse_outdir_inside_bundle
    assert _refuse_outdir_inside_bundle("out") is False           # the documented first-run layout
    assert _refuse_outdir_inside_bundle(str(ROOT / "runs" / "x")) is False
    assert _refuse_outdir_inside_bundle(str(ROOT / "tools")) is True
    assert _refuse_outdir_inside_bundle(str(ROOT)) is True


def test_bundle_write_override_env_is_honoured(tmp_path):
    from mamey.cli import _refuse_outdir_inside_bundle
    assert _refuse_outdir_inside_bundle(str(ROOT / "mamey" / "x")) is True
    os.environ["MAMEY_ALLOW_BUNDLE_WRITES"] = "1"
    try:
        assert _refuse_outdir_inside_bundle(str(ROOT / "mamey" / "x")) is False
    finally:
        del os.environ["MAMEY_ALLOW_BUNDLE_WRITES"]
    assert _refuse_outdir_inside_bundle(str(tmp_path / "fine")) is False
