"""P2 engine-side test: master append idempotency via drop_existing_strain."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from mamey.workbook_dedup import drop_existing_strain
import openpyxl

def _wb_with(strains_rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "B1_BGC_Master"
    ws.append(["strain","BGC_ID","x"])
    for s,b in strains_rows: ws.append([s,b,"data"])
    # add A2 too
    a2 = wb.create_sheet("A2_Strain_Registry"); a2.append(["strain","tax"])
    for s in {s for s,_ in strains_rows}: a2.append([s,"Streptomyces"])
    return wb

def test_drop_removes_all_rows_for_strain():
    wb = _wb_with([("X","B1"),("X","B2"),("Y","B1")])
    removed = drop_existing_strain(wb, "X")
    assert removed == 2 + 1  # 2 B1 rows + 1 A2 row
    b1 = wb["B1_BGC_Master"]
    remaining = [r[0] for r in b1.iter_rows(min_row=2, values_only=True) if r[0]]
    assert remaining == ["Y"]

def test_drop_noop_for_absent_strain():
    wb = _wb_with([("X","B1")])
    assert drop_existing_strain(wb, "Z") == 0

def test_reappend_after_drop_is_idempotent():
    wb = _wb_with([("X","B1"),("X","B2")])
    drop_existing_strain(wb, "X")
    b1 = wb["B1_BGC_Master"]
    b1.append(["X","B1","new"]); b1.append(["X","B2","new"])
    from collections import Counter
    keys = Counter((r[0],r[1]) for r in b1.iter_rows(min_row=2, values_only=True) if r[0])
    assert all(n==1 for n in keys.values())
