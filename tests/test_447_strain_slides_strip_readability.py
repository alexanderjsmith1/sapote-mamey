"""tools/strain_slides.py: the multi-contig strip stays readable and never presents a weaker region as the core.

The owner, 2026-10-02 (late), on two strips: labels piled up on four partner panels, reference genes printed as bare protein
accessions (AFV52203.1), and a "core" region with almost no GECCO signal while its "partner" held most reference genes.
The locks: a found gene's label is its reference name, or its first Pfam name + "-like", never an accession; labels are
thinned so neighbours keep a fixed distance, named genes first; at most two partner panels are drawn (the rest are
counted under the strip); a partner region holding more reference genes than this region (at least 3) is flagged.
"""
import sys
from pathlib import Path

import pytest

pytest.importorskip("matplotlib")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import strain_slides as ss  # noqa: E402


def _D():
    return {"gfeat": {"ctg2_5": [{"domain": "PF00001", "i_evalue": "1e-30"}]}, "pfam": {"PF00001": "Trp_halogenase"}}


def test_a_reference_accession_is_never_the_label():
    D = _D()
    assert ss.ref_label(D, {"name": "AFV52203.1", "best_locus": "ctg2_5"}) == "Trp_halogen-like"
    assert ss.ref_label(D, {"name": "QKO28672.1", "best_locus": "ctg2_9"}) == "ctg2_9"
    assert ss.ref_label(D, {"name": "atH", "best_locus": "ctg2_5"}) == "atH"


def test_labels_are_thinned_with_names_first():
    items = [(1.0, "ctg1_1", 2), (1.05, "atH", 3), (5.0, "", 1), (9.0, "pfam", 1)]
    kept = ss.thin_labels(items, 0, 10, inches=2.0, gap_in=0.17)
    assert [k[1] for k in kept] == ["atH", "pfam"]


def _panel(c, n, supported):
    return {"contig": c, "regions": [f"BGC{c}"],
            "matched": {f"t{c}_{i}": {"partner_verdict": "SUPPORTED" if i < supported else ""} for i in range(n)}}


def test_partners_rank_by_supported_finds():
    ps = [_panel("a", 5, 0), _panel("b", 2, 2), _panel("c", 3, 1)]
    assert [p["contig"] for p in sorted(ps, key=ss.partner_rank, reverse=True)[:2]] == ["b", "c"]


def test_a_partner_with_more_reference_genes_is_flagged(monkeypatch):
    st = {"table": [{"status": "PRESENT_IN_CORE"}] * 2}
    monkeypatch.setattr(ss, "rescue_panels", lambda *a: [_panel("x", 3, 1)])
    ml = ss.main_locus({}, "BGC1", st, {})
    assert ml and ml[1:] == (3, 2)
    monkeypatch.setattr(ss, "rescue_panels", lambda *a: [_panel("x", 2, 1)])
    assert ss.main_locus({}, "BGC1", {"table": [{"status": "PRESENT_IN_CORE"}]}, {}) is None  # fewer than 3


def test_audit_fit_flags_text_past_the_slide(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(ss.SW), Inches(ss.SH)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ss.text(s, 8.6, 0.95, 4.4, 1.0, [[("word " * 900, False, ss.INK)]], size=9)
    ss.text(s, 0.5, 0.95, 4.0, 1.0, [[("fits", False, ss.INK)]], size=9)
    prs.save(tmp_path / "d.pptx")
    bad = ss.audit_fit(tmp_path / "d.pptx")
    assert len(bad) == 1 and "past the slide bottom" in bad[0]


def test_fit_paras_moves_overflow_and_never_ends_on_a_heading():
    P = [[("x " * 300, False, ss.INK)], [("Heading", True, ss.ACCENT)], [("y " * 300, False, ss.INK)]]
    size, kept, rest = ss.fit_paras(P, 4.4, 1.2)
    assert size == 7.5 and kept == P[:1] and rest == P[1:]


def test_a_draft_relative_is_consistent_not_decisive():
    """A deck review: ten slides said "Two layers agree" on a draft relative, which cannot decide a split."""
    sc = {"layers_verdict": "CONSISTENT", "relative": "Streptomyces sp. QL37", "locus_gap_genes": "0"}
    assert "a draft cannot settle it" in ss.position_text(dict(sc, relative_assembly="draft"))
    assert "cannot settle" not in ss.position_text(dict(sc, relative_assembly="complete"))


def test_star_labels_never_overlap(tmp_path):
    """Codex Tasks 323/324: locus-tag labels on close PCoA stars overlapped. Labels now take the first free spot."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import strain_slides_pcoa as sp
    fig, ax = plt.subplots(figsize=(4, 3), dpi=100)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    pts = [(0.5, 0.5), (0.505, 0.5), (0.5, 0.505), (0.51, 0.51), (0.495, 0.495)]
    skipped = sp.place_labels(ax, pts, [f"ctg1_{i}" for i in range(5)])
    fig.canvas.draw()
    boxes = [t.get_window_extent() for t in ax.texts]
    assert len(boxes) + skipped == 5
    assert not any(a.overlaps(b) for i, a in enumerate(boxes) for b in boxes[i + 1:])
    plt.close(fig)
