"""tools/strain_slides.py and tools/strain_slides_pcoa.py: per-strain slide decks from a sources file.

2026-10-02: "for each strain slide deck, we want everything related to that strain", then "you still have loci that
need to be reversed" and "this shows two nodes ... as a contig rescue, but then under RGGMCI it lists a different rescue?".
The locks below hold the rules those messages produced, on a synthetic strain (no cohort data ships in tests):
- every antiSMASH region gets a slide; GECCO-only clusters are found by NODE_n_length_L even when GECCO rewrote the name;
- a gene strip is drawn in the orientation the map renderer recorded, not a guess;
- an RG-GMCI HIGH link the two-proof check rates WEAK is named as such and does not raise the region's rank;
- every region slide carries the review-draft line until a rescue verdict table is supplied, then the verdict instead;
- a deck that exists is never overwritten.
"""
import csv
import json
from pathlib import Path

import pytest

pytest.importorskip("pptx")
pytest.importorskip("matplotlib")

import sys
TOOLS = Path(__file__).resolve().parents[1] / "tools"
sys.path.insert(0, str(TOOLS))
import strain_slides as ss  # noqa: E402

S = "XS-001"
C1, C2 = "NODE_1_length_6000_cov_10.5", "NODE_2_length_3000_cov_9.5"


def _csv(path, rows, delim=","):
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]), delimiter=delim)
        w.writeheader()
        w.writerows(rows)


def _package(tmp):
    pkg = tmp / "package"
    pkg.mkdir(parents=True)
    (pkg / "manifest.json").write_text(json.dumps({"assembly": {"genome_bp": 9000, "contigs": 3, "n50": 6000, "gc_pct": 70.1},
                                                   "bgc_counts": {"raw": 2, "interior": 1, "edge": 0, "full_contig": 1}}))
    inv = lambda b, c, s, e, bd, prod, kcb: {"BGC_ID": b, "Contig": c, "antiSMASH_Region": "region001", "Start": s, "End": e,
                                           "Length_kb": f"{(e - s) / 1000:.1f}", "Boundary": bd, "Products": prod,
                                           "KCB_top": kcb, "KCB_score": "100.0"}
    _csv(pkg / f"{S}_2_inventory.csv", [inv("BGC001", C1, 1000, 4000, "Interior", "NRPS", "BGC0000001.1 | examplin | knownclusterblast #1"),
                                        inv("BGC002", C2, 0, 3000, "Full-contig", "NRPS", "")])
    genes = []
    for b, c, n in (("BGC001", C1, 1), ("BGC002", C2, 2)):
        for k in range(3):
            s = (1000 if n == 1 else 0) + k * 900
            genes.append({"bgc_id": b, "locus_tag": f"ctg{n}_{k + 1}", "bgc_start": 1000 if n == 1 else 0,
                          "bgc_end": 4000 if n == 1 else 3000, "cds_start": s, "cds_end": s + 800, "strand": "+",
                          "gene_function_inference": "core biosynthetic" if k == 1 else "biosynthetic context",
                          "sec_met_domains": "AMP-binding" if k == 1 else "", "aa_length": 266,
                          "product_qualifier": "AMP-binding (E-value: 1e-50, bitscore: 200, seeds: 1)" if k == 1 else ""})
    _csv(pkg / f"{S}_gene_by_gene_all_bgcs.csv", genes)
    _csv(pkg / f"{S}_domains.csv", [{"locus_tag": "ctg1_3", "feature_type": "PFAM_domain", "domain": "Thioesterase",
                                     "description": "Thioesterase domain", "evalue": "1e-20"}])
    _csv(pkg / f"{S}_4A_RGGMCI_ranked_pairs.csv", [{"bgc_a": "BGC002", "bgc_b": "BGC001", "contig_a": C2, "contig_b": C1,
                                                    "rggmci_confidence": "HIGH_RG_GMCI_RESCUE", "rggmci_score": "33",
                                                    "supporting_references": "30", "complementary_disjoint_refs": "0",
                                                    "overlapping_subject_refs": "30",
                                                    "terminus_override_note": "OVERRODE_OVERLAPPING_PARALOG_x"}])
    _csv(pkg / f"{S}_4D_two_proof_rescue.csv", [{"bgc_a": "BGC002", "bgc_b": "BGC001", "verdict": "WEAK"}])
    return pkg


def _gecco(tmp):
    g = tmp / "gecco"
    g.mkdir(parents=True)
    # GECCO rewrites the coverage part of the contig name; matching must use NODE_n_length_L
    gen = [{"sequence_id": "NODE_1_length_6000_cov_10.49", "protein_id": f"ctg1_{k + 1}", "start": 1000 + k * 900,
            "end": 1800 + k * 900, "strand": "+", "average_p": "0.95"} for k in range(3)]
    gen += [{"sequence_id": "NODE_3_length_5000_cov_8.1", "protein_id": f"ctg3_{k + 1}", "start": 100 + k * 900,
             "end": 900 + k * 900, "strand": "-", "average_p": "0.9"} for k in range(3)]
    _csv(g / f"{S}.genes.tsv", gen, "\t")
    _csv(g / f"{S}.features.tsv", [{"sequence_id": "NODE_3_length_5000_cov_8.1", "protein_id": "ctg3_2", "domain": "PF00109"}], "\t")
    _csv(g / f"{S}.clusters.tsv", [{"sequence_id": "NODE_1_length_6000_cov_10.49", "cluster_id": "c1", "start": 1000, "end": 3600,
                                    "type": "NRP", "average_p": "0.95"},
                                   {"sequence_id": "NODE_3_length_5000_cov_8.1", "cluster_id": "c2", "start": 100, "end": 2700,
                                    "type": "Polyketide", "average_p": "0.90"}], "\t")
    return g


def _sources(tmp, **extra):
    src = {"strain": S, "package": str(_package(tmp)), "gecco_dir": str(_gecco(tmp))}
    src.update(extra)
    return src


def _texts(deck):
    from pptx import Presentation
    return [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in Presentation(deck).slides]


def test_every_region_and_the_gecco_only_cluster_get_slides(tmp_path):
    r = ss.build(_sources(tmp_path), tmp_path / "out", "t1")
    t = _texts(tmp_path / "out" / r["deck"])
    assert r["regions"] == 2 and r["gecco_only"] == 1
    assert sum("BGC001 ·" in x for x in t) == 1 and sum("BGC002 ·" in x for x in t) == 1
    assert any("GECCO-only candidate on NODE_3_length_5000" in x and "KS N" in x for x in t)
    assert (tmp_path / "out" / f"{S}_gene_tables_t1_{r['deck'].split('_')[-1].replace('.pptx', '')}.pptx").exists()


def test_four_component_identity_is_copied_from_the_inventory(tmp_path):
    r = ss.build(_sources(tmp_path), tmp_path / "out", "t2")
    t = _texts(tmp_path / "out" / r["deck"])
    assert any(f"{S} / {C1} / region001 / BGC001" in x for x in t)


def test_a_weak_two_proof_link_is_named_and_does_not_raise_rank(tmp_path):
    D = ss.load(_sources(tmp_path))
    high, _ = ss.rgg_for(D, "BGC002")
    assert high == []                     # HIGH by score, WEAK by two-proof: not counted
    r = ss.build(_sources(tmp_path / "b"), tmp_path / "out", "t3")
    t = " ".join(_texts(tmp_path / "out" / r["deck"]))
    assert "1 high-confidence link by score" in t and "rates 1 as weak" in t and "two similar loci, not one split cluster" in t


def test_review_draft_line_until_verdicts_are_supplied(tmp_path):
    r = ss.build(_sources(tmp_path), tmp_path / "out", "t4")
    assert any("REVIEW DRAFT" in x for x in _texts(tmp_path / "out" / r["deck"]))
    v = tmp_path / "verdicts.tsv"
    _csv(v, [{"strain": S, "core identity": f"{S} / {C1} / region001 / BGC001", "partner contig": C2, "method": "RG-GMCI",
              "verdict": "TWO_SIMILAR_LOCI", "rule": "same reference genes in both"}], "\t")
    r = ss.build(_sources(tmp_path / "v", rescue_verdicts_tsv=str(v)), tmp_path / "out", "t5")
    t = _texts(tmp_path / "out" / r["deck"])
    assert not any("REVIEW DRAFT" in x for x in t)
    assert any("Gene-level review: two similar loci, not one cluster" in x for x in t)


def test_strip_orientation_comes_from_the_renderers_record(tmp_path, monkeypatch):
    src = _sources(tmp_path)
    gap = tmp_path / "gap"
    (gap / "BGC001_vs_BGC0000001").mkdir(parents=True)
    _csv(gap / "SUMMARY.tsv", [{"bgc": f"{S} / {C1} / region001 / BGC001", "check": "run", "reference": "BGC0000001",
                                "reference_name": "examplin", "folder": "BGC001_vs_BGC0000001", "partner_verdicts": "",
                                "split_genes": ""}], "\t")
    _csv(gap / "BGC001_vs_BGC0000001" / "gap_rescue.tsv",
         [{"name": f"g{k}", "status": "PRESENT_IN_CORE", "best_identity_pct": "60", "best_locus": f"ctg1_{3 - k}",
           "reference_product": "x"} for k in range(3)], "\t")
    out = tmp_path / "out"
    lay = out / f"{S}_t6_assets" / "maps" / "BGC001_vs_BGC0000001"
    lay.mkdir(parents=True)
    from PIL import Image
    Image.new("RGB", (40, 20), "white").save(lay / "gap_rescue.png")
    (lay / "map_layout.json").write_text(json.dumps({"flipped": {"NODE_1_length_6000": True}}))
    seen = {}
    real = ss.strip

    def spy(*a, **k):
        seen.setdefault("flip", []).append(k.get("flip"))
        return real(*a, **k)
    monkeypatch.setattr(ss, "strip", spy)
    ss.build(dict(src, gap_rescue_dir=str(gap)), out, "t6")
    assert True in seen["flip"]           # BGC001's strip follows the recorded reversal


def test_an_existing_deck_is_never_overwritten(tmp_path):
    src = _sources(tmp_path)
    ss.build(src, tmp_path / "out", "t7")
    with pytest.raises(SystemExit):
        ss.build(src, tmp_path / "out", "t7")


def test_template_lists_every_source_key(capsys):
    assert ss.main(["template"]) == 0
    keys = json.loads(capsys.readouterr().out)
    for k in ("strain", "package", "gecco_dir", "gap_rescue_dir", "rescue_verdicts_tsv", "trees", "family_figures"):
        assert k in keys


def test_pcoa_panels_mark_the_strain_and_count_its_tiers(tmp_path):
    import strain_slides_pcoa as sp
    kit = tmp_path / "kit" / "out_KS"
    kit.mkdir(parents=True)
    rows = [{"id": f"p{i}", "PC1": i / 10, "PC2": (i % 3) / 10, "n_represented": 1, "source": src, "group": g,
             "strain": st, "origin": ""}
            for i, (src, g, st) in enumerate([("reference_held", "", ""), ("MIBiG", "", ""), ("isolate", "ga", S),
                                              ("isolate", "ga", S), ("isolate", "gb", "XS-002")])]
    _csv(kit / "PCOA_KS.tsv", rows, "\t")
    _csv(kit / "NEAREST_KS.tsv", [{"id": "p2", "nearest_pident": "50"}, {"id": "p3", "nearest_pident": "90"}], "\t")
    (kit / "RUN_KS.json").write_text(json.dumps({"pct_axes": [20.0, 10.0]}))
    assert sp.main(["--kit", str(tmp_path / "kit"), "--strain", S, "--out", str(tmp_path / "p"),
                    "--groups", "ga:group A,gb:group B", "--sets-a", "KS", "--sets-b", "KS"]) == 0
    c = list(csv.DictReader(open(tmp_path / "p" / f"{S}_PCOA_COUNTS.tsv"), delimiter="\t"))[0]
    assert (c["n"], c["below_70"], c["from_70_to_85"], c["n_other_gb"]) == ("2", "1", "0", "1")
    assert (tmp_path / "p" / f"{S}_PCOA_biosynthetic_core.png").exists()


def test_biosynthetic_logic_names_found_and_unseen_steps_from_the_genes(tmp_path):
    """The owner, 2 Oct: slides need "gene function and what we can say about biosynthetic logic". The block lists each
    class step as found (with its genes) or not seen in the annotations, and never names a product."""
    D = ss.load(_sources(tmp_path))
    text = " ".join(t for para in ss.biosynthetic_logic(D, "BGC001") for t, _, _ in para)
    assert "NRPS: " in text and "adenylation (A) (ctg1_2)" in text
    assert "Not seen in the annotations" in text and "condensation (C)" in text
    assert "not a product call" in text or "unconfirmed" in text


def _kit_two_sets(tmp_path):
    """KS holds two genes of BGC001 (ctg1_2, ctg1_3) and one other gene of the strain; RES01, a CARD family, holds a
    supported partner gene (ctg9_4) on another contig."""
    for T, pts in (("KS", [("isolate", S, "ctg1_2", "50"), ("isolate", S, "ctg1_3", "90"), ("isolate", S, "ctg7_1", "95"),
                           ("isolate", "XS-002", "ctg1_2", "60"), ("MIBiG", "", "", ""), ("reference_held", "", "", "")]),
                   ("RES01", [("isolate", S, "ctg9_4", "75"), ("reference_held", "", "", "")])):
        d = tmp_path / "kit" / f"out_{T}"
        d.mkdir(parents=True)
        rows = [{"id": f"{T}{i}", "PC1": i / 10, "PC2": (i % 3) / 10, "n_represented": 1, "source": src, "group": "g",
                 "strain": st, "origin": "", "locus_tag": tag} for i, (src, st, tag, _) in enumerate(pts)]
        _csv(d / f"PCOA_{T}.tsv", rows, "\t")
        _csv(d / f"NEAREST_{T}.tsv", [{"id": f"{T}{i}", "nearest_pident": p} for i, (_, _, _, p) in enumerate(pts) if p], "\t")
        (d / f"RUN_{T}.json").write_text(json.dumps({"pct_axes": [20.0, 10.0]}))
    with open(tmp_path / "kit" / "RES_SETS.tsv", "w") as fh:
        fh.write("set\tamr_gene_family\tproteins\tisolate_proteins\nRES01\texample efflux pump\t2\t1\n")
    return tmp_path / "kit"


def test_a_bgc_figure_stars_only_that_bgcs_genes_and_its_supported_partners(tmp_path):
    """The owner, 3 Oct: "for each BGC I want the PCA plots that show where the genes from that specific BGC show up in
    the PCA plot (giant star at those node or something similar)", and the CARD plots too."""
    import strain_slides_pcoa as sp
    kit = _kit_two_sets(tmp_path)
    img, summ = sp.bgc_figure(kit, S, {"ctg1_2", "ctg1_3", "ctg9_4"}, tmp_path / "b.png", {})
    assert img.exists()
    got = {T: (n, nl, tags) for T, n, nl, tags in summ}
    assert got["KS"] == (2, 1, ["ctg1_2", "ctg1_3"])        # the strain's other KS gene and XS-002's same tag are not starred
    assert got["RES01"] == (1, 0, ["ctg9_4"])               # a CARD family set is read like any other
    assert sp.NAME["RES01"][0] == "CARD: example efflux pump"
    assert sp.bgc_figure(kit, S, {"ctg5_5"}, tmp_path / "c.png", {}) == (None, [])


def test_a_gene_with_no_identity_is_its_own_tier_not_85_or_more(tmp_path, monkeypatch):
    """A hostile review: a BGC gene with no best reference/MIBiG identity was filled white, the "85% or more"
    tier. It is now grey, with its own legend entry."""
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.axes import Axes
    import strain_slides_pcoa as sp
    kit = _kit_two_sets(tmp_path)
    near = kit / "out_KS" / "NEAREST_KS.tsv"
    near.write_text("\n".join(l for l in near.read_text().splitlines() if not l.startswith("KS1\t")) + "\n")  # ctg1_3
    fills, real = [], Axes.scatter

    def spy(self, x, y, *a, **k):
        if k.get("marker") == "*":
            fills.extend(k.get("c") or [])
        return real(self, x, y, *a, **k)
    monkeypatch.setattr(Axes, "scatter", spy)
    sp.bgc_figure(kit, S, {"ctg1_2", "ctg1_3"}, tmp_path / "b.png", {})
    assert sorted(fills) == sorted([sp.STRAIN, sp.NONE])  # ctg1_2 at 50%; ctg1_3 with no identity is grey, not white
