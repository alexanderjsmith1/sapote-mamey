"""test_e1_judgment_writeback.py — v9.7.67

Tests for update_e1_from_judgment: judgment_store → master workbook write path.
Verifies that completed Mode B records populate E1_Mode_B_Index and update
A4_Completeness_Audit in the master workbook.
"""
import csv
import json
import os
import tempfile
from pathlib import Path

import pytest

from mamey.judgment_store import init_register, record_mode_b, record_batch_complete
from mamey.master_workbook import update_e1_from_judgment, CANONICAL_V1_HEADERS


# ── Fixtures ──────────────────────────────────────────────────────────────────

@pytest.fixture
def pkg(tmp_path):
    (tmp_path / "manifest.json").write_text(
        json.dumps({"strain_id": "AS-TEST", "mode": "gold", "bgcs": []})
    )
    # Minimal triage board so update_e1 can read product/length metadata
    tb = tmp_path / "AS-TEST_4_triage_board.csv"
    with open(tb, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["BGC_ID", "Products", "Arch_Capacity", "Arch", "Start", "End",
                    "Boundary", "Lead_tier_auto"])
        w.writerow(["BGC001", "RiPP; azole-containing-RiPP", "NRPS", "A", "0", "18000", "Interior", "High"])
        w.writerow(["BGC002", "T1PKS", "PKS", "B", "20000", "60000", "Edge", "Medium"])
        w.writerow(["BGC003", "terpene", "Terpene", "C", "62000", "72000", "Interior", "Inventory"])
    return tmp_path


def _make_minimal_workbook(path: Path):
    """Write a minimal master workbook with canonical headers."""
    from openpyxl import Workbook as OWB
    wb = OWB()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    for sheet_name, headers in CANONICAL_V1_HEADERS.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(headers)
    # Pre-populate A4 for AS-TEST
    a4 = wb["A4_Completeness_Audit"]
    a4.append(["AS-TEST", "PASS", "PASS", "PASS", "PASS", "PASS", "PASS", "JUDGMENT_PENDING",
               "PASS", "PASS_EXTRACTION", "Load manifest for Sapote judgment"])
    # Pre-populate E1 with the "ALL" placeholder row
    e1 = wb["E1_Mode_B_Index"]
    e1.append(["AS-TEST", "ALL", "", "", "", "JUDGMENT_PENDING", "Sapote/Claude",
               "2026-06-17", "", "", ""])
    wb.save(str(path))


# ── Basic functionality ───────────────────────────────────────────────────────

def test_e1_rows_written_per_bgc(pkg, tmp_path):
    """After Mode B for BGC001, E1 has one row per BGC (not "ALL")."""
    wb_path = tmp_path / "master.xlsx"
    _make_minimal_workbook(wb_path)
    init_register(pkg, "AS-TEST", ["BGC001", "BGC002", "BGC003"])
    record_mode_b(pkg, "BGC001", "# §1\nContent.", "Lay text for BGC001.", "Ferm text.", "batch_1")

    result = update_e1_from_judgment(pkg, wb_path)

    from openpyxl import load_workbook as _lw
    wb = _lw(wb_path)
    e1 = wb["E1_Mode_B_Index"]
    # Collect all E1 rows for AS-TEST
    rows = [e1.cell(row=r, column=1).value for r in range(2, e1.max_row + 1)
            if e1.cell(row=r, column=1).value == "AS-TEST"]
    assert len(rows) == 3, f"Expected 3 per-BGC rows, got {len(rows)}"


def test_e1_no_all_placeholder_after_update(pkg, tmp_path):
    """The 'ALL' placeholder row must be replaced, not duplicated."""
    wb_path = tmp_path / "master.xlsx"
    _make_minimal_workbook(wb_path)
    init_register(pkg, "AS-TEST", ["BGC001", "BGC002"])
    record_mode_b(pkg, "BGC001", "# §1", "Lay.", "Ferm.", "batch_1")
    update_e1_from_judgment(pkg, wb_path)

    from openpyxl import load_workbook as _lw
    wb = _lw(wb_path)
    e1 = wb["E1_Mode_B_Index"]
    bgc_ids = [e1.cell(row=r, column=2).value for r in range(2, e1.max_row + 1)
               if e1.cell(row=r, column=1).value == "AS-TEST"]
    assert "ALL" not in bgc_ids, f"'ALL' placeholder still present: {bgc_ids}"


def test_complete_bgc_has_layperson_summary(pkg, tmp_path):
    """COMPLETE BGC row in E1 must carry layperson_summary."""
    wb_path = tmp_path / "master.xlsx"
    _make_minimal_workbook(wb_path)
    init_register(pkg, "AS-TEST", ["BGC001"])
    record_mode_b(pkg, "BGC001", "# §1\nContent.", "BGC001 encodes azole-RiPP capacity.", "Grow at 28C.", "batch_1")
    update_e1_from_judgment(pkg, wb_path)

    from openpyxl import load_workbook as _lw
    wb = _lw(wb_path)
    e1 = wb["E1_Mode_B_Index"]
    headers = [e1.cell(row=1, column=c).value for c in range(1, e1.max_column + 1)]
    lay_col = headers.index("layperson_summary") + 1
    for r in range(2, e1.max_row + 1):
        if e1.cell(row=r, column=2).value == "BGC001":
            val = e1.cell(row=r, column=lay_col).value
            assert val and "azole-RiPP" in val, f"layperson_summary missing or wrong: {val!r}"
            break


def test_pending_bgc_has_no_layperson_summary(pkg, tmp_path):
    """PENDING BGC row must not carry layperson_summary (nothing was written)."""
    wb_path = tmp_path / "master.xlsx"
    _make_minimal_workbook(wb_path)
    init_register(pkg, "AS-TEST", ["BGC001", "BGC002"])
    record_mode_b(pkg, "BGC001", "# §1", "BGC001 lay.", "BGC001 ferm.", "batch_1")
    update_e1_from_judgment(pkg, wb_path)

    from openpyxl import load_workbook as _lw
    wb = _lw(wb_path)
    e1 = wb["E1_Mode_B_Index"]
    headers = [e1.cell(row=1, column=c).value for c in range(1, e1.max_column + 1)]
    lay_col = headers.index("layperson_summary") + 1
    for r in range(2, e1.max_row + 1):
        if e1.cell(row=r, column=2).value == "BGC002":
            val = e1.cell(row=r, column=lay_col).value
            assert not val, f"PENDING BGC002 should not have layperson_summary, got {val!r}"
            break


def test_a4_updated_to_in_progress(pkg, tmp_path):
    """A4 E1_mode_b must reflect partial completion."""
    wb_path = tmp_path / "master.xlsx"
    _make_minimal_workbook(wb_path)
    init_register(pkg, "AS-TEST", ["BGC001", "BGC002", "BGC003"])
    record_mode_b(pkg, "BGC001", "# §1", "Lay.", "Ferm.", "batch_1")
    result = update_e1_from_judgment(pkg, wb_path)

    assert "IN_PROGRESS" in result["e1_audit_status"]
    from openpyxl import load_workbook as _lw
    wb = _lw(wb_path)
    a4 = wb["A4_Completeness_Audit"]
    headers = [a4.cell(row=1, column=c).value for c in range(1, a4.max_column + 1)]
    e1_col = headers.index("E1_mode_b") + 1
    for r in range(2, a4.max_row + 1):
        if a4.cell(row=r, column=1).value == "AS-TEST":
            val = a4.cell(row=r, column=e1_col).value
            assert "IN_PROGRESS" in str(val), f"A4 E1_mode_b should be IN_PROGRESS, got {val!r}"
            break


def test_a4_updated_to_complete_when_all_done(pkg, tmp_path):
    """A4 E1_mode_b must show COMPLETE when all BGCs have Mode B."""
    wb_path = tmp_path / "master.xlsx"
    _make_minimal_workbook(wb_path)
    init_register(pkg, "AS-TEST", ["BGC001", "BGC002", "BGC003"])
    for bgc in ["BGC001", "BGC002", "BGC003"]:
        record_mode_b(pkg, bgc, "# §1", f"{bgc} lay.", f"{bgc} ferm.", "batch_1")
    result = update_e1_from_judgment(pkg, wb_path)
    assert result["e1_audit_status"] == "COMPLETE"


def test_idempotent_second_call_replaces_rows(pkg, tmp_path):
    """Calling update_e1_from_judgment twice must not duplicate rows."""
    wb_path = tmp_path / "master.xlsx"
    _make_minimal_workbook(wb_path)
    init_register(pkg, "AS-TEST", ["BGC001", "BGC002"])
    record_mode_b(pkg, "BGC001", "# §1", "Lay1.", "Ferm1.", "batch_1")
    update_e1_from_judgment(pkg, wb_path)
    # Batch 2: BGC002 now complete
    record_mode_b(pkg, "BGC002", "# §1", "Lay2.", "Ferm2.", "batch_2")
    update_e1_from_judgment(pkg, wb_path)

    from openpyxl import load_workbook as _lw
    wb = _lw(wb_path)
    e1 = wb["E1_Mode_B_Index"]
    rows = [e1.cell(row=r, column=2).value for r in range(2, e1.max_row + 1)
            if e1.cell(row=r, column=1).value == "AS-TEST"]
    # Exactly 2 rows, not 4
    assert len(rows) == 2, f"Expected 2 rows after 2 batches, got {len(rows)}: {rows}"
    assert "BGC001" in rows and "BGC002" in rows


def test_missing_workbook_raises(pkg, tmp_path):
    """FileNotFoundError when workbook doesn't exist."""
    init_register(pkg, "AS-TEST", ["BGC001"])
    with pytest.raises(FileNotFoundError):
        update_e1_from_judgment(pkg, tmp_path / "nonexistent.xlsx")


def test_result_summary_counts(pkg, tmp_path):
    """Return dict has correct counts."""
    wb_path = tmp_path / "master.xlsx"
    _make_minimal_workbook(wb_path)
    init_register(pkg, "AS-TEST", ["BGC001", "BGC002", "BGC003"])
    record_mode_b(pkg, "BGC001", "# §1", "Lay.", "Ferm.", "batch_1")
    result = update_e1_from_judgment(pkg, wb_path)
    assert result["bgcs_written"] == 3
    assert result["bgcs_complete"] == 1
    assert result["bgcs_pending"] == 2
    assert result["strain"] == "AS-TEST"
