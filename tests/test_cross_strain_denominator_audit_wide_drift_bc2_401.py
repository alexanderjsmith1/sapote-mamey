"""BC2 .401 audit: tools/cross_strain_denominator_audit.py's own docstring documents a real
historical incident -- an 18-strain cohort snapshot survived, undetected, inside a 24-strain
master (drift = 6). The gate's own catch window was `n_reg - 12` to `n_reg + 12`: symmetric,
and calibrated for that drift magnitude. But a stale cohort denominator is, BY CONSTRUCTION,
always smaller than the live registry length (the cohort only grows) -- so as the cohort keeps
growing and a stale number sits unnoticed for longer, the drift only gets LARGER, and the old
symmetric window gets WORSE at its actual job exactly when it matters most.

Reproduced live against the real pristine script: a genuine stale denominator (17/18) inside a
44-strain registry (drift 26, well outside the old +/-12 window) reported a clean PASS -- the
identical defect class the tool's own docstring cites as its origin, just with more elapsed
drift than the original 24-strain incident.

Fixed by widening the lower bound to a small fixed floor (2) instead of `n_reg - 12`; the upper
bound is UNCHANGED. No existing test exercised the lower bound as protecting anything real --
the one legitimate "different ratio, not a cohort count" case this project's own test suite
documents (a BGC-total ratio, e.g. 164/1131) is already excluded by the exact-match `allowed`
set (`{n_reg, n_bgc}`), not by window proximity.

Honest trade-off, explicitly tested rather than hidden: this widening can now flag a genuinely
small, unrelated per-category ratio (e.g. "3/8 diagnostics") as a suspected stale cohort count,
where the old (narrower) window would have silently passed it. This is a deliberate choice
consistent with the tool's own stated "fail-closed invariant" design: a false positive here
costs a human a few seconds of review; the false negative this fix closes already shipped once
as a real incident. `test_small_category_ratio_now_over_flags_by_design` below documents this
trade-off explicitly so it is understood, not discovered by surprise.
"""
from __future__ import annotations

import os
import sys
import tempfile
import pathlib

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from cross_strain_denominator_audit import audit_denominators  # noqa: E402


def _make_master(path, findings_rows, n_strains=24, n_bgc=100):
    wb = openpyxl.Workbook()
    reg = wb.active
    reg.title = "A2_Strain_Registry"
    reg.append(["strain", "taxonomy"])
    for i in range(n_strains):
        reg.append([f"AS-{900 + i}", "sp."])
    b1 = wb.create_sheet("B1_BGC_Master")
    b1.append(["bgc_uid"])
    for i in range(n_bgc):
        b1.append([f"AS-9900:BGC{i}"])
    f = wb.create_sheet("Cross_Strain_Findings")
    f.append(["#", "finding", "value"])
    for r in findings_rows:
        f.append(r)
    wb.save(path)


def test_wide_drift_stale_denominator_now_caught():
    """The actual regression this fix closes: a stale denominator whose drift from the live
    registry length exceeds the old +/-12 window."""
    d = tempfile.mkdtemp()
    p = os.path.join(d, "wide_drift.xlsx")
    # a real, historically-plausible shape: an 18-strain-era stale denominator surviving into
    # a 44-strain master (drift 26 -- well outside the old +/-12 window).
    _make_master(p, [[3, "ectoine", "17/18"]], n_strains=44, n_bgc=500)
    violations = audit_denominators(p)
    assert len(violations) == 1, f"expected the wide-drift stale denominator to be caught; got {violations!r}"
    assert "Stale cohort count" in violations[0]
    assert "M=18" in violations[0]
    assert "cohort size is 44" in violations[0]


def test_narrow_drift_stale_denominator_still_caught():
    """No regression: the ORIGINAL documented incident shape (small drift, within the old
    window too) is still caught."""
    d = tempfile.mkdtemp()
    p = os.path.join(d, "narrow_drift.xlsx")
    _make_master(p, [[3, "universal", "18/18"], [4, "ectoine", "17/18 / 15/18"]], n_strains=24)
    violations = audit_denominators(p)
    assert len(violations) == 3
    assert all("Stale cohort count" in v for v in violations)


def test_small_category_ratio_now_over_flags_by_design():
    """The explicit, accepted trade-off: a genuine small per-category ratio unrelated to
    cohort size (not equal to n_reg or n_bgc) is now flagged too, where the old narrower
    window would have silently passed it. Documented here so this is a known, deliberate
    consequence of the fail-closed widening, not a silent surprise."""
    d = tempfile.mkdtemp()
    p = os.path.join(d, "small_category.xlsx")
    _make_master(p, [[9, "diagnostics", "3/8 diagnostics passed"]], n_strains=44, n_bgc=500)
    violations = audit_denominators(p)
    assert len(violations) == 1, (
        "expected the widened lower bound to flag this unrelated small ratio too -- if this "
        "assertion fails, the trade-off documented in this file's own docstring has changed "
        "and should be re-verified, not silently accepted"
    )


def test_bgc_total_ratio_still_not_flagged_even_far_from_registry():
    """No regression: a ratio over the BGC total (e.g. 164/1131) is still correctly excluded
    by the exact-match `allowed` set, regardless of how far it sits from the registry length --
    this protection was never provided by the window in the first place."""
    d = tempfile.mkdtemp()
    p = os.path.join(d, "bgc.xlsx")
    _make_master(p, [[7, "dark", "164/1131 (15%)"]], n_strains=24, n_bgc=1131)
    assert audit_denominators(p) == []
