"""Regression test for atomic workbook I/O (ledger F2).

An in-place `wb.save(path)` interrupted mid-write corrupts the workbook into an unreadable BadZipFile.
atomic_save writes to a sibling .tmp and os.replace()s it in, so an interrupted save never touches the original.
"""
import os, sys, tempfile
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tools"))
from _wbio import atomic_save


def test_original_survives_interrupted_save():
    d = tempfile.mkdtemp(); path = os.path.join(d, "wb.xlsx")
    wb = openpyxl.Workbook(); wb["Sheet"]["A1"] = "original"; wb.save(path)

    class Boom:                              # a workbook whose save fails mid-write
        def save(self, p):
            open(p, "w").write("half-written garbage")   # leaves a corrupt .tmp
            raise IOError("interrupted")

    try:
        atomic_save(Boom(), path)
    except IOError:
        pass

    # the original must still be a valid, readable xlsx with its original content
    assert openpyxl.load_workbook(path)["Sheet"]["A1"].value == "original"


def test_successful_save_replaces_content():
    d = tempfile.mkdtemp(); path = os.path.join(d, "wb.xlsx")
    wb = openpyxl.Workbook(); wb["Sheet"]["A1"] = "v1"; wb.save(path)
    wb2 = openpyxl.Workbook(); wb2["Sheet"]["A1"] = "v2"
    atomic_save(wb2, path)
    assert openpyxl.load_workbook(path)["Sheet"]["A1"].value == "v2"


if __name__ == "__main__":
    test_original_survives_interrupted_save()
    test_successful_save_replaces_content()
    print("atomic-save regression: both tests pass")
