"""Known metadata-only source counts and independent row-authority boundaries."""
import copy
import csv
import hashlib
import importlib.util
import json
import sys
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from mamey import antismash_evidence as ae
from mamey.antismash_tables import build_structured_tables
from mamey.release_policy import resolve_run_release
from mamey.report_release_policy import policy_for_master_rows, policy_for_records, row_policy

CONTIG = "CP123456.1"


def hit(location="[120:150](+)"):
    return {"location": location, "locus_tag": "invented_gene", "domain": "RRE",
            "identifier": "invented_family", "evalue": 1e-20, "score": 50.0}


def record(locations=("[120:150](+)",)):
    return {"id": CONTIG, "modules": {"antismash.modules.rrefinder": {
        "schema_version": 1, "record_id": CONTIG, "bitscore_cutoff": 25.0,
        "hits_by_protocluster": {}, "hits_by_cds": {"invented_gene": [hit(x) for x in locations]}}}}


def evidence(tmp_path, layout):
    first = record()
    members = {"one.json": [first]}
    if layout == "two_files": members["two.json"] = [copy.deepcopy(first)]
    elif layout == "two_records": members["one.json"].append(copy.deepcopy(first))
    elif layout == "coordinate_aliases": members["one.json"] = [record(("[120:150](+)", "[0120:0150](+)"))]
    elif layout == "empty": members["one.json"] = [record(())]
    path = tmp_path / "source.zip"
    with ZipFile(path, "w") as archive:
        for name, records in members.items():
            archive.writestr(name, json.dumps({"version": "8.0.4", "records": records}))
    return path, ae.parse_antismash_evidence(path, json_mode="off", include_structured=True)


def tables(path, ev):
    bgc = SimpleNamespace(bgc_id="BGC001", contig=CONTIG, start=100, end=1000,
                          products=["RiPP"], source_gbk="")
    return build_structured_tables(path, ev, [bgc])


@pytest.mark.parametrize("layout", ["two_files", "two_records", "coordinate_aliases"])
def test_valid_raw_duplicate_sources_are_complete_with_one_retained_row(tmp_path, layout):
    path, ev = evidence(tmp_path, layout)
    assert len(ev["rrefinder_hits"]) == ev["record_extraction"]["rrefinder_hit_count"] == 2
    out = tables(path, ev)
    assert len(out["rrefinder"]) == 1
    state = out["rrefinder_evidence_state"]
    assert out["rrefinder_status"] == state["state"] == "SOURCE_OBSERVED_HITS"
    assert (state["raw_source_rows"], state["retained_rows"], state["source_rows"]) == (2, 1, 1)
    assert state["source_completeness"] == "COMPLETE_FOR_SELECTED_SOURCE_RECORDS"
    assert state["upstream_search"] == "UNKNOWN"
    assert state["biological_absence"] == "NOT_ESTABLISHED"
    assert out["rrefinder"][0]["record_id"] == CONTIG
    assert out["rrefinder"][0]["mapping_status"] == "MAPPED"
    assert out["rrefinder"][0]["bgc_id"] == "BGC001"


@pytest.mark.parametrize("corrupt_count", [1, 3])
def test_raw_receipt_corruption_refuses_even_when_one_retained_row_matches(tmp_path, corrupt_count):
    path, ev = evidence(tmp_path, "two_files")
    ev["record_extraction"]["rrefinder_hit_count"] = corrupt_count
    out = tables(path, ev)
    assert len(out["rrefinder"]) == 1
    state = out["rrefinder_evidence_state"]
    assert state["state"] == "SOURCE_CONFLICT"
    assert (state["raw_source_rows"], state["retained_rows"]) == (2, 1)
    assert state["extraction_receipt"]["rrefinder_hit_count"] == corrupt_count


@pytest.mark.parametrize("layout,expected,count", [("single", "SOURCE_OBSERVED_HITS", 1),
                                                   ("empty", "SOURCE_OBSERVED_EMPTY", 0)])
def test_unrepeated_source_and_observed_empty_counts_remain_distinct(tmp_path, layout, expected, count):
    path, ev = evidence(tmp_path, layout)
    state = tables(path, ev)["rrefinder_evidence_state"]
    assert state["state"] == expected
    assert (state["raw_source_rows"], state["retained_rows"], state["source_rows"]) == (count, count, count)
    assert state["biological_absence"] == "NOT_ESTABLISHED"


def public_row(strain="reference"):
    decision = resolve_run_release(strain, "PUBLIC").to_dict()
    return {"strain": strain, **row_policy(policy_for_records([
        {"strain_id": strain, "release": "PUBLIC", "release_decision": decision}], strain))}


@pytest.mark.parametrize("legacy_first", [False, True])
def test_current_same_strain_authority_cannot_clear_legacy_master_row(legacy_first):
    current = public_row()
    legacy = {"strain": "reference", "release": "PUBLIC"}
    rows = [legacy, current] if legacy_first else [current, legacy]
    policy = policy_for_master_rows(rows)
    assert (policy["effective_release"], policy["privacy_assignment_state"]) == ("PRIVATE", "HOLD")


def test_each_explicit_same_strain_row_is_admitted_without_erasing_scope_conflicts():
    assert policy_for_master_rows([public_row(), public_row()])["effective_release"] == "PUBLIC"
    conflicting = public_row()
    decision = json.loads(conflicting["release_decision_json"])
    decision["reason"] = "different_unadmitted_authority"
    conflicting["release_decision_json"] = json.dumps(decision)
    assert policy_for_master_rows([public_row(), conflicting])["effective_release"] == "PRIVATE"


def exporter():
    root = Path(__file__).resolve().parents[1]
    sys.path.insert(0, str(root / "tools"))
    spec = importlib.util.spec_from_file_location("scope_count_figure_export", root / "tools/export_figure_ready.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("sheet", ["TIGRFAM_Check", "Cross_Strain_Findings"])
@pytest.mark.parametrize("scope,expected", [("explicit", "PUBLIC"), ("same_strain_legacy", "PRIVATE"),
                                           ("foreign_legacy", "PRIVATE"), ("missing_identity", "PRIVATE")])
def test_optional_export_data_has_independent_authority_and_whole_export_scope(tmp_path, sheet, scope, expected):
    public = public_row()
    wb = Workbook()
    a2 = wb.active; a2.title = "A2_Strain_Registry"
    policy_keys = list(row_policy(policy_for_records([], "unknown")))
    a2.append(["strain", "taxonomy", "bgc_count"] + policy_keys)
    a2.append(["reference", "Invented reference", 1] + [public[k] for k in policy_keys])
    b1 = wb.create_sheet("B1_BGC_Master")
    b1.append(["strain", "BGC_ID", "contig", "region", "products"] + policy_keys)
    b1.append(["reference", "BGC001", CONTIG, "region001", "RiPP"] + [public[k] for k in policy_keys])
    optional = wb.create_sheet(sheet)
    data = {"strain": "reference", **public}
    if scope == "same_strain_legacy": data = {"strain": "reference", "release": "PUBLIC"}
    if scope == "foreign_legacy": data = {"strain": "foreign_reference", "release": "PRIVATE"}
    if scope == "missing_identity": data = {k: v for k, v in public.items() if k != "strain"}
    fields = ["strain", "TIGR01454"] if sheet == "TIGRFAM_Check" else ["strain", "#", "finding", "metric", "value", "claim_status", "note"]
    data.update(TIGR01454=2, finding="Invented metadata finding", metric="count", value=2,
                claim_status="METHOD CAVEAT", note="Synthetic scope test")
    optional.append(fields + policy_keys)
    optional.append([data.get(k) for k in fields + policy_keys])
    source = tmp_path / "master.xlsx"; wb.save(source)
    before = source.read_bytes(); out = tmp_path / "fresh_export"
    exporter().main(str(source), str(out))
    assert source.read_bytes() == before
    policy = json.loads((out / "release_policy.json").read_text())
    assert policy["policy"]["effective_release"] == expected
    optional_name = "diagnostics_long.csv" if sheet == "TIGRFAM_Check" else "cross_strain_findings.csv"
    assert (out / optional_name).is_file()
    paths = sorted(out.glob("*.csv"))
    assert len(paths) == 6  # Five mandatory tables plus the selected optional table.
    for path in paths:
        with path.open(newline="") as handle:
            rows = list(csv.DictReader(handle))
        assert rows, path.name
        assert all(row["release"] == expected for row in rows)
        assert all(row["source_master_sha256"] == hashlib.sha256(before).hexdigest() for row in rows)
    with (out / optional_name).open(newline="") as handle:
        retained = list(csv.DictReader(handle))
    if sheet == "TIGRFAM_Check": assert retained[0]["n_hits"] == "2"
    else: assert retained[0]["finding"] == "Invented metadata finding"
